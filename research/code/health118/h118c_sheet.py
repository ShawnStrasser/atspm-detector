"""Note 118c: draw the v4 example charts from the saved v4c plot data (h118c_plotdata.py) and write
review/health_review_v4.xlsx FRESH (stale PNGs deleted first).

    python h118c_sheet.py [--charts-only]
"""
from __future__ import annotations

import json
import os
import sys
from pathlib import Path

import pandas as pd

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import health_charts as HC  # noqa: E402

DCW = Path(os.environ.get("DC_WORK") or (Path.home() / "dc_work"))
PD = DCW / "s118c" / "plot"
REPO = HERE.parents[2]
XL = REPO / "review" / "health_review_v4.xlsx"
CH = REPO / "review" / "health_review_v4_charts"
DC = REPO / "review" / "health_review_v4_day_charts"
CAT = {"stuck": "Stuck on", "dropout": "Goes silent", "level": "Count drops", "night_drop": "Misses vehicles at night",
       "choppy": "Erratic counts", "rapid": "Too-fast actuations", "volume": "Too many for the traffic",
       "chatter": "Chattering", "occspk": "Erratic time ON", "prof": "Busy at night",
       "shape24": "Unusual daily pattern", "count_on": "Count zone held ON", "occ_hi": "ON longer than its kind",
       "C1": "Extension time on a count zone", "C2": "Set to pulse but holds ON"}


def _dropped(dd, t, hours=4):
    """phase mates whose counts fell by >= 70 % across time t (mean per 15 min, 4 h each side, >= 5 before)."""
    t = pd.Timestamp(t)
    m = dd[~dd["self"]]
    b = m[(m.t >= t - pd.Timedelta(hours=hours)) & (m.t < t)].groupby("detector")["count"].mean()
    a = m[(m.t >= t) & (m.t < t + pd.Timedelta(hours=hours))].groupby("detector")["count"].mean()
    x = pd.concat([b.rename("b"), a.rename("a")], axis=1).dropna()
    x = x[x.b >= 5]
    return int((x.a <= 0.3 * x.b).sum()), float((x.a / x.b).clip(upper=1).sum())


def phase_note(r, dd, mk, ee):
    """does the whole phase do the same thing at the same time? (an event such as a closure, not one bad detector)
    Read from the saved day data only. Returns (short note, n mates the same, n mates) or None."""
    m = dd[~dd["self"]]
    nm = m.detector.nunique()
    if not r["flagged"] or nm == 0:
        return None
    c = r["category"]
    if c == "stuck":
        sp = [q for q in (mk or {}).get("spans", []) if q["label"] == "counted"]
        if not sp:
            return None
        q = max(sp, key=lambda z: z["dur_s"])
        a, b = pd.Timestamp(q["t0"]).ceil("15min"), pd.Timestamp(q["t1"])
        x = m[(m.t >= a) & (m.t + pd.Timedelta(minutes=15) <= b)]
        if x.empty:
            return None
        n = int((x.groupby("detector").pct_on.mean() >= 95).sum())
        if 2 * n >= nm:
            w = f"{HC.hm(q['t0'])}-{HC.hm(q['t1'])}"
            return (f"all {nm} phase mates ON too {w}" if n == nm else f"{n} of {nm} phase mates ON too {w}"), n, nm
    if c == "dropout":
        a, b = pd.Timestamp(r["silent_t0"]), pd.Timestamp(r["silent_t1"])
        x = m[(m.t >= a) & (m.t < b)].groupby("label")["count"].sum()
        sil = [k.split(":")[0] for k in x[x == 0].index]
        if len(sil) >= 2 or 2 * len(sil) >= nm:
            return "silent too: " + ", ".join(sil), len(sil), nm
    if c in ("level", "count_on"):
        if c == "level":
            t = r["drop_at"]
        else:
            h = ee[(ee.series == "hi") & ee.value.notna()]
            if h.empty:
                return None
            t0 = (h.t.min() - pd.Timedelta(minutes=7.5)).floor("15min")    # first 15 min held ON; the mates' sharpest
            sm = dd[~dd["self"]].groupby("t")["count"].sum()                     # drop within +-1 h
            step = sm.rolling(2).mean().shift(1) - sm[::-1].rolling(2).mean()[::-1]  # 30 min before - 30 min after
            step = step[(step.index >= t0 - pd.Timedelta(hours=1)) & (step.index <= t0 + pd.Timedelta(hours=1))]
            t = (step.idxmax() if step.notna().any() else t0) + pd.Timedelta(minutes=7.5)   # end of the first low bin
        n = _dropped(dd, t)[0]
        if 2 * n >= nm:
            who = f"all {nm} phase mates" if n == nm else f"{n} of {nm} phase mates"
            return f"{who} drop at {pd.Timestamp(t):%H:%M}" + (" too" if c == "level" else ""), n, nm
    return None


def reason(r):
    """ONE short line (<= 12 words), numbers that are on the chart."""
    c, f = r["category"], r["flagged"]
    pn = r.get("phase_n")
    if pn:                                               # the whole phase did it too: say so
        n, nm = pn
        if c == "stuck":
            w = r["phase_note"].split()[-1]
            return f"Held ON {r['n_ep']} times, {HC.dur_txt(r['tot_s'])}; " + ("whole phase ON " if n == nm else f"{n} of {nm} mates ON ") + w
        if c == "dropout":
            return f"No actuations {HC.hm(r['silent_t0'])}-{HC.hm(r['silent_t1'])}; {r['phase_note']}"
        if c == "level":
            return f"{r['own_before']:.0f} to {r['own_after']:.0f} per 15 min; " + ("all" if n == nm else f"{n} of {nm}") + " mates dropped too"
        if c == "count_on":
            return f"ON up to {r['hi_occ_max']:.0f} %, {r['hi_min']:.0f} min; " + ("all" if n == nm else f"{n} of {nm}") + f" mates dropped at {r['phase_note'].split()[-1]}"
    if c == "stuck":
        if f:
            return (f"Held ON {r['n_ep']} times, {HC.dur_txt(r['tot_s'])} in total" if r["n_ep"] > 1
                    else f"Held ON {HC.dur_txt(r['tot_s'])}")
        return "Long ONs while its phase mates were queued too"
    if c == "dropout":
        if f:
            return f"No actuations {HC.hm(r['silent_t0'])}-{HC.hm(r['silent_t1'])}; ~{r['expected']:,.0f} expected"
        return f"Quiet at night; ~{r['expected']:.0f} expected once faulty mates are left out"
    if c == "level":
        return f"{r['own_before']:.0f} to {r['own_after']:.0f} per 15 min; ~{r['exp_after']:.0f} expected"
    if c == "night_drop":
        if r.get("silent_night"):
            return f"No actuations {HC.hm(r['silent_t0'])}-{HC.hm(r['silent_t1'])}; ~{r['expected']:.0f} expected"
        return f"Night {r['night_n']:.0f} actuations vs ~{r['night_exp']:.0f} expected"
    if c == "choppy":
        return (f"{r['n_off']} of {r['n_sc']} periods outside the range" if f
                else f"All {r['n_sc']} periods inside the expected range")
    if c == "rapid":
        if f and r["x_spk"] >= r["x_zf"]:
            return f"{r['n_spk']} bursts; normal at most {int(r['lim_n_spk'])}"
        normal = r["fem_all"] + r["lim_zf"] * (r["fem_all"] + 1) ** 0.5
        if f:
            return f"{r['fo_all']:,.0f} back-to-back; normal up to {normal:,.0f}"
        return f"{r['fo_all']:,.0f} back-to-back vehicle ONs; normal up to {normal:,.0f}"
    if c == "volume":
        return f"Busiest {r['q5']:,.0f} veh/h/lane vs limit {r['lim_vol']:,.0f}"
    if c == "chatter":
        return f"{100 * r['chat_frac']:.1f} % chatter vs limit {100 * r['lim_chat']:.1f} %"
    if c == "occspk":
        return f"{r['n3_exc']:.0f} unexplained ON minutes vs limit {r['lim_n3']:.0f}"
    if c == "prof":
        if not f:
            return f"Night {100 * r['night_ratio']:.0f} % of busiest hours vs normal {100 * r['exp_ratio']:.0f} %; under limit"
        if r.get("occ_drove") and r.get("night_on") is not None:
            eo = max(r["exp_on"], r.get("typ_on_night") or 0)
            return f"ON {r['night_on']:.0f} % at night; " + ("normal under 1 %" if eo < 1 else f"its type about {eo:.0f} %")
        if r.get("sig_wide"):
            return f"Whole phase busy at night; {r['sig_odd_n']} of {r['sig_n']} signal detectors unusual"
        return f"Night {100 * r['night_ratio']:.0f} % of busiest hours vs normal {100 * r['exp_ratio']:.0f} %"
    if c == "shape24":
        span = f"{r['ev_h0']:02d}-{r['ev_h1'] + 1:02d} h"
        if not f:
            return f"Fewer counts {span} while ON {r['pct_run']:.0f} %: queue"
        who = "mates" if r["ev_from"] == "mates" else "type"
        return f"{r['ev_obs']:,.0f} actuations {span}; its {who} ~{r['ev_exp']:,.0f}"
    if c == "count_on":
        if f:
            return f"ON up to {r['hi_occ_max']:.0f} % in light traffic, {r['hi_min']:.0f} min"
        return "Held ON by extension (ONs again without OFF): config note"
    if c == "occ_hi":
        return f"{r['hi_min']:.0f} min ON longer than similar zones"
    if c == "C1":
        return f"{r['n_relog']:,} of {r['n_start'] + r['n_relog']:,} ONs logged again without OFF"
    if c == "C2":
        return f"{r['n_ge5']} ONs over 5 s, longest {HC.dur_txt(r['long_max'])}"
    return ""


def main(charts_only=False):
    rows = json.loads((PD / "rows.json").read_text())
    marks = json.loads((PD / "marks.json").read_text())
    day = pd.read_parquet(PD / "day15.parquet")
    ev = pd.read_parquet(PD / "ev.parquet")
    for p in (CH, DC):
        p.mkdir(parents=True, exist_ok=True)
        for f in p.glob("*.png"):
            f.unlink()
    out = []
    for r in rows:
        i = r["row"]
        # a config note is not a fault and not a healthy look-alike; keep the scorer's status when it says more
        if r["category"] in ("C1", "C2") and r["status"] == "ok":
            r = dict(r, status="config note")
        if r["category"] == "count_on" and not r["flagged"] and r.get("c1") and r["status"] == "ok":
            r = dict(r, status="config note")
        dd, ee = day[day.row == i], ev[ev.row == i]
        if r["category"] == "prof":  # the type's typical night % ON as drawn on the chart (01-05 h)
            o = ee[(ee.series == "otype_med") & ee.t.dt.hour.between(1, 4)]
            if len(o):
                r = dict(r, typ_on_night=float(o.value.mean()))
        pn = phase_note(r, dd, marks.get(str(i)), ee)
        if pn:
            r = dict(r, phase_note=pn[0], phase_n=(pn[1], pn[2]))
        stem = f"{i:02d}_{r['category']}_{r['signal']}_d{r['detector']}.png"
        title = HC.draw_category(r, dd, ee, marks.get(str(i)), CH / stem)
        HC.draw_day(r, dd, DC / stem)
        rs = reason(r)
        assert len(rs.split()) <= 12, (i, rs)
        out.append(dict(r, title=title, reason=rs, png=stem))
        print(i, title, "|", rs, flush=True)
    (PD / "sheet_rows.json").write_text(json.dumps(out, indent=1, default=str))
    if not charts_only:
        write_xlsx(out)


def write_xlsx(out):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Examples"
    head = [
        "Health v4: example charts per problem type. Is the status right and the chart clear? Y / N / ?",
        "ok = fine · watch = note only · suspect = probably a problem · bad = clearly a problem · config note = setup.",
        "Rows with status ok are healthy look-alikes the checks leave alone on purpose.",
    ]
    for k, h in enumerate(head, start=1):
        ws.cell(k, 1, h).font = Font(bold=(k == 1), size=11)
    cols = ["#", "Signal", "Detector", "Category", "Status", "Reason", "Chart", "Day chart", "Your answer", "Comment"]
    h0 = len(head) + 2
    for j, c in enumerate(cols, start=1):
        x = ws.cell(h0, j, c)
        x.font = Font(bold=True)
        x.fill = PatternFill("solid", fgColor="DDDDDD")
    yel = PatternFill("solid", fgColor="FFF2B3")
    for k, r in enumerate(out, start=1):
        rr = h0 + k
        vals = [k, r["signal"], r["label"], CAT[r["category"]], r["status"].replace("not_enough_data", "too little data"),
                r["reason"]]
        for j, v in enumerate(vals, start=1):
            ws.cell(rr, j, v).alignment = Alignment(vertical="top")
        for j, (folder, txt) in enumerate(((CH.name, "chart"), (DC.name, "day chart")), start=7):
            c = ws.cell(rr, j, txt)
            c.hyperlink = f"{folder}/{r['png']}"
            c.font = Font(color="0563C1", underline="single")
        for j in (9, 10):
            ws.cell(rr, j).fill = yel
    for col, w in zip("ABCDEFGHIJ", (4, 9, 24, 28, 11, 58, 8, 10, 12, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(h0 + 1, 1)
    wb.save(XL)
    print("wrote", XL, len(out), "rows")


if __name__ == "__main__":
    main(charts_only="--charts-only" in sys.argv)
