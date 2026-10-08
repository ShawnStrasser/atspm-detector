"""Note 118b: draw the v4 example charts from the saved plot data and write review/health_review_v4.xlsx (fresh sheet).

    python h118b_sheet.py [--charts-only]
"""
from __future__ import annotations

import json
import os
import sys
from pathlib import Path

import pandas as pd

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import health_charts as HC  # noqa: E402

DCW = Path(os.environ.get("DC_WORK") or (Path.home() / "dc_work"))
PD = DCW / "s118b" / "plot"
REPO = HERE.parents[2]
XL = REPO / "review" / "health_review_v4.xlsx"
CH = REPO / "review" / "health_review_v4_charts"
DC = REPO / "review" / "health_review_v4_day_charts"
CAT = {"stuck": "Stuck on", "dropout": "Goes silent", "level": "Count drops", "night_drop": "Misses vehicles at night",
       "choppy": "Erratic counts", "rapid": "Too-fast actuations", "volume": "Too many for the traffic",
       "chatter": "Chattering", "occspk": "Erratic time ON", "prof": "Busy at night", "occ_hi": "ON longer than its kind",
       "C1": "Extension time on a count zone", "C2": "Set to pulse but holds ON"}


def reason(r):
    """ONE short line, numbers that are on the chart."""
    c, f = r["category"], r["flagged"]
    if c == "stuck":
        if f:
            return (f"Stuck on {r['n_ep']} times, {HC.dur_txt(r['tot_s'])} in total" if r["n_ep"] > 1
                    else f"Stuck on for {HC.dur_txt(r['tot_s'])}")
        return "Long ON during a phase-wide queue; not counted"
    if c == "dropout":
        return f"No actuations {HC.hm(r['silent_t0'])}-{HC.hm(r['silent_t1'])}; ~{r['expected']:,.0f} expected"
    if c == "level":
        return f"{r['own_before']:.0f} to {r['own_after']:.0f} per 15 min; ~{r['exp_after']:.0f} expected"
    if c == "night_drop":
        return f"Night {r['night_n']:.0f} actuations vs ~{r['night_exp']:.0f} expected"
    if c == "choppy":
        return (f"{r['n_off']} of {r['n_sc']} periods outside the range" if f
                else f"All {r['n_sc']} periods inside the expected range")
    if c == "rapid":
        if f and r["x_spk"] >= r["x_zf"]:
            return f"{r['n_spk']} bursts; normal at most {int(r['lim_n_spk'])}"
        normal = r["fem_all"] + r["lim_zf"] * (r["fem_all"] + 1) ** 0.5
        return f"{r['fo_all']:,.0f} back-to-back; normal up to {normal:,.0f}" + ("" if f else "; heavy traffic")
    if c == "volume":
        return f"Busiest {r['q5']:,.0f} veh/h/lane vs limit {r['lim_vol']:,.0f}"
    if c == "chatter":
        return f"{100 * r['chat_frac']:.1f} % chatter vs limit {100 * r['lim_chat']:.1f} %"
    if c == "occspk":
        return f"{r['n3_exc']:.0f} unexplained ON minutes vs limit {r['lim_n3']:.0f}"
    if c == "prof":
        return (f"Night {100 * r['night_ratio']:.0f} % of busiest hours vs normal {100 * r['exp_ratio']:.0f} %"
                + ("" if f else "; just under limit"))
    if c == "occ_hi":
        return f"{r['hi_min']:.0f} min ON longer than similar zones"
    if c == "C1":
        return f"{r['n_relog']:,} of {r['n_start'] + r['n_relog']:,} ONs logged again without OFF"
    if c == "C2":
        return f"{r['n_ge5']} ONs over 5 s, longest {HC.dur_txt(r['long_max'])}"
    return ""


def main(charts_only=False):
    rows = json.loads((PD / "rows.json").read_text())
    marks = json.loads((PD / "marks.json").read_text())
    day = pd.read_parquet(PD / "day15.parquet")
    ev = pd.read_parquet(PD / "ev.parquet")
    for p in (CH, DC):
        p.mkdir(parents=True, exist_ok=True)
        for f in p.glob("*.png"):
            f.unlink()
    out = []
    for r in rows:
        i = r["row"]
        if r["category"] in ("C1", "C2"):  # a config note is not a fault and not a healthy look-alike
            r = dict(r, status="config note")
        dd, ee = day[day.row == i], ev[ev.row == i]
        stem = f"{i:02d}_{r['category']}_{r['signal']}_d{r['detector']}.png"
        title = HC.draw_category(r, dd, ee, marks.get(str(i)), CH / stem)
        HC.draw_day(r, dd, DC / stem)
        out.append(dict(r, title=title, reason=reason(r), png=stem))
        print(i, title, "|", reason(r), flush=True)
    (PD / "sheet_rows.json").write_text(json.dumps(out, indent=1, default=str))
    if not charts_only:
        write_xlsx(out)


def write_xlsx(out):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Examples"
    head = [
        "Health v4: one example chart per problem type. Is the chart clear, and is the status right? Answer Y / N / ?.",
        "Status: ok = fine · watch = note only · suspect = probably a problem · bad = clearly a problem. "
        "Config note = setup issue, not a fault.",
        "Rows with status ok are healthy look-alikes the check correctly leaves alone.",
    ]
    for k, h in enumerate(head, start=1):
        ws.cell(k, 1, h).font = Font(bold=(k == 1), size=11)
    cols = ["#", "Signal", "Detector", "Category", "Status", "Reason", "Chart", "Day chart", "Your answer", "Comment"]
    h0 = len(head) + 2
    for j, c in enumerate(cols, start=1):
        x = ws.cell(h0, j, c)
        x.font = Font(bold=True)
        x.fill = PatternFill("solid", fgColor="DDDDDD")
    yel = PatternFill("solid", fgColor="FFF2B3")
    for k, r in enumerate(out, start=1):
        rr = h0 + k
        vals = [k, r["signal"], r["label"], CAT[r["category"]], r["status"].replace("not_enough_data", "too little data"),
                r["reason"]]
        for j, v in enumerate(vals, start=1):
            ws.cell(rr, j, v).alignment = Alignment(vertical="top")
        for j, (folder, txt) in enumerate(((CH.name, "chart"), (DC.name, "day chart")), start=7):
            c = ws.cell(rr, j, txt)
            c.hyperlink = f"{folder}/{r['png']}"
            c.font = Font(color="0563C1", underline="single")
        for j in (9, 10):
            ws.cell(rr, j).fill = yel
    for col, w in zip("ABCDEFGHIJ", (4, 9, 24, 28, 9, 52, 8, 10, 12, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(h0 + 1, 1)
    wb.save(XL)
    print("wrote", XL, len(out), "rows")


if __name__ == "__main__":
    main(charts_only="--charts-only" in sys.argv)
