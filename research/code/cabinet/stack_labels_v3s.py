"""Label table variant v3s (note 54): the user's 2026-09-30 rule "ATSPM classes, one per lane" on the label side.

    python stack_labels_v3s.py [--only-yes]
        -> research/labels/function_labels_v3s.parquet   (git-ignored, like every label file)
           %DC_WORK%/cabinet/dec_role_changed_v3s.parquet  (note-52 table recomputed on the v3s labels)
           %DC_WORK%/cabinet/stacked/v3s_summary.json

Starts from function_labels_v3 (unchanged) and applies, each with its own flag column:
  1. stack_relabel   every member of a stacked group (stacked_review.py groups.parquet: same timing phase, different
                     technology, ONs starting together, labels agreeing on the role or one plain Other) carries the
                     group's role: `function`, `label_print_first`, `print_function` -> role; `truth_v3s` -> role
                     (a confirmed proposal counts as a user ruling). Answers in review/stacked_detectors_review.xlsx
                     are honoured: "no" -> group left as it is; "?" -> members out of training AND scoring
                     (`exclude_train_score`); "yes" or blank -> applied (blank flagged `stack_unreviewed`;
                     --only-yes applies "yes" rows only). `stack_group` / `stack_role` are set for every member,
                     relabelled or not, so the scorer can give stack-aware credit.
  2. yr_count_identical   a Yellow_Red zone whose actuations are identical to a Count zone on the same phase
                     (counts within 1 %, >= 98 % of ONs both ways within 0.3 s; stacked_review.yr_count) ->
                     `exclude_train_score` (left out of training and scoring; listed for staff).
  3. readmit_radar_over_loops   signals flagged unusual_layout ONLY because live radar (or video) zones sit over live
                     loops -> unusual_layout False, train_use = labelled & not dead. Signals flagged for any other
                     reason (rewired / stale print, two intersections, railroad pre-signal, radar reusing loop
                     channels) stay unusual. `unusual_layout_v3` keeps the old flag.
  4. Dec-role filter, function/phase rule only (orchestrator 2026-09-30; the 'absent' rule is NOT used): the note-52
     table is recomputed against the v3s labels (a relabelled loop now carries the radar's class) and written to
     dec_role_changed_v3s.parquet with `dec_role_changed_*` = status in {function, phase}; v3_retrain.py
     --labels v3s --dec-role reads it.
Cleansing only (prints, technology, card structure are allowed here); nothing is a model input. locked_v2 asserted absent.
"""
from __future__ import annotations

import argparse
import json
import re
import sys

import numpy as np
import pandas as pd

import os

from cab_common import CAB, DC_WORK, REPO, STORE

# note 80: with DC_CAB_STORE set, input = cab_final's function_labels_<tag>_base, output = function_labels_<tag> (v4l)
TAG = os.environ.get("DC_LABELS_TAG", "v4l")
V3 = REPO / (f"research/labels/function_labels_{TAG}_base.parquet" if STORE else "research/labels/function_labels_v3.parquet")
V3S = REPO / (f"research/labels/function_labels_{TAG}.parquet" if STORE else "research/labels/function_labels_v3s.parquet")
ST = CAB / "stacked"
XL = REPO / "review/stacked_detectors_review.xlsx"
DEC_V3S = CAB / (f"dec_role_changed_{TAG}.parquet" if STORE else "dec_role_changed_v3s.parquet")
C7 = ["Advance", "Presence", "Count", "Yellow_Red", "Mid", "Bike", "Other"]
RADAR_OVER = re.compile(r"(radar|video).{0,80}(over|alongside) (the |a )?live|live loops.{0,60}under the radar", re.I)
NOT_ONLY = re.compile(r"rewired|two intersections|2 intersections|print (does not|no longer|outdated)|predates|"
                      r"transitional|pre-signal|reuse channel|added after the print|same channels", re.I)
# flagged in a record note only (v3_retrain EXTRA_UNUSUAL); reasons read from the notes (note 28 / records)
NOTE_ONLY = {"04073": "live radar advance zones over live advance loops on ph2",
             "13025": "live radar count-advance zones over live advance loops on the same lanes",
             "2C023": "hybrid radar stop-bar zones + live loops (only ph1 shares a lane); not flagged in the record",
             "01064": "railroad pre-signal stop lines", "10086": "rewired: print MT numbers moved to other channels",
             "12032": "heavily rewired", "2C042": "radar zones reuse channel numbers of live loops; stale config"}


def locked() -> set[str]:
    return set(pd.read_csv(DC_WORK / "official/locked_v2.csv").DeviceId.astype(str).str.lower())


def unusual_reasons(names) -> pd.DataFrame:
    rows = []
    for dn in sorted(names):
        # released NEWTEST signals (note 36) keep their records in the locked store
        f = next((x for x in ((CAB / "signals" / f"{dn}.json",) if STORE else
                              (DC_WORK / "cabinet/signals" / f"{dn}.json", DC_WORK / "cabinet_locked/signals" / f"{dn}.json"))
                  if x.exists()), None)
        s = json.loads(f.read_text()) if f else {}
        why = s.get("signal_flags_reason") or ""
        note = s.get("note") or ""
        if not why and re.search(r"\bunusual_layout\b", note):
            why = note
        why = why or NOTE_ONLY.get(dn, "")
        ok = bool(RADAR_OVER.search(why)) and not NOT_ONLY.search(why)
        if dn == "2C023":
            ok = True        # hybrid: loops and radar overlap only on ph1 (record: "not flagged unusual_layout")
        rows.append(dict(DeviceName=dn, reason=why[:200], readmit=ok))
    return pd.DataFrame(rows)


def honest_truth(v3: pd.DataFrame) -> pd.Series:
    """Scoring truth of notes 45 / 49 / 51: user ruling > high-confidence print > config label where no print."""
    pf, src = v3.print_function, v3.source.astype("string")
    user = src.eq("user_ruling").fillna(False)
    high = v3.print_source.eq("print") & v3.print_confidence.eq("high") & pf.isin(C7)
    noprint = ~v3.print_source.eq("print").fillna(False) | pf.isna()
    return pd.Series(np.where(user, v3.function, np.where(high, pf, np.where(noprint & v3.func7_v2.notna(),
                                                                              v3.func7_v2, None))), index=v3.index)


def review_answers() -> dict:
    """(DeviceName, phase, members tuple) -> 'yes' | 'no' | '?' | '' from the user's review workbook."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(XL, read_only=True, data_only=True)
    except Exception as e:                                       # noqa: BLE001 (open in Excel, missing ...)
        print(f"review workbook not readable ({e}); every group treated as unanswered", file=sys.stderr)
        return {}
    ws = wb["stacked groups"]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    ia, ig, ip, isg = (hdr.index("Your answer (yes / no / ?)"), hdr.index("Detectors in the group"),
                       hdr.index("Phase"), hdr.index("Signal"))
    out = {}
    for r in it:
        if r[isg] is None:
            continue
        mem = tuple(sorted(int(x) for x in re.findall(r"(\d+) \(", str(r[ig]))))
        a = str(r[ia] or "").strip().lower()
        a = "yes" if a.startswith("y") else "no" if a.startswith("n") else "?" if a.startswith("?") else a
        out[(str(r[isg]), str(r[ip]), mem)] = a
    return out


def main(only_yes: bool):
    v3 = pd.read_parquet(V3)
    lk = locked()
    assert not v3.DeviceId.str.lower().isin(lk).any(), "locked signal in v3"
    d = v3.copy()
    d["dev"] = d.DeviceId.str.lower()
    for c in ("function", "label_print_first", "print_function", "train_use", "unusual_layout"):
        d[f"{c}_v3"] = d[c]
    d["truth_v3"] = honest_truth(d)
    d["truth_v3s"] = d.truth_v3.copy()
    for c in ("stack_relabel", "stack_unreviewed", "yr_count_identical", "readmit_radar_over_loops",
              "exclude_train_score"):
        d[c] = False
    d["stack_group"], d["stack_role"], d["stack_answer"] = None, None, None
    idx = {k: i for i, k in enumerate(zip(d.dev, d.detector.astype(int)))}
    # ---- 1. stacked groups
    G = pd.read_parquet(ST / "groups.parquet")
    assert not G.dev.isin(lk).any()
    ans = review_answers()
    n = dict(groups=len(G), applied=0, no=0, q=0, unreviewed=0, rows_relabelled=0)
    for g in G.itertuples():
        mem = tuple(sorted(int(m) for m in g.members))
        a = ans.get((str(g.DeviceName), str(g.phase), mem), "")
        gid = f"{g.DeviceName}_{g.phase}_{g.role}_{'-'.join(map(str, mem))}"
        if any((g.dev, m) not in idx for m in mem):   # a member no longer in the table (store copy, note 80)
            n["missing_member"] = n.get("missing_member", 0) + 1
            continue
        rows = [idx[(g.dev, m)] for m in mem]
        d.iloc[rows, d.columns.get_loc("stack_answer")] = a or "blank"
        if a == "no" or (only_yes and a != "yes"):
            n["no"] += a == "no"
            continue
        if a == "?":
            d.iloc[rows, d.columns.get_loc("exclude_train_score")] = True
            n["q"] += 1
            continue
        n["applied"] += 1
        n["unreviewed"] += a != "yes"
        for i in rows:
            d.iat[i, d.columns.get_loc("stack_group")] = gid
            d.iat[i, d.columns.get_loc("stack_role")] = g.role
            d.iat[i, d.columns.get_loc("stack_unreviewed")] = a != "yes"
            ch = d.function.iat[i] != g.role or d.truth_v3.iat[i] != g.role
            d.iat[i, d.columns.get_loc("stack_relabel")] = bool(ch)
            n["rows_relabelled"] += bool(ch)
            for c in ("function", "label_print_first", "print_function", "truth_v3s"):
                d.iat[i, d.columns.get_loc(c)] = g.role
    # ---- 2. YR identical to a Count zone
    Y = pd.read_parquet(ST / "yr_count_pairs.parquet")
    Y = Y[Y.identical]
    for dev, yr in zip(Y.dev, Y.yr.astype(int)):
        if (dev, yr) not in idx:
            continue
        i = idx[(dev, yr)]
        d.iat[i, d.columns.get_loc("yr_count_identical")] = True
        d.iat[i, d.columns.get_loc("exclude_train_score")] = True
    # ---- 3. radar-over-loops signals back into training
    un = d[d.unusual_layout.astype(bool)].DeviceName.dropna().astype(str).unique()
    R = unusual_reasons(set(un) | {"2C023", "04073", "13025"})
    re_names = set(R.DeviceName[R.readmit])
    m = d.DeviceName.astype(str).isin(re_names)
    d.loc[m, "readmit_radar_over_loops"] = True
    d.loc[m, "unusual_layout"] = False
    d.loc[m, "train_use"] = d.loc[m, "function"].notna() & ~d.loc[m, "dead"].eq(True)
    # train_use for relabelled rows on usual signals follows the new label
    rl = d.stack_relabel & ~d.unusual_layout.astype(bool)
    d.loc[rl, "train_use"] = ~d.loc[rl, "dead"].eq(True)
    d.loc[d.exclude_train_score, "train_use"] = False
    d["train_use_validated"] = d.train_use & d.validated.eq("pass").fillna(False).astype(bool)
    # ---- 4. Dec-role, function / phase rule only, on the v3s labels
    import cab_final as CF
    a_ = CF.dec_role_changed(d, d.function)
    b_ = CF.dec_role_changed(d, d.truth_v3s)
    dr = a_.rename(columns={"label": "label_train", "dec_status": "dec_status_train"}).drop(columns="dec_role_changed")
    dr["dec_role_changed_train"] = a_.dec_status.isin(["function", "phase"]).to_numpy() & a_.label.notna().to_numpy()
    dr["label_truth"], dr["dec_status_truth"] = b_.label.to_numpy(), b_.dec_status.to_numpy()
    dr["dec_role_changed_truth"] = b_.dec_status.isin(["function", "phase"]).to_numpy() & b_.label.notna().to_numpy()
    dr["rule"] = "function_phase"
    assert not dr.DeviceId.isin(lk).any()
    dr.to_parquet(DEC_V3S, index=False)
    d["dec_fp_train"] = dr.dec_role_changed_train.to_numpy()
    d["dec_fp_truth"] = dr.dec_role_changed_truth.to_numpy()
    d = d.drop(columns="dev")
    d.to_parquet(V3S, index=False)
    s = dict(stack=n, stack_rows_by_change=d[d.stack_group.notna()].groupby("stack_relabel").size().to_dict(),
             yr_identical=int(d.yr_count_identical.sum()), exclude_train_score=int(d.exclude_train_score.sum()),
             readmitted_signals=sorted(re_names), readmitted_rows=int(m.sum()),
             still_unusual=sorted(set(R.DeviceName[~R.readmit])),
             train_use=[int(v3.train_use.sum()), int(d.train_use.sum())],
             train_use_validated=[int(v3.train_use_validated.sum()), int(d.train_use_validated.sum())],
             dec_fp_train=int(d.dec_fp_train.sum()), dec_fp_truth=int(d.dec_fp_truth.sum()),
             unusual_reasons=R.to_dict("records"))
    json.dump(s, open(ST / (f"{TAG}_summary.json" if STORE else "v3s_summary.json"), "w"), indent=1, default=str)
    print(json.dumps({k: v for k, v in s.items() if k != "unusual_reasons"}, indent=1, default=str))
    print(R.to_string())


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--only-yes", action="store_true", help="apply only the groups the user answered 'yes'")
    main(ap.parse_args().only_yes)
