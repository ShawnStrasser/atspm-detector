"""Stacked-detector review, step 2: groups -> chart data -> charts -> review/stacked_detectors_review.xlsx.

    python stacked_review.py          (needs stacked_pairs.py output)

A STACKED group = 2+ detectors of different technology (loop / radar / video) on the same timing phase whose ONs
start together (the lower-count member's ONs start within +-1.5 s of the other's, well above chance) and whose
labels say the same role (or one says Other). New rule under review (user, 2026-09-30): every member carries the real
class; the model outputs at most one of each class per lane and the other member counts as Other, scored correct.
Charts are drawn from saved sample data (%DC_WORK%/cabinet/stacked/chart_data/). locked_v2 signals never enter.
Cleansing / review only: nothing here is a model input.
"""
from __future__ import annotations

import sys
import urllib.parse
from pathlib import Path

import duckdb
import numpy as np
import pandas as pd

from cab_common import DC_WORK, REPO
from stacked_pairs import IVF, OUT as SOUT, load_labels, TOL

XL = REPO / "review/stacked_detectors_review.xlsx"
CHARTS = REPO / "review/stacked_detectors_charts"
DATA = SOUT / "chart_data"
PDFS = [DC_WORK / "cabinet/pdf", DC_WORK / "cabinet_locked/pdf"]
PM = ["Advance", "Presence", "Count", "Yellow_Red", "Mid"]
TECH = {"loop", "radar", "video"}
C = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100"]
INK, INK2, GRID, SURF = "#0b0b0b", "#52514e", "#e4e3df", "#fcfcfb"
ROLE = {"Presence": "stop-bar presence", "Count": "stop-bar count", "Yellow_Red": "yellow-red (stop bar)",
        "Advance": "advance", "Mid": "mid (between advance and stop bar)"}
# stacked edge thresholds (lower-count member into the higher-count one)
EXCESS_MIN, RATIO_MIN, LAG_MAX, OVL_MAX, OVL_MIN, MIN_ACT = 0.35, 0.35, 1.0, 0.70, 0.30, 200
# Yellow_Red vs Count, sheet 2
IDENT = dict(ratio=0.01, match=0.98)  # identical: counts within 1 %, >= 98 % of ONs both ways start within 0.3 s
NEAR = dict(ratio=0.05, match=0.90)   # near-identical: counts within 5 %, >= 90 % both ways within 1.5 s


def role_txt(role: str) -> str:
    return ROLE.get(role) or (" or ".join(ROLE.get(x, x).replace("stop-bar ", "") for x in role.split(" or ")) + "?")


def lane_txt(r) -> str:
    li, ls, lt = r.lane_index, r.lanes_spanned, r.lane_type
    s = ""
    if pd.notna(ls) and int(ls) >= 2:
        s = f"spans {int(ls)} lanes"
    elif pd.notna(li):
        s = f"lane {int(li)}"
    if lt and isinstance(lt, str):
        s = (s + f" ({lt})").strip()
    return s or "-"


def edges(pairs: pd.DataFrame, lab: pd.DataFrame) -> pd.DataFrame:
    k = lab.set_index(["dev", "detector"])
    cols = ["technology", "function", "print_subtype", "unusual_layout"]
    p = pairs.copy()
    for side in ("a", "b"):
        x = k.reindex(list(zip(p.dev, p[side])))[cols].reset_index(drop=True)
        x.columns = [f"{c}_{side}" for c in cols]
        p = pd.concat([p.reset_index(drop=True), x], axis=1)
    # orient: s = smaller count, g = bigger count
    sw = p.n_a > p.n_b
    p["s"], p["g"] = np.where(sw, p.b, p.a), np.where(sw, p.a, p.b)
    p["m_sg"] = np.where(sw, p.match_ba, p.match_ab)
    p["m_gs"] = np.where(sw, p.match_ab, p.match_ba)
    p["chance_sg"] = np.where(sw, p.chance_ba, p.chance_ab)
    p["ratio"] = np.minimum(p.n_a, p.n_b) / np.maximum(p.n_a, p.n_b)
    p["excess"] = (p.m_sg - p.chance_sg) / (1 - p.chance_sg).clip(lower=1e-6)
    return p


PROMOTABLE = {"superseded_by_radar", "unexplained", "explained", None}  # Other rows that are no zone type of their own


def role_of(fa, fb, pa, pb, sa, sb):
    """Shared role of a pair from its labels, or None. An Other member must be a plain Other (e.g. a loop superseded by
    the radar), never a zone type the user ruled Other (advance-presence, long zone, 20-75 ft, extension ...) or Bike.
    Presence / Count / Yellow_Red roles need compatible ON durations (a pulse zone and a holding zone differ)."""
    oa = fa not in PM
    ob = fb not in PM
    if oa and ob:
        return None
    if (oa and (fa not in ("Other", None) or (sa if isinstance(sa, str) else None) not in PROMOTABLE)) or        (ob and (fb not in ("Other", None) or (sb if isinstance(sb, str) else None) not in PROMOTABLE)):
        return None
    if not oa and not ob and fa != fb:
        return None
    r = fb if oa else fa
    if r in ("Presence", "Count", "Yellow_Red") and abs(pa - pb) > 0.5:
        return None
    return r


def build_groups(p: pd.DataFrame, lab: pd.DataFrame) -> pd.DataFrame:
    e = p[p.technology_a.isin(TECH) & p.technology_b.isin(TECH) & (p.technology_a != p.technology_b)
          & (p.ratio >= RATIO_MIN) & (np.minimum(p.n_a, p.n_b) >= MIN_ACT)].copy()
    e["role"] = [role_of(*v) for v in zip(e.function_a, e.function_b, e.pulse_a, e.pulse_b, e.print_subtype_a, e.print_subtype_b)]
    e = e[e.role.notna()]
    # similarity by role: presence zones hold, so ON-time overlap; pulse / advance zones: ONs starting together
    pres = e.role == "Presence"
    e["sim"] = np.where(pres, np.minimum(e.ovl_a, e.ovl_b) * 0.5 + np.maximum(e.ovl_a, e.ovl_b) * 0.5, e.excess)
    ok_p = pres & (np.maximum(e.ovl_a, e.ovl_b) >= OVL_MAX) & (np.minimum(e.ovl_a, e.ovl_b) >= OVL_MIN)
    ok_o = ~pres & (e.excess >= EXCESS_MIN) & (e.lag_ab.abs() <= LAG_MAX)
    e = e[ok_p | ok_o]
    # each kept edge must be the best partner (highest similarity, other technology) of at least one of its ends
    best_a = e.groupby(["dev", "phase", "a"]).sim.transform("max")
    best_b = e.groupby(["dev", "phase", "b"]).sim.transform("max")
    e = e[(e.sim >= best_a) | (e.sim >= best_b)].copy()
    # one role per detector: a detector in several role groups keeps Presence if it holds (median ON > 1 s),
    # otherwise its highest-similarity role
    long_on = {}
    for side in ("a", "b"):
        for dv, d_, dm in zip(e.dev, e[side], e[f"dmed_{side}"]):
            long_on[(dv, d_)] = dm > 1.0
    cand = pd.concat([e[["dev", "phase", "a", "role", "sim"]].rename(columns={"a": "det"}),
                      e[["dev", "phase", "b", "role", "sim"]].rename(columns={"b": "det"})])
    cand["pref"] = cand.sim + np.where((cand.role == "Presence").to_numpy() & np.array([long_on[(x, y)] for x, y in zip(cand.dev, cand.det)]), 10, 0)
    keep = cand.sort_values("pref", ascending=False).drop_duplicates(["dev", "det"]).set_index(["dev", "det"]).role
    e = e[[keep.get((x, a)) == r and keep.get((x, b)) == r for x, a, b, r in zip(e.dev, e.a, e.b, e.role)]].copy()
    # groups: connected components per (signal, phase, role)
    groups = []
    for (dev, ph, role), g in e.groupby(["dev", "phase", "role"]):
        par = {}

        def f(x):
            while par.setdefault(x, x) != x:
                x = par[x]
            return x
        for a, b in zip(g.a, g.b):
            par[f(a)] = f(b)
        comps = {}
        for x in set(g.a) | set(g.b):
            comps.setdefault(f(x), set()).add(x)
        for mem in comps.values():
            ge = g[g.a.isin(mem) & g.b.isin(mem)]
            groups.append(dict(dev=dev, DeviceName=ge.DeviceName.iloc[0], phase=ph, role=role, members=sorted(mem),
                               best=ge.sort_values("sim", ascending=False).iloc[0].to_dict(),
                               min_sim=ge.sim.min(), mean_m=ge.m_sg.mean(), min_ratio=ge.ratio.min(),
                               corr=ge.corr15.min()))
    G = pd.DataFrame(groups)
    k = lab.set_index(["dev", "detector"])
    info = []
    for r in G.itertuples():
        L = k.loc[[(r.dev, m) for m in r.members]]
        techs = sorted(set(L.technology))
        funcs = list(L.function.fillna("unlabelled"))
        info.append(dict(tech="+".join(techs), n_members=len(r.members),
                         change=any(f != r.role for f in funcs), labels=funcs,
                         unusual=bool(L.unusual_layout.fillna(False).astype(bool).any()),
                         superseded=bool((L.print_subtype == "superseded_by_radar").any())))
    G = pd.concat([G, pd.DataFrame(info)], axis=1)
    G["question"] = False
    return G


def borderline(p: pd.DataFrame, lab: pd.DataFrame, G: pd.DataFrame, n_max: int = 8) -> pd.DataFrame:
    """Pairs of different technology that behave as one detector but whose labels do not agree on a role (two
    different classes, or a class + a zone type the user ruled Other). Listed as questions, never relabelled."""
    used = {(r.dev, m) for r in G.itertuples() for m in r.members}
    e = p[p.technology_a.isin(TECH) & p.technology_b.isin(TECH) & (p.technology_a != p.technology_b)
          & (p.ratio >= 0.6) & (np.minimum(p.n_a, p.n_b) >= MIN_ACT)].copy()
    e["role"] = [role_of(*v) for v in zip(e.function_a, e.function_b, e.pulse_a, e.pulse_b, e.print_subtype_a, e.print_subtype_b)]
    fa, fb = e.function_a.fillna("Other"), e.function_b.fillna("Other")
    e = e[e.role.isna() & (fa != fb) & ~fa.isin(["Bike"]) & ~fb.isin(["Bike"]) & (fa.isin(PM) | fb.isin(PM))]
    e = e[[(x, a) not in used and (x, b) not in used for x, a, b in zip(e.dev, e.a, e.b)]]
    e["sim"] = np.maximum(e.excess, np.minimum(e.ovl_a, e.ovl_b))
    e = e[((e.excess >= 0.6) & (e.lag_ab.abs() <= LAG_MAX)) | (np.minimum(e.ovl_a, e.ovl_b) >= 0.6)]
    e = e.sort_values("sim", ascending=False)
    rows, per = [], {}
    for r in e.itertuples():
        if per.get(r.dev, 0) >= 2:
            continue
        per[r.dev] = per.get(r.dev, 0) + 1
        roles = sorted({f if f in PM else "Other" for f in (r.function_a, r.function_b)})
        rows.append(dict(dev=r.dev, DeviceName=r.DeviceName, phase=r.phase, role=" or ".join(roles), members=[r.a, r.b],
                         best=e.loc[r.Index].to_dict(), min_sim=r.sim, mean_m=r.m_sg, min_ratio=r.ratio, corr=r.corr15,
                         tech="+".join(sorted({r.technology_a, r.technology_b})), n_members=2, change=True,
                         labels=[r.function_a, r.function_b], unusual=bool(r.unusual_layout_a or r.unusual_layout_b),
                         superseded=False, question=True))
        if len(rows) >= n_max:
            break
    return pd.DataFrame(rows)


def yr_count(p: pd.DataFrame) -> pd.DataFrame:
    y = p[((p.function_a == "Yellow_Red") & (p.function_b == "Count")) | ((p.function_b == "Yellow_Red") & (p.function_a == "Count"))].copy()
    y["yr"] = np.where(y.function_a == "Yellow_Red", y.a, y.b)
    y["cnt"] = np.where(y.function_a == "Yellow_Red", y.b, y.a)
    both = np.minimum(y.match_ab, y.match_ba)
    both_t = np.minimum(y.tight_ab, y.tight_ba)
    dr = 1 - y.ratio
    y["identical"] = (dr <= IDENT["ratio"]) & (both_t >= IDENT["match"])
    y["near"] = (dr <= NEAR["ratio"]) & (both >= NEAR["match"])
    return y


def save_chart_data(sel: pd.DataFrame):
    DATA.mkdir(parents=True, exist_ok=True)
    cn = duckdb.connect()
    cn.execute("SET memory_limit='10GB'")
    cn.execute("SET threads=8")
    for r in sel.itertuples():
        f = DATA / f"{r.DeviceName}_{r.phase}_{'-'.join(map(str, r.members))}"
        if Path(str(f) + "_counts.parquet").exists():
            continue
        dl = ",".join(map(str, r.members))
        iv = cn.sql(f"""SELECT Detector det, t_on, t_off FROM '{IVF}' WHERE lower(DeviceId) = '{r.dev}'
                        AND Detector IN ({dl}) ORDER BY t_on""").df()
        iv["t"] = iv.t_on.dt.floor("15min")
        grid = pd.date_range(iv.t.min(), iv.t.max(), freq="15min")
        cnt = (iv.groupby(["t", "det"]).size().unstack(fill_value=0).reindex(grid, fill_value=0)
               .reindex(columns=r.members, fill_value=0))
        cnt.index.name = "t"
        cnt.columns = [str(c) for c in cnt.columns]
        cnt.reset_index().to_parquet(str(f) + "_counts.parquet", index=False)
        # lag of each s-member ON to the nearest g-member ON (best pair), all ONs
        b = r.best
        s, g = int(b["s"]), int(b["g"])
        ep = pd.Timestamp("2026-01-01")
        a_on = (iv[iv.det == s].t_on - ep).dt.total_seconds().to_numpy()
        g_on = np.sort((iv[iv.det == g].t_on - ep).dt.total_seconds().to_numpy())
        i = np.searchsorted(g_on, a_on)
        lo, hi = g_on[np.clip(i - 1, 0, len(g_on) - 1)] - a_on, g_on[np.clip(i, 0, len(g_on) - 1)] - a_on
        lag = np.where(np.abs(lo) <= np.abs(hi), lo, hi)
        pd.DataFrame({"lag": lag, "s": s, "g": g}).to_parquet(str(f) + "_lags.parquet", index=False)
        # a 10-minute ON timeline in a medium-busy daytime period of the s member
        c10 = iv[iv.det == s].set_index("t_on").resample("10min").size()
        c10 = c10[(c10.index.hour >= 7) & (c10.index.hour < 19)]
        t0 = c10[c10 >= c10.quantile(0.6)].index[len(c10[c10 >= c10.quantile(0.6)]) // 2] if len(c10) else iv.t_on.min()
        t1 = t0 + pd.Timedelta(minutes=10)
        iv[(iv.t_off >= t0) & (iv.t_on <= t1)][["det", "t_on", "t_off"]].assign(w0=t0, w1=t1).to_parquet(
            str(f) + "_timeline.parquet", index=False)


def _style(ax):
    ax.set_facecolor(SURF)
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color(INK2)
    ax.tick_params(colors=INK2, labelsize=8)
    ax.grid(axis="y", color=GRID, lw=0.8)
    ax.set_axisbelow(True)


def draw(sel: pd.DataFrame, lab: pd.DataFrame) -> dict:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.dates as mdates
    import matplotlib.pyplot as plt
    CHARTS.mkdir(parents=True, exist_ok=True)
    k = lab.set_index(["dev", "detector"])
    out = {}
    for r in sel.itertuples():
        stem = f"{r.DeviceName}_{r.phase}_{'-'.join(map(str, r.members))}"
        cnt = pd.read_parquet(DATA / f"{stem}_counts.parquet")
        lags = pd.read_parquet(DATA / f"{stem}_lags.parquet")
        tl = pd.read_parquet(DATA / f"{stem}_timeline.parquet")

        def nm(d):
            x = k.loc[(r.dev, d)]
            return f"det {d} ({x.technology}, label {x.function if isinstance(x.function, str) else 'none'})"
        fig = plt.figure(figsize=(10, 6.4), facecolor=SURF)
        gs = fig.add_gridspec(2, 2, height_ratios=[1.15, 1], width_ratios=[1.6, 1])
        ax = fig.add_subplot(gs[0, :])
        _style(ax)
        for i, d in enumerate(r.members[:4]):
            ax.plot(cnt.t, cnt[str(d)], color=C[i], lw=1.8 if i == 0 else 1.3,
                    label=f"{nm(d)}: {int(cnt[str(d)].sum()):,} in 66 h")
        ax.set_ylabel("actuations per 15 min", color=INK2, fontsize=9)
        ax.legend(frameon=False, fontsize=8, loc="lower left", bbox_to_anchor=(0, 1.0), ncol=2)
        ax.xaxis.set_major_formatter(mdates.DateFormatter("%a %H:%M"))
        # timeline
        a2 = fig.add_subplot(gs[1, 0])
        _style(a2)
        w0, w1 = (tl.w0.iloc[0], tl.w1.iloc[0]) if len(tl) else (None, None)
        for i, d in enumerate(r.members[:4]):
            s = tl[tl.det == d]
            a2.barh([len(r.members) - i] * len(s), (s.t_off - s.t_on).dt.total_seconds().fillna(0).clip(lower=1.0) / 86400, left=s.t_on,
                    height=0.6, color=C[i])
        a2.set_yticks([len(r.members) - i for i in range(len(r.members[:4]))])
        a2.set_yticklabels([f"det {d}" for d in r.members[:4]], fontsize=8)
        if w0 is not None:
            a2.set_xlim(w0, w1)
        a2.xaxis.set_major_formatter(mdates.DateFormatter("%H:%M"))
        a2.grid(axis="y", visible=False)
        a2.grid(axis="x", color=GRID, lw=0.8)
        a2.set_title("10 daytime minutes: each bar = detector ON (ONs under 1 s drawn 1 s wide)", fontsize=9, color=INK2, loc="left")
        # lag histogram
        a3 = fig.add_subplot(gs[1, 1])
        _style(a3)
        s, g = int(lags.s.iloc[0]), int(lags.g.iloc[0])
        lg = lags.lag.clip(-5, 5)
        a3.hist(lg, bins=np.arange(-5, 5.01, 0.25), color=C[r.members.index(g) if g in r.members[:4] else 0])
        a3.axvspan(-TOL, TOL, color=GRID, alpha=0.5, zorder=0)
        a3.set_xlabel(f"seconds from det {s} ON to nearest det {g} ON", color=INK2, fontsize=8)
        a3.set_title(f"{(lags.lag.abs() <= TOL).mean():.0%} of det {s} ONs within {TOL:g} s", fontsize=9,
                     color=INK2, loc="left")
        fig.suptitle(f"{r.DeviceName} {r.phase}: {role_txt(r.role)} - detectors {', '.join(map(str, r.members))}, "
                     f"Sept 2026 log", x=0.01, ha="left", fontsize=10, color=INK)
        fig.tight_layout()
        png = CHARTS / f"{stem}.png"
        fig.savefig(png, dpi=105, facecolor=SURF)
        plt.close(fig)
        out[stem] = png
    return out


def pdf_for(sig):
    if not isinstance(sig, str):
        return None
    for d in PDFS:
        f = sorted(d.glob(f"{sig}_*.pdf"))
        if f:
            return f[0]
    return None


def sentence(r, lab) -> str:
    k = lab.set_index(["dev", "detector"])
    b = r.best
    s, g = int(b["s"]), int(b["g"])
    ts, tg = k.loc[(r.dev, s)].technology, k.loc[(r.dev, g)].technology
    ns, ng = int(min(b["n_a"], b["n_b"])), int(max(b["n_a"], b["n_b"]))
    lag = b["lag_ab"] if s == int(b["a"]) else -b["lag_ab"]
    ov = b["ovl_a"] if s == int(b["a"]) else b["ovl_b"]
    ovg = b["ovl_b"] if s == int(b["a"]) else b["ovl_a"]
    txt = (f"{ts.capitalize()} det {s} counted {ns:,} and {tg} det {g} {ng:,} actuations in the 66 h "
           f"(ratio {ns / ng:.2f}); 15-min counts correlate r = {b['corr15']:.2f}. ")
    if r.role == "Presence":
        txt += (f"They hold ON together: det {s} is ON while det {g} is ON for {ov:.0%} of its ON time, and det {g} "
                f"for {ovg:.0%} of its own; {b['m_sg']:.0%} of det {s}'s ONs start within {TOL:g} s of a det {g} ON "
                f"(chance alone {b['chance_sg']:.0%}).")
    else:
        txt += (f"{b['m_sg']:.0%} of det {s}'s ONs start within {TOL:g} s of a det {g} ON (median lag {lag:+.1f} s; "
                f"chance alone {b['chance_sg']:.0%}); ON at the same time for {ov:.0%} of det {s}'s ON time.")
    if len(r.members) > 2:
        txt += f" ({len(r.members)} detectors in the group; the numbers are for the closest pair.)"
    return txt


def main():
    lab = load_labels()
    pairs = pd.read_parquet(SOUT / "pairs.parquet")
    locked = set(pd.read_csv(DC_WORK / "official/locked_v2.csv").DeviceId.str.lower())
    assert not set(pairs.dev) & locked
    p = edges(pairs, lab)
    G = build_groups(p, lab)
    Y = yr_count(p)
    Y.to_parquet(SOUT / "yr_count_pairs.parquet", index=False)
    print(G.groupby(["role", "tech", "change"]).size())
    Q = borderline(p, lab, G)
    # --- sheet rows: every clear group (<= 32, labels that would change first), then <= 8 questions;
    # signals ordered by their clearest group, a signal's groups kept together
    G["score"] = G.min_sim + 0.5 * G.min_ratio + 0.3 * G["corr"].fillna(0)
    clear = G.sort_values(["change", "score"], ascending=False).head(32)
    Q["score"] = Q.min_sim if len(Q) else []
    sel = pd.concat([clear, Q], ignore_index=True)
    sel["sig_rank"] = sel.groupby(["question", "DeviceName"]).score.transform("max")
    sel = sel.sort_values(["question", "sig_rank", "DeviceName", "phase", "score"],
                          ascending=[True, False, True, True, False]).reset_index(drop=True)
    save_chart_data(sel)
    charts = draw(sel, lab)
    k = lab.set_index(["dev", "detector"])
    rows = []
    for i, r in enumerate(sel.itertuples(), 1):
        L = k.reindex([(r.dev, m) for m in r.members])
        stem = f"{r.DeviceName}_{r.phase}_{'-'.join(map(str, r.members))}"
        cur = "; ".join(f"det {m}: {f}" + (f" ({st.replace('_', ' ')})" if f == "Other" and isinstance(st, str) else "")
                        for m, f, st in zip(r.members, L.function.fillna("unlabelled"), L.print_subtype))
        if r.question:
            prop = (f"? They behave as one detector, but the labels disagree ({r.role}). If they cover the same lane "
                    f"and role, which class should both carry?")
        else:
            prop = "; ".join(f"det {m}: {r.role}" for m in r.members)
            prop += (f". The model may call only one of them {r.role} for this lane (the higher-probability one); "
                     f"the other is then Other and scored correct.")
            if not r.change:
                prop = "No label changes (both already " + r.role + "). " + prop
        lanes = sorted({lane_txt(x) for x in L.itertuples()} - {"-"})
        what = sentence(r, lab)
        if r.unusual:
            what += " The signal is flagged unusual layout (live radar over live loops), so today it is left out of training."
        rows.append({"#": i, "Signal": r.DeviceName, "Phase": r.phase, "Lane (print)": " / ".join(lanes) or "-",
                     "Role": role_txt(r.role),
                     "Detectors in the group": ", ".join(f"{m} ({t})" for m, t in zip(r.members, L.technology)),
                     "Current labels": cur, "Proposed labels under the new rule": prop, "What the data shows": what,
                     "Chart": charts[stem], "Cabinet print": pdf_for(r.DeviceName),
                     "Your answer (yes / no / ?)": "", "Comment": ""})
    G.drop(columns=["best"]).to_parquet(SOUT / "groups.parquet", index=False)
    Q.drop(columns=["best"], errors="ignore").to_parquet(SOUT / "questions.parquet", index=False)
    write_xlsx(rows, G, Y, lab, Q)


def write_xlsx(rows, G, Y, lab, Q):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "stacked groups"
    cols = list(rows[0])
    ws.append(cols)
    link = Font(color="0563C1", underline="single")
    for r in rows:
        ws.append([("open chart" if k == "Chart" else ("open print" if v else "no print on file"))
                   if k in ("Chart", "Cabinet print") else v for k, v in r.items()])
        n = ws.max_row
        c = ws.cell(n, cols.index("Chart") + 1)
        c.hyperlink = f"{CHARTS.name}/{Path(r['Chart']).name}"
        c.font = link
        if r["Cabinet print"]:
            c = ws.cell(n, cols.index("Cabinet print") + 1)
            c.hyperlink = "file:///" + urllib.parse.quote(str(r["Cabinet print"]), safe=":\\/()_-.,'")
            c.font = link
    widths = (4, 8, 7, 16, 16, 22, 34, 40, 70, 11, 11, 12, 35)
    # --- sheet 2: counts
    w2 = wb.create_sheet("counts")
    tot = G.groupby(["role", "tech"]).agg(groups=("dev", "size"), signals=("dev", "nunique"),
                                         labels_would_change=("change", "sum")).reset_index()
    w2.append(["Stacked groups found (2+ detectors of different technology, same phase, ONs starting together, "
               "labels agreeing on the role or one of them Other). All training / dev signals, Sept 2026 log."])
    w2.append([])
    w2.append(["Role", "Technology", "Groups", "Signals", "Groups where a label would change"])
    hdr_rows = [3]
    for r in tot.sort_values(["role", "tech"]).itertuples():
        w2.append([role_txt(r.role), r.tech, int(r.groups), int(r.signals), int(r.labels_would_change)])
    w2.append(["All", "", int(len(G)), int(G.dev.nunique()), int(G.change.sum())])
    w2.append(["Also listed at the end of sheet 1 as questions: pairs of different technology that behave as one "
               f"detector but whose labels name different roles: {len(Q)}"])
    w2.append([])
    ny = lab[lab.function == "Yellow_Red"]
    yr_act = Y.groupby(["dev", "yr"]).agg(ident=("identical", "any"), near=("near", "any")).reset_index()
    w2.append(["Yellow_Red zones whose actuations match a Count zone on the same phase. Identical = counts within 1 % "
               "and at least 98 % of the ONs of each start within 0.3 s of an ON of the other. Near-identical = counts "
               "within 5 % and at least 90 % of the ONs of each start within 1.5 s of the other (includes identical). "
               "A Yellow_Red zone that misses the first car of green would differ by more than 5 %."])
    w2.append([])
    n_yr_act = int(Y.groupby(["dev", "yr"]).ngroups)
    w2.append(["Yellow_Red zones labelled (training / dev)", int(len(ny))])
    w2.append(["... with 50+ actuations and a Count zone on the same phase", n_yr_act])
    w2.append(["... identical to a Count zone", int(yr_act.ident.sum())])
    w2.append(["... near-identical to a Count zone", int(yr_act.near.sum())])
    w2.append([])
    w2.append(["Signal", "Yellow_Red zones (with a Count zone on the phase)", "Identical to a Count zone",
               "Near-identical", "Which (YR det = Count det)"])
    hdr_rows.append(w2.max_row)
    nn = lab[lab.DeviceName.notna()]
    names = dict(zip(nn.dev, nn.DeviceName))
    Yp = Y[Y.near].sort_values(["dev", "yr"])
    for dev, g in sorted(yr_act.groupby("dev"), key=lambda t: str(names.get(t[0]))):
        if not g.near.any():
            continue
        which = ", ".join(sorted({f"{int(a)} = {int(b)}" + (" (identical)" if i else "")
                                  for a, b, i in zip(Yp[Yp.dev == dev].yr, Yp[Yp.dev == dev].cnt, Yp[Yp.dev == dev].identical)}))
        w2.append([names[dev], int(len(g)), int(g.ident.sum()), int(g.near.sum()), which])
    for sh, ws_ in ((ws, widths), (w2, (28, 22, 16, 14, 60))):
        for j, w in enumerate(ws_):
            sh.column_dimensions[chr(65 + j)].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for rr in hdr_rows:
        for c in w2[rr]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for row in w2.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for rr in (1, hdr_rows[0] + len(tot) + 3, hdr_rows[0] + len(tot) + 5):
        w2.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=5)
        w2.row_dimensions[rr].height = 48
    ws.freeze_panes = "A2"
    try:
        wb.save(XL)
    except PermissionError:
        wb.save(XL.with_name(XL.stem + "_new.xlsx"))
        print("workbook open in Excel -> wrote _new", file=sys.stderr)
    print("saved", XL, len(rows), "rows; groups", len(G), "; YR near", int(yr_act.near.sum()), "ident", int(yr_act.ident.sum()))


if __name__ == "__main__":
    main()
