"""Note 99: review sheet of phases where the lane model's OUT-OF-FOLD lane count disagrees with the print lane count.

Lane output = lanes D (note 58 / 77: engineered pair cues + LightGBM same-lane model + per-lane role grouping), the
six-fold OOF decode `%DC_WORK%/final_v3_work/f77/ln8/lanes_D.func{,_ph}.parquet` (every frame signal / window >= 30 min).
Truth = print lanes per phase (ln1_cues.truth rules) rebuilt on the current label table (rpath.LABELS_CURRENT).
Scope: Sept-2026 log (period stg), samples >= 30 min (m30 / h1 / h3 / h6 / h24 / full66); a phase's model answer = the
majority lane count over its samples. Ranked by (share of samples giving the majority answer) x (mean n_lanes_conf of
those samples), then by the number of agreeing samples. locked_v2 asserted absent. Charts from SAVED sample data only.

    python lane_review99.py rows    -> %DC_WORK%/rev99/{phase_majority,rows99}.parquet + profile99.json
    python lane_review99.py data    -> %DC_WORK%/rev99/review_data/*.parquet (full66 ONs of each row's detectors)
    python lane_review99.py write   -> review/lane_count_review.xlsx + review/lane_count_review_charts/
"""
from __future__ import annotations

import os
for _v in ("OMP_NUM_THREADS", "MKL_NUM_THREADS", "OPENBLAS_NUM_THREADS"):
    os.environ.setdefault(_v, "4")
import json  # noqa: E402
import sys  # noqa: E402
import urllib.parse  # noqa: E402
from pathlib import Path  # noqa: E402

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

CODE = next(p for p in Path(__file__).resolve().parents if p.name == "code")
sys.path.insert(0, str(CODE))
sys.path.insert(0, str(CODE / "lanes"))
sys.path.insert(0, str(CODE / "evaluation"))
import rpath  # noqa: F401,E402
import ln1_cues as L1  # noqa: E402
import review64 as R  # noqa: E402

DCW = L1.DCW
REPO = L1.REPO
OUT = DCW / "rev99"
DATA = OUT / "review_data"
LN = DCW / "final_v3_work" / "f77" / "ln8"
STG = DCW / "official" / "stg" / "cache"
FULL = ("2026-09-18 16:15:00", 237600)
N_ROWS = 25
MIN_OVER = 8
MAX_PER_SIGNAL = 3
MIN_ACT = 10
OUTNAME = "lane_count_review"
QUESTION = ("Is the print's lane count right for this phase? Y (print right) / N (model right) / ? "
            "- add a comment if you see why the model got it wrong.")
HEAD = ["#", "Signal", "Phase", "Print says (lanes)", "Model says (lanes)", "How sure",
        "Detectors on the phase (print function - model lane, print lane)", "What this row asks", "Chart",
        "Print link", "Answer", "Comment"]
WIDTH = (4, 8, 6, 8, 8, 12, 46, 50, 7, 7, 10, 30)
FN = {"Advance": "Advance", "Presence": "Presence", "Count": "Stop-bar Count", "Yellow_Red": "Yellow-Red",
      "Other": "Other", "Mid": "Mid", "Bike": "Bike"}
COL = ["#2a78d6", "#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#008300", "#4a3aa7", "#7a7a7a"]
INK, INK2, GRID, SURF = R.INK, R.INK2, R.GRID, R.SURF


def log(m):
    print(m, flush=True)


def lanes_txt(s: str) -> str:
    return "+".join(s.split(",")) if isinstance(s, str) and s else ""


def plane_txt(li, sp) -> str:
    if pd.isna(li):
        return ""
    a, n = int(li), int(sp) if pd.notna(sp) else 1
    return str(a) if n <= 1 else "+".join(str(a + k) for k in range(n))


# ------------------------------------------------------------------------------------------------ rows
def stage_rows():
    OUT.mkdir(parents=True, exist_ok=True)
    lk = L1.locked()
    L1.V3 = rpath.LABELS_CURRENT
    ph, det = L1.truth()
    assert not ph.DeviceId.isin(lk).any()
    ph["ph_num"] = ph.target.str[1:].astype(int)
    P = pd.read_parquet(LN / "lanes_D.func_ph.parquet")
    P = P[P.period == "stg"]
    assert not P.DeviceId.isin(lk).any()
    x = ph.merge(P, left_on=["DeviceId", "ph_num"], right_on=["DeviceId", "phase"], suffixes=("", "_p"))
    g = x.groupby(["DeviceId", "target"])
    M = g.agg(n=("win", "size"), truth=("n_lanes", "first"), ph_num=("ph_num", "first"),
              maj=("n_lanes_p", lambda s: s.value_counts().sort_index().idxmax())).reset_index()
    x = x.merge(M[["DeviceId", "target", "maj"]], on=["DeviceId", "target"])
    a = x[x.n_lanes_p == x.maj].groupby(["DeviceId", "target"]).agg(
        k=("win", "size"), conf=("n_lanes_conf", "mean"),
        wins=("win", lambda s: ",".join(sorted(s)))).reset_index()
    M = M.merge(a, on=["DeviceId", "target"])
    M["wrong"] = M.maj != M.truth
    M["dir"] = np.where(M.maj > M.truth, "over", np.where(M.maj < M.truth, "under", "ok"))
    M["share"] = M.k / M.n
    M["score"] = M.share * M.conf
    # phase facts for the profile
    L = pd.read_parquet(rpath.LABELS_CURRENT, columns=["DeviceId", "DeviceName", "detector", "technology",
                                                       "stack_group", "phase_target"])
    L["DeviceId"] = L.DeviceId.str.lower()
    L = L.rename(columns={"detector": "det"}).drop_duplicates(["DeviceId", "det"])
    L["det"] = L.det.astype(int)
    d = det.merge(L[["DeviceId", "det", "technology", "stack_group"]], on=["DeviceId", "det"], how="left")
    veh = d[d.high_veh]
    D = pd.read_parquet(LN / "lanes_D.func.parquet")
    D = D[(D.period == "stg") & (D.win == "full66")]
    pp = D.groupby(["DeviceId", "Detector"]).phase.first()
    non = D.groupby(["DeviceId", "Detector"]).n_on.first()
    d["pred_phase"] = [pp.get((s, x_), np.nan) for s, x_ in zip(d.DeviceId, d.det)]
    d["n_on66"] = [non.get((s, x_), 0) for s, x_ in zip(d.DeviceId, d.det)]
    d["n_on66"] = d.n_on66.fillna(0)
    gp = d.groupby(["DeviceId", "target"])
    # lanes the log can show: distinct print lanes holding an ACTIVE (>= 10 ONs in 66 h) single-lane vehicle detector
    lv = d.lane_index.notna() & ~d.lane_type.isin(L1.NONVEH)
    act = d[lv & (d.span == 1) & (d.n_on66 >= MIN_ACT)]
    vis = act.groupby(["DeviceId", "target"]).lane_index.nunique()
    dead = d[lv & (d.n_on66 < MIN_ACT)].groupby(["DeviceId", "target"]).size()
    tech = veh.groupby(["DeviceId", "target"]).technology.agg(
        lambda s: s.mode().iloc[0] if s.notna().any() else "unknown")
    mixed = veh.groupby(["DeviceId", "target"]).technology.nunique() > 1
    F = pd.DataFrame({
        "n_print_dets": gp.size(),
        "n_veh_dets": veh.groupby(["DeviceId", "target"]).size(),
        "has_span": veh.assign(s=veh.span > 1).groupby(["DeviceId", "target"]).s.any(),
        "has_stack": d.assign(s=d.stack_group.notna()).groupby(["DeviceId", "target"]).s.any(),
        "dup_role": veh.groupby(["DeviceId", "target", "lane_index", "print_function"]).size().gt(1)
                       .groupby(level=[0, 1]).any(),
        "moved": d.assign(m=d.pred_phase.notna() & (d.pred_phase != d.ph_num)).groupby(["DeviceId", "target"]).m.any(),
        "tech": tech, "mixed_tech": mixed, "visible": vis, "n_dead": dead}).reset_index()
    M = M.merge(F, on=["DeviceId", "target"], how="left")
    # model group on the phase holds a detector the print/label puts on ANOTHER phase (phase error leaking in)
    lp = dict(zip(zip(L.DeviceId, L.det), L.phase_target))
    Dq = D.assign(lab_ph=[lp.get((s_, int(x_))) for s_, x_ in zip(D.DeviceId, D.Detector)])
    Dq = Dq[Dq.lab_ph.astype(str).str.startswith("P") & (Dq.lab_ph.astype(str) != "P" + Dq.phase.astype(int).astype(str))
            & (Dq.lanes != "")]
    mi = set(zip(Dq.DeviceId, "P" + Dq.phase.astype(int).astype(str)))
    M["moved_in"] = [(s_, t_) in mi for s_, t_ in zip(M.DeviceId, M.target)]
    M["tech"] = np.where(M.mixed_tech.fillna(False), "mixed", M.tech.fillna("unknown"))
    for c in ("has_span", "has_stack", "dup_role", "moved"):
        M[c] = M[c].fillna(False).astype(bool)
    M["few"] = M.n_veh_dets.fillna(0) <= 2
    M["visible"] = M.visible.fillna(0).astype(int)
    M["n_dead"] = M.n_dead.fillna(0).astype(int)
    # blind = the log cannot show the missing lanes (only spanning zones / dead loops there): not a question for the user
    M["blind"] = (M.dir == "under") & (M.visible.clip(lower=1) <= M.maj)
    M["blind_dead"] = M.blind & (M.n_dead > 0)
    names = L.drop_duplicates("DeviceId").set_index("DeviceId").DeviceName
    M["DeviceName"] = M.DeviceId.map(names)
    M.to_parquet(OUT / "phase_majority.parquet", index=False)

    def rate(m):
        return {"n": int(m.sum()), "wrong_pct": round(100 * M[m].wrong.mean(), 1),
                "over": int((M[m].dir == "over").sum()), "under": int((M[m].dir == "under").sum())}
    W = M[M.wrong]
    prof = {"phases": int(len(M)), "signals": int(M.DeviceId.nunique()), "wrong": int(len(W)),
            "wrong_pct": round(100 * M.wrong.mean(), 1),
            "over": int((W.dir == "over").sum()), "under": int((W.dir == "under").sum()),
            "truth_x_model": {f"{t}->{m}": int(n) for (t, m), n in W.groupby(["truth", "maj"]).size().items()},
            "by_truth": {int(t): rate(M.truth == t) for t in sorted(M.truth.unique())},
            "by_tech": {t: rate(M.tech == t) for t in M.tech.value_counts().index},
            "span": rate(M.has_span), "no_span": rate(~M.has_span),
            "stack": rate(M.has_stack), "no_stack": rate(~M.has_stack),
            "dup_role_same_lane": rate(M.dup_role), "no_dup_role": rate(~M.dup_role),
            "moved_det": rate(M.moved), "no_moved": rate(~M.moved),
            "few_le2": rate(M.few), "more_than_2": rate(~M.few),
            "under_blind": int(W.blind.sum()), "under_blind_dead_det": int(W.blind_dead.sum()),
            "under_blind_span_only": int((W.blind & ~W.blind_dead).sum()),
            "under_not_blind": int(((W.dir == "under") & ~W.blind).sum()),
            "wrong_not_blind": int((~W.blind).sum()),
            "over_with_det_from_other_print_phase": int((W[W.dir == "over"].moved_in).sum()),
            "confident_wrong_score_ge_.8": int((W.score >= .8).sum()),
            "sample_acc": round(float((x.n_lanes == x.n_lanes_p).mean()), 4)}
    for t in (1, 2, 3):
        mm = M.truth == t
        prof[f"few_truth{t}"] = rate(mm & M.few)
        prof[f"many_truth{t}"] = rate(mm & ~M.few)
    json.dump(prof, open(OUT / "profile99.json", "w"), indent=1)
    # real disagreements (blind under-counts out) by phase trait
    M["err"] = M.wrong & ~M.blind

    def rr(m):
        z = M[m]
        return {"n": int(len(z)), "err_pct": round(100 * float(z.err.mean()), 1),
                "over": int((z.err & (z.dir == "over")).sum()), "under": int((z.err & (z.dir == "under")).sum()),
                "blind": int(z.blind.sum())}
    nb = {"all": rr(M.wrong | ~M.wrong)}
    nb.update({f"tech_{t}": rr(M.tech == t) for t in M.tech.value_counts().index})
    for c in ("has_span", "has_stack", "dup_role", "moved", "moved_in", "few"):
        nb[c], nb[f"not_{c}"] = rr(M[c]), rr(~M[c])
    json.dump(nb, open(OUT / "profile99_nonblind.json", "w"), indent=1)
    log(json.dumps(prof, indent=1))
    # selection: top over / under by score, then k; at most MAX_PER_SIGNAL per signal; blind phases left out
    W = W[~W.blind]
    W = W.sort_values(["score", "k", "conf"], ascending=False)
    n_over = max(MIN_OVER, round(N_ROWS * (W.dir == "over").sum() / max(1, len(W))))
    pick, per = [], {}
    for direction, n in (("over", n_over), ("under", N_ROWS - n_over)):
        c = 0
        for r in W[W.dir == direction].itertuples():
            if c >= n:
                break
            if per.get(r.DeviceId, 0) >= MAX_PER_SIGNAL:
                continue
            per[r.DeviceId] = per.get(r.DeviceId, 0) + 1
            pick.append(r.Index)
            c += 1
    S = W.loc[pick].copy()
    # representative sample: full66 if it gives the majority answer, else the longest majority window
    order = ["full66", "h24_a", "h24_b", "h6_a", "h6_b", "h3_a", "h3_b", "h1_a", "h1_c", "h1_b",
             "m30_a", "m30_b", "m30_d", "m30_c"]
    S["rep_win"] = S.wins.map(lambda w: next(o for o in order if o in w.split(",")))
    # detectors: print detectors on the phase + the model's group on that phase in rep_win
    Dall = pd.read_parquet(LN / "lanes_D.func.parquet")
    Dall = Dall[Dall.period == "stg"]
    LF = pd.read_parquet(rpath.LABELS_CURRENT, columns=["DeviceId", "detector", "phase_target", "print_function",
                                                        "truth_v3s", "lane_index", "lanes_spanned", "technology"])
    LF["DeviceId"] = LF.DeviceId.str.lower()
    LF = LF.rename(columns={"detector": "det"}).drop_duplicates(["DeviceId", "det"])
    LF["det"] = LF.det.astype(int)
    dl = []
    for r in S.itertuples():
        dw = Dall[(Dall.DeviceId == r.DeviceId) & (Dall.win == r.rep_win)].set_index("Detector")
        pr = det[(det.DeviceId == r.DeviceId) & (det.target == r.target)].set_index("det")
        grp = set(dw.index[dw.phase == r.ph_num].astype(int))
        dets = sorted(set(pr.index.astype(int)) | grp)
        lf = LF[LF.DeviceId == r.DeviceId].set_index("det")
        items = []
        for dd in dets:
            fn = pr.print_function.get(dd) if dd in pr.index else (lf.print_function.get(dd) if dd in lf.index else None)
            if not isinstance(fn, str):
                fn = lf.truth_v3s.get(dd) if dd in lf.index else None
            on_print = dd in pr.index
            p_l = plane_txt(pr.lane_index.get(dd), pr.span.get(dd)) if on_print else ""
            other_ph = None if on_print else (lf.phase_target.get(dd) if dd in lf.index else None)
            if dd in dw.index:
                mph, ml, note, n_on = dw.phase.get(dd), dw.lanes.get(dd), dw.lane_note.get(dd), dw.n_on.get(dd)
            else:
                mph, ml, note, n_on = np.nan, "", "", 0
            items.append(dict(det=int(dd), fn=fn if isinstance(fn, str) else None, on_print=on_print, p_lane=p_l,
                              other_ph=other_ph if isinstance(other_ph, str) else None,
                              m_phase=None if pd.isna(mph) else int(mph), m_lanes=ml if isinstance(ml, str) else "",
                              m_note=note if isinstance(note, str) else "", n_on=int(n_on) if pd.notna(n_on) else 0,
                              tech=lf.technology.get(dd) if dd in lf.index else None))
        dl.append(json.dumps(items))
    S["dets_json"] = dl
    S = S.sort_values(["DeviceName", "ph_num"]).reset_index(drop=True)
    S.to_parquet(OUT / "rows99.parquet", index=False)
    log(f"rows {len(S)}: over {(S.dir == 'over').sum()} / under {(S.dir == 'under').sum()}, signals "
        f"{S.DeviceId.nunique()}, score min {S.score.min():.3f}")


# ------------------------------------------------------------------------------------------------ chart data
def stage_data():
    import duckdb
    DATA.mkdir(parents=True, exist_ok=True)
    S = pd.read_parquet(OUT / "rows99.parquet")
    cn = duckdb.connect()
    cn.execute("SET memory_limit='10GB'")
    cn.execute("SET threads=4")
    cn.execute(f"SET temp_directory='{(DCW / 'tmp').as_posix()}'")
    t0 = pd.Timestamp(FULL[0])
    t1 = t0 + pd.Timedelta(seconds=int(FULL[1]))
    for r in S.itertuples():
        f = DATA / f"{r.DeviceName}_{r.target}_on.parquet"
        if f.exists():
            continue
        dets = [i["det"] for i in json.loads(r.dets_json)]
        iv = cn.sql(f"""SELECT DISTINCT Detector::INT det, t_on, t_off FROM '{(STG / 'det_intervals.parquet').as_posix()}'
                        WHERE lower(DeviceId) = '{r.DeviceId}' AND Detector IN ({','.join(map(str, dets))})
                          AND t_on >= TIMESTAMP '{t0}' AND t_on < TIMESTAMP '{t1}' ORDER BY t_on""").df()
        iv.to_parquet(f, index=False)
    log(f"chart data for {len(S)} rows in {DATA}")


# ------------------------------------------------------------------------------------------------ charts + sheet
def det_label(i) -> str:
    fn = FN.get(i["fn"], i["fn"]) if i["fn"] else "no print function"
    return fn


def det_line(i, ph_num) -> str:
    fn = det_label(i)
    if not i["on_print"]:
        fn = f"{fn} (print puts it on {i['other_ph']})" if i["other_ph"] else f"{fn} (not on this phase in the print)"
    if i["m_phase"] is None:
        m = "no actuations in the sample" if i["n_on"] == 0 else "model: no answer"
    elif i["m_phase"] != ph_num:
        m = f"model puts it on P{i['m_phase']}"
    elif i["m_lanes"]:
        m = ("model lanes " if "," in i["m_lanes"] else "model lane ") + lanes_txt(i["m_lanes"])
    else:
        nt = i["m_note"]
        m = ("model calls it Bike (no lane)" if nt == "bike detector" else
             "model: no lane (too few actuations)" if nt.startswith("too few") else
             "model: no lane" + (f" ({nt})" if nt else ""))
    p = (("print lanes " if "+" in i["p_lane"] else "print lane ") + i["p_lane"]) if i["p_lane"] else "print: no lane"
    return f"det {i['det']}: {fn} - {m}, {p}"


def _join(xs):
    xs = list(xs)
    return xs[0] if len(xs) == 1 else ", ".join(xs[:-1]) + " and " + xs[-1]


def ask_txt(r, items) -> str:
    ln = lambda n: f"{n} lane" + ("" if n == 1 else "s")  # noqa: E731
    on = [i for i in items if i["on_print"] and i["p_lane"] and i["m_phase"] == r.ph_num and i["m_lanes"]]
    why = ""
    if r.dir == "under":
        for a in on:
            for b in on:
                if a["det"] < b["det"] and not set(a["p_lane"].split("+")) & set(b["p_lane"].split("+")) \
                        and set(a["m_lanes"].split(",")) & set(b["m_lanes"].split(",")):
                    why = (f" It puts det {a['det']} and det {b['det']} (print lanes {a['p_lane']} and {b['p_lane']}) "
                           f"in the same lane.")
                    break
            if why:
                break
        moved = [i for i in items if i["on_print"] and i["m_phase"] is not None and i["m_phase"] != r.ph_num]
        if not why and moved:
            why = f" It puts {_join(['det ' + str(i['det']) for i in moved])} on another phase."
    else:
        for a in on:
            for b in on:
                if a["det"] < b["det"] and a["p_lane"] == b["p_lane"] \
                        and not set(a["m_lanes"].split(",")) & set(b["m_lanes"].split(",")):
                    why = (f" It splits det {a['det']} and det {b['det']} (both print lane {a['p_lane']}) "
                           f"into lanes {lanes_txt(a['m_lanes'])} and {lanes_txt(b['m_lanes'])}.")
                    break
            if why:
                break
        extra = [i for i in items if not i["on_print"] and i["m_phase"] == r.ph_num and i["m_lanes"]]
        if not why and extra:
            why = (f" It adds {_join(['det ' + str(i['det']) for i in extra])}, which the print does "
                   f"not put on {r.target}.")
    return (f"Print shows {r.target} with {ln(int(r.truth))}; the model counts {ln(int(r.maj))} on {int(r.k)} of "
            f"{int(r.n)} samples.{why} Is the print's lane count right?")


def sure_txt(r) -> str:
    return f"{int(r.k)} of {int(r.n)} samples, {r.conf:.0%} sure"


def same_moment(iv: pd.DataFrame, dets: list, tol=0.5) -> np.ndarray:
    """M[a, b] = share of det a's actuations (ON starts) with det b turning ON within +-tol s."""
    st = {d: np.sort(iv.t_on[iv.det == d].to_numpy("datetime64[ns]").astype("int64")) / 1e9 for d in dets}
    M = np.full((len(dets), len(dets)), np.nan)
    for i, a in enumerate(dets):
        for j, b in enumerate(dets):
            if i == j or len(st[a]) == 0 or len(st[b]) == 0:
                continue
            lo = np.searchsorted(st[b], st[a] - tol, side="left")
            hi = np.searchsorted(st[b], st[a] + tol, side="right")
            M[i, j] = float(np.mean(hi > lo))
    return M


def draw(r, items, path: Path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    iv = pd.read_parquet(DATA / f"{r.DeviceName}_{r.target}_on.parquet")
    dets = [i["det"] for i in items]
    t0 = pd.Timestamp(FULL[0])
    bins = pd.date_range(t0, t0 + pd.Timedelta(seconds=int(FULL[1])), freq="15min")
    ln = lambda n: f"{int(n)} lane" + ("" if int(n) == 1 else "s")  # noqa: E731
    fig = plt.figure(figsize=(11, 7.8), facecolor=SURF)
    gs = fig.add_gridspec(2, 2, height_ratios=[1.2, 1], width_ratios=[1.35, 1], hspace=.55, wspace=.28,
                          top=.89 - .025 * ((len(items) - 1) // 4))
    ax = fig.add_subplot(gs[0, :])
    R._style(ax)
    for k, i in enumerate(items):
        t = iv.t_on[iv.det == i["det"]]
        c = np.histogram(t.to_numpy("datetime64[ns]").astype("int64"), bins.asi8)[0]
        p = f"print lane {i['p_lane']}" if i["p_lane"] else ("print: other phase" if not i["on_print"] else "print: no lane")
        m = (f"model lane {lanes_txt(i['m_lanes'])}" if i["m_phase"] == r.ph_num and i["m_lanes"]
             else (f"model: P{i['m_phase']}" if i["m_phase"] not in (None, r.ph_num) else "model: no lane"))
        ax.plot(bins[:-1], c, color=COL[k % len(COL)], lw=1.6, ls="-" if k < len(COL) else "--",
                label=f"det {i['det']} {det_label(i)}: {len(t):,}")
    ax.set_ylabel("actuations per 15 min", color=INK2, fontsize=9)
    ax.legend(loc="lower left", bbox_to_anchor=(0, 1.0), fontsize=8, frameon=False, ncol=min(4, len(items)),
              labelcolor=INK, borderaxespad=0.2)
    import matplotlib.dates as mdates
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%a %H:%M"))
    fig.suptitle(f"{r.DeviceName} {r.target} - Print: {ln(r.truth)}, model: {ln(r.maj)}   (Sept 2026 log, 66 h; "
                 f"lanes in the table at the bottom right)", x=0.06, ha="left", color=INK, fontsize=12)
    # same-moment heatmap
    M = same_moment(iv, dets)
    ax2 = fig.add_subplot(gs[1, 0])
    ax2.imshow(np.nan_to_num(M, nan=0), cmap="Blues", vmin=0, vmax=max(.4, float(np.nanmax(M)) if np.isfinite(M).any() else .4), aspect="auto")
    ax2.set_xticks(range(len(dets)), [f"d{d}" for d in dets], fontsize=8, color=INK2)
    ax2.set_yticks(range(len(dets)), [f"d{d}" for d in dets], fontsize=8, color=INK2)
    for a in range(len(dets)):
        for b in range(len(dets)):
            if a != b and not np.isnan(M[a, b]):
                ax2.text(b, a, f"{M[a, b]:.0%}", ha="center", va="center", fontsize=7.5 if len(dets) < 9 else 6,
                         color="white" if M[a, b] > .6 * max(.4, float(np.nanmax(M))) else INK)
    ax2.set_title("same moment: % of the ROW detector's actuations with the\nCOLUMN detector turning on within 0.5 s",
                  loc="left", fontsize=9, color=INK)
    for s in ax2.spines.values():
        s.set_visible(False)
    # small lane table
    ax3 = fig.add_subplot(gs[1, 1])
    ax3.axis("off")
    cell = [[f"d{i['det']}", det_label(i)[:14], i["p_lane"] or ("other ph." if not i["on_print"] else "-"),
             (lanes_txt(i["m_lanes"]) if i["m_phase"] == r.ph_num and i["m_lanes"] else
              (f"on P{i['m_phase']}" if i["m_phase"] not in (None, r.ph_num) else "-"))] for i in items]
    tb = ax3.table(cellText=cell, colLabels=["det", "print function", "print lane", "model lane"], loc="upper left",
                   cellLoc="left", colLoc="left")
    tb.auto_set_font_size(False)
    tb.set_fontsize(8)
    tb.scale(1, 1.25)
    for (rr, cc), c in tb.get_celld().items():
        c.set_width([.12, .36, .24, .24][cc])
    for (rr, cc), c in tb.get_celld().items():
        c.set_edgecolor(GRID)
        c.set_facecolor("#ddebf7" if rr == 0 else SURF)
    fig.savefig(path, dpi=100, facecolor=SURF, bbox_inches="tight")
    plt.close(fig)


def stage_write():
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    S = pd.read_parquet(OUT / "rows99.parquet")
    out = REPO / "review"
    xl = out / f"{OUTNAME}.xlsx"
    assert not xl.exists(), f"{xl} exists - never overwrite a sheet the user may have typed in"
    chd = out / f"{OUTNAME}_charts"
    chd.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "lane counts to check"
    ws.append([QUESTION])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEAD))
    ws.row_dimensions[1].height = 34
    ws.append(HEAD)
    link = Font(color="0563C1", underline="single")
    cc, cp = HEAD.index("Chart") + 1, HEAD.index("Print link") + 1
    links = []
    for n, r in enumerate(S.itertuples(), 1):
        items = json.loads(r.dets_json)
        png = chd / f"{r.DeviceName}_{r.target}.png"
        draw(r, items, png)
        pdf = R.pdf_for(r.DeviceName)
        pl = "file:///" + urllib.parse.quote(str(pdf).replace("\\", "/"), safe=":/()_-.,'") if pdf else None
        ws.append([n, r.DeviceName, r.target, int(r.truth), int(r.maj), sure_txt(r),
                   "\n".join(det_line(i, r.ph_num) for i in items), ask_txt(r, items), "chart",
                   "print" if pl else "none", "", ""])
        k = ws.max_row
        ws.cell(k, cc).hyperlink = f"{OUTNAME}_charts/{png.name}"
        ws.cell(k, cc).font = link
        if pl:
            ws.cell(k, cp).hyperlink = pl
            ws.cell(k, cp).font = link
        links.append({"row": n, "signal": r.DeviceName, "phase": r.target, "chart": png.name,
                      "print": str(pdf) if pdf else None})
    for j, w in enumerate(WIDTH):
        ws.column_dimensions[chr(65 + j)].width = w
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"
    wb.save(xl)
    json.dump(links, open(OUT / "links99.json", "w"), indent=1)
    log(f"saved {xl} ({len(S)} rows), charts in {chd}")


if __name__ == "__main__":
    {"rows": stage_rows, "data": stage_data, "write": stage_write}[sys.argv[1]]()
