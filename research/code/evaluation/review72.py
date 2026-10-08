"""Note 72: the user's LABEL review sheet (likely-wrong labels only), built on note 64's review builder.

Candidate = cand64 trees + fj blend (six folds x 3 seeds), Sept-2026 log, samples >= 30 min. Starts from review64's items
(one detector, wrong after review64's exclusions in >= half its samples; phase errors take precedence) and keeps only the
items where the LABEL is the likely culprit:
  * note-70 bucket (a): config-only label the model contradicts (A3 confident, A4 unsure at how-sure >= A4_MIN), or
  * a confident contradiction: model's mean probability on its own answer >= CONF over the wrong samples.
Dropped (counted): definitional rows pending the user's open question (note-70 B1-B5 rules evaluated on every wrong
window, not first-match: Other subtypes acting like a PM class, radar long Advance called Other, YR/Count twins, Mid vs
Advance / Presence, stacked extra lane), unhealthy-only (A2 / health bad), labels the user ruled on already
(source user_ruling: never ask twice), plus everything review64 already excludes (misconfigured, field issue, switch /
additional call phase, overlap phase, known print-phase mislabels, YR/Count twin swaps).
Same signal + same pattern (kind, label, model answer) collapse into ONE row listing the detectors.
Charts are drawn from SAVED sample data (`%DC_WORK%/cand64/review_data72/`), nothing re-scored.

    python review72.py            # dry run: counts -> %DC_WORK%/cand64/review72_dryrun.json
    python review72.py --write    # review/function_label_review_v1.xlsx + review/function_label_review_v1_charts/
"""
from __future__ import annotations

import os
os.environ.setdefault("OMP_NUM_THREADS", "4")
import argparse  # noqa: E402
import json  # noqa: E402
import sys  # noqa: E402
import time  # noqa: E402
import urllib.parse  # noqa: E402
from pathlib import Path  # noqa: E402

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

sys.path.insert(0, str(next(p for p in Path(__file__).resolve().parents if p.name == "code")))
import rpath  # noqa: F401,E402
from common import DC_WORK  # noqa: E402
import review64 as R  # noqa: E402
import err51_attrib as E51  # noqa: E402

C64 = DC_WORK / "cand64"
DATA = C64 / "review_data72"
GE30 = R.GE30
CONF = 0.9
A4_MIN = 0.7
MIN_N = 3
MAX_DETS_ROW = 12
OUTNAME = "function_label_review_v1"
NAME = R.NAME
QUESTION = ("Is the label right? Each row is a detector (or a group of detectors at one signal behaving the same way) "
            "where the model disagrees with the label on most 30-min-plus samples of the Sept 2026 log, and the label is "
            "the likely culprit. Answer L (label right), M (model right) or ? (can't tell). Optional: write the correct "
            "class or phase if both are wrong.")
INK, INK2, GRID, SURF = R.INK, R.INK2, R.GRID, R.SURF
C = R.C
ERR_FILE = DC_WORK / "err70" / "errors_everything.parquet"   # note 81 (review81.py) points it at the champion's errors


def log(m):
    print(f"[{time.strftime('%H:%M:%S')}] {m}", flush=True)


# ------------------------------------------------------------------------------------------------ selection
def err70_flags() -> pd.DataFrame:
    """per (DeviceId, Detector): share of its wrong >= 30-min Sept windows matching a definitional (B) rule, the unhealthy
    rule, and the dominant note-70 bucket."""
    e = pd.read_parquet(ERR_FILE)
    e = e[(e.period == "stg") & e.wgroup.isin(GE30)].copy()
    t, p, st = e.truth, e.pred, e.print_subtype.fillna("none")
    yrc = ((t == "Yellow_Red") & (p == "Count")) | ((t == "Count") & (p == "Yellow_Red"))
    e["defn"] = (e.err.eq("stack_extra")
                 | (t.eq("Other") & st.isin(E51.OTHER_LIKE) & p.isin(E51.PM | {"Mid"}))
                 | (t.eq("Advance") & st.isin(E51.RADAR_LONG_ADV) & p.eq("Other"))
                 | (yrc & ((e.twin < .05) | e.twin.isna()))
                 | (t.eq("Mid") & p.isin(["Advance", "Presence"])) | (p.eq("Mid") & t.isin(["Advance", "Presence"])))
    e["unh"] = e.validated.isin(["unhealthy", "no_data"]) | e.health.eq("bad")
    g = e.groupby(["DeviceId", "Detector"])
    out = g.agg(defn=("defn", "mean"), unh=("unh", "mean"), n70=("ok", "size")).reset_index()
    out["cat"] = g.cat.agg(lambda s: s.value_counts().index[0]).to_numpy()
    out["truth_src"] = g.truth_src.first().to_numpy()
    return out


def select(items: pd.DataFrame):
    lf = R.label_facts()[["DeviceId", "Detector"]]
    v = pd.read_parquet(rpath.LABELS_CURRENT,  # note 81: v4l
                        columns=["DeviceId", "detector", "source", "print_confidence", "print_subtype"])
    v["DeviceId"] = v.DeviceId.str.lower()
    v = v.rename(columns={"detector": "Detector"}).drop_duplicates(["DeviceId", "Detector"])
    del lf
    it = items.merge(v, on=["DeviceId", "Detector"], how="left")
    fl = err70_flags()
    it = it.merge(fl.astype({"Detector": it.Detector.dtype}), on=["DeviceId", "Detector"], how="left")
    why = np.full(len(it), "", object)
    fn = it.kind.eq("function")
    rules = [
        (fn & it.defn.fillna(0).ge(0.5), "definitional, pending your open question"),
        (fn & it.unh.fillna(0).ge(0.5), "unhealthy detector"),
        (it.source.eq("user_ruling"), "you already ruled on this label"),
        (it.n.lt(MIN_N), "fewer than 3 samples"),
    ]
    for m, r in rules:
        why[(why == "") & m.to_numpy()] = r
    bucket_a = fn & it.cat.isin(["A3 config-only label, model confident"]) | \
        (fn & it.cat.isin(["A4 config-only label, model unsure"]) & it.conf.ge(A4_MIN))
    conf = it.conf.ge(CONF)
    it["why_drop"] = why
    it["type"] = np.where(bucket_a, "config-only label contradicted", np.where(conf, "confident contradiction", ""))
    it.loc[it.why_drop == "", "why_drop"] = np.where(it.loc[it.why_drop == "", "type"] == "", "not confident enough", "")
    return it


def collapse(it: pd.DataFrame) -> pd.DataFrame:
    """one row per (signal, kind, label, model answer)."""
    k = it[it.why_drop == ""].copy()
    k["lab_key"] = np.where(k.kind == "phase", k.phase_target.astype(str), k.label_function.astype(str))
    k["mod_key"] = np.where(k.kind == "phase", k.model_phase.astype("Int64").astype(str), k.model_function.astype(str))
    k = k.sort_values("conf", ascending=False)
    rows = []
    for (d, kind, lk, mk), g in k.groupby(["DeviceId", "kind", "lab_key", "mod_key"], sort=False):
        r = g.iloc[0].to_dict()
        r["dets"] = sorted(int(x) for x in g.Detector)
        r["n_dets"] = len(g)
        r["phases"] = sorted({str(x)[1:] for x in g.phase_target if str(x)[1:].isdigit()}, key=int)
        r["conf_max"] = float(g.conf.max())
        r["conf_mean"] = float(g.conf.mean())
        r["n_err_sum"], r["n_sum"] = int(g.n_err.sum()), int(g.n.sum())
        r["types"] = ", ".join(sorted(set(g.type)))
        rows.append(r)
    out = pd.DataFrame(rows)
    out["sig_rank"] = out.groupby("DeviceId").conf_max.transform("max")
    out = out.sort_values(["sig_rank", "DeviceId", "conf_max"], ascending=[False, True, False]).reset_index(drop=True)
    return out


# ------------------------------------------------------------------------------------------------ chart data + charts
def stem(r) -> str:
    return f"{r.DeviceName or r.DeviceId[:8]}_d{int(r.Detector)}_{r.kind}"


def mates_of(L, r) -> dict:
    """detector -> label text, for the labelled phase and the model's phase at this signal (lane-mates for the chart)."""
    s = L[L.DeviceId == r.DeviceId]
    phs = {str(r.phase_target)} | ({f"P{int(r.model_phase)}"} if pd.notna(r.model_phase) else set())
    s = s[s.phase_target.astype(str).isin(phs)]
    return {int(d): f"{NAME.get(t, t) if isinstance(t, str) else 'unlabelled'} {p}"
            for d, t, p in zip(s.Detector, s.truth_v3s, s.phase_target)}


def save_chart_data(rows: pd.DataFrame):
    import duckdb
    DATA.mkdir(parents=True, exist_ok=True)
    cn = duckdb.connect()
    cn.execute("SET memory_limit='10GB'")
    cn.execute("SET threads=4")
    cn.execute(f"SET temp_directory='{(DC_WORK / 'tmp').as_posix()}'")
    L = R.label_facts()
    for r in rows.itertuples():
        st = stem(r)
        if (DATA / f"{st}_on.parquet").exists():
            continue
        dets = sorted(set(mates_of(L, r)) | set(r.dets))
        iv = cn.sql(f"""SELECT Detector::INT det, t_on, t_off FROM '{(R.STG / 'det_intervals.parquet').as_posix()}'
                        WHERE lower(DeviceId) = '{r.DeviceId}' AND Detector IN ({','.join(map(str, dets))})
                        ORDER BY t_on""").df().drop_duplicates()
        iv.to_parquet(DATA / f"{st}_on.parquet", index=False)
        cy = cn.sql(f"""SELECT Phase::INT phase, green_start, yellow_start, red_start, next_green
                        FROM '{(R.STG / 'phase_cycles.parquet').as_posix()}' WHERE lower(DeviceId) = '{r.DeviceId}'
                        ORDER BY green_start""").df()
        cy.to_parquet(DATA / f"{st}_cycles.parquet", index=False)


def cycle_profile(iv: pd.DataFrame, cy: pd.DataFrame, phase: int, lo=-40, hi=60, step=1.0):
    """share of cycles with the detector ON, and actuation starts per cycle per second, vs seconds from begin of green
    of `phase` (negative = the red before it); plus median green length."""
    g = cy[cy.phase == phase].dropna(subset=["green_start"])
    if len(g) < 5 or len(iv) == 0:
        return None
    gs = g.green_start.to_numpy("datetime64[ns]").astype("int64")
    glen = float(np.nanmedian((g.yellow_start - g.green_start).dt.total_seconds()))
    t = np.arange(lo, hi + step, step)
    on = iv.t_on.to_numpy("datetime64[ns]").astype("int64")
    off = iv.t_off.fillna(iv.t_on).to_numpy("datetime64[ns]").astype("int64")
    off = np.maximum(off, on)
    # occupancy: for each grid time per cycle, is there an interval covering it? use sorted on + running max of off
    order = np.argsort(on)
    on, off = on[order], np.maximum.accumulate(off[order])
    occ = np.zeros(len(t))
    starts = np.zeros(len(t))
    for j, s in enumerate(t):
        q = gs + int(s * 1e9)
        i = np.searchsorted(on, q, side="right") - 1
        occ[j] = np.mean((i >= 0) & (off[np.clip(i, 0, None)] >= q))
        a = np.searchsorted(on, q, side="left")
        b = np.searchsorted(on, q + int(step * 1e9), side="left")
        starts[j] = np.mean(b - a) / step
    return t, occ, starts, glen


def draw(rows: pd.DataFrame, out_dir: Path) -> dict:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.dates as mdates
    import matplotlib.pyplot as plt
    out_dir.mkdir(parents=True, exist_ok=True)
    L = R.label_facts()
    res = {}
    for r in rows.itertuples():
        st = stem(r)
        iv = pd.read_parquet(DATA / f"{st}_on.parquet")
        cy = pd.read_parquet(DATA / f"{st}_cycles.parquet")
        det = int(r.Detector)
        mates = mates_of(L, r)
        me = iv[iv.det == det]
        lab_ph = int(str(r.phase_target)[1:]) if str(r.phase_target)[1:].isdigit() else None
        mod_ph = int(r.model_phase) if pd.notna(r.model_phase) else None
        fig = plt.figure(figsize=(10.5, 9.2), facecolor=SURF)
        gsp = fig.add_gridspec(3, 2, height_ratios=[1, 1, 0.95], width_ratios=[1.5, 1])
        # (1) top: phase item -> share of ONs starting in each phase's green; function -> 15-min counts vs mates
        ax = fig.add_subplot(gsp[0, :])
        R._style(ax)
        if r.kind == "phase":
            sh = R.green_share(me.t_on.to_numpy("datetime64[ns]").astype("int64"), cy)
            col = [C[1] if p == mod_ph else C[0] if p == lab_ph else "#c9c8c3" for p in sh.index]
            ax.bar([f"P{p}" for p in sh.index], 100 * sh.to_numpy(), color=col)
            ax.set_ylabel("% of its actuations\nstarting in that green", color=INK2, fontsize=9)
            ax.set_title(f"Which phase's green its actuations fall in: blue = label P{lab_ph}, orange = model P{mod_ph}",
                         fontsize=9, color=INK2, loc="left")
        else:
            iv2 = iv.assign(t=iv.t_on.dt.floor("15min"))
            cnt = iv2.groupby(["t", "det"]).size().unstack(fill_value=0)
            grp = [d for d in r.dets if d != det][:2]
            oth = [d for d in cnt.columns if d != det and d not in grp][:4]
            for i, d in enumerate([det] + grp + oth):
                if d not in cnt:
                    continue
                tag = "this one" if d == det else ("same row" if d in grp else mates.get(d, ""))
                ax.plot(cnt.index, cnt[d], color=C[i % len(C)], lw=1.9 if d == det else 1.0,
                        ls="-" if d == det or d not in grp else "--", label=f"det {d} ({tag})")
            ax.set_ylabel("actuations per 15 min", color=INK2, fontsize=9)
            ax.legend(frameon=False, fontsize=8, loc="lower left", bbox_to_anchor=(0, 1.0), ncol=4)
            ax.xaxis.set_major_formatter(mdates.DateFormatter("%a %H:%M"))
        # (2) middle: cycle profile around begin of green (label phase, and model phase for phase items)
        a4 = fig.add_subplot(gsp[1, 0])
        a5 = fig.add_subplot(gsp[1, 1])
        for a in (a4, a5):
            R._style(a)
        phs = [p for p in dict.fromkeys((lab_ph, mod_ph)) if p is not None]
        if r.kind == "function":
            phs = phs[:1]
        glen = None
        for i, p in enumerate(phs):
            pr = cycle_profile(me, cy, p)
            if pr is None:
                continue
            t, occ, starts, gl = pr
            glen = gl if glen is None else glen
            colr = C[0] if p == lab_ph else C[1]
            nm = f"P{p} green" + (" (label)" if p == lab_ph else " (model)")
            a4.plot(t, 100 * occ, color=colr, lw=1.8, label=nm)
            a5.plot(t, starts, color=colr, lw=1.6, label=nm)
            if i == 0:
                for a in (a4, a5):
                    a.axvspan(0, gl, color=C[2], alpha=0.13, lw=0)
        for a in (a4, a5):
            a.axvline(0, color=C[2], lw=1)
            a.set_xlabel("s from begin of green (shaded = median green, left of 0 = red)", color=INK2, fontsize=7.5)
            if len(phs) > 1:
                a.legend(frameon=False, fontsize=7.5, loc="upper right")
        a4.set_ylabel("% of cycles ON", color=INK2, fontsize=9)
        a4.set_ylim(0, 100)
        a4.set_title("Is it ON during red, through green, or not at all?", fontsize=9, color=INK2, loc="left")
        a5.set_ylabel("starts per cycle per s", color=INK2, fontsize=9)
        a5.set_title("When new actuations start", fontsize=9, color=INK2, loc="left")
        # (3) bottom left: 10 busy daytime minutes with green / yellow / red of the labelled phase
        a2 = fig.add_subplot(gsp[2, 0])
        R._style(a2)
        c10 = me.set_index("t_on").resample("10min").size()
        c10 = c10[(c10.index.hour >= 7) & (c10.index.hour < 19)]
        if len(c10):
            busy = c10[c10 >= c10.quantile(0.6)]
            t0 = busy.index[len(busy) // 2]
        else:
            t0 = me.t_on.min() if len(me) else cy.green_start.min()
        t0 = pd.Timestamp(t0)
        t1 = t0 + pd.Timedelta(minutes=10)
        ylab = []
        for p in phs:
            g = cy[(cy.phase == p) & (cy.next_green >= t0) & (cy.green_start <= t1)]
            y = len(ylab)
            for c0, c1, colr in (("green_start", "yellow_start", C[2]), ("yellow_start", "red_start", C[3]),
                                 ("red_start", "next_green", "#d9534f")):
                w = (g[c1] - g[c0]).dt.total_seconds() / 86400
                a2.barh([y] * len(g), w, left=g[c0], height=0.6, color=colr, alpha=0.75 if colr != C[2] else 0.9)
            ylab.append(f"P{p} signal" + (" (label)" if p == lab_ph else " (model)"))
        dd = [det] + [d for d in r.dets if d != det][:1] + \
             [d for d in sorted(iv.det.unique()) if d != det and d not in r.dets][:2]
        for d in dd:
            s = iv[(iv.det == d) & (iv.t_off >= t0) & (iv.t_on <= t1)]
            y = len(ylab)
            a2.barh([y] * len(s), (s.t_off - s.t_on).dt.total_seconds().fillna(0).clip(lower=1.0) / 86400,
                    left=s.t_on, height=0.6, color=C[0] if d == det else C[5])
            ylab.append(f"det {d}" + ("" if d == det else f" ({'same row' if d in r.dets else mates.get(d, '')})"))
        a2.set_yticks(range(len(ylab)))
        a2.set_yticklabels(ylab, fontsize=7.5)
        a2.set_xlim(t0, t1)
        a2.xaxis.set_major_formatter(mdates.DateFormatter("%H:%M"))
        a2.grid(axis="y", visible=False)
        a2.grid(axis="x", color=GRID, lw=0.8)
        a2.set_title("10 busy daytime minutes (green / yellow / red; ONs under 1 s drawn 1 s wide)", fontsize=9,
                     color=INK2, loc="left")
        # (4) bottom right: ON-duration histogram (equal-width bins on a log scale)
        a3 = fig.add_subplot(gsp[2, 1])
        R._style(a3)
        du = (me.t_off - me.t_on).dt.total_seconds().dropna().clip(0.1, 120)
        bins = np.geomspace(0.1, 120, 22)
        a3.hist(du, bins=bins, color=C[0])
        a3.set_xscale("log")
        a3.set_xticks([0.1, 0.3, 1, 3, 10, 30, 100])
        a3.set_xticklabels(["0.1", "0.3", "1", "3", "10", "30", "100"])
        a3.set_xlabel("seconds ON per actuation", color=INK2, fontsize=8)
        a3.set_ylabel("actuations", color=INK2, fontsize=8)
        a3.set_title(f"ON length: {(du <= 0.15).mean():.0%} one-tick, median {du.median():.1f} s",
                     fontsize=9, color=INK2, loc="left")
        fl = NAME.get(r.label_function, r.label_function) if isinstance(r.label_function, str) else ""
        fm = NAME.get(r.model_function, r.model_function) if isinstance(r.model_function, str) else ""
        lab = (f"phase {lab_ph} " if lab_ph else "") + fl
        mod = (f"phase {mod_ph} " if mod_ph else "") + fm
        extra = f"; same pattern: det {', '.join(map(str, [d for d in r.dets if d != det]))}" if r.n_dets > 1 else ""
        fig.suptitle(f"{r.DeviceName} det {det} ({r.technology or 'technology unknown'}): label {lab}, model says {mod}"
                     f"\nSept 2026 log, {len(me):,} actuations{extra}", x=0.01, ha="left", fontsize=10, color=INK)
        fig.tight_layout()
        png = out_dir / f"{st}.png"
        fig.savefig(png, dpi=95, facecolor=SURF)
        plt.close(fig)
        res[st] = png
    return res


# ------------------------------------------------------------------------------------------------ sheet (note 91 layout)
# The user found "P5, P8 stop-bar presence" ambiguous (phase list + one function). Now each detector's phase and function
# are written together, one detector per line: "det 15: P5 Presence" / "det 23: P8 Presence" (user, 2026-10-05).
FN_FULL = {"Advance": "Advance", "Presence": "Presence", "Count": "Stop-bar Count", "Yellow_Red": "Yellow-Red",
           "Other": "Other", "Mid": "Mid", "Bike": "Bike"}
HEAD = ["#", "Signal", "Detector(s)", "Label", "Model says", "How sure", "What this row asks", "Chart", "Print", "Answer"]
WIDTH = (4, 8, 10, 24, 24, 7, 60, 7, 7, 30)
ARM = "fj"                                     # review81 sets "champ" (prediction column of C64 / function_rows.parquet)


def phase_txt(p) -> str:
    s = str(p).strip()
    s = s[1:] if s[:1] in "Pp" else s
    try:
        return f"P{int(float(s))}"
    except ValueError:
        return "(no phase)"


def fn_txt(t, model: bool = False) -> str:
    if not isinstance(t, str) or not t or t == "None":
        return "(no function answer)" if model else "(function not labelled)"
    return FN_FULL.get(t, t)


def pf_txt(p, t, model: bool = False) -> str:
    """'P5 Presence'."""
    return f"{phase_txt(p) if p is not None else '(no phase)'} {fn_txt(t, model)}"


def _mode(s):
    s = s.dropna()
    return s.value_counts().index[0] if len(s) else None


def det_facts(rows: pd.DataFrame, frows: Path | None = None, arm: str | None = None) -> pd.DataFrame:
    """per detector of each row: labelled phase / function and the model's phase / function, as lists aligned with dets.
    Label phase = label table phase_target; model phase = the model's most common phase over the detector's >= 30-min
    Sept samples. Function rows: label / model function are the row's (identical for every detector by the collapse key).
    Phase rows: phase from the row (identical by the key); each detector's own build-time label function and most common
    model function."""
    frows = frows or C64 / "function_rows.parquet"
    arm = arm or ARM
    L = R.label_facts()
    lp = {(d, int(x)): p for d, x, p in zip(L.DeviceId, L.Detector, L.phase_target)}
    lt = {(d, int(x)): t for d, x, t in zip(L.DeviceId, L.Detector, L.truth_v3s)}
    f = pd.read_parquet(frows, columns=["DeviceId", "Detector", "period", "wgroup", "pred_phase", "truth_v3s",
                                        f"pred_{arm}"])
    f = f[(f.period == "stg") & f.wgroup.isin(GE30) & f.DeviceId.isin(set(rows.DeviceId))]
    g = f.groupby(["DeviceId", "Detector"])
    M = pd.DataFrame({"mph": g.pred_phase.agg(_mode), "mfn": g[f"pred_{arm}"].agg(_mode), "tfn": g.truth_v3s.agg(_mode)})
    M = {(d, int(x)): v for (d, x), v in zip(M.index, M.itertuples(index=False))}
    cols = {k: [] for k in ("d_lab_ph", "d_lab_fn", "d_mod_ph", "d_mod_fn")}
    for r in rows.itertuples():
        a = {k: [] for k in cols}
        for d in _dets(r):
            k = (r.DeviceId, int(d))
            m = M.get(k)
            if r.kind == "phase":
                a["d_lab_ph"].append(str(r.phase_target))
                a["d_mod_ph"].append(int(r.model_phase))
                lf = r.label_function if int(d) == int(r.Detector) else (m.tfn if m is not None else lt.get(k))
                a["d_lab_fn"].append(lf if isinstance(lf, str) else None)
                a["d_mod_fn"].append(m.mfn if m is not None and isinstance(m.mfn, str) else None)
            else:
                a["d_lab_ph"].append(lp.get(k))
                a["d_mod_ph"].append(int(m.mph) if m is not None and pd.notna(m.mph) else None)
                a["d_lab_fn"].append(r.label_function)
                a["d_mod_fn"].append(r.model_function)
        for k in cols:
            cols[k].append(a[k])
    rows = rows.copy()
    for k, v in cols.items():
        rows[k] = v
    return rows


def _dets(r) -> list:
    return [int(x) for x in r.dets] if isinstance(r.dets, (list, tuple, np.ndarray)) else \
        [int(x) for x in str(r.dets).split(",")]


def _per_det(dets, vals) -> str:
    """one value, or one 'det N: value' line per detector when there are several."""
    if len(dets) == 1:
        return vals[0]
    return "\n".join(f"det {d}: {v}" for d, v in zip(dets, vals))


def _join(xs) -> str:
    xs = [str(x) for x in xs]
    return xs[0] if len(xs) == 1 else ", ".join(xs[:-1]) + " and " + xs[-1]


def row_cells(r) -> dict:
    """Label / Model says cells ('P5 Presence', one 'det N: P5 Presence' line per detector when several) + the one-line
    question of a review row (needs det_facts columns)."""
    dets = _dets(r)
    lab = [pf_txt(p, t) for p, t in zip(r.d_lab_ph, r.d_lab_fn)]
    mod = [pf_txt(p, t, model=True) for p, t in zip(r.d_mod_ph, r.d_mod_fn)]
    if r.kind == "phase":
        lph, mph = phase_txt(r.d_lab_ph[0]), phase_txt(r.d_mod_ph[0])
        who = f"det {dets[0]} is" if len(dets) == 1 else f"dets {_join(dets)} are"
        ask = f"Label says {who} on {lph}; model says {mph}. Which phase is right?"
    elif len(dets) == 1:
        ask = f"Label says det {dets[0]} is {lab[0]}; model says {mod[0]}. Which is right?"
    else:
        ask = (f"Label says {_join([f'det {d} is {x}' for d, x in zip(dets, lab)])}; "
               f"model says {_join([f'det {d} is {x}' for d, x in zip(dets, mod)])}. Which is right?")
    return {"lab": _per_det(dets, lab), "mod": _per_det(dets, mod), "ask": ask}


def det_text(r) -> str:
    d = _dets(r)
    if len(d) > MAX_DETS_ROW:
        return ", ".join(map(str, d[:MAX_DETS_ROW])) + f" +{len(d) - MAX_DETS_ROW} more"
    return ", ".join(map(str, d))


def read_answers(xl: Path) -> dict:
    """(signal, detector tuple) -> the user's non-empty Answer cell, from either sheet layout (Answer = last header)."""
    from openpyxl import load_workbook
    wb = load_workbook(xl, read_only=True, data_only=True)
    ws = wb["labels to check"]
    head = list(next(ws.iter_rows(min_row=2, max_row=2, values_only=True)))
    ia, idt = head.index("Answer"), next(i for i, h in enumerate(head) if str(h).startswith("Detector"))
    out = {}
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[0] is None or r[ia] is None or str(r[ia]).strip() == "":
            continue
        key = (str(r[1]), tuple(int(x) for x in str(r[idt]).split(",") if x.strip().isdigit()))
        assert key not in out, key
        out[key] = r[ia]
    wb.close()
    return out


def write_sheet(rows: pd.DataFrame, charts: dict | None, xl: Path, chart_rel: str, answers: list | None = None,
                links: list | None = None):
    """answers: per row, the user's Answer text to carry over (note 91); links: per row (chart target, print target or
    None) from an earlier build of the same sheet, else computed from charts / R.pdf_for."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    if "d_lab_ph" not in rows:
        rows = det_facts(rows)
    if answers is None and xl.exists():          # never drop what the user typed: carry it to the same detector row
        old = read_answers(xl)
        answers = [old.pop((str(r.DeviceName), tuple(_dets(r))), None) for r in rows.itertuples()]
        if old:
            raise RuntimeError(f"{xl}: answers on rows that are not in the new build, not overwriting: {old}")
    wb = Workbook()
    ws = wb.active
    ws.title = "labels to check"
    ws.append([QUESTION])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEAD))
    ws.row_dimensions[1].height = 48
    ws.append(HEAD)
    link = Font(color="0563C1", underline="single")
    cc, cp = HEAD.index("Chart") + 1, HEAD.index("Print") + 1
    for i, r in enumerate(rows.itertuples(), 1):
        x = row_cells(r)
        if links is not None:
            ch, pdf = links[i - 1]
        else:
            ch = f"{chart_rel}/{Path(charts[stem(r)]).name}"
            p = R.pdf_for(r.DeviceName)
            pdf = "file:///" + urllib.parse.quote(str(p).replace("\\", "/"), safe=":/()_-.,'") if p else None
        ans = answers[i - 1] if answers is not None else None
        ws.append([i, r.DeviceName, det_text(r), x["lab"], x["mod"], f"{r.conf_max:.0%}", x["ask"], "chart",
                   "print" if pdf else "none", "" if ans is None else ans])
        n = ws.max_row
        c = ws.cell(n, cc)
        c.hyperlink = ch
        c.font = link
        if pdf:
            c = ws.cell(n, cp)
            c.hyperlink = pdf
            c.font = link
    for j, w in enumerate(WIDTH):
        ws.column_dimensions[chr(65 + j)].width = w
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"
    xl.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xl)
    log(f"saved {xl} ({len(rows)} rows)")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--write", action="store_true")
    ap.add_argument("--limit", type=int, default=0, help="charts + sheet for the first N rows only (preview)")
    a = ap.parse_args()
    t0 = time.time()
    meta = json.load(open(C64 / "function_meta.json"))
    assert len(meta["fj_folds"]) == 6, "fj arm needs all six folds"
    items, dry64 = R.build_items("fj")
    assert not items.DeviceId.isin(R.locked()).any()
    it = select(items)
    rows = collapse(it)
    dry = {"review64_items": int(len(items)), "review64_by_kind": items.kind.value_counts().to_dict(),
           "dropped_items": it.why_drop[it.why_drop != ""].value_counts().to_dict(),
           "kept_items": int((it.why_drop == "").sum()),
           "would_add_items_if_conf_ge_0.8": int(((it.why_drop == "not confident enough") & it.conf.ge(0.8)).sum()),
           "kept_items_by_type": {f"{k} / {t}": int(n) for (k, t), n in
                                  it[it.why_drop == ""].groupby(["kind", "type"]).size().items()},
           "rows": int(len(rows)), "signals": int(rows.DeviceId.nunique()),
           "rows_by_type": rows.types.value_counts().to_dict(), "rows_by_kind": rows.kind.value_counts().to_dict(),
           "rows_multi_detector": int((rows.n_dets > 1).sum()),
           "rows_by_pattern": (rows.kind + ": " + rows.label_function.astype(str) + " -> " +
                               rows.model_function.astype(str)).value_counts().head(15).to_dict(),
           "rows_by_technology": rows.technology.fillna("unknown").value_counts().to_dict(),
           "rows_by_label_source": rows.source.fillna("none").value_counts().to_dict(),
           "rows_with_print": int(sum(R.pdf_for(n) is not None for n in rows.DeviceName)),
           "rows_conf_ge": {f"{t:.2f}": int((rows.conf_max >= t).sum()) for t in (0.7, 0.8, 0.9, 0.95)}}
    log(json.dumps(dry, indent=1, default=str))
    json.dump(dry, open(C64 / "review72_dryrun.json", "w"), indent=1, default=str)
    rows.drop(columns=[c for c in rows.columns if rows[c].map(lambda v: isinstance(v, (set, list))).any()]) \
        .assign(dets=rows.dets.map(lambda d: ",".join(map(str, d)))).to_parquet(C64 / "review72_rows.parquet", index=False)
    if a.write or a.limit:
        sel = rows.head(a.limit) if a.limit else rows
        save_chart_data(sel)
        log("chart data saved")
        if a.write:
            out = R.REPO / "review"
            ch = draw(sel, out / f"{OUTNAME}_charts")
            write_sheet(sel, ch, out / f"{OUTNAME}.xlsx", f"{OUTNAME}_charts")
        else:
            d = C64 / "review72_preview"
            ch = draw(sel, d / "charts")
            write_sheet(sel, ch, d / f"{OUTNAME}.xlsx", "charts")
    log(f"done in {time.time() - t0:.0f} s")


if __name__ == "__main__":
    main()
