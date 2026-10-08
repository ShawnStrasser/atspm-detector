"""Note 102: lane-count review sheet v2 -- ONE ROW PER SIGNAL, after the wide-right-lane fix.

Model = lanes from research/code/lanes/ln102_wide.py (six-fold OOF, 2026 recipe) for the chosen variant (--tag) and the
retrained baseline (--base, to find errors the fix introduced).  Model phase = frame OOF predicted phase; model function =
the v5 function OOF decoded through the same lanes (ln102_wide atspm -> f102/ln8/func_<out>.parquet).  Truth = print lanes
v4q + the user's lane answers of review v1.  Scope = Sept-2026 log, the 14 samples >= 30 min; a phase's model answer = the
majority over its samples (note 99).  Left out: phases the user answered in review v1, under-counts the log cannot show
(note 99 'blind'), phases with how-sure < MIN_SCORE.  locked_v2 asserted absent.  Charts from SAVED sample data only.

    python lane_review102.py rows  --tag W --base base --func main
    python lane_review102.py data
    python lane_review102.py write          -> review/lane_count_review_v2.xlsx + review/lane_count_review_v2_charts/
"""
from __future__ import annotations

import os
for _v in ("OMP_NUM_THREADS", "MKL_NUM_THREADS", "OPENBLAS_NUM_THREADS"):
    os.environ.setdefault(_v, "4")
import argparse  # noqa: E402
import json  # noqa: E402
import sys  # noqa: E402
import urllib.parse  # noqa: E402
from collections import namedtuple  # noqa: E402
from pathlib import Path  # noqa: E402

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

CODE = next(p for p in Path(__file__).resolve().parents if p.name == "code")
for _d in ("", "lanes", "evaluation"):
    sys.path.insert(0, str(CODE / _d) if _d else str(CODE))
import rpath  # noqa: E402
import ln1_cues as L1  # noqa: E402
import ln102_wide as M  # noqa: E402
import lane_review99 as R99  # noqa: E402
import review64 as R  # noqa: E402

DCW = L1.DCW
REPO = L1.REPO
OUT = DCW / "rev102"
DATA = OUT / "review_data"
V1 = REPO / "review" / "lane_count_review.xlsx"
OUTNAME = "lane_count_review_v2"
MIN_SCORE = 0.70
MIN_SCORE_NEW = 0.60
N_KEEP = 17
N_NEW = 3
MIN_ACT = 10
ATS = {"Advance", "Presence", "Count", "Yellow_Red"}
FN = {"Advance": "Advance", "Presence": "Presence", "Count": "Count", "Yellow_Red": "Yellow-Red", "Other": "Other",
      "Mid": "Mid", "Bike": "Bike", None: "no print function"}
ORDER = ["full66", "h24_a", "h24_b", "h6_a", "h6_b", "h3_a", "h3_b", "h1_a", "h1_c", "h1_b",
         "m30_a", "m30_b", "m30_d", "m30_c"]
QUESTION = ("One row per signal. For each phase listed: is the print's lane count right? Answer Y (print right) / "
            "N (model right) / ? per phase, e.g. 'P4 Y, P8 N'. Comment if you see why the model is wrong.")
HEAD = ["#", "Signal", "Phases affected (print lanes vs model lanes)", "Detectors (print | model)", "Likely cause",
        "Question", "Print link", "Chart", "Answer", "Comment"]
WIDTH = (4, 8, 26, 70, 22, 46, 7, 7, 14, 30)


def log(m):
    print(m, flush=True)


def reviewed_v1() -> set:
    """(signal, phase) the user answered in review v1 (read only; the file is never written)."""
    from openpyxl import load_workbook
    ws = load_workbook(V1, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    h = list(rows[1])
    i_s, i_p, i_a, i_c = h.index("Signal"), h.index("Phase"), h.index("Answer"), h.index("Comment")
    return {(str(r[i_s]), str(r[i_p])) for r in rows[2:] if r and (r[i_a] not in (None, "") or r[i_c] not in (None, ""))}


def majority(tag: str) -> pd.DataFrame:
    x = M.sample_rows(tag)
    g = x.groupby(["DeviceId", "target"])
    Mj = g.agg(n=("win", "size"), truth=("n_lanes", "first"), ph_num=("ph_num", "first"),
               maj=("n_lanes_p", lambda s: s.value_counts().sort_index().idxmax())).reset_index()
    x = x.merge(Mj[["DeviceId", "target", "maj"]], on=["DeviceId", "target"])
    a = x[x.n_lanes_p == x.maj].groupby(["DeviceId", "target"]).agg(
        k=("win", "size"), conf=("n_lanes_conf", "mean"), wins=("win", lambda s: ",".join(sorted(s)))).reset_index()
    Mj = Mj.merge(a, on=["DeviceId", "target"])
    Mj["wrong"] = Mj.maj != Mj.truth
    Mj["dir"] = np.where(Mj.maj > Mj.truth, "over", np.where(Mj.maj < Mj.truth, "under", "ok"))
    Mj["score"] = Mj.k / Mj.n * Mj.conf
    return Mj


def blind_flags(Mj: pd.DataFrame, tag: str) -> pd.DataFrame:
    """note 99: an under-count is blind when the active single-lane vehicle detectors cover no more print lanes than
    the model counted (the log cannot show the missing lane)."""
    det = pd.read_parquet(M.OUT / "truth_det.parquet")
    D = pd.read_parquet(M.lanes_file(tag))
    D = D[(D.period == "stg") & (D.win == "full66")]
    non = D.groupby(["DeviceId", "Detector"]).n_on.first()
    det["n_on66"] = [non.get((s, x), 0) for s, x in zip(det.DeviceId, det.det)]
    det["n_on66"] = det.n_on66.fillna(0)
    lv = det.lane_index.notna() & ~det.lane_type.isin(L1.NONVEH)
    act = det[lv & (det.span == 1) & (det.n_on66 >= MIN_ACT)]
    vis = act.groupby(["DeviceId", "target"]).lane_index.nunique()
    Mj["visible"] = [int(vis.get((s, t), 0)) for s, t in zip(Mj.DeviceId, Mj.target)]
    Mj["blind"] = (Mj.dir == "under") & (Mj.visible.clip(lower=1) <= Mj.maj)
    return Mj


def det_items(dev, target, ph_num, win, tag, func_col, F, det, LF, Dall):
    dw = Dall[(Dall.DeviceId == dev) & (Dall.win == win)].set_index("Detector")
    pr = det[(det.DeviceId == dev) & (det.target == target)].set_index("det")
    grp = set(dw.index[(dw.phase == ph_num) & (dw.lanes != "")].astype(int))
    dets = sorted(set(pr.index.astype(int)) | grp)
    lf = LF[LF.DeviceId == dev].set_index("det")
    fw = F[(F.DeviceId == dev) & (F.win == win)].set_index("Detector")
    items = []
    for d in dets:
        on_print = d in pr.index
        fn = pr.print_function.get(d) if on_print else (lf.print_function.get(d) if d in lf.index else None)
        if not isinstance(fn, str):
            fn = lf.truth_v3s.get(d) if d in lf.index else None
        p_l = R99.plane_txt(pr.lane_index.get(d), pr.span.get(d)) if on_print else ""
        other_ph = None if on_print else (lf.phase_target.get(d) if d in lf.index else None)
        if d in dw.index:
            mph, ml, note, n_on = dw.phase.get(d), dw.lanes.get(d), dw.lane_note.get(d), dw.n_on.get(d)
        else:
            mph, ml, note, n_on = np.nan, "", "", 0
        mf = fw[func_col].get(d) if d in fw.index else None
        fph = fw.pred_phase.get(d) if d in fw.index else np.nan
        if pd.isna(mph) and pd.notna(fph):
            mph = fph
        items.append(dict(det=int(d), fn=fn if isinstance(fn, str) else None, on_print=on_print, p_lane=p_l,
                          other_ph=other_ph if isinstance(other_ph, str) else None,
                          m_phase=None if pd.isna(mph) else int(mph), m_lanes=ml if isinstance(ml, str) else "",
                          m_note=note if isinstance(note, str) else "", n_on=int(n_on) if pd.notna(n_on) else 0,
                          m_func=mf if isinstance(mf, str) else None))
    return items


def det_line(i, target) -> str:
    fn = FN.get(i["fn"], i["fn"])
    if i["on_print"]:
        p = f"print {target} {fn} " + (("lanes " if "+" in i["p_lane"] else "lane ") + i["p_lane"] if i["p_lane"]
                                        else "(no lane)")
    else:
        p = f"print {i['other_ph'] or 'not on ' + target} {fn}"
    if i["m_phase"] is None:
        m = "model: no actuations in the sample" if i["n_on"] == 0 else "model: no phase"
    else:
        mf = FN.get(i["m_func"], i["m_func"] or "?")
        if i["m_lanes"] and i["m_phase"] == int(target[1:]):
            ml = ("lanes " if "," in i["m_lanes"] else "lane ") + R99.lanes_txt(i["m_lanes"])
        elif i["m_note"].startswith("too few") or (i["n_on"] < MIN_ACT and i["m_phase"] == int(target[1:])):
            ml = "no lane (too few actuations)"
        elif i["m_note"] == "bike detector" and i["m_func"] != "Bike":
            ml = "no lane (the lane step saw it as Bike)"
        elif i["m_func"] == "Bike" or i["m_note"] == "bike detector":
            ml = "no lane"
        else:
            ml = ""
        m = f"model P{i['m_phase']} {mf}" + (f" {ml}" if ml else "")
    flags = []
    tp = int(target[1:])
    if i["on_print"] and i["m_phase"] is not None and i["m_phase"] != tp:
        flags.append("wrong phase")
    if not i["on_print"] and i["m_phase"] == tp and i["m_lanes"]:
        flags.append("wrong phase")
    if i["m_phase"] is not None and i["m_func"]:
        pf, mf = i["fn"], i["m_func"]
        if (pf in ATS or mf in ATS) and pf != mf and i["on_print"]:
            flags.append("wrong function")
    return f"det {i['det']}: {p} | {m}" + (f"   <- {', '.join(flags)}" if flags else "")


def cause_of(r, items, pat) -> str:
    tp = r.ph_num
    moved_out = [i for i in items if i["on_print"] and i["p_lane"] and i["m_phase"] not in (None, tp)
                 and i["fn"] not in ("Bike", "Other", None)]
    moved_in = [i for i in items if not i["on_print"] and i["m_phase"] == tp and i["m_lanes"]]
    if (r.dir == "over" and moved_in) or (r.dir == "under" and moved_out):
        return "wrong phase"
    nolane = [i for i in items if i["on_print"] and i["p_lane"] and i["m_phase"] == tp and not i["m_lanes"]
              and i["n_on"] >= MIN_ACT and i["fn"] in ATS]
    anchor_from_non = [i for i in items if i["on_print"] and i["fn"] not in ATS and i["m_func"] in ATS
                       and i["m_phase"] == tp and i["m_lanes"]]
    demoted = [i for i in items if i["on_print"] and i["p_lane"] and i["fn"] in ATS and i["m_phase"] == tp
               and i["m_func"] not in ATS]
    if (r.dir == "under" and (nolane or demoted)) or (r.dir == "over" and anchor_from_non and not pat):
        return "wrong function"
    if pat:
        return "wide right lane"
    if r.dir == "over" and len({i["m_func"] for i in items if i["m_lanes"]}) and any(
            i["on_print"] and i["fn"] in ATS and i["m_func"] in ATS and i["fn"] != i["m_func"] for i in items):
        return "wrong function"
    return "lane grouping"


def art(w) -> str:
    return ("an " if str(w)[:1].lower() in "aeiou" else "a ") + str(w)


def ask(r, items, cause) -> str:
    ln = lambda n: f"{int(n)} lane" + ("" if int(n) == 1 else "s")  # noqa: E731
    tp = r.ph_num
    head = f"{r.target}: print {ln(r.truth)}, model {ln(r.maj)}"
    on = [i for i in items if i["on_print"] and i["p_lane"] and i["m_phase"] == tp and i["m_lanes"]]
    if cause == "wrong phase":
        mi = [i for i in items if not i["on_print"] and i["m_phase"] == tp and i["m_lanes"]]
        mo = [i for i in items if i["on_print"] and i["p_lane"] and i["m_phase"] not in (None, tp)]
        if r.dir == "over" and mi:
            i = mi[0]
            return f"{head} (adds det {i['det']}, print {i['other_ph'] or 'other phase'}). Is det {i['det']} on " \
                   f"{i['other_ph'] or 'another phase'}?"
        if mo:
            i = mo[0]
            return f"{head} (moves det {i['det']} to P{i['m_phase']}). Is det {i['det']} on {r.target}?"
    if r.dir == "over":
        for a in on:
            for b in on:
                if a["det"] < b["det"] and a["p_lane"] == b["p_lane"] and \
                        not set(a["m_lanes"].split(",")) & set(b["m_lanes"].split(",")):
                    if cause == "wide right lane":
                        return f"{head}. Are det {a['det']} and det {b['det']} in one (wide) lane?"
                    return f"{head} (splits det {a['det']} and det {b['det']}). Same lane, as the print shows?"
    else:
        for a in on:
            for b in on:
                if a["det"] < b["det"] and not set(a["p_lane"].split("+")) & set(b["p_lane"].split("+")) \
                        and set(a["m_lanes"].split(",")) & set(b["m_lanes"].split(",")):
                    return f"{head} (puts det {a['det']} and det {b['det']} in one lane). Separate lanes, as the " \
                           f"print shows?"
        nl = [i for i in items if i["on_print"] and i["p_lane"] and i["m_phase"] == tp and not i["m_lanes"]
              and i["n_on"] >= MIN_ACT]
        if nl:
            i = nl[0]
            return f"{head} (det {i['det']} gets no lane). " \
                   f"Is det {i['det']} {art(FN.get(i['fn'], i['fn']))} in its own lane?"
    return f"{head}. Is the print's lane count right?"


def stage_rows(a):
    OUT.mkdir(parents=True, exist_ok=True)
    lk = L1.locked()
    rev = reviewed_v1()
    names = pd.read_parquet(rpath.LABELS_CURRENT, columns=["DeviceId", "DeviceName"]).dropna().drop_duplicates()
    names["DeviceId"] = names.DeviceId.str.lower()
    nm = dict(zip(names.DeviceId, names.DeviceName))
    pat = M.pattern_phases()
    pk = set(zip(pat.DeviceId, pat.target))
    Mt = blind_flags(majority(a.tag), a.tag)
    Mb = blind_flags(majority(a.base), a.base)
    assert not Mt.DeviceId.isin(lk).any()
    for X in (Mt, Mb):
        X["DeviceName"] = X.DeviceId.map(nm)
        X["reviewed"] = [(s, t) in rev for s, t in zip(X.DeviceName, X.target)]
        X["pat"] = [(s, t) in pk for s, t in zip(X.DeviceId, X.target)]
    j = Mt.merge(Mb[["DeviceId", "target", "maj", "wrong", "score"]], on=["DeviceId", "target"], suffixes=("", "_b"))
    j["new"] = j.wrong & ~j.wrong_b
    elig = j.wrong & ~j.blind & ~j.reviewed & (j.score >= MIN_SCORE)
    prof = {"phases": int(len(j)), "wrong_after": int(j.wrong.sum()), "wrong_before": int(j.wrong_b.sum()),
            "fixed": int((j.wrong_b & ~j.wrong).sum()), "new_wrong": int(j.new.sum()),
            "blind_after": int((j.wrong & j.blind).sum()), "reviewed_left_out": int((j.wrong & j.reviewed).sum()),
            "confident_nonblind_unreviewed": int(elig.sum()),
            "signals_eligible": int(j[elig].DeviceId.nunique()),
            "new_eligible": int((elig & j.new).sum()),
            "pattern_wrong_after": int((j.wrong & j.pat).sum()), "pattern_wrong_before": int((j.wrong_b & j.pat).sum())}
    log(json.dumps(prof))
    E = j[elig].copy()
    sg = E.groupby("DeviceId").agg(best=("score", "max"), n=("target", "size"), new=("new", "any"),
                                   allnew=("new", "all")).reset_index()
    sg = sg.sort_values(["best", "n"], ascending=False)
    keep = list(sg[~sg.allnew].DeviceId[:N_KEEP])
    # NEW = wrong now, right for the review-v1 model; looser how-sure floor (few are confident)
    En = j[j.new & ~j.blind & ~j.reviewed & (j.score >= MIN_SCORE_NEW)].sort_values("score", ascending=False)
    keep_new = [d for d in dict.fromkeys(En.DeviceId) if d not in keep][:N_NEW]
    E = pd.concat([E, En[En.DeviceId.isin(keep + keep_new)]]).drop_duplicates(["DeviceId", "target"])
    S = E[E.DeviceId.isin(keep + keep_new)].copy()
    S["sig_new"] = S.DeviceId.isin(keep_new)
    prof["new_rows"] = [nm.get(d) for d in keep_new]
    S["rep_win"] = S.wins.map(lambda w: next(o for o in ORDER if o in w.split(",")))
    det = pd.read_parquet(M.OUT / "truth_det.parquet")
    LF = pd.read_parquet(rpath.LABELS_CURRENT, columns=["DeviceId", "detector", "phase_target", "print_function",
                                                        "truth_v3s"])
    LF["DeviceId"] = LF.DeviceId.str.lower()
    LF = LF.rename(columns={"detector": "det"}).drop_duplicates(["DeviceId", "det"])
    LF["det"] = LF.det.astype(int)
    Dall = pd.read_parquet(M.lanes_file(a.tag))
    Dall = Dall[Dall.period == "stg"]
    F = pd.read_parquet(M.OUT / f"func_{a.func}.parquet")
    F["DeviceId"] = F.DeviceId.str.lower()
    F["Detector"] = F.Detector.astype(int)
    fcol = f"func_{a.tag}"
    js, causes, asks = [], [], []
    for r in S.itertuples():
        it = det_items(r.DeviceId, r.target, r.ph_num, r.rep_win, a.tag, fcol, F, det, LF, Dall)
        c = cause_of(r, it, r.pat)
        js.append(json.dumps(it))
        causes.append(c)
        asks.append(ask(r, it, c))
    S["dets_json"], S["cause"], S["ask"] = js, causes, asks
    S = S.sort_values(["sig_new", "DeviceName", "ph_num"]).reset_index(drop=True)
    S.to_parquet(OUT / "rows102.parquet", index=False)
    j.to_parquet(OUT / "phase_majority102.parquet", index=False)
    prof["rows_signals"] = int(S.DeviceId.nunique())
    prof["rows_phases"] = int(len(S))
    prof["causes"] = S.cause.value_counts().to_dict()
    json.dump(prof, open(OUT / "profile102.json", "w"), indent=1)
    log(f"rows: {S.DeviceId.nunique()} signals / {len(S)} phases; causes {prof['causes']}")


def stage_data(a):
    import duckdb
    DATA.mkdir(parents=True, exist_ok=True)
    S = pd.read_parquet(OUT / "rows102.parquet")
    S = S.sort_values("score", ascending=False).drop_duplicates("DeviceId")
    cn = duckdb.connect()
    cn.execute("SET memory_limit='10GB'")
    cn.execute("SET threads=4")
    cn.execute(f"SET temp_directory='{(DCW / 'tmp').as_posix()}'")
    t0 = pd.Timestamp(R99.FULL[0])
    t1 = t0 + pd.Timedelta(seconds=int(R99.FULL[1]))
    for r in S.itertuples():
        f = DATA / f"{r.DeviceName}_{r.target}_on.parquet"
        if f.exists():
            continue
        dets = [i["det"] for i in json.loads(r.dets_json)]
        iv = cn.sql(f"""SELECT DISTINCT Detector::INT det, t_on, t_off FROM '{(R99.STG / 'det_intervals.parquet').as_posix()}'
                        WHERE lower(DeviceId) = '{r.DeviceId}' AND Detector IN ({','.join(map(str, dets))})
                          AND t_on >= TIMESTAMP '{t0}' AND t_on < TIMESTAMP '{t1}' ORDER BY t_on""").df()
        iv.to_parquet(f, index=False)
    log(f"chart data for {len(S)} signals")


def stage_write(a):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    S = pd.read_parquet(OUT / "rows102.parquet")
    out = REPO / "review"
    xl = out / f"{OUTNAME}.xlsx"
    assert not xl.exists(), f"{xl} exists - never overwrite a sheet the user may have typed in"
    chd = out / f"{OUTNAME}_charts"
    chd.mkdir(parents=True, exist_ok=True)
    R99.DATA = DATA
    wb = Workbook()
    ws = wb.active
    ws.title = "lane counts by signal"
    ws.append([QUESTION])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(HEAD))
    ws.row_dimensions[1].height = 34
    ws.append(HEAD)
    link = Font(color="0563C1", underline="single")
    cc, cp = HEAD.index("Chart") + 1, HEAD.index("Print link") + 1
    ln = lambda n: f"{int(n)}"  # noqa: E731
    links = []
    order = S.groupby("DeviceId", sort=False).first().reset_index()
    for n, s in enumerate(order.itertuples(), 1):
        g = S[S.DeviceId == s.DeviceId].sort_values("ph_num")
        new = bool(g.sig_new.iloc[0])
        phases = "\n".join(f"{r.target}: print {ln(r.truth)}, model {ln(r.maj)} ({int(r.k)} of {int(r.n)} samples)" + (" NEW" if r.new else "")
                           for r in g.itertuples())
        dl = "\n".join(f"{r.target}:\n" + "\n".join("  " + R99_line for R99_line in
                                                   [det_line(i, r.target) for i in json.loads(r.dets_json)])
                       for r in g.itertuples())
        cause = "; ".join(dict.fromkeys(f"{r.target} {r.cause}" if len(g) > 1 else r.cause for r in g.itertuples()))
        if new:
            cause = "NEW since review v1 (right then, wrong now): " + cause
        q = "\n".join(r.ask for r in g.itertuples())
        top = g.sort_values("score", ascending=False).iloc[0]
        png = chd / f"{s.DeviceName}_{top.target}.png"
        Row = namedtuple("Row", "DeviceName target truth maj ph_num")
        items = json.loads(top.dets_json)
        for i in items:
            i.setdefault("tech", None)
        try:
            R99.draw(Row(s.DeviceName, top.target, int(top.truth), int(top.maj), int(top.ph_num)), items, png)
            chart = "chart"
        except Exception as e:  # noqa: BLE001
            log(f"chart failed {s.DeviceName}: {e}")
            chart, png = "", None
        pdf = R.pdf_for(s.DeviceName)
        pl = "file:///" + urllib.parse.quote(str(pdf).replace("\\", "/"), safe=":/()_-.,'") if pdf else None
        ws.append([n, s.DeviceName, phases, dl, cause, q, "print" if pl else "none", chart, "", ""])
        k = ws.max_row
        if png is not None:
            ws.cell(k, cc).hyperlink = f"{OUTNAME}_charts/{png.name}"
            ws.cell(k, cc).font = link
        if pl:
            ws.cell(k, cp).hyperlink = pl
            ws.cell(k, cp).font = link
        if new:
            for c in ws[k]:
                c.fill = PatternFill("solid", fgColor="FFF2CC")
        links.append({"row": n, "signal": s.DeviceName, "phases": list(g.target), "chart": png.name if png else None,
                      "print": str(pdf) if pdf else None, "new": new})
    for jx, w in enumerate(WIDTH):
        ws.column_dimensions[chr(65 + jx)].width = w
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"
    wb.save(xl)
    json.dump(links, open(OUT / "links102.json", "w"), indent=1)
    log(f"saved {xl} ({len(order)} rows), charts in {chd}")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("stage", choices=["rows", "data", "write"])
    ap.add_argument("--tag", default="W")
    ap.add_argument("--base", default="base")
    ap.add_argument("--func", default="main")
    a = ap.parse_args()
    {"rows": stage_rows, "data": stage_data, "write": stage_write}[a.stage](a)
