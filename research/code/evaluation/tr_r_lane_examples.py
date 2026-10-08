"""Note 102b: how the print-lane truth (v4q + user lane answers, f102/ln8) treats a loop approach whose stop bar has a
through/right loop plus an extra loop on the right (corner widening).  Counts by how it was counted (1 or 2 lanes) and print
confidence; writes review/tr_r_lane_examples.xlsx (10 examples).  Labels/analysis only, locked signals asserted absent.

    python tr_r_lane_examples.py
"""
from __future__ import annotations

import json
import os
import urllib.parse
from pathlib import Path

import pandas as pd

DCW = Path(os.environ.get("DC_WORK") or (Path.home() / "dc_work"))
REPO = next(p for p in Path(__file__).resolve().parents if p.name == "research").parent
W = DCW / "final_v3_work" / "f102" / "ln8"
PDF = DCW / "cabinet" / "pdf"
XL = REPO / "review" / "tr_r_lane_examples.xlsx"
SB = ["stopbar_presence", "stopbar_count", "stopbar_secondary", "presence_20_75"]
USER_FIXED = {("01030", "P8"), ("01072", "P4")}  # lane_truth_corrections_user_v1.csv
EXAMPLES = [("04059", "P8"), ("01074", "P4"), ("01017", "P8"), ("03044", "P4"),
            ("01072", "P8"), ("01063", "P8"), ("03022", "P4"), ("2B337", "P8"), ("03054", "P4"), ("04013", "P4")]


def load():
    T = pd.read_parquet(W / "truth_det.parquet")
    P = pd.read_parquet(W / "truth_phase.parquet")
    L = pd.read_parquet(REPO / "research" / "labels" / "function_labels_v4q.parquet")
    L["DeviceId"] = L.DeviceId.str.lower()
    L = L.rename(columns={"detector": "det"}).drop_duplicates(["DeviceId", "det"])
    lk = pd.read_csv(DCW / "official" / "locked_v2.csv", dtype=str)
    d = T.merge(L[["DeviceId", "det", "DeviceName", "print_subtype", "technology", "description"]],
                on=["DeviceId", "det"], how="left")
    assert not d.DeviceId.isin(set(lk.DeviceId.str.lower())).any()
    assert not d.DeviceName.isin(set(lk.DeviceName)).any()
    return d, P


def classify(d, P):
    d = d[d.lane_index.notna() & (d.span == 1)].copy()
    d["lsb"] = d.print_subtype.isin(SB) & (d.technology == "loop")
    rows = []
    for (dev, tg), g in d.groupby(["DeviceId", "target"]):
        name = g.DeviceName.iloc[0]
        sb = g[g.lsb]
        if len(sb) < 2:
            continue
        if (name, tg) not in USER_FIXED and (g[g.print_subtype.isin(SB)].technology != "loop").any():
            continue  # stacked radar / video zones, not a side-by-side loop split
        types = {k: set(x.lane_type.dropna()) for k, x in g.groupby("lane_index")}
        sbk = {k: x for k, x in sb.groupby("lane_index")}
        advk = set(g[g.print_function == "Advance"].lane_index)
        hit = None
        for k, x in sbk.items():
            t = types[k]
            if len(x) >= 2 and any("R" in s for s in t) and any("T" in s for s in t):
                hit = ("1 lane", x)
                break
        if hit is None:
            for k in sbk:
                if (k + 1 in sbk and set(sbk[k + 1].lane_type) <= {"R", "TR"} and any("R" in s for s in types[k])
                        and (k + 1) not in advk):
                    stem = not any("T" in s for s in types[k])
                    hit = ("2 lanes (L/R stem)" if stem else "2 lanes", pd.concat([sbk[k], sbk[k + 1]]))
                    break
        if hit is None:
            continue
        kind, x = hit
        c = x.print_confidence
        conf = "high" if (c == "high").all() else ("low" if (c == "low").any() else "medium")
        rows.append(dict(DeviceId=dev, DeviceName=name, target=tg, kind=kind, conf=conf,
                         user_fixed=(name, tg) in USER_FIXED, ph_kept=bool(g.phase_kept.all())))
    return pd.DataFrame(rows)


def det_line(r):
    role = {"stopbar_secondary": "extra stop-bar loop (labelled Other)", "advance": "advance loop"}.get(
        r.print_subtype, f"stop-bar {r.print_function}" if r.print_subtype in SB else f"{r.print_function}")
    desc = str(r.description).replace("%20", " ").replace("%2c", ",").replace("&amp;", "&") if pd.notna(r.description) else ""
    desc = f" ({desc if desc.lower().startswith('loop') else ('loops ' if ',' in desc else 'loop ') + desc})" if desc else ""
    return f"det {r.det}: {r.lane_type} {role}{desc} - our lane {int(r.lane_index)}"


def why(kind, user_fixed, g):  # noqa: C901
    if user_fixed:
        return "1 lane - you told us in the last review (extra loop is in the same lane)"
    adv = g[g.print_function == "Advance"]
    up = f"; upstream there is one advance loop ({'/'.join(sorted(set(adv.lane_type)))})" if len(adv) == 1 else ""
    if kind == "1 lane":
        nl = int(g.lane_index.max())
        tot = "" if nl == 1 else f" (phase total {nl} lanes: the through/right side counted once)"
        return "1 lane - the extra right loop was read as a second stop-bar loop in the same lane" + tot
    sbt = g[g.print_subtype.isin(SB)].sort_values("lane_index").lane_type.tolist()
    return f"2 lanes - the two stop-bar loops were read as two lanes ({' + '.join(sbt)}){up}"


def main():
    d, P = load()
    R = classify(d, P)
    core = R[R.kind != "2 lanes (L/R stem)"]
    summ = {"phases": int(len(core)), "signals": int(core.DeviceId.nunique()),
            "by_kind_conf": {k: g.conf.value_counts().to_dict() for k, g in R.groupby("kind")},
            "one_lane_user_fixed": int(R.user_fixed.sum()), "not_in_lane_truth": R[~R.ph_kept].kind.value_counts().to_dict(),
            "rows": R[["DeviceName", "target", "kind", "conf", "ph_kept"]].values.tolist()}
    print(json.dumps(summ, indent=1))
    n1 = int((core.kind == "1 lane").sum())
    n2 = core[core.kind == "2 lanes"]
    nst = int((R.kind == "2 lanes (L/R stem)").sum())
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    assert not XL.exists(), "never overwrite a sheet the user may have typed in"
    wb = Workbook()
    ws = wb.active
    ws.title = "TR + R examples"
    a1 = (f"Partly. Of {len(core)} approaches whose print shows a through/right stop-bar loop plus an extra loop on the right, "
          f"{n1} are counted as one lane ({n1 - int(core.user_fixed.sum())} as read from the print, "
          f"{int(core.user_fixed.sum())} from your last answers).")
    a2 = (f"{len(n2)} are still counted as two lanes ({int((n2.conf == 'high').sum())} of them high-confidence print readings), "
          f"e.g. 01072 P8, the twin of the 01072 P4 you corrected; plus {nst} similar left/right splits on a side street.")
    head = ["#", "Signal", "Phase", "Detectors on that approach (print position, label - our lane)",
            "How we counted it", "Print link", "Question: Is this one lane or two? 1 / 2 / ?", "Answer"]
    ws.append([a1])
    ws.append([a2])
    ws.append([])
    ws.append(head)
    link = Font(color="0563C1", underline="single")
    for n, (s, tg) in enumerate(EXAMPLES, 1):
        r = R[(R.DeviceName == s) & (R.target == tg)]
        assert len(r) == 1, (s, tg)
        r = r.iloc[0]
        g = d[(d.DeviceName == s) & (d.target == tg) & d.lane_index.notna()].sort_values(["lane_index", "det"])
        g = g[g.print_subtype.isin(SB + ["advance"])]
        dets = "\n".join(det_line(x) for x in g.itertuples())
        pdfs = sorted(PDF.glob(f"{s}_*.pdf"))
        ws.append([n, s, tg, dets, why(r.kind, r.user_fixed, g), "print" if pdfs else "none",
                   "Is this one lane or two? 1 / 2 / ?", ""])
        if pdfs:
            c = ws.cell(ws.max_row, 6)
            c.hyperlink = "file:///" + urllib.parse.quote(str(pdfs[0]).replace("\\", "/"), safe=":/()_-.,'")
            c.font = link
    for k, w in zip("ABCDEFGH", [4, 8, 7, 62, 44, 8, 22, 10]):
        ws.column_dimensions[k].width = w
    for row in ws.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for rr in (1, 2):
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=len(head))
        ws.row_dimensions[rr].height = 32
        ws.cell(rr, 1).font = Font(bold=True)
    for c in ws[4]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    wb.save(XL)
    print("wrote", XL)


if __name__ == "__main__":
    main()
