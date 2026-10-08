"""Note 81: the user's label review sheet v1b (second batch), champion vs v4l truth, built with note 72's builder.

Champion = note-77 OOF (trees arm c + siba -> context stacker f77, 3 seeds -> D lanes, gate .9, stack pick, twin decode),
scored on the v4l labels (rpath.LABELS_CURRENT; note 80 + note 81 rules / user corrections). Phase rows are cand64's
(the phase champion has not changed since note 64). Selection = review72.select with every threshold at p >= .8:
  * likely label error: config-only label contradicted (note-70 A3 / A4) or a confident contradiction, model's mean
    probability on its own answer over the wrong samples >= .8;
  * dropped as in v1: definitional rows pending the user's radar / long-zone question (note-70 B1-B5 on every wrong
    window), unhealthy, already ruled (source user_ruling), < 3 samples, plus review64's own exclusions;
  * NEW for v1b: every detector already on sheet v1 (read-only), every detector the user answered (corrections list),
    and the R4 / R5 / R6 exclusions (no truth -> never scored anyway).
Same narrow sheet as v1 (question on top, 9 columns, same-signal same-pattern rows collapsed, chart + print link).
Charts drawn from SAVED sample data (`%DC_WORK%/rev81/review_data81/`), nothing re-scored. locked_v2 asserted absent.

    python review81.py rows       # champion rows on v4l + champion error attribution -> %DC_WORK%/rev81/
    python review81.py review     # dry run: counts -> %DC_WORK%/rev81/review81_dryrun.json
    python review81.py review --write   # review/function_label_review_v1b.xlsx + review/function_label_review_v1b_charts/
"""
from __future__ import annotations

import os
os.environ["F76_ARM"] = "c"
for _v in ("OMP_NUM_THREADS", "MKL_NUM_THREADS", "OPENBLAS_NUM_THREADS"):
    os.environ.setdefault(_v, "4")
import argparse  # noqa: E402
import json  # noqa: E402
import shutil  # noqa: E402
import sys  # noqa: E402
import time  # noqa: E402
from pathlib import Path  # noqa: E402

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

CODE = next(p for p in Path(__file__).resolve().parents if p.name == "code")
for _p in (CODE, CODE / "evaluation", CODE / "final77"):
    sys.path.insert(0, str(_p))
import rpath  # noqa: E402,F401
from common import DC_WORK  # noqa: E402

OUT = DC_WORK / "rev81"
V1 = rpath.REPO / "review" / "function_label_review_v1.xlsx"          # READ ONLY (the user's answers live there)
CORR = rpath.REPO / "research" / "labels" / "function_label_corrections_v4l.csv"
OUTNAME = "function_label_review_v1b"
P_MIN = 0.8
KEY = ["DeviceId", "Detector", "period", "win"]
GE30 = ["m30", "h1", "h3", "h6", "h24", "full"]
C7 = ["Advance", "Presence", "Count", "Yellow_Red", "Mid", "Bike", "Other"]


def log(m):
    print(f"[{time.strftime('%H:%M:%S')}] {m}", flush=True)


def locked():
    return set(pd.read_csv(DC_WORK / "official" / "locked_v2.csv").DeviceId.str.lower())


# ================================================================================================ rows
def stage_rows():
    import atspm_score as AS
    assert Path(AS.V3S) == rpath.LABELS_CURRENT and rpath.LABELS_CURRENT.name == "function_labels_v4l.parquet"
    import err78 as E78
    x = E78.env()
    E, fr = x["E"], x["E"]["fr"]
    assert not fr.DeviceId.isin(locked()).any()
    Pst = np.mean([np.load(x["F77"] / "s74" / "f77" / f"stack_ctx_s{s}.npy").astype(float) for s in (0, 1, 2)], 0)
    pr, cr = E78.decode(Pst)
    ok = E78.frame_ok(cr)
    d = fr[KEY + ["wgroup", "fold", "pred_phase", "det_n_on", "validated", "truth_v3s", "stack_group",
                  "stack_role"]].copy()
    d["validated"] = d.validated.astype(str)
    d["pred_champ"] = np.array(C7, object)[pr]
    d["p_pred_champ"] = Pst[np.arange(len(Pst)), pr]
    d["p_max"] = Pst.max(1)
    for s in "ER":
        d[f"ok_{s}_champ"] = ok[s]
    for s, r in E["rows"].items():
        e = np.full(len(fr), None, object)
        e[r] = cr[s].err.to_numpy(object)
        d[f"err_{s[0].upper()}_champ"] = e
    g30 = d.wgroup.isin(GE30).to_numpy()
    head = {}
    for s in "ER":
        m = g30 & ~np.isnan(ok[s])
        head[s] = [round(float(np.mean(ok[s][m])), 4), int(m.sum())]
        log(f"champion on v4l, >= 30 min {s}: {head[s][0]:.4f} on {head[s][1]:,} rows")
    OUT.mkdir(parents=True, exist_ok=True)
    d.to_parquet(OUT / "function_rows.parquet", index=False)
    shutil.copy2(DC_WORK / "cand64" / "phase_rows.parquet", OUT / "phase_rows.parquet")
    # champion error attribution (note-70 rules, everything set) for the definitional / unhealthy / config-only flags
    a = d.rename(columns={"pred_champ": "pred", "ok_E_champ": "ok_E", "err_E_champ": "err_E"})
    df, e = E78.attrib_rows(a, "E")
    e = e.copy()
    e["bucket"] = e.cat.str[0]
    e.to_parquet(OUT / "errors_everything.parquet", index=False)
    json.dump({"champion_v4l_ge30": head, "errors_everything": int(len(e)),
               "buckets_ge30": e[e.wgroup.isin(GE30)].bucket.value_counts().to_dict()},
              open(OUT / "rows81.json", "w"), indent=1)


# ================================================================================================ review
def _cls_txt(t: str) -> str:
    """'P4, P8 stop-bar presence' -> 'stop-bar presence'; 'phase 7' stays."""
    t = str(t).strip()
    if t.startswith("phase "):
        return t
    w = t.split(" ")
    while w and w[0].rstrip(",").startswith("P") and w[0].rstrip(",")[1:].isdigit():
        w = w[1:]
    return " ".join(w)


def v1_sheet():
    """sheet v1 (opened read-only, never saved): (DeviceName, detector) set and (signal, label, model) pattern set."""
    from openpyxl import load_workbook
    wb = load_workbook(V1, read_only=True, data_only=True)
    ws = wb["labels to check"]
    head = [c for c in next(ws.iter_rows(min_row=2, max_row=2, values_only=True))]
    new = "What this row asks" in head                     # note-91 layout ('P5 Presence', one det per line)
    first = lambda v: str(v).split("\n")[0].split(": ", 1)[-1].strip().lower()   # noqa: E731  'p5 presence'
    ph_of = lambda v: "phase " + first(v).split(" ")[0][1:]                      # noqa: E731
    fn_of = lambda v: first(v).split(" ", 1)[1]                                  # noqa: E731
    dets, pats = set(), set()
    for r in ws.iter_rows(min_row=3, values_only=True):
        if r[0] is None or r[1] is None:
            continue
        for x in str(r[2]).split(","):
            x = x.strip().split(" ")[0]
            if x.isdigit():
                dets.add((str(r[1]), int(x)))
        if new:
            h = {k: v for k, v in zip(head, r)}
            if str(h["What this row asks"]).endswith("Which phase is right?"):
                pats.add((str(r[1]), ph_of(h["Label"]), ph_of(h["Model says"])))
            else:
                pats.add((str(r[1]), fn_of(h["Label"]), fn_of(h["Model says"])))
        else:
            pats.add((str(r[1]), _cls_txt(r[3]).lower(), _cls_txt(r[4]).lower()))
    wb.close()
    return dets, pats


def row_pattern(r, NAME) -> tuple:
    """the (signal, label, model) text of a collapsed v1b row without the phase list, lower case; matches both the old
    v1 layout ('phase 7' / NAME text) and the note-91 layout ('P7 Stop-bar Count' / review72.FN_FULL text)."""
    import review72 as R72
    if r.kind == "phase":
        lab, mod = f"phase {str(r.phase_target)[1:]}", f"phase {int(r.model_phase)}"
        return (str(r.DeviceName), lab, mod)
    if _V1_NEW:
        return (str(r.DeviceName), R72.fn_txt(r.label_function).lower(), R72.fn_txt(r.model_function, True).lower())
    return (str(r.DeviceName), NAME.get(r.label_function, r.label_function).lower(),
            NAME.get(r.model_function, r.model_function).lower())


def _v1_is_new() -> bool:
    from openpyxl import load_workbook
    wb = load_workbook(V1, read_only=True)
    h = next(wb["labels to check"].iter_rows(min_row=2, max_row=2, values_only=True))
    wb.close()
    return "What this row asks" in h


_V1_NEW = False


def stage_review(write: bool, limit: int = 0):
    import review64 as R
    import review72 as R72
    R.C64 = OUT                                   # build_items reads phase_rows / function_rows from here
    R72.C64, R72.DATA, R72.ERR_FILE = OUT, OUT / "review_data81", OUT / "errors_everything.parquet"
    R72.CONF, R72.A4_MIN, R72.OUTNAME, R72.ARM = P_MIN, P_MIN, OUTNAME, "champ"
    items, dry64 = R.build_items("champ")
    assert not items.DeviceId.isin(locked()).any()
    # review64 shows a phase item's MOST COMMON predicted phase over all its windows, which is the label itself when the
    # detector is wrong in exactly half of them (14037 d39 / d50): show the phase of its errors instead
    ph = items.kind.eq("phase") & items.pred_phase_item.notna()
    items.loc[ph, "model_phase"] = items.loc[ph, "pred_phase_item"]
    it = R72.select(items)
    # v1b exclusions on top of v1's rules
    global _V1_NEW
    _V1_NEW = _v1_is_new()
    v1, v1_pats = v1_sheet()
    corr = pd.read_csv(CORR, dtype={"DeviceName": str})
    answered = {(n, int(d)) for n, d in zip(corr.DeviceName, corr.detector)}
    lab = pd.read_parquet(rpath.LABELS_CURRENT, columns=["DeviceId", "detector", "exclude_train_score",
                                                        "phase_user_confirmed", "user_review", "rule_v4l"])
    lab["DeviceId"] = lab.DeviceId.str.lower()
    lab = lab.rename(columns={"detector": "Detector"}).astype({"Detector": it.Detector.dtype})
    it = it.merge(lab, on=["DeviceId", "Detector"], how="left")
    nd = list(zip(it.DeviceName.astype(str), it.Detector.astype(int)))
    # note-70 B2 keys on the print subtype; radar advance zones tagged plain 'advance' (channels 49-52 far radar) that
    # the model calls Other are the same open radar / long-zone question
    radv = (it.kind.eq("function") & it.technology.eq("radar") & it.label_function.eq("Advance")
            & it.model_function.eq("Other")).to_numpy()
    extra = [
        (radv, "definitional, pending your open question"),
        (np.array([k in v1 for k in nd]), "already on sheet v1"),
        (np.array([k in answered for k in nd]) | it.user_review.notna().to_numpy(), "user already answered"),
        (it.exclude_train_score.eq(True).to_numpy() | it.phase_user_confirmed.eq(True).to_numpy(),
         "label excluded / confirmed (v4l rules)"),
        (it.conf.lt(P_MIN).to_numpy(), "not confident enough"),
    ]
    why = it.why_drop.to_numpy(object).copy()
    for m, r in extra:
        why[(why == "") & m] = r
    it["why_drop"] = why
    rows = R72.collapse(it)
    # same signal + same pattern as a v1 row: the user answers that question there (never ask it twice); listed so his
    # v1 answer can be carried over to these detectors
    same = np.array([row_pattern(r, R.NAME) in v1_pats for r in rows.itertuples()], bool)
    sp = rows[same].assign(dets=rows[same].dets.map(lambda d: ",".join(map(str, d))))
    sp = sp[["DeviceName", "kind", "dets", "phase_target", "label_function", "model_function", "model_phase", "conf_max"]]
    sp.to_csv(OUT / "review81_same_pattern_as_v1.csv", index=False)
    log(f"{int(same.sum())} rows repeat a v1 pattern at the same signal -> left out (rev81/review81_same_pattern_as_v1.csv)")
    rows = rows[~same].reset_index(drop=True)
    assert (rows.conf_max >= P_MIN).all() and not rows.DeviceId.isin(locked()).any()
    assert not any((n, int(d)) in v1 for n, ds in zip(rows.DeviceName, rows.dets) for d in ds)
    dry = {"review64_items": int(len(items)), "review64_by_kind": items.kind.value_counts().to_dict(),
           "review64_dry": dry64,
           "dropped_items": it.why_drop[it.why_drop != ""].value_counts().to_dict(),
           "kept_items": int((it.why_drop == "").sum()),
           "kept_items_by_type": {f"{k} / {t}": int(n) for (k, t), n in
                                  it[it.why_drop == ""].groupby(["kind", "type"]).size().items()},
           "rows_same_pattern_as_v1_left_out": int(same.sum()),
           "rows": int(len(rows)), "signals": int(rows.DeviceId.nunique()),
           "rows_by_type": rows.types.value_counts().to_dict(), "rows_by_kind": rows.kind.value_counts().to_dict(),
           "rows_multi_detector": int((rows.n_dets > 1).sum()),
           "rows_by_pattern": (rows.kind + ": " + rows.label_function.astype(str) + " -> " +
                               rows.model_function.astype(str)).value_counts().head(15).to_dict(),
           "rows_by_technology": rows.technology.fillna("unknown").value_counts().to_dict(),
           "rows_by_label_source": rows.source.fillna("none").value_counts().to_dict(),
           "rows_with_print": int(sum(R.pdf_for(n) is not None for n in rows.DeviceName)),
           "rows_conf_ge": {f"{t:.2f}": int((rows.conf_max >= t).sum()) for t in (0.8, 0.9, 0.95)}}
    log(json.dumps(dry, indent=1, default=str))
    json.dump(dry, open(OUT / "review81_dryrun.json", "w"), indent=1, default=str)
    rows.drop(columns=[c for c in rows.columns if rows[c].map(lambda v: isinstance(v, (set, list))).any()]) \
        .assign(dets=rows.dets.map(lambda d: ",".join(map(str, d)))).to_parquet(OUT / "review81_rows.parquet", index=False)
    if write or limit:
        sel = rows.head(limit) if limit else rows
        R72.save_chart_data(sel)
        log("chart data saved")
        if write:
            out = rpath.REPO / "review"
            xl = out / f"{OUTNAME}.xlsx"
            assert xl != V1
            ch = R72.draw(sel, out / f"{OUTNAME}_charts")
            R72.write_sheet(sel, ch, xl, f"{OUTNAME}_charts")
        else:
            d = OUT / "preview"
            ch = R72.draw(sel, d / "charts")
            R72.write_sheet(sel, ch, d / f"{OUTNAME}.xlsx", "charts")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("stage", choices=["rows", "review"])
    ap.add_argument("--write", action="store_true")
    ap.add_argument("--limit", type=int, default=0)
    a = ap.parse_args()
    t0 = time.time()
    if a.stage == "rows":
        stage_rows()
    else:
        stage_review(a.write, a.limit)
    log(f"done in {time.time() - t0:.0f} s")
