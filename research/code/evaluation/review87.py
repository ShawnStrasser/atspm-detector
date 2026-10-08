"""Note 87: the user's LABEL-DISCRETION review sheet -- every training / dev label we left out or changed by a rule or an
agent's judgment (never the user's own rulings), so the user can say whether each decision was right.

Scope: the v4l label table (rpath.LABELS_CURRENT; training + dev signals only, locked_v2 asserted absent; the locked
answer key is never opened). One ITEM = one labelled detector with at least one of these decisions (primary = first):
  loop_count_excl   R1: config Count / Yellow_Red on a print loop, no print reading -> out of training and scoring
  loop_count_print  R1: same, print reading used instead (label changed)
  yr_identical      R3 / v3s: Yellow_Red zone whose actuations equal a Count zone -> out of training and scoring
  beh_tied_low      R5: low-confidence print reading tied / classed by hi-res behaviour -> out
  noise_input       R6: unused input (< 15 ONs) -> unclassified
  expl_sweep        note 80: an 'explained away' Other entry that could hide a real detector removed -> unlabelled
  reread_flip       note 80: print re-read changed the class (training label or scoring truth)
  reread_removed    note 80: print re-read removed the scoring truth
  stack_relabel     v3s: member of a stacked group given the group's role
  misconfigured     label check: label kept, field set-up wrong (pulse presence ...) -> out of training
  check_fail        label check (behaviour) failed -> out of training (still scored in the 'everything' set)
  unhealthy         label check: unhealthy by actuations -> out of training
  dq_check          data-quality check (prints' location; not the fault-event part) -> out of training (--clean)
  dec_role          Dec-2024 config export disagrees -> Dec-2024 samples dropped
  unusual_signal    whole signal flagged unusual_layout -> all its labels out of training (ONE row per signal)
Counted but NOT listed (reported to the orchestrator): the user's own decisions (R2 2B091, R4, user_ruling), dq_suspect
rows whose only DQ evidence is detector fault events 84-88 (the user banned those, 2026-09-28), not_checkable rows not
re-admitted, no_data / dead rows (no evidence either way).
Model answer = champion note-77 OOF on v4l (`%DC_WORK%/rev81/function_rows.parquet`), Sept-2026 samples >= 30 min
(Dec-2024 samples for dec_role). Charts from SAVED data (`%DC_WORK%/rev87/data/`), nothing re-scored.

    python review87.py items            # items + counts -> %DC_WORK%/rev87/
    python review87.py data             # chart data (DuckDB, 4 threads)
    python review87.py write            # review/label_discretion_review_v1.xlsx + _charts/
"""
from __future__ import annotations

import os
for _v in ("OMP_NUM_THREADS", "MKL_NUM_THREADS", "OPENBLAS_NUM_THREADS"):
    os.environ.setdefault(_v, "4")
import argparse  # noqa: E402
import json  # noqa: E402
import re  # noqa: E402
import sys  # noqa: E402
import textwrap  # noqa: E402
import time  # noqa: E402
import urllib.parse  # noqa: E402
from pathlib import Path  # noqa: E402

import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402

CODE = next(p for p in Path(__file__).resolve().parents if p.name == "code")
for _p in (CODE, CODE / "evaluation", CODE / "trackA"):
    sys.path.insert(0, str(_p))
import rpath  # noqa: E402
from common import DC_WORK  # noqa: E402

OUT = DC_WORK / "rev87"
DATA = OUT / "data"
OUTNAME = "label_discretion_review_v1"
CAB = DC_WORK / "cabinet_v4l"
CHANGES = rpath.REPO / "research" / "labels" / "function_label_changes_v4l.csv"
PDFS = [DC_WORK / "cabinet" / "pdf", CAB / "pdf"]
CACHE = {"stg": DC_WORK / "official" / "stg" / "cache", "dec": DC_WORK / "cache"}
GE30 = ["m30", "h1", "h3", "h6", "h24", "full"]
ATS = {"Advance", "Presence", "Count", "Yellow_Red"}
NAME = {"Advance": "advance", "Presence": "stop-bar presence", "Count": "stop-bar count", "Yellow_Red": "yellow-red",
        "Mid": "mid", "Bike": "bike", "Other": "other"}
ORDER = ["unusual_signal", "loop_count_excl", "loop_count_print", "yr_identical", "beh_tied_low", "noise_input",
         "expl_sweep", "reread_flip", "reread_removed", "stack_relabel", "misconfigured", "check_fail", "unhealthy",
         "dq_check", "dec_role"]
PLAIN = {"unusual_signal": "whole signal left out (unusual layout)",
         "loop_count_excl": "loop called Count / YR by config: left out",
         "loop_count_print": "loop called Count / YR by config: print reading used",
         "yr_identical": "yellow-red identical to a count zone: left out",
         "beh_tied_low": "low-confidence print reading tied by behaviour: left out",
         "noise_input": "unused input with < 15 actuations: left out",
         "expl_sweep": "'explained-away' Other label removed",
         "reread_flip": "print re-read changed the class",
         "reread_removed": "print re-read removed it from scoring",
         "stack_relabel": "stacked detector given the group's role",
         "misconfigured": "misconfigured (label kept): left out of training",
         "check_fail": "behaviour check failed: left out of training",
         "unhealthy": "unhealthy (by actuations): left out of training",
         "dq_check": "data-quality check: left out of training",
         "dec_role": "Dec-2024 config disagrees: Dec-2024 data dropped"}
QUESTION = ("Was our decision right? Answer Y (right), N (wrong - use the original label) or ? (can't tell). "
            "Each row is a label on a training or dev signal that we left out of training / scoring, or changed, by an "
            "automatic rule or an agent's judgment (your own rulings are not here). Wrongly dropped labels starve or skew "
            "training. First come labels of the four ATSPM classes where the model AGREES with the label we dropped, "
            "then other ATSPM decisions, then the rest; rows are grouped by signal.")
MAX_DETS_ROW = 12


def log(m):
    print(f"[{time.strftime('%H:%M:%S')}] {m}", flush=True)


def locked() -> set:
    return set(pd.read_csv(DC_WORK / "official" / "locked_v2.csv").DeviceId.str.lower())


def nm(c) -> str:
    return NAME.get(c, c) if isinstance(c, str) else "unlabelled"


def short(s, n=150) -> str:
    s = re.sub(r"\s+", " ", str(s)).strip()
    return s if len(s) <= n else s[: n - 3].rsplit(" ", 1)[0] + "..."


def pct(x) -> str:
    return f"{x:.0%}"


# ================================================================================================ plain reasons
def why_check(fc: str, reason: str) -> str:
    """one plain sentence for a failed behaviour check (label_check codes), numbers taken from the reason text."""
    r = str(reason)
    def f(p):
        m = re.search(p, r)
        return m.groups() if m else None
    c = fc.split(",")[0].replace("misconfigured:", "")
    mis = fc.startswith("misconfigured")
    if c == "pulse_presence":
        return "Presence zone set to pulse (every ON <= 0.3 s), so it cannot hold a call; label kept, not trained on."
    if c == "count_off_in_red":
        g = f(r"occupied (\d+%) of red")
        s = f"A count zone should be quiet in red; this one is occupied {g[0] if g else 'much'} of red"
        return s + (" (zone not set to pulse); label kept, not trained on." if mis else " (acts like presence).")
    if c == "count_short_ons":
        g = f(r"typical actuation ([\d.]+) s")
        s = f"Count actuations should be short pulses; typical actuation here is {g[0] if g else '?'} s"
        return s + (" (not set to pulse); label kept, not trained on." if mis else " (acts like presence).")
    if c == "advance_leads_lane":
        g = f(r"before (\d+%) of det (\d+)'s .*?chance level (\d+%)")
        if g:
            return (f"An advance should fire 1.5-9 s before the stop-bar zone in its lane; it led only {g[0]} of det "
                    f"{g[1]}'s off-peak cars ({g[2]} by chance).")
        return "An advance should fire 1.5-9 s before the stop-bar zone in its lane; it does not."
    if c == "advance_red_arrivals":
        g = f(r"(\d+%) of actuations on red while red was (\d+%)")
        if g:
            return (f"Running free, an advance gets cars on red as often as on green; this one got {g[0]} of its "
                    f"actuations on red while red was {g[1]} of the time (acts like a stop-bar count).")
        return "Running free, it is quiet on red (acts like a stop-bar count, not an advance)."
    if c == "advance_misses":
        g = f(r"misses detections: ([\d,]+) actuations vs ([\d,]+) on its same-lane stop-bar loop det (\d+)")
        if g:
            return f"Advance loop counts far fewer cars than its same-lane stop-bar loop: {g[0]} vs {g[1]} on det {g[2]}."
        return "Advance loop counts far fewer cars than its same-lane stop-bar loop."
    if c == "presence_holds_red":
        g = f(r"(\d+%) of ([\d,]+) vehicles arriving on red")
        if g:
            return (f"A presence zone stays on while a car waits on red; this one held only {g[0]} of {g[1]} red "
                    f"arrivals until green (acts like a count / advance zone).")
        return "A presence zone stays on while a car waits on red; this one does not."
    if c == "bike_far_below":
        g = f(r"([\d,]+) actuations vs ([\d,]+) actuations")
        return (f"A bike loop counts far less than the car detectors; it counted {g[0]} vs {g[1]} for a typical car "
                f"detector of its phase." if g else "A bike loop counts far less than the car detectors; this one does not.")
    if c == "mid_tracks_sum":
        g = f(r"correlation ([\d.]+)")
        return (f"A mid loop across both lanes should track the sum of the lane advance loops; correlation only "
                f"{g[0]}." if g else "A mid loop should track the sum of the lane advance loops; it does not.")
    if c == "role_order_occupancy":
        return "The count zone is more occupied than the presence zone on its phase (roles look swapped)."
    if c in ("health", "card_fault", "saturation"):
        parts = [p.strip() for p in r.split("|")]
        p = next((p for p in parts if p.startswith(("detector health", "card fault")) or "actuations in 5 minutes" in p),
                 parts[0])
        p = re.sub(r"^detector health: ", "", p)
        return "Unhealthy by its actuations: " + short(p, 120) + "."
    return short(r, 150)


DQ_TXT = {"pair_sb": "its counts do not follow the same-lane advance loop off-peak",
          "order": "cars over the same-lane advance do not show up here a few seconds later",
          "excl": "it fires together with the side-by-side loop (should fire separately)",
          "bike": "the bike loop counts like a vehicle loop",
          "near_dead": "almost no actuations compared with its phase",
          "span": "the spanning loop does not track the sum of the lane loops",
          "sat": "more actuations per 5 min than its lanes can carry"}


def why_dq(flags: str, reasons: str) -> str:
    fl = [f for f in str(flags).split(",") if f and f not in ("health", "dead")]
    r = str(reasons)
    bits = []
    for f in fl:
        t = DQ_TXT.get(f, f)
        m = re.search(r"excl: (\d+%) of ONs coincide with adv(\d+)", r) if f == "excl" else None
        if m:
            t = f"{m.group(1)} of its ONs coincide with det {m.group(2)}'s (side-by-side loops should fire separately)"
        m = re.search(r"near-dead: (\d+) ONs vs phase median (\d+)", r) if f == "near_dead" else None
        if m:
            t = f"almost no actuations ({m.group(1)} vs {m.group(2)} for its phase)"
        m = re.search(r"bike: counts like a vehicle loop \(([\d.]+) of phase median\)", r) if f == "bike" else None
        if m:
            t = f"the bike loop counts like a vehicle loop ({float(m.group(1)):.0%} of its phase's median)"
        bits.append(t)
    h = re.search(r"health: ([^;]*)", r)
    if h:
        m = re.search(r"stuck-on ([\d.]+) min", h.group(1))
        if m:
            bits.append(f"stuck on for {float(m.group(1)):.0f} min")
        m = re.search(r"chatter ([\d.]+%)", h.group(1))
        if m:
            bits.append(f"chatters ({m.group(1)} of re-triggers within 0.3 s)")
    if not bits:
        return "Data-quality check flagged it (no detail kept)."
    s = "; ".join(bits)
    return "Data-quality check: " + s[0].lower() + s[1:] + "."


def dq_fault_only(flags: str, reasons: str) -> bool:
    """dq_suspect only because of detector fault events 84-88 (health part = 'faults ...' alone, no other flag)."""
    fl = [f for f in str(flags).split(",") if f]
    if fl != ["health"]:
        return False
    h = re.search(r"health: ([^;]*)", str(reasons))
    return bool(h) and h.group(1).strip().startswith("faults") and not re.search(r"stuck|chatter", h.group(1))


# ================================================================================================ items
def model_summary() -> pd.DataFrame:
    f = pd.read_parquet(DC_WORK / "rev81" / "function_rows.parquet",
                        columns=["DeviceId", "Detector", "period", "wgroup", "pred_champ", "p_pred_champ"])
    f = f[f.wgroup.isin(GE30)]
    out = []
    for (d, det, per), g in f.groupby(["DeviceId", "Detector", "period"], sort=False):
        vc = g.pred_champ.value_counts()
        top = vc.index[0]
        out.append((d, int(det), per, top, int(vc.iloc[0]), len(g), float(g.p_pred_champ[g.pred_champ == top].mean())))
    return pd.DataFrame(out, columns=["DeviceId", "Detector", "period", "m_pred", "m_n_top", "m_n", "m_p"])


def split_of() -> dict:
    dev = set(pd.read_csv(DC_WORK / "data" / "splits" / "device_id_valid.csv").DeviceId.str.lower())
    rel = pd.read_csv(DC_WORK / "official" / "newtest_released.csv")
    rel = set(rel.DeviceId.str.lower()) if "DeviceId" in rel else set()
    return {"dev": dev, "rel": rel}


def stage_items():
    lk = locked()
    L = pd.read_parquet(rpath.LABELS_CURRENT)
    assert rpath.LABELS_CURRENT.name == "function_labels_v4l.parquet"
    L["DeviceId"] = L.DeviceId.str.lower()
    assert not L.DeviceId.isin(lk).any(), "locked_v2 in v4l"
    # names: label table, else config export
    cfgp = DC_WORK / "data" / "labels" / "detector_config_current.parquet"
    names = L.dropna(subset=["DeviceName"]).drop_duplicates("DeviceId").set_index("DeviceId").DeviceName.astype(str).to_dict()
    if cfgp.exists():
        c = pd.read_parquet(cfgp)
        ncol = next((x for x in ("DeviceName", "Name", "SignalName") if x in c), None)
        if ncol:
            c["DeviceId"] = c.DeviceId.astype(str).str.lower()
            for k, v in c.drop_duplicates("DeviceId").set_index("DeviceId")[ncol].astype(str).items():
                names.setdefault(k, v)
    L["DeviceName"] = L.DeviceId.map(names).fillna(L.DeviceId.str[:8])
    dq = pd.read_parquet(CAB / "dq_print.parquet", columns=["DeviceId", "detector", "dq_reasons"])
    dq["DeviceId"] = dq.DeviceId.str.lower()
    L = L.merge(dq.drop_duplicates(["DeviceId", "detector"]), on=["DeviceId", "detector"], how="left")
    dr = pd.read_parquet(CAB / "dec_role_changed_v4l.parquet")
    dr["DeviceId"] = dr.DeviceId.str.lower()
    L = L.merge(dr[["DeviceId", "detector", "dec_function", "dec_phase", "dec_status_train", "dec_status_truth"]]
                .drop_duplicates(["DeviceId", "detector"]), on=["DeviceId", "detector"], how="left")
    yr = pd.read_parquet(CAB / "stacked" / "yr_count_pairs.parquet")
    yr = yr.sort_values(["identical", "m_sg"], ascending=False).drop_duplicates(["dev", "yr"])
    yrp = {(d.lower(), int(a)): (int(b), float(r)) for d, a, b, r in zip(yr.dev, yr.yr, yr.cnt, yr.ratio)}
    C = pd.read_csv(CHANGES, dtype={"DeviceName": str})
    C["key"] = C.DeviceName + "_" + C.detector.astype(int).astype(str)
    ex = pd.read_csv(DC_WORK / "lab80" / "expl_removed.csv", dtype=str)
    exo = {f"{a}_{int(b)}": o for a, b, o in zip(ex.DeviceName, ex.detector, ex.old)}
    import v3_retrain as V
    V.LABELS_V3, V.DEC_ROLE_FILE = V.LABEL_SETS["v3s"]
    nch = V.nc_high_mask(L.rename(columns={"detector": "Detector"}))
    L["key"] = L.DeviceName.astype(str) + "_" + L.detector.astype(int).astype(str)
    b = lambda s: s.fillna(False).astype(bool)
    rule = L.rule_v4l.astype("string").fillna("")
    labelled = L.function.notna()
    usr = rule.str.startswith(("R2", "R4")) | L.user_review.notna()
    items = []

    def add(m, key, had, action, why, new=None, period="stg", scope="train+score"):
        for i in np.flatnonzero(m.to_numpy()):
            r = L.iloc[i]
            items.append(dict(DeviceId=r.DeviceId, DeviceName=r.DeviceName, Detector=int(r.detector),
                              phase=str(r.phase_target) if pd.notna(r.phase_target) else "", technology=r.technology,
                              key=key, had=had(r) if callable(had) else had, new=new(r) if callable(new) else new,
                              action=action(r) if callable(action) else action, why=why(r), period=period, scope=scope))

    unu = b(L.unusual_layout) & labelled
    base = ~unu & ~usr
    # R1
    m = base & rule.eq("R1_loop_config_count->excluded")
    add(m, "loop_count_excl", lambda r: r.func7_v2, "left out of training and scoring",
        lambda r: f"Config calls it {nm(r.func7_v2)}, but the print shows a loop, and a loop is never Count / Yellow_Red; "
                  f"the print gives no other reading, so it was left out.")
    m = base & rule.eq("R1_loop_config_count->print")
    add(m, "loop_count_print", lambda r: r.func7_v2, lambda r: f"changed to {nm(r.function)}",
        lambda r: f"Config calls it {nm(r.func7_v2)}, but the print shows a loop (never Count / Yellow_Red); the print's "
                  f"own reading ({nm(r.function)}, {r.print_confidence} confidence) was used instead.", new=lambda r: r.function)
    # R3 / YR identical
    m = base & b(L.yr_count_identical)

    def yr_why(r):
        p = yrp.get((r.DeviceId, int(r.detector)))
        t = f" det {p[0]} (count ratio {p[1]:.2f})" if p else ""
        return (f"Its actuations are the same as the count zone{t} on the same phase, so the two cannot be told apart; "
                f"left out of training and scoring.")
    add(m, "yr_identical", "Yellow_Red", "left out of training and scoring", yr_why)
    # R5
    m = base & rule.eq("R5_behaviour_tied_low")
    add(m, "beh_tied_low", lambda r: r.function, "left out of training and scoring",
        lambda r: "Low-confidence print reading whose channel or class was decided from its own hi-res behaviour "
                  "(circular as a label). Reader: " + short(r.confidence_reason, 110))
    # R6
    m = base & rule.eq("R6_noise_input")
    add(m, "noise_input", lambda r: r.function, "left unclassified",
        lambda r: f"Input not on the print and with no config function, only {int(r.n_on_new)} actuations (< 15) in "
                  f"the Sept log: treated as noise.")
    # note-80 change list: explained-other sweep, re-read flips / truth removals
    Cx = C[C.cause.eq("explained_other_sweep")]
    for r in Cx.itertuples():
        row = L[L.key.eq(r.key)]
        if len(row) == 0:
            continue
        i = row.index[0]
        items.append(dict(DeviceId=L.at[i, "DeviceId"], DeviceName=r.DeviceName, Detector=int(r.detector),
                          phase=str(L.at[i, "phase_target"]) if pd.notna(L.at[i, "phase_target"]) else "",
                          technology=L.at[i, "technology"], key="expl_sweep", had=r.train_label_v3s or r.function_v3s,
                          new=None, action="label removed (now unlabelled)",
                          why="The print reader had explained it away as an unused channel ('"
                              + short(exo.get(r.key, ""), 90) + "'), but it may be a real detector added after the "
                              "print, so the Other label was removed.", period="stg", scope="train+score"))
    kind = lambda a, b_: ("added" if pd.isna(a) and pd.notna(b_) else "removed" if pd.notna(a) and pd.isna(b_)
                          else "flip" if pd.notna(a) and a != b_ else "")
    for r in C[C.cause.isin(["print_reread", "pipeline_rerun", "released_newtest_pipeline"])].itertuples():
        kt, kr = kind(r.train_label_v3s, r.train_label_v4l), kind(r.score_truth_v3s, r.score_truth_v4l)
        row = L[L.key.eq(r.key)]
        if len(row) == 0 or (kt != "flip" and kr != "flip" and not (r.cause == "print_reread" and kr == "removed")):
            continue
        i = row.index[0]
        if usr.loc[i] or unu.loc[i]:
            continue
        if kt == "flip" or kr == "flip":
            had = r.train_label_v3s if kt == "flip" else r.score_truth_v3s
            new = r.train_label_v4l if kt == "flip" else r.score_truth_v4l
            items.append(dict(DeviceId=L.at[i, "DeviceId"], DeviceName=r.DeviceName, Detector=int(r.detector),
                              phase=str(L.at[i, "phase_target"]), technology=L.at[i, "technology"], key="reread_flip",
                              had=had, new=new, action=f"changed to {nm(new)}",
                              why="A second reading of the print changed it: " + short(L.at[i, "confidence_reason"], 120),
                              period="stg", scope="train+score"))
        else:
            items.append(dict(DeviceId=L.at[i, "DeviceId"], DeviceName=r.DeviceName, Detector=int(r.detector),
                              phase=str(L.at[i, "phase_target"]), technology=L.at[i, "technology"],
                              key="reread_removed", had=r.score_truth_v3s, new=None,
                              action="left out of scoring" + ("" if pd.notna(r.train_label_v4l) else " and training"),
                              why=f"After the print re-read its {r.source_v4l} label no longer counts as scoring truth "
                                  f"(print {r.print_function_v4l if pd.notna(r.print_function_v4l) else 'silent'}, "
                                  f"{r.print_confidence_v4l if pd.notna(r.print_confidence_v4l) else 'no'} confidence).",
                              period="stg", scope="score"))
    # stacked groups
    chg = (L.function_v3.astype("string").fillna("-") != L.function.astype("string").fillna("-")) | \
          (L.truth_v3.astype("string").fillna("-") != L.truth_v3s.astype("string").fillna("-"))
    m = base & b(L.stack_relabel) & chg & ~rule.str.startswith("R")

    def st_why(r):
        mates = str(r.stack_group).rsplit("_", 1)[-1].replace("-", ", ")
        return (f"Stacked with det {mates} on {r.phase_target} (same lane and role, another technology, actuations start "
                f"together): every member carries the group's role {nm(r.stack_role)}.")
    add(m, "stack_relabel", lambda r: r.function_v3 if pd.notna(r.function_v3) else r.truth_v3,
        lambda r: f"changed to {nm(r.function)}", st_why, new=lambda r: r.function)
    # label checks (training only); exclude_train_score rows are already out
    tu = b(L.train_use) & ~b(L.exclude_train_score)
    for key, v in (("misconfigured", "misconfigured"), ("check_fail", "fail"), ("unhealthy", "unhealthy")):
        m = base & labelled & tu & L.validated.eq(v).fillna(False)
        add(m, key, lambda r: r.function, "left out of training" + (" (label kept)" if v == "misconfigured" else ""),
            lambda r: why_check(str(r.failed_checks), r.validation_reason), scope="train")
    # dq (--clean), only where nothing else keeps it out of training
    okval = L.validated.eq("pass").fillna(False) | nch
    dqs = base & labelled & tu & okval & L.dq_suspect.eq(True).fillna(False) & ~b(L.dead)
    fo = pd.Series([dq_fault_only(f, r) for f, r in zip(L.dq_flags, L.dq_reasons)], index=L.index)
    add(dqs & ~fo, "dq_check", lambda r: r.function, "left out of training",
        lambda r: why_dq(r.dq_flags, r.dq_reasons), scope="train")
    # Dec role
    m = base & labelled & (b(L.dec_fp_train) | b(L.dec_fp_truth))

    def dec_why(r):
        st = r.dec_status_train if r.dec_status_train in ("function", "phase") else r.dec_status_truth
        dp = f"phase {int(r.dec_phase)}" if pd.notna(r.dec_phase) else "no phase"
        if st == "phase":
            return (f"The Dec-2024 config export puts it on {dp}, not {r.phase_target}; its Dec-2024 samples were left "
                    f"out (Sept kept).")
        return (f"The Dec-2024 config export lists it as {nm(r.dec_function)} ({dp}), not {nm(r.function)}; its Dec-2024 "
                f"samples were left out (Sept kept).")
    add(m, "dec_role", lambda r: r.function, "Dec-2024 data dropped", dec_why, period="dec", scope="dec")
    I = pd.DataFrame(items)
    # whole-signal rows (unusual layout)
    sig = []
    for d, g in L[unu].groupby("DeviceId"):
        n = g.DeviceName.iloc[0]
        p = CAB / "signals" / f"{n}.json"
        why = ""
        if p.exists():
            rec = json.loads(p.read_text(encoding="utf-8"))
            why = rec.get("signal_flags_reason") or re.sub(r"^\s*unusual_layout[:.;]?\s*", "", str(rec.get("note") or ""))
        sig.append(dict(DeviceId=d, DeviceName=n, dets_all=sorted(int(x) for x in g.detector), labels=g.function.tolist(),
                        dets_label=dict(zip(g.detector.astype(int), g.function)), phases=dict(zip(g.detector.astype(int),
                        g.phase_target.astype(str))), why=short("Unusual layout: " + (why or "flagged by the print reader"), 170)))
    S = pd.DataFrame(sig)
    # model answers
    M = model_summary()
    I = I.merge(M.rename(columns={"Detector": "Detector"}), left_on=["DeviceId", "Detector", "period"],
                right_on=["DeviceId", "Detector", "period"], how="left")
    I["agree"] = I.m_pred.notna() & I.m_pred.eq(I.had) & (I.m_n_top * 2 > I.m_n)
    I["atspm"] = I.had.isin(ATS) | I.new.isin(ATS)
    I["tier"] = np.where(I.had.isin(ATS) & I.agree, 1, np.where(I.atspm, 2, 3))
    # one row per detector: primary decision = first in ORDER, others named in 'also'
    I["ord"] = I.key.map({k: i for i, k in enumerate(ORDER)})
    I = I.sort_values(["DeviceId", "Detector", "ord"])
    also = I.groupby(["DeviceId", "Detector"]).key.agg(lambda s: [PLAIN[k] for k in list(s)[1:]])
    alsow = I.groupby(["DeviceId", "Detector"]).why.agg(lambda s: list(s)[1:])
    P = I.drop_duplicates(["DeviceId", "Detector"]).set_index(["DeviceId", "Detector"])
    P["also"] = also
    P["also_why"] = alsow
    P = P.reset_index()
    # signal-row model answers
    Ms = M[M.period == "stg"].set_index(["DeviceId", "Detector"])
    if len(S):
        ag, n_at, best = [], [], []
        for r in S.itertuples():
            a = t = 0
            bb = (None, -1.0)
            for det, lab in r.dets_label.items():
                if lab not in ATS or (r.DeviceId, det) not in Ms.index:
                    continue
                mm = Ms.loc[(r.DeviceId, det)]
                t += 1
                if mm.m_pred == lab and mm.m_n_top * 2 > mm.m_n:
                    a += 1
                    if mm.m_p > bb[1]:
                        bb = (det, float(mm.m_p))
            ag.append(a)
            n_at.append(t)
            best.append(bb[0] if bb[0] is not None else next((d for d, l in r.dets_label.items() if l in ATS),
                                                               r.dets_all[0]))
        S["n_atspm_scored"], S["n_atspm_agree"], S["lead"] = n_at, ag, best
    # counts (labelled detectors, by reason, split) -------------------------------------------------
    sp = split_of()
    grp = lambda d: "dev" if d in sp["dev"] else ("released" if d in sp["rel"] or False else "train")
    L["split"] = L.DeviceId.map(grp)
    L.loc[b(L.released_from_newtest), "split"] = "released"
    I["split"] = I.DeviceId.map(dict(zip(L.DeviceId, L.split)))
    cnt = {}
    for k in ORDER[1:]:
        x = I[I.key == k]
        cnt[k] = dict(detectors=int(len(x)), **{s: int((x.split == s).sum()) for s in ("train", "dev", "released")},
                      atspm_label=int(x.had.isin(ATS).sum()), model_agrees_atspm=int(((x.tier == 1)).sum()))
    xu = L[unu]
    cnt["unusual_signal"] = dict(detectors=int(len(xu)), signals=int(xu.DeviceId.nunique()),
                                 **{s: int((xu.split == s).sum()) for s in ("train", "dev", "released")},
                                 atspm_label=int(xu.function.isin(ATS).sum()))
    # not listed
    nl = {}
    nl["user_own_R2_R4"] = int((labelled & usr).sum())
    nl["user_ruling_source"] = int((L.source == "user_ruling").sum())
    fo_rows = base & labelled & tu & okval & L.dq_suspect.eq(True).fillna(False) & ~b(L.dead) & fo
    nl["dq_fault_events_only_training_dropped"] = int(fo_rows.sum())
    nl["dq_fault_events_only_by_class"] = L[fo_rows].function.value_counts().to_dict()
    ncx = base & labelled & tu & L.validated.eq("not_checkable").fillna(False) & ~nch
    nl["not_checkable_not_readmitted"] = int(ncx.sum())
    nl["not_checkable_by_class"] = L[ncx].function.value_counts().to_dict()
    nl["not_checkable_by_source"] = L[ncx].source.value_counts().to_dict()
    nl["no_data_labelled"] = int((base & labelled & L.validated.eq("no_data").fillna(False)).sum())
    nl["dead_labelled"] = int((base & labelled & b(L.dead)).sum())
    tot = dict(label_rows=len(L), labelled=int(labelled.sum()), train_use_validated=int(b(L.train_use_validated).sum()),
               detectors_with_decision=int(len(P)) + int(len(xu)),
               detector_items_any_reason=int(len(I)))
    OUT.mkdir(parents=True, exist_ok=True)
    for c in ("also", "also_why"):
        P[c] = P[c].map(lambda v: json.dumps(v))
    P.to_parquet(OUT / "items.parquet", index=False)
    I.drop(columns=["ord"]).to_parquet(OUT / "items_all.parquet", index=False)
    S.assign(dets_all=S.dets_all.map(json.dumps), labels=S.labels.map(json.dumps),
             dets_label=S.dets_label.map(lambda d: json.dumps({str(k): v for k, v in d.items()})),
             phases=S.phases.map(lambda d: json.dumps({str(k): v for k, v in d.items()}))).to_parquet(
        OUT / "signals_unusual.parquet", index=False)
    res = dict(totals=tot, by_reason=cnt, not_listed=nl,
               primary_by_reason=P.key.value_counts().to_dict(),
               tiers_detector=P.tier.value_counts().sort_index().to_dict())
    json.dump(res, open(OUT / "counts.json", "w"), indent=1, default=str)
    log(json.dumps(res, indent=1, default=str))


# ================================================================================================ rows
def build_rows() -> pd.DataFrame:
    P = pd.read_parquet(OUT / "items.parquet")
    S = pd.read_parquet(OUT / "signals_unusual.parquet")
    P["also"] = P.also.map(json.loads)
    P["also_why"] = P.also_why.map(json.loads)
    P["mkey"] = P.m_pred.fillna("none")
    P["newk"] = P.new.fillna("-")
    P["conf"] = P.m_p.fillna(0)
    L = pd.read_parquet(rpath.LABELS_CURRENT, columns=["DeviceId", "detector", "failed_checks", "dq_flags"])
    L["DeviceId"] = L.DeviceId.str.lower()
    P = P.merge(L.rename(columns={"detector": "Detector"}).drop_duplicates(["DeviceId", "Detector"]),
                on=["DeviceId", "Detector"], how="left")
    P["sub"] = np.where(P.key.isin(["check_fail", "misconfigured", "unhealthy"]),
                        P.failed_checks.fillna("").str.split(",").str[0], "")
    P.loc[P.key.eq("dq_check"), "sub"] = P.dq_flags.fillna("")
    P.loc[P.key.eq("dec_role"), "sub"] = P.why.str.extract(r"(lists it as [^(]*|puts it on phase \d+)")[0].fillna("")
    rows = []
    for (d, k, had, act, sub, tier), g in P.sort_values("conf", ascending=False).groupby(
            ["DeviceId", "key", "had", "action", "sub", "tier"], sort=False, dropna=False):
        r = g.iloc[0].to_dict()
        r["dets"] = sorted(int(x) for x in g.Detector)
        r["lead"] = int(g.Detector.iloc[0])
        r["phases"] = sorted({p for p in g.phase if isinstance(p, str) and p}, key=lambda p: (len(p), p))
        r["n_dets"] = len(g)
        vc = g.m_pred.value_counts()
        r["m_mode"] = vc.index[0] if len(vc) else None
        r["m_mode_n"] = int(vc.iloc[0]) if len(vc) else 0
        r["m_mode_p"] = float(g.m_p[g.m_pred == r["m_mode"]].mean()) if len(vc) else None
        r["conf"] = float(g.conf.max())
        r["why_lead"] = g.why.iloc[0]
        r["also_lead"] = g.also.iloc[0]
        r["kind"] = "det"
        rows.append(r)
    for s in S.itertuples():
        lab = json.loads(s.dets_label)
        labs = pd.Series(list(lab.values())).value_counts()
        rows.append(dict(DeviceId=s.DeviceId, DeviceName=s.DeviceName, key="unusual_signal", kind="signal",
                         dets=json.loads(s.dets_all), lead=int(s.lead), n_dets=len(lab), phases=[],
                         had=", ".join(f"{n} {nm(c)}" for c, n in labs.items()), new=None,
                         action="whole signal left out of training", why_lead=s.why, also_lead=[],
                         m_pred=None, m_p=None, n_atspm=int(s.n_atspm_scored), n_agree=int(s.n_atspm_agree),
                         tier=1 if s.n_atspm_agree > 0 else 2,
                         # agreement share shrunk towards 0 for few scored labels (2 of 2 must not outrank everything)
                         conf=s.n_atspm_agree / (s.n_atspm_scored + 2), period="stg",
                         technology=None, phase=json.loads(s.phases).get(str(int(s.lead)), "")))
    R = pd.DataFrame(rows)
    R["sig_tier"] = R.groupby("DeviceId").tier.transform("min")
    R["sig_conf"] = R.apply(lambda r: r.conf if r.tier == r.sig_tier else -1, axis=1)
    R["sig_conf"] = R.groupby("DeviceId").sig_conf.transform("max")
    R["kord"] = R.key.map({k: i for i, k in enumerate(ORDER)})
    R = R.sort_values(["sig_tier", "sig_conf", "DeviceId", "tier", "conf", "kord"],
                      ascending=[True, False, True, True, False, True]).reset_index(drop=True)
    return R


def stem(r) -> str:
    return f"{r.DeviceName}_d{int(r.lead)}_{r.key}"


# ================================================================================================ chart data
def stage_data():
    import duckdb
    R = build_rows()
    DATA.mkdir(parents=True, exist_ok=True)
    cn = duckdb.connect()
    cn.execute("SET memory_limit='10GB'")
    cn.execute("SET threads=4")
    cn.execute(f"SET temp_directory='{(DC_WORK / 'tmp').as_posix()}'")
    for per in ("stg", "dec"):
        devs = sorted(set(R.DeviceId[R.period.eq(per)]) | (set(R.DeviceId) if per == "stg" else set()))
        devs = [d for d in devs if not (DATA / f"{d}_{per}_on.parquet").exists()]
        if not devs:
            continue
        lst = ",".join(f"'{d}'" for d in devs)
        tmp_on, tmp_cy = OUT / f"_tmp_{per}_on.parquet", OUT / f"_tmp_{per}_cy.parquet"
        cn.execute(f"""COPY (SELECT lower(DeviceId) dev, Detector::INT det, t_on, t_off
                             FROM '{(CACHE[per] / 'det_intervals.parquet').as_posix()}'
                             WHERE lower(DeviceId) IN ({lst}) ORDER BY dev, t_on) TO '{tmp_on.as_posix()}'""")
        cn.execute(f"""COPY (SELECT lower(DeviceId) dev, Phase::INT phase, green_start, yellow_start, red_start, next_green
                             FROM '{(CACHE[per] / 'phase_cycles.parquet').as_posix()}'
                             WHERE lower(DeviceId) IN ({lst}) ORDER BY dev, green_start) TO '{tmp_cy.as_posix()}'""")
        for d in devs:
            cn.sql(f"SELECT det, t_on, t_off FROM '{tmp_on.as_posix()}' WHERE dev = '{d}'").df().drop_duplicates() \
                .to_parquet(DATA / f"{d}_{per}_on.parquet", index=False)
            cn.sql(f"SELECT phase, green_start, yellow_start, red_start, next_green FROM '{tmp_cy.as_posix()}' "
                   f"WHERE dev = '{d}'").df().to_parquet(DATA / f"{d}_{per}_cycles.parquet", index=False)
        tmp_on.unlink()
        tmp_cy.unlink()
        log(f"{per}: chart data for {len(devs)} signals")


# ================================================================================================ charts + sheet
def label_map(dev: str, L: pd.DataFrame) -> dict:
    s = L[L.DeviceId == dev]
    return {int(d): (f, str(p)) for d, f, p in zip(s.detector, s.function, s.phase_target)}


def partners(r) -> list:
    """detectors the reason names (count partner, same-lane loop, stack mates), for the counts panel."""
    t = str(r.why_lead)
    out = [int(x) for x in re.findall(r"det (\d+)", t)]
    m = re.search(r"Stacked with det ([\d, ]+) on", t)
    if m:
        out += [int(x) for x in m.group(1).split(",") if x.strip().isdigit()]
    return [x for x in dict.fromkeys(out) if x != int(r.lead)]


def draw(R: pd.DataFrame, out_dir: Path) -> dict:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.dates as mdates
    import matplotlib.pyplot as plt
    import review64 as R64
    import review72 as R72
    C, INK, INK2, GRID, SURF = R64.C, R64.INK, R64.INK2, R64.GRID, R64.SURF
    out_dir.mkdir(parents=True, exist_ok=True)
    L = pd.read_parquet(rpath.LABELS_CURRENT, columns=["DeviceId", "detector", "function", "phase_target"])
    L["DeviceId"] = L.DeviceId.str.lower()
    res = {}
    for r in R.itertuples():
        st = stem(r)
        png = out_dir / f"{st}.png"
        res[st] = png
        if png.exists():
            continue
        per = r.period
        f_on, f_cy = DATA / f"{r.DeviceId}_{per}_on.parquet", DATA / f"{r.DeviceId}_{per}_cycles.parquet"
        iv = pd.read_parquet(f_on) if f_on.exists() else pd.DataFrame(columns=["det", "t_on", "t_off"])
        cy = pd.read_parquet(f_cy) if f_cy.exists() else pd.DataFrame(columns=["phase", "green_start"])
        det = int(r.lead)
        lm = label_map(r.DeviceId, L)
        ph_txt = r.phase if isinstance(r.phase, str) else ""
        lab_ph = int(ph_txt[1:]) if ph_txt[1:].isdigit() else None
        tag = lambda d: (f"{nm(lm[d][0])} {lm[d][1]}" if d in lm else "no label")
        me = iv[iv.det == det]
        fig = plt.figure(figsize=(10.5, 9.4), facecolor=SURF)
        gsp = fig.add_gridspec(3, 2, height_ratios=[1, 1, 0.95], width_ratios=[1.5, 1])
        # (1) 15-min counts: this one, the detectors the reason names, the same row, the same phase
        ax = fig.add_subplot(gsp[0, :])
        R64._style(ax)
        if len(iv):
            iv2 = iv.assign(t=iv.t_on.dt.floor("15min"))
            cnt = iv2.groupby(["t", "det"]).size().unstack(fill_value=0)
            same_ph = [d for d, (f, p) in lm.items() if p == ph_txt and d != det]
            pick = [det] + partners(r) + [d for d in r.dets if d != det][:2] + same_ph
            pick = [d for d in dict.fromkeys(pick) if d in cnt][:7]
            for i, d in enumerate(pick):
                lbl = "this one" if d == det else ("same row" if d in r.dets else tag(d))
                ax.plot(cnt.index, cnt[d], color=C[i % len(C)], lw=1.9 if d == det else 1.0,
                        ls="--" if (d in r.dets and d != det) else "-", label=f"det {d} ({lbl})")
            ax.legend(frameon=False, fontsize=7.5, loc="lower left", bbox_to_anchor=(0, 1.0), ncol=4)
            ax.xaxis.set_major_formatter(mdates.DateFormatter("%a %H:%M"))
        ax.set_ylabel("actuations per 15 min", color=INK2, fontsize=9)
        # (2) cycle profile around begin of green of the labelled phase
        a4, a5 = fig.add_subplot(gsp[1, 0]), fig.add_subplot(gsp[1, 1])
        for a in (a4, a5):
            R64._style(a)
        pr = R72.cycle_profile(me, cy, lab_ph) if (lab_ph is not None and len(me) and len(cy)) else None
        if pr is not None:
            t, occ, starts, gl = pr
            a4.plot(t, 100 * occ, color=C[0], lw=1.8)
            a5.plot(t, starts, color=C[0], lw=1.6)
            for a in (a4, a5):
                a.axvspan(0, gl, color=C[2], alpha=0.13, lw=0)
        for a in (a4, a5):
            a.axvline(0, color=C[2], lw=1)
            a.set_xlabel(f"s from begin of {ph_txt or 'its phase'} green (shaded = median green, left of 0 = red)",
                         color=INK2, fontsize=7.5)
        a4.set_ylabel("% of cycles ON", color=INK2, fontsize=9)
        a4.set_ylim(0, 100)
        a4.set_title("Is it ON during red, through green, or not at all?", fontsize=9, color=INK2, loc="left")
        a5.set_ylabel("starts per cycle per s", color=INK2, fontsize=9)
        a5.set_title("When new actuations start", fontsize=9, color=INK2, loc="left")
        # (3) 10 busy daytime minutes
        a2 = fig.add_subplot(gsp[2, 0])
        R64._style(a2)
        if len(me):
            c10 = me.set_index("t_on").resample("10min").size()
            c10 = c10[(c10.index.hour >= 7) & (c10.index.hour < 19)]
            t0 = c10[c10 >= c10.quantile(0.6)].index[len(c10[c10 >= c10.quantile(0.6)]) // 2] if len(c10) else me.t_on.min()
        else:
            t0 = cy.green_start.min() if len(cy) else pd.Timestamp("2026-09-01 08:00")
        t0 = pd.Timestamp(t0)
        t1 = t0 + pd.Timedelta(minutes=10)
        ylab = []
        if lab_ph is not None and len(cy):
            g = cy[(cy.phase == lab_ph) & (cy.next_green >= t0) & (cy.green_start <= t1)]
            for c0, c1, colr in (("green_start", "yellow_start", C[2]), ("yellow_start", "red_start", C[3]),
                                 ("red_start", "next_green", "#d9534f")):
                w = (g[c1] - g[c0]).dt.total_seconds() / 86400
                a2.barh([0] * len(g), w, left=g[c0], height=0.6, color=colr, alpha=0.85)
            ylab.append(f"{ph_txt} signal")
        dd = [det] + partners(r)[:2] + [d for d in r.dets if d != det][:1]
        dd += [d for d, (f, p) in lm.items() if p == ph_txt and d not in dd][:2]
        for d in list(dict.fromkeys(dd))[:5]:
            s = iv[(iv.det == d) & (iv.t_off >= t0) & (iv.t_on <= t1)]
            y = len(ylab)
            a2.barh([y] * len(s), (s.t_off - s.t_on).dt.total_seconds().fillna(0).clip(lower=1.0) / 86400,
                    left=s.t_on, height=0.6, color=C[0] if d == det else C[5])
            ylab.append(f"det {d}" + ("" if d == det else f" ({tag(d)})"))
        a2.set_yticks(range(len(ylab)))
        a2.set_yticklabels(ylab, fontsize=7.5)
        a2.set_xlim(t0, t1)
        a2.xaxis.set_major_formatter(mdates.DateFormatter("%H:%M"))
        a2.grid(axis="y", visible=False)
        a2.grid(axis="x", color=GRID, lw=0.8)
        a2.set_title("10 busy daytime minutes (green / yellow / red; ONs under 1 s drawn 1 s wide)", fontsize=9,
                     color=INK2, loc="left")
        # (4) ON length
        a3 = fig.add_subplot(gsp[2, 1])
        R64._style(a3)
        du = (me.t_off - me.t_on).dt.total_seconds().dropna().clip(0.1, 120) if len(me) else pd.Series(dtype=float)
        a3.hist(du, bins=np.geomspace(0.1, 120, 22), color=C[0])
        a3.set_xscale("log")
        a3.set_xticks([0.1, 0.3, 1, 3, 10, 30, 100])
        a3.set_xticklabels(["0.1", "0.3", "1", "3", "10", "30", "100"])
        a3.set_xlabel("seconds ON per actuation", color=INK2, fontsize=8)
        a3.set_ylabel("actuations", color=INK2, fontsize=8)
        a3.set_title(f"ON length: {(du <= 0.15).mean():.0%} one-tick, median {du.median():.1f} s" if len(du)
                     else "ON length: no actuations", fontsize=9, color=INK2, loc="left")
        log_txt = "Dec 2024 log" if per == "dec" else "Sept 2026 log"
        head = (f"{r.DeviceName} det {det} ({r.technology if isinstance(r.technology, str) else 'technology unknown'}): "
                f"label {ph_txt} {nm(r.had) if r.kind == 'det' else '(whole signal)'}; we: {r.action}")
        why = textwrap.fill("Why: " + str(r.why_lead), 150)
        extra = f"   other detectors on this row: {', '.join(map(str, [d for d in r.dets if d != det][:12]))}" \
            if r.n_dets > 1 else ""
        fig.suptitle(f"{head}\n{why}\n{log_txt}, {len(me):,} actuations{extra}", x=0.01, ha="left", fontsize=9.5,
                     color=INK)
        fig.tight_layout()
        fig.savefig(png, dpi=95, facecolor=SURF)
        plt.close(fig)
    return res


def pdf_for(name):
    if not isinstance(name, str):
        return None
    for d in PDFS:
        fs = sorted(d.glob(f"{name}_*.pdf"))
        if fs:
            return fs[0]
    return None


def det_text(d) -> str:
    return ", ".join(map(str, d[:MAX_DETS_ROW])) + (f" +{len(d) - MAX_DETS_ROW} more" if len(d) > MAX_DETS_ROW else "")


def row_text(r):
    ph = (", ".join(r.phases) + " ") if r.kind == "det" and r.phases else ""
    if r.kind == "signal":
        had = f"{r.n_dets} labels: {r.had}"
        mod = f"agrees on {int(r.n_agree)} of {int(r.n_atspm)} ATSPM labels" if r.n_atspm else "no prediction"
    else:
        had = ph + nm(r.had)
        if r.n_dets > 1 and isinstance(r.m_mode, str):
            mod = f"{nm(r.m_mode)} {r.m_mode_p:.0%} on {int(r.m_mode_n)} of {int(r.n_dets)} detectors"
        elif isinstance(r.m_pred, str):
            mod = f"{nm(r.m_pred)} {r.m_p:.0%} ({int(r.m_n_top)} of {int(r.m_n)} samples)"
        else:
            mod = "no prediction (no 30-min+ sample in that log)"
        if r.period == "dec" and isinstance(r.m_pred, str):
            mod += ", Dec 2024"
    why = str(r.why_lead)
    if r.kind == "det" and r.n_dets > 1:
        why = f"det {int(r.lead)}: {why} Others on this row: same decision."
    if r.kind == "det" and r.also_lead:
        why += " Also: " + "; ".join(r.also_lead) + "."
    return had, mod, why


def write_sheet(R: pd.DataFrame, charts: dict, xl: Path, chart_rel: str, counts: dict):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "decisions to check"
    ws.append([QUESTION])
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=10)
    ws.row_dimensions[1].height = 78
    ws.append(["#", "Signal", "Detector(s)", "Label we had", "What we did", "Why", "Model says", "Chart", "Print",
               "Answer"])
    link = Font(color="0563C1", underline="single")
    for i, r in enumerate(R.itertuples(), 1):
        had, mod, why = row_text(r)
        pdf = pdf_for(r.DeviceName)
        ws.append([i, r.DeviceName, det_text(list(r.dets)), had, r.action, why, mod, "chart", "print" if pdf else "none",
                   ""])
        n = ws.max_row
        c = ws.cell(n, 8)
        c.hyperlink = f"{chart_rel}/{Path(charts[stem(r)]).name}"
        c.font = link
        if pdf:
            c = ws.cell(n, 9)
            c.hyperlink = "file:///" + urllib.parse.quote(str(pdf).replace("\\", "/"), safe=":/()_-.,'")
            c.font = link
    for j, w in enumerate((5, 8, 12, 16, 16, 48, 17, 7, 7, 10)):
        ws.column_dimensions[chr(65 + j)].width = w
    for c in ws[2]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A3"
    # counts tab (information only)
    w2 = wb.create_sheet("counts")
    w2.append(["Detectors per decision (training + dev signals; a detector can have more than one)", "", "", "", "", ""])
    w2.append(["Decision", "Detectors", "Training signals", "Dev signals", "Released test half", "ATSPM label"])
    for k in ORDER:
        c = counts["by_reason"].get(k)
        if c:
            w2.append([PLAIN[k] + (f" ({c['signals']} signals)" if "signals" in c else ""), c["detectors"],
                       c["train"], c["dev"], c["released"], c["atspm_label"]])
    for c in w2[2]:
        c.font = Font(bold=True)
    for j, w in enumerate((52, 10, 10, 10, 12, 10)):
        w2.column_dimensions[chr(65 + j)].width = w
    xl.parent.mkdir(parents=True, exist_ok=True)
    wb.save(xl)
    log(f"saved {xl} ({len(R)} rows)")


def stage_write(limit: int = 0, preview: bool = False):
    R = build_rows()
    assert not R.DeviceId.isin(locked()).any()
    if limit:
        R = R.head(limit)
    counts = json.load(open(OUT / "counts.json"))
    R.assign(dets=R.dets.map(lambda d: ",".join(map(str, d))), phases=R.phases.map(lambda d: ",".join(d)),
             also_lead=R.also_lead.map(lambda v: "; ".join(v))) \
        [["DeviceId", "DeviceName", "key", "kind", "dets", "lead", "phases", "had", "new", "action", "why_lead",
          "also_lead", "m_pred", "m_p", "tier", "period"]].to_parquet(OUT / "rows.parquet", index=False)
    if preview:
        base = OUT / "preview"
        ch = draw(R, base / "charts")
        write_sheet(R, ch, base / f"{OUTNAME}.xlsx", "charts", counts)
        return
    out = rpath.REPO / "review"
    ch = draw(R, out / f"{OUTNAME}_charts")
    write_sheet(R, ch, out / f"{OUTNAME}.xlsx", f"{OUTNAME}_charts", counts)
    log(json.dumps(dict(rows=len(R), signals=int(R.DeviceId.nunique()), by_tier=R.tier.value_counts().sort_index()
                        .to_dict(), by_key=R.key.value_counts().to_dict(),
                        multi=int((R.n_dets > 1).sum()), with_print=int(sum(pdf_for(n) is not None for n in R.DeviceName))),
                   default=str))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("stage", choices=["items", "data", "write"])
    ap.add_argument("--limit", type=int, default=0)
    ap.add_argument("--preview", action="store_true")
    a = ap.parse_args()
    t0 = time.time()
    {"items": stage_items, "data": stage_data}.get(a.stage, lambda: stage_write(a.limit, a.preview))()
    log(f"done in {time.time() - t0:.0f} s")
