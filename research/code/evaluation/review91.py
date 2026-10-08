"""Note 91: re-lay the three open review sheets in explicit columns (the user found "P5, P8 stop-bar presence" ambiguous).

Nothing is re-selected or re-scored: the rows are the saved build rows of each sheet, the chart / print links are copied
from the sheet as it stands, and every Answer the user already typed is carried to the same (signal, detector) row.
Layout (review72.write_sheet / sanity89.write_sheet): Label | Model says | How sure | What this row asks; each detector's
phase and function written together ("P5 Presence", "P4 Stop-bar Count"), one "det N: P5 Presence" line per detector in
grouped rows (user correction 2026-10-05: never "P5, P8 Presence").

    python review91.py            # dry run: builds into %DC_WORK%/rev91/ and verifies, review/ untouched
    python review91.py --write    # backup to review/_backup/, then rewrite in place (or <name>_v2.xlsx if locked)
"""
from __future__ import annotations

import argparse
import shutil
import sys
import time
from pathlib import Path

import pandas as pd

CODE = next(p for p in Path(__file__).resolve().parents if p.name == "code")
for _p in (CODE, CODE / "evaluation"):
    sys.path.insert(0, str(_p))
import rpath  # noqa: E402
from common import DC_WORK  # noqa: E402
import review72 as R72  # noqa: E402
import sanity89 as S89  # noqa: E402

REV = rpath.REPO / "review"
WORK = DC_WORK / "rev91"
SHEETS = {
    "function_label_review_v1": dict(rows=DC_WORK / "cand64" / "review72_rows.parquet",
                                     frows=DC_WORK / "cand64" / "function_rows.parquet", arm="fj",
                                     question=R72.QUESTION),
    "function_label_review_v1b": dict(rows=DC_WORK / "rev81" / "review81_rows.parquet",
                                      frows=DC_WORK / "rev81" / "function_rows.parquet", arm="champ",
                                      question=R72.QUESTION),
    "label_sanity_20": dict(rows=DC_WORK / "x89" / "v4n" / "label_sanity_20_rows.parquet"),
}


def log(m):
    print(f"[{time.strftime('%H:%M:%S')}] {m}", flush=True)


def snapshot(xl: Path) -> dict:
    """everything a later check needs: question, per-row signal / detector text / answer / chart and print links."""
    from openpyxl import load_workbook
    wb = load_workbook(xl)
    ws = wb["labels to check"]
    head = [c.value for c in ws[2]]
    ia = head.index("Answer") + 1
    idt = next(i for i, h in enumerate(head) if str(h).startswith("Detector")) + 1
    ic, ip = head.index("Chart") + 1, head.index("Print") + 1
    rows = []
    for n in range(3, ws.max_row + 1):
        if ws.cell(n, 1).value is None:
            continue
        ch, pr = ws.cell(n, ic).hyperlink, ws.cell(n, ip).hyperlink
        rows.append(dict(sig=str(ws.cell(n, 2).value),
                         dets=tuple(int(x) for x in str(ws.cell(n, idt).value).split(",") if x.strip().isdigit()),
                         answer=ws.cell(n, ia).value, chart=ch.target if ch else None, print=pr.target if pr else None,
                         comment=ws.cell(n, ia).comment))
    extra = [(c.coordinate, c.value) for row in ws.iter_rows(min_row=3) for c in row
             if c.column > len(head) and c.value is not None]          # anything typed right of the table
    other = [s for s in wb.sheetnames if s != "labels to check"]
    out = dict(question=ws.cell(1, 1).value, head=head, rows=rows, extra=extra, other_sheets=other,
               comments=[(c.coordinate, c.comment.text) for row in ws.iter_rows() for c in row if c.comment])
    wb.close()
    return out


def locked_for_write(xl: Path) -> bool:
    try:
        with open(xl, "r+b"):
            pass
    except PermissionError:
        return True
    return (xl.parent / f"~${xl.name}").exists()


def build(name: str, cfg: dict, xl_in: Path, xl_out: Path) -> dict:
    old = snapshot(xl_in)
    assert not old["extra"] and not old["other_sheets"] and not old["comments"], (name, old["extra"],
                                                                                  old["other_sheets"], old["comments"])
    rows = pd.read_parquet(cfg["rows"])
    links = [(r["chart"], r["print"]) for r in old["rows"]]
    answers = [r["answer"] for r in old["rows"]]
    assert len(rows) == len(old["rows"]), (name, len(rows), len(old["rows"]))
    if name == "label_sanity_20":
        for r, o in zip(rows.itertuples(), old["rows"]):
            assert (str(r.DeviceName), (int(r.lead),)) == (o["sig"], o["dets"]), (r.DeviceName, r.lead, o)
        assert old["question"] == S89.QUESTION
        S89.write_sheet(rows, links, xl_out, answers=answers)
    else:
        for r, o in zip(rows.itertuples(), old["rows"]):
            assert (str(r.DeviceName), tuple(R72._dets(r))) == (o["sig"], o["dets"]), (r.DeviceName, r.dets, o)
        assert old["question"] == cfg["question"]
        rows = R72.det_facts(rows, cfg["frows"], cfg["arm"])
        R72.write_sheet(rows, None, xl_out, "", answers=answers, links=links)
    new = snapshot(xl_out)
    # ---- verification: nothing the user typed is lost, links unchanged and pointing at real files
    assert new["question"] == old["question"]
    assert len(new["rows"]) == len(old["rows"])
    for o, n in zip(old["rows"], new["rows"]):
        assert (o["sig"], o["dets"], o["answer"], o["chart"], o["print"]) == \
               (n["sig"], n["dets"], n["answer"], n["chart"], n["print"]), (o, n)
        assert (REV / n["chart"]).exists(), n["chart"]
    oa = {(o["sig"], o["dets"]): o["answer"] for o in old["rows"] if o["answer"] not in (None, "")}
    na = {(o["sig"], o["dets"]): o["answer"] for o in new["rows"] if o["answer"] not in (None, "")}
    assert oa == na
    log(f"{name}: {len(new['rows'])} rows, {len(oa)} answers carried and identical, links identical -> {xl_out}")
    return dict(rows=len(new["rows"]), answers=len(oa), out=str(xl_out))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--write", action="store_true")
    a = ap.parse_args()
    WORK.mkdir(parents=True, exist_ok=True)
    res = {}
    for name, cfg in SHEETS.items():
        src = REV / f"{name}.xlsx"
        if not a.write:
            res[name] = build(name, cfg, src, WORK / f"{name}.xlsx")
            continue
        bk = REV / "_backup"
        bk.mkdir(exist_ok=True)
        b = bk / f"{name}_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        shutil.copy2(src, b)
        assert b.read_bytes() == src.read_bytes()
        tmp = WORK / f"{name}.xlsx"
        res[name] = build(name, cfg, b, tmp)                 # built and verified from the backup copy
        dst = REV / f"{name}_v2.xlsx" if locked_for_write(src) else src
        shutil.copyfile(tmp, dst)
        res[name].update(out=str(dst), backup=str(b), locked=dst != src)
        log(f"{name}: backup {b.name}; written {dst.name}" + (" (original LOCKED, wrote _v2)" if dst != src else ""))
    print(res)


if __name__ == "__main__":
    main()
