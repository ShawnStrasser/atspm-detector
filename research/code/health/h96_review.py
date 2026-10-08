"""Note 96: short review sheet review/health_review_v2.xlsx (+ _charts/, _day_charts/): the cases where the proposed
context-aware health logic (h96_resolve) and the current package disagree, or where it is unsure - at most ~20 rows,
none of the 70 detectors already answered in v1 except v1 row 53 (the user asked to see its occupancy).

Charts from the saved w40 events (Sat 26 16:15 - Mon 28 24:00) only, same style as v1 (note 82c / 82d): line charts,
this detector blue, comparison orange, flagged period shaded, >= 2 h of context; plus an occupancy panel.
  Chart     : top = actuations of this detector and the traffic on its phase (predicted Advance + Count detectors
              combined, scaled to the same total so the shapes compare); bottom = % of time ON, this detector vs
              the average of the other detectors on its phase.
  Day chart : the calendar day of the flagged period: actuations per 15 min (this detector vs its best-matching
              same-phase detector) and % of time ON of both.
Selection, rows, sentences: hi-res log + the classifier's outputs only (inputs as note 79); seed 96.

    python h96_review.py
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

os.environ.setdefault("OMP_NUM_THREADS", "4")
sys.path.insert(0, str(Path(__file__).resolve().parent))
import h82_charts3 as C3  # noqa: E402
import h96_occ as O  # noqa: E402

import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
import pyarrow.dataset as ds  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

hc = O.hc
DCW, OUT = O.DCW, O.OUT
REPO = Path(__file__).resolve().parents[3]
XL = REPO / "review" / "health_review_v2.xlsx"
CH = REPO / "review" / "health_review_v2_charts"
DC = REPO / "review" / "health_review_v2_day_charts"
THIS, PART, INK, MUTED = C3.THIS, C3.PART, C3.INK, C3.MUTED
FN = {"Advance": "Advance", "Presence": "Presence", "Count": "Count", "Yellow_Red": "Yellow-red", "Mid": "Mid",
      "Bike": "Bike", "Other": "Other"}
CHK = {"stuck": "stuck on", "dropout": "goes silent", "chatter": "chattering", "rapid": "too-fast actuations",
       "volume": "too many in 5 min", "level": "count drops", "choppy": "erratic counts",
       "night_drop": "misses vehicles at night", "night_day": "busier at night", "corr": "doesn't follow traffic"}
PLAN = [("R1", 3), ("R2", 3), ("R3", 2), ("R4", 2), ("R5", 1), ("R7", 2), ("R9", 2), ("N1", 3)]
WHAT = {"R1": "Stuck on, but traffic was heavy", "R2": "Counts don't follow traffic, but its time ON does",
        "R3": "Count drops, but its time ON kept up", "R4": "Too many in 5 min, but it spans 2+ lanes",
        "R5": "Too-fast actuations, few actuations", "R7": "Just over a limit, nothing else wrong",
        "R9": "Flagged only because its neighbours are broken",
        "N1": "NEW: ON much longer than this kind of detector", "V1": "v1 row 53 again, with time ON"}
NEWTXT = {"R1": "ok (a queue in heavy traffic)", "R2": "ok (long zone: its time ON follows traffic)",
          "R3": "ok (its time ON kept up with traffic)", "R4": "ok (busy multi-lane detector)",
          "R5": "watch only (too few actuations to judge)", "R7": "watch only (just over the limit, no other finding)",
          "R9": "ok (judged against its healthy neighbours)", "N1": "watch only (new note)"}
ST = {"ok": "ok", "suspect": "suspect", "bad": "bad", "not_enough_data": "too little data", "watch": "watch"}


def names():
    n = pd.concat([pd.read_parquet(DCW / "official" / "labels_official.parquet", columns=["DeviceId", "DeviceName"]),
                   pd.read_parquet(REPO / "research" / "labels" / "function_labels_v3.parquet",
                                   columns=["DeviceId", "DeviceName"])]).drop_duplicates().dropna(subset=["DeviceName"])
    return dict(zip(n.DeviceId.str.lower(), n.DeviceName.astype(str)))


def select():
    R = pd.read_parquet(OUT / "resolved.parquet")
    v1 = pd.read_csv(DCW / "health82" / "rows.csv")
    seen = set(zip(v1.dev.str.lower(), v1.det))
    R = R[~pd.Series(list(zip(R.DeviceId, R.detector))).isin(seen).to_numpy()]
    R = R[R.wg.isin(["h3", "h24"])]
    old = R.status.isin(["suspect", "bad"])
    rng = np.random.default_rng(96)
    R = R.assign(rnd=rng.random(len(R)))
    chosen, used = [], set()
    for rule, k in PLAN:
        if rule == "N1":
            c = R[R.rules.str.contains("N1") & R.status.eq("ok")]
        else:
            c = R[R.rules.str.split(",").str[0].eq(rule) & old & ~R.new_status.isin(["suspect", "bad"])]
        c = c.sort_values("rnd")
        got, wgs = 0, []
        for _, r in c.iterrows():
            if r.DeviceId in used:
                continue
            if rule == "R7" and any(c.get("rule") == "R7" and c.watch == r.watch for c in chosen):
                continue                                                 # one per check
            if rule == "R1" and not np.isfinite(r.ep_phx):
                continue                                                 # needs phase peers to show
            if got < k - 1 and wgs.count(r.wg) >= (k + 1) // 2:     # spread 3 h / 24 h samples
                continue
            rr = r.copy()
            rr["rule"] = rule
            chosen.append(rr)
            used.add(r.DeviceId)
            wgs.append(r.wg)
            got += 1
            if got == k:
                break
    S = pd.DataFrame(chosen)
    # v1 row 53 (2B530 d60, 24 h Mon): the user asked for occupancy
    r53 = v1[v1.n == 53].iloc[0]
    A = pd.read_parquet(OUT / "resolved.parquet")
    x = A[(A.DeviceId == r53.dev.lower()) & (A.window == r53.window) & (A.detector == r53.det)].iloc[0].copy()
    x["rule"] = "V1"
    return pd.concat([S, x.to_frame().T], ignore_index=True)


def hm(t):
    return pd.Timestamp(t).strftime("%a %H:%M")


def pass1(r):
    h = pd.read_parquet(OUT / "health.parquet", filters=[("DeviceId", "==", r.DeviceId), ("window", "==", r.window)])
    return h[h.detector == r.detector].iloc[0]


def main_check(r):
    s = {k: float(r[f"s_{k}"]) for k in CHK if pd.notna(r.get(f"s_{k}")) and float(r[f"s_{k}"]) >= .35}
    return max(s, key=s.get) if s else None


def sentence(r):
    rule = r.rule
    occf = f"its time ON follows the traffic on its phase (correlation {r.c_occ_ref:.2f})"
    if rule == "R1":
        return (f"Held ON for {float(r.ep_dur) / 60:.0f} min ({hm(r.ep_t0)} - {pd.Timestamp(r.ep_t1):%H:%M}); meanwhile "
                f"the other detectors on its phase were ON {r.ep_phx:.1f}x as much as usual and the traffic on its phase "
                f"kept moving ({r.ep_refx:.1f}x its usual count).")
    if rule == "R2":
        k = "corr" if float(r.s_corr or 0) >= .35 else "choppy"
        what = (f"its counts do not rise and fall with the signal (correlation {float(r['corr']):.2f}, normally about "
                f"{float(r.corr_exp):.2f})" if k == "corr" else
                f"its 15-min counts jump around {float(r.chop15):.0f}x more than the others explain (limit "
                f"{float(r.chop_lim):.0f}x)")
        return f"Check: {what}. But it is a long zone and {occf}."
    if rule == "R3":
        return (f"Check: from {hm(pd.Timestamp(O.WIN[r.window][0]) + pd.Timedelta(minutes=5 * int(float(r.level_b))))} "
                f"it counted only {float(r.level_ratio):.0%} of its earlier share. But its time ON kept up with traffic "
                f"({r.lv_occx:.0%} of before, relative to traffic).")
    if rule == "R4":
        return (f"Check: {float(r.max5):.0f} actuations in one 5 min (limit 150). The classifier says it spans "
                f"{int(r.lanes)} lanes (healthy 2-lane detectors reach about 140) and its counts follow traffic "
                f"(correlation {r.c_cnt_ref:.2f}).")
    if rule == "R5":
        return (f"Check: {float(r.ioi_lt1):.0%} of its actuations start within 1 s of the previous one "
                f"({float(r.rapid):.2f}x the limit), but it had only {int(r.n_on)} actuations in the sample.")
    if rule == "R7":
        k = r.watch.split(",")[0]
        num = {"night_drop": f"at night it counted {float(r.night_ratio):.1%} of what was expected (limit 19 %)",
               "level": f"it dropped to {float(r.level_ratio):.1%} of its earlier share (limit 15 %)",
               "choppy": f"its counts jump {float(r.chop15):.1f}x more than the others explain (limit "
                         f"{float(r.chop_lim):.1f}x)",
               "rapid": f"{float(r.ioi_lt1):.0%} of actuations within 1 s ({float(r.rapid):.2f}x the limit)",
               "corr": f"correlation with the signal {float(r['corr']):.2f}, normally about {float(r.corr_exp):.2f}",
               "dropout": f"silent {(float(r.drop_b1) - float(r.drop_b0)) * 5:.0f} min while about "
                          f"{float(r.drop_lam):.0f} actuations were expected",
               "night_day": f"{float(r.night_day):.1f}x its day rate at night"}.get(k, k)
        return f"Check '{CHK.get(k, k)}' just over its limit: {num}. No other check fired."
    if rule == "R9":
        A = pd.read_parquet(OUT / "resolved.parquet", filters=[("DeviceId", "==", r.DeviceId), ("window", "==", r.window)])
        bad = A[(A.status == "bad") & (A.phase == r.phase) & (A.detector != r.detector)].detector.astype(int).tolist()
        mc = main_check(pass1(r))
        who = ", ".join(f"det {b}" for b in bad) if bad else "other detectors"
        verb = "is itself" if len(bad) == 1 else "are themselves"
        return (f"The current check '{CHK.get(mc, mc)}' compared it with {who} on its phase, which {verb} "
                f"flagged bad. Compared with the healthy detectors only, nothing is wrong with it.")
    if rule == "N1":
        return (f"ON {float(r.occ):.0%} of the sample; for {int(r.hi_n) * 15} min it was ON far longer than "
                f"{FN.get(r.fn, r.fn)} zones are at that traffic, and its time ON does not follow traffic "
                f"(correlation {r.c_occ_ref:.2f}). No current check fired.")
    if rule == "V1":
        return (f"You asked to see its time ON (v1 row 53). Counts do not follow traffic (correlation "
                f"{float(r['corr']):.2f}), but {occf}. It is still flagged because it also held ON for "
                f"{float(r.ep_dur) / 60:.0f} min ({hm(r.ep_t0)}).")
    return ""


def title_of(r, d):
    return {"R1": f"Det {d} held ON {float(r.ep_dur) / 60:.0f} min while its phase was busy",
            "R2": f"Det {d}'s counts wander, but its time ON follows traffic",
            "R3": f"Det {d} counted less from {hm(pd.Timestamp(O.WIN[r.window][0]) + pd.Timedelta(minutes=5 * int(float(r.level_b))))}, but its time ON kept up",
            "R4": f"Det {d} had {float(r.max5):.0f} actuations in one 5 min",
            "R5": f"Det {d} has many quick repeats, but only {int(r.n_on)} actuations",
            "R7": f"Det {d} is just over one limit",
            "N1": f"Det {d} is ON {float(r.occ):.0%} of the time",
            "R9": f"Det {d} looks fine next to its healthy neighbours",
            "V1": f"Det {d}: counts and time ON over the weekend"}[r.rule]


# ---------------------------------------------------------------- data + charts
def events(dev):
    p = O.EVD / f"DeviceId={dev}"
    if not p.is_dir():
        p = next(q for q in O.EVD.iterdir() if q.name.lower() == f"deviceid={dev}")
    e = ds.dataset(p).to_table(filter=ds.field("EventId").isin(list(hc.ALLOWED))).to_pandas()
    return e[["Timestamp", "EventId", "Parameter"]].drop_duplicates()


def bins(ev, z0, z1, bs):
    B = hc.events_to_bins(ev, z0, z1, bin_s=bs)
    n = B["n_on"].astype(float)
    occ = B["occ"].astype(float) / bs * 100
    n[:, ~B["cov"]] = np.nan
    occ[:, ~B["cov"]] = np.nan
    tt = z0 + pd.to_timedelta(np.arange(n.shape[1]) * bs + bs / 2, unit="s")
    return {int(d): i for i, d in enumerate(B["dets"])}, tt, n, occ


def inputs(r):
    x = pd.read_parquet(OUT / "feat.parquet", filters=[("DeviceId", "==", r.DeviceId), ("window", "==", r.window)])
    return dict(zip(x.detector, x.phase)), dict(zip(x.detector, x.fn))


def flagged(r, t0, t1, ev, d):
    if r.rule in ("R1", "V1") and pd.notna(pd.Timestamp(r.ep_t0)):
        return pd.Timestamp(r.ep_t0), pd.Timestamp(r.ep_t1), "the long ON"
    if r.rule == "R3":
        return t0 + pd.Timedelta(minutes=5 * int(float(r.level_b))), t1, "after the drop (in the sample)"
    if r.rule == "R4":
        ix, tt, n, _ = bins(ev, t0, t1, 300)
        k = int(np.nanargmax(n[ix[d]]))
        a = t0 + pd.Timedelta(minutes=5 * k)
        return a, a + pd.Timedelta(minutes=5), "the busiest 5 minutes"
    return t0, t1, "the sample the check looked at"


def two_panel(path, title, sub, z0, z1, bs, ix, tt, n, occ, d, top2, top2_lab, low2, low2_lab, shade, sample,
              top_ylabel):
    fig, (a1, a2) = plt.subplots(2, 1, figsize=(13, 8.2), sharex=True, gridspec_kw=dict(height_ratios=[1, 1]))
    ser = [(f"det {d}", tt, n[ix[d]], THIS, 2.6)]
    if top2 is not None:
        ser.append((top2_lab, tt, top2, PART, 2.0))
    C3.line_chart(a1, z0, z1, ser, top_ylabel, shade=shade, sample=sample)
    ser = [(f"det {d}", tt, occ[ix[d]], THIS, 2.6)]
    if low2 is not None:
        ser.append((low2_lab, tt, low2, PART, 2.0))
    C3.line_chart(a2, z0, z1, ser, f"% of each {C3.bin_word(bs).rstrip('s')} ON", shade=shade, sample=sample,
                  ylim=(0, 102))
    a1.set_xlabel("")
    for ax in (a1, a2):
        lg = ax.get_legend()
        h_, l_ = (lg.legend_handles, [t.get_text() for t in lg.get_texts()]) if lg else ax.get_legend_handles_labels()
        ax.legend(h_, l_, ncol=2, fontsize=10.5, frameon=False, loc="lower left", bbox_to_anchor=(0, 1.0))
    fig.suptitle(title, x=0.01, ha="left", fontsize=16, fontweight="bold", color=INK)
    fig.text(0.01, 0.94, sub, ha="left", va="top", fontsize=12, color=MUTED)
    fig.tight_layout(rect=(0, 0, 1, 0.92 - 0.025 * sub.count("\n")))
    fig.savefig(path, dpi=100)
    plt.close(fig)


def charts(r, sig, d, png, day_png):
    ev = events(r.DeviceId)
    s0, h = O.WIN[r.window]
    t0 = pd.Timestamp(s0)
    t1 = t0 + pd.Timedelta(hours=h)
    cov0, cov1 = ev.Timestamp.min().floor("15min"), ev.Timestamp.max().ceil("15min")
    ph, fn = inputs(r)
    A = pd.read_parquet(OUT / "resolved.parquet", filters=[("DeviceId", "==", r.DeviceId), ("window", "==", r.window)])
    bad = set(A.loc[A.status == "bad", "detector"].astype(int)) - {d}      # broken detectors are nobody's yardstick
    p = ph.get(d, np.nan)
    same = [k for k in ph if k != d and k not in bad and np.isfinite(p) and ph[k] == p]
    traf = [k for k in same if fn.get(k) in ("Advance", "Count")]
    tlab = "traffic on its phase (advance + count detectors, scaled)"
    if not traf:
        traf = [k for k in ph if k != d and k not in bad and fn.get(k) in ("Advance", "Count")]
        tlab = "traffic at the signal (advance + count detectors, scaled)"
    a, b, lab = flagged(r, t0, t1, ev, d)
    # main chart: the sample + >= 3 h either side (24 h samples: all saved data)
    if h >= 24:
        z0, z1 = cov0, cov1
    else:
        z0, z1 = max(cov0, t0 - pd.Timedelta(hours=3)), min(cov1, t1 + pd.Timedelta(hours=3))
    bs = 900
    ix, tt, n, occ = bins(ev, z0, z1, bs)
    if d not in ix:
        return None
    tr = np.nansum(n[[ix[k] for k in traf if k in ix]], 0) if traf else None
    if tr is not None:
        tr = tr * np.nansum(n[ix[d]]) / max(np.nansum(tr), 1)
    po = np.nanmean(occ[[ix[k] for k in same if k in ix]], 0) if same else None
    model = f"P{int(p)} {FN.get(fn.get(d), '?')}" if np.isfinite(p) else f"no phase {FN.get(fn.get(d), '?')}"
    sub = f"Signal {sig}   |   det {d}: {model}   |   {WHAT[r.rule]}"
    title = title_of(r, d)
    two_panel(png, title, sub, z0, z1, bs, ix, tt, n, occ, d, tr, tlab, po,
              "other detectors on its phase (average)", [(a, b, lab)], (t0, t1), f"actuations per 15 minutes")
    # day chart: calendar day of the flagged start, raw counts, best-matching same-phase detector
    day = a.normalize()
    y0, y1 = day, day + pd.Timedelta(days=1)
    ix2, tt2, n2, occ2 = bins(ev, y0, y1, 900)
    part, best = None, -2
    for k in same:
        if k in ix2:
            x1, x2 = n2[ix2[d]], n2[ix2[k]]
            m = np.isfinite(x1) & np.isfinite(x2)
            if m.sum() > 4 and np.std(x1[m]) > 0 and np.std(x2[m]) > 0:
                c = np.corrcoef(x1[m], x2[m])[0, 1]
                if c > best:
                    part, best = k, c
    pl = f"det {part} (same phase)" if part is not None else None
    sub2 = f"Signal {sig}   |   det {d}: {model}   |   {WHAT[r.rule]}"
    if cov1 < y1 - pd.Timedelta(minutes=15):
        sub2 += f"\nThe saved data end at {cov1:%H:%M}; nothing after that is shown."
    two_panel(day_png, f"Det {d} - actuations and time ON per 15 min, {day:%a %d %b}", sub2, y0, y1, 900, ix2, tt2,
              n2, occ2, d, n2[ix2[part]] if part is not None else None, pl,
              occ2[ix2[part]] if part is not None else None, pl, [(max(a, y0), min(b, y1), lab)], (t0, t1),
              "actuations per 15 minutes")
    return title, model


def main():
    S = select()
    nm = names()
    for p_ in (CH, DC):
        p_.mkdir(parents=True, exist_ok=True)
    rows = []
    for i, r in enumerate(S.itertuples(index=False), start=1):
        r = pd.Series(r._asdict())
        d = int(r.detector)
        sig = nm.get(r.DeviceId, r.DeviceId[:8])
        png = CH / f"{i:02d}_{r.rule}_{sig}_d{d}.png"
        dpng = DC / png.name
        out = charts(r, sig, d, png, dpng)
        if out is None:
            continue
        title, model = out
        mc = main_check(pass1(r) if r.rule == "R9" else r)
        old = ST.get(r.status, r.status) + (f" ({CHK[mc]})" if mc and r.status in ("suspect", "bad") else "")
        new = NEWTXT.get(r.rule, ST.get(r.new_status, r.new_status))
        if r.rule == "V1":
            new = ST.get(r.new_status, r.new_status) + " (still flagged: stuck on)"
        s0, h = O.WIN[r.window]
        rows.append(dict(n=i, what=WHAT[r.rule], signal=sig, det=f"det {d}: {model}",
                         sample=f"{'24 h' if h >= 24 else '3 h' if h >= 3 else '30 min'}, {pd.Timestamp(s0):%a %d %b %H:%M}",
                         old=old, new=new, saw=sentence(r), png=png, dpng=dpng, title=title, rule=r.rule,
                         dev=r.DeviceId, window=r.window, detector=d))
        print(i, r.rule, sig, d, r.window, "|", old, "->", new, "|", rows[-1]["saw"])
    pd.DataFrame(rows).to_csv(OUT / "review_rows.csv", index=False)
    write_xl(rows)


def write_xl(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cases"
    ws["A1"] = "Is this a real problem? Answer Y (yes, real problem), N (no, the detector is fine) or ? (can't tell)."
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Data: Sat 26 - Mon 28 Sep 2026. Each row = one detector in one sample. The proposed new health logic "
                "and the current one disagree on these (or it is a new note). 'Old' = what the current package says; "
                "'New' = what the proposed logic would say. Chart: top = this detector's actuations (blue) vs the "
                "traffic on its phase (orange; detectors flagged bad left out); bottom = % of time ON (blue) vs the other detectors on its phase "
                "(orange). Flagged period shaded; hours before and after shown. Day chart = the whole day.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:L2")
    ws.row_dimensions[2].height = 62
    head = ["#", "What the new logic does", "Signal", "Det", "Sample", "Old result", "New result",
            "What the data shows", "Chart", "Day chart", "Answer (Y / N / ?)", "Comment"]
    for j, h in enumerate(head, start=1):
        c = ws.cell(row=3, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2A78D6")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for k, r in enumerate(rows, start=4):
        vals = [r["n"], r["what"], r["signal"], r["det"], r["sample"], r["old"], r["new"], r["saw"], "chart",
                "day chart", None, None]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=k, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=k, column=9).hyperlink = f"{CH.name}/{r['png'].name}"
        ws.cell(row=k, column=10).hyperlink = f"{DC.name}/{r['dpng'].name}"
        for j in (9, 10):
            ws.cell(row=k, column=j).font = Font(color="0563C1", underline="single")
        ws.cell(row=k, column=11).fill = PatternFill("solid", fgColor="FFF8DC")
    for col, w in zip("ABCDEFGHIJKL", (4, 26, 8, 20, 18, 22, 24, 70, 8, 10, 14, 40)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    XL.parent.mkdir(exist_ok=True)
    wb.save(XL)
    print("saved", XL, len(rows), "rows")


if __name__ == "__main__":
    main()
