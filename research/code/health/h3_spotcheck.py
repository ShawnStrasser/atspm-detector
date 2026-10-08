"""Note 43: short follow-up spot-check of the v3 health scorer (review/spotcheck_health.xlsx).

<= 10 rows, one per detector, 66-h Sept-2026 window, training signals only (never locked_v2):
the user's "maybe" rows re-judged, his stuck-loop example, and new cases showing the choppiness
rule, stuck-period handling (bad periods + recovery check) and the phase/function-aware calls.
Chart per row (review/spotcheck_health_charts/v3_<signal>_d<det>.png): 15-min counts of the
detector vs the other detectors on its PREDICTED phase and the signal total, bad periods shaded;
second panel = the evidence with a healthy baseline (a healthy same-phase detector, or the typical
healthy detector of the same predicted function from h3_baseline.py).  Phase / function are the
OOF predictions the scorer used.  The first sheet's answers of 2026-09-28 move to a second sheet.

    python h3_spotcheck.py
"""
from __future__ import annotations

import json
import sys
import urllib.parse
import warnings
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.dates as mdates  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
import pyarrow.dataset as ds  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
import hb_build as HB  # noqa: E402
import hb_data as H  # noqa: E402
import health_core as hc  # noqa: E402

warnings.filterwarnings("ignore")
XL = H.REPO / "review" / "spotcheck_health.xlsx"
CH = H.REPO / "review" / "spotcheck_health_charts"
PDF = H.DCW / "cabinet" / "pdf"
EV = H.DCW / "official" / "stg" / "cache" / "events"
T0, T1 = pd.Timestamp("2026-09-18 16:15"), pd.Timestamp("2026-09-21 10:25")
SIBC = ["#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#4a3aa7", "#008300"]
DETC, TOTC, INK, MUTED, BASEC, BADC = "#2a78d6", "#d9d8d3", "#0b0b0b", "#52514e", "#8a8a85", "#fde2d6"


def fmt(t):
    return pd.Timestamp(t).strftime("%a %d %b %H:%M")


def style(ax):
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color("#bbbbbb")
    ax.tick_params(colors=MUTED, labelsize=8)
    ax.grid(axis="y", color="#ececec", lw=0.6)


def load(dev):
    ev = ds.dataset(EV / f"DeviceId={dev}").to_table(filter=ds.field("EventId").isin(list(hc.ALLOWED))).to_pandas()
    pf = pd.read_parquet(H.DCW / "health3" / "oof_pf.parquet")
    pf = pf[(pf.period == "stg") & (pf.wgroup == "full") & (pf.DeviceId == dev)]
    ph, fn = dict(zip(pf.detector, pf.pred_phase)), dict(zip(pf.detector, pf.pred_function))
    h = hc.health(ev, T0, T1, HB.expected_channels().get(dev), ph, fn)
    B = hc.events_to_bins(ev, T0, T1, HB.expected_channels().get(dev))
    return ev, h, B, ph, fn


def on_table(ev, d):
    x = ev[ev.EventId.isin([81, 82]) & (ev.Parameter == d)].drop_duplicates().sort_values(["Timestamp", "EventId"],
                                                                                          ascending=[True, False])
    # continuous ON (note 47): from an ON to the next OFF; ONs inside it (ON -> ON, no OFF between) are normal
    # extension re-calls and are not plotted separately
    e = x.EventId.to_numpy()
    first = (e == 82) & np.r_[True, e[:-1] != 82]
    on = x.Timestamp.to_numpy()[first]
    off = x[x.EventId == 81].Timestamp.to_numpy()
    k = np.searchsorted(off, on, "left")
    end = np.where(k < len(off), off[np.minimum(k, len(off) - 1)], np.datetime64("NaT"))
    return pd.DataFrame({"t": on, "dur": (end - on) / np.timedelta64(1, "s")})


def zres(B, i, ref):
    """standardised 15-min residual of detector i around the local share of `ref` (the chop model)."""
    n15 = hc._agg(B["n_on"].astype(float))
    ok = hc._agg(B["cov"][None].astype(float))[0] == 3
    x, r = np.where(ok, n15[i], 0), np.where(ok, n15[ref].sum(0), 0)
    h = int(hc.H_REF * 4)
    X, R = hc._movsum(x, h) - x, hc._movsum(r, h) - r
    s = X / np.maximum(R, 1e-9)
    e = s * r
    z = (x - e) / np.sqrt(np.maximum(e * (1 + s), 1))
    return np.where(ok & (R > 0), z, np.nan)


def chart(row, ev, h, B, ph, fn, kind, path):
    dets = list(B["dets"])
    d = int(row.detector)
    i = dets.index(d)
    p = ph.get(d)
    sib = [k for k in dets if k != d and ph.get(k) == p and B["n_on"][dets.index(k)].sum() > 0]
    sib = sorted(sib, key=lambda k: -B["n_on"][dets.index(k)].sum())[:5]
    stat = dict(zip(h.detector, h.status))
    healthy = [k for k in sib if stat.get(k) == "ok"]
    n15 = hc._agg(B["n_on"].astype(float))
    cov15 = hc._agg(B["cov"][None].astype(float))[0] > 0
    tt = T0 + pd.to_timedelta(np.arange(n15.shape[1]) * 15, unit="min")
    msk = lambda y: np.where(cov15, y, np.nan)  # noqa: E731
    fig, (ax, a2) = plt.subplots(2, 1, figsize=(12, 7.4), gridspec_kw={"height_ratios": [1.25, 1]})
    tot = n15.sum(0)
    ref = max([n15[i].max()] + [n15[dets.index(k)].max() for k in sib] + [1])
    kk = ref / max(tot.max(), 1)
    ax.fill_between(tt, 0, msk(tot * kk), color=TOTC, lw=0, step="mid", label=f"all {len(dets)} detectors (scaled x{kk:.2f})")
    periods = json.loads(row.bad_periods) if isinstance(row.bad_periods, str) else list(row.bad_periods or [])
    for P in periods:
        for a_ in (ax, a2 if kind in ("stuck", "silent") else None):
            if a_ is not None:
                a_.axvspan(pd.Timestamp(P["start"]), pd.Timestamp(P["end"]), color=BADC, lw=0, zorder=0)
    for j, k in enumerate(sib):
        ax.plot(tt, msk(n15[dets.index(k)]), color=SIBC[j], lw=1.0, label=f"d{k} (same predicted phase, {stat.get(k)})")
    ax.plot(tt, msk(n15[i]), color=DETC, lw=2.4, label=f"d{d} (this detector)")
    ax.set_ylabel("actuations per 15 min", fontsize=9, color=MUTED)
    ax.xaxis.set_major_locator(mdates.HourLocator(byhour=[0, 6, 12, 18]))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%a %H:%M"))
    ax.set_xlim(tt[0], tt[-1])
    ax.legend(fontsize=7.5, loc="upper left", frameon=False, ncol=2)
    if periods:
        ax.text(1.0, 1.01, "shaded = bad periods reported by the scorer", transform=ax.transAxes, ha="right",
                fontsize=7.5, color=MUTED)
    style(ax)
    fig.suptitle(f"{row.signal}   detector {d}   predicted phase {int(p) if p else '-'}, predicted {fn.get(d, '-')}"
                 f"   -   {row.status.upper()}", fontsize=11.5, color=INK, x=0.01, ha="left", fontweight="bold")
    if kind in ("stuck", "silent"):
        base = healthy[:1] or sib[:1]
        for k in base:
            o = on_table(ev, k)
            a2.scatter(o.t, o.dur.clip(lower=0.1), s=4, color=BASEC, alpha=.5, label=f"d{k} (healthy, same phase)")
        o = on_table(ev, d)
        a2.scatter(o.t, o.dur.clip(lower=0.1), s=6, color=DETC, label=f"d{d}")
        a2.set_yscale("log")
        a2.axhline(900, color="#e34948", lw=0.8, ls="--")
        a2.set_ylabel("length of each ON (s, log)", fontsize=9, color=MUTED)
        a2.set_xlim(tt[0], tt[-1])
        a2.xaxis.set_major_formatter(mdates.DateFormatter("%a %H:%M"))
        a2.set_title("how long each actuation stayed ON (dashed = 15 min); shaded = bad period", fontsize=9,
                     loc="left", color=MUTED)
    elif kind == "rapid":
        bl = pd.read_parquet(H.DCW / "health3" / "ioi_baseline.parquet")
        b = bl[bl.pred_function == fn.get(d)]
        if len(b):
            b = b.iloc[0]
            e = np.asarray(b.edges)
            a2.fill_between(e[:-1], b.lo, b.hi, step="post", color=TOTC, lw=0,
                            label=f"typical healthy {fn.get(d)} detector (middle 80 % of {int(b.n)})")
            a2.step(e[:-1], b.med, where="post", color=BASEC, lw=1.2, label="its median")
        e = np.logspace(-1, np.log10(600), 40)
        for k, c, lw in ([(healthy[0], SIBC[0], 1.2)] if healthy else []) + [(d, DETC, 2.4)]:
            t = np.sort(ev[(ev.EventId == 82) & (ev.Parameter == k)].Timestamp.drop_duplicates().to_numpy())
            ioi = np.diff(t) / np.timedelta64(1, "s")
            hh, _ = np.histogram(ioi, e)
            a2.step(e[:-1], hh / max(len(ioi), 1), where="post", color=c, lw=lw,
                    label=f"d{k}" + (" (this)" if k == d else " (healthy, same phase)"))
        a2.axvspan(0.1, 1.0, color=BADC, lw=0, zorder=0)
        a2.set_xscale("log")
        a2.set_xlabel("time from one actuation to the next (s, log); shaded < 1 s", fontsize=9, color=MUTED)
        a2.set_ylabel("share of intervals", fontsize=9, color=MUTED)
        a2.set_title("gaps between successive ONs, against a healthy baseline", fontsize=9, loc="left", color=MUTED)
    else:                                               # choppiness: residuals around the phase's own pattern
        refi = [dets.index(k) for k in sib] or [j for j in range(len(dets)) if j != i]
        a2.axhspan(-3, 3, color="#f1f1ee", lw=0, zorder=0)
        for j, k in enumerate(sib):
            others = [dets.index(q) for q in sib if q != k] + [i]
            a2.plot(tt, zres(B, dets.index(k), others), color=SIBC[j], lw=0.8, alpha=.8, label=f"d{k}")
        a2.plot(tt, zres(B, i, refi), color=DETC, lw=2.0, label=f"d{d} (this)")
        a2.set_ylabel("spike size (standard units)", fontsize=9, color=MUTED)
        a2.set_xlim(tt[0], tt[-1])
        a2.xaxis.set_major_formatter(mdates.DateFormatter("%a %H:%M"))
        a2.set_title("15-min count minus what the other detectors on the phase predict (grey band = +-3, where "
                     "healthy detectors stay); a spike the whole phase shares cancels out", fontsize=9, loc="left",
                     color=MUTED)
    a2.legend(fontsize=7.5, frameon=False, ncol=3)
    style(a2)
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    fig.savefig(path, dpi=110)
    plt.close(fig)


def main():
    rows = pd.read_csv(H.DCW / "health3" / "spotcheck_v3_rows.csv")
    CH.mkdir(parents=True, exist_ok=True)
    out = []
    for r in rows.itertuples():
        ev, h, B, ph, fn = load(r.DeviceId)
        x = h[h.detector == r.detector].iloc[0]
        x = x.copy()
        x["signal"] = r.signal
        x["bad_periods"] = x.bad_periods
        png = CH / f"v3_{r.signal}_d{int(r.detector)}.png"
        chart(x, ev, h, B, ph, fn, r.kind, png)
        pdfs = sorted(PDF.glob(f"{r.signal}_*.pdf"))
        p = ph.get(int(r.detector))
        out.append(dict(signal=r.signal, detector=int(r.detector), phase=int(p) if p else None,
                        function=fn.get(int(r.detector), "-"), status=x.status, what=r.what, png=png.name,
                        pdf=pdfs[0] if pdfs else None))
        print(r.signal, r.detector, x.status, "|", x.reason[:200])
    # keep the earlier answers as a second sheet
    old = openpyxl.load_workbook(XL)
    ws0 = old.worksheets[0]
    prev = [[c.value for c in row] for row in ws0.iter_rows()]
    wb = Workbook()
    ws = wb.active
    ws.title = "health follow-up"
    ws.append(["Signal", "Detector", "Phase (predicted)", "Predicted function", "Health status",
               "What was found (log Fri 18 Sep 16:15 - Mon 21 Sep 10:25, 2026)", "Chart", "Cabinet print",
               "Your answer (yes / no / ?)", "Comment"])
    link = Font(color="0563C1", underline="single")
    for n, r in enumerate(out, start=2):
        ws.append([r["signal"], r["detector"], r["phase"], r["function"], r["status"], r["what"], "open chart",
                   "open print" if r["pdf"] is not None else "no print on file", "", ""])
        ws.cell(n, 7).hyperlink = f"spotcheck_health_charts/{r['png']}"
        ws.cell(n, 7).font = link
        if r["pdf"] is not None:
            ws.cell(n, 8).hyperlink = "file:///" + urllib.parse.quote(str(r["pdf"]), safe=":\\/()_-.,'")
            ws.cell(n, 8).font = link
    for col, w in zip("ABCDEFGHIJ", (9, 9, 9, 12, 13, 95, 12, 14, 14, 40)):
        ws.column_dimensions[col].width = w
    w2 = wb.create_sheet("your answers 2026-09-28")
    for row in prev:
        w2.append(row)
    for n in range(2, len(prev) + 1):                    # re-point the old links
        for c, attr in ((6, "chart"), (7, "print")):
            src = ws0.cell(n, c)
            if src.hyperlink is not None:
                w2.cell(n, c).hyperlink = src.hyperlink.target
                w2.cell(n, c).font = link
    for col, w in zip("ABCDEFGHI", (9, 9, 7, 13, 90, 12, 14, 14, 50)):
        w2.column_dimensions[col].width = w
    for sh in (ws, w2):
        for c in sh[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
        for row in sh.iter_rows(min_row=1):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"
    wb.save(XL)


if __name__ == "__main__":
    main()
