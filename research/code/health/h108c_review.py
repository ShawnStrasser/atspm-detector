"""Note 108 (v3c): review/health_review_v3.xlsx rebuilt so that EVERY chart shows what its row's one-line explanation
claims (user, 2026-10-06: "your charts don't match the explanations, I can't follow this").

Same 30 rows as h108_review (dc_work/health108/review_rows108.csv).  For every row the statistic behind each finding
is RE-COMPUTED from the saved hi-res events with the package's own functions (health_core.health pass 1 / pass 2, the
same phase dicts, the same references), the explanation is written FROM those numbers, and the same numbers are drawn
on the chart (evidence panel per finding: limit as a dashed line with its value, the flagged value labelled, the
flagged period shaded, >= 3 h of context either side where data exist).  Every recomputed number is checked against
the resolver's saved value (dc_work/health108/v3c/auto_check.csv).  Below the evidence panels: every detector on the
predicted phase as its own line ('det 12 P4 Advance L1', flagged detector thick, mates' own status in the legend),
real 15-min counts and % ON, nothing scaled.

The user's entries (Answer / Comment / earlier comment) are carried cell by cell from the backup given as argument.

    python h108c_review.py <backup xlsx to read the user's entries from>
"""
from __future__ import annotations

import json
import os
import sys
from pathlib import Path

os.environ.setdefault("OMP_NUM_THREADS", "4")
HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import h104b_review as V  # noqa: E402
import h96_health as HH  # noqa: E402
from h108_review import KIND  # noqa: E402

import matplotlib  # noqa: E402
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from matplotlib.lines import Line2D  # noqa: E402
from matplotlib.patches import Patch  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

V3, RV = V.V3, V.RV
O = RV.O
hc = O.hc
DCW = O.DCW
H8 = DCW / "health108"
WK = H8 / "v3c"
REPO = HERE.parents[2]
XL = REPO / "review" / "health_review_v3.xlsx"
CH = REPO / "review" / "health_review_v3_charts"
DC = REPO / "review" / "health_review_v3_day_charts"
H1 = pd.Timedelta(hours=1)
M15 = pd.Timedelta(minutes=15)
FN = V.FN
THIS, INK, MUTED, MARK, LIMC, SHADE = "#1f4e9c", "#1a1a1a", "#555555", "#c62828", "#b8860b", "#fff1c2"
PAL = ["#eb6834", "#2a9d8f", "#a03ca0", "#7a7a75", "#d1495b", "#5c8001", "#8c564b", "#17becf", "#e377c2", "#3b3b98",
       "#9e9e2a", "#c49a00"]
ST = {"ok": "ok", "suspect": "suspect", "bad": "bad", "not_enough_data": "too little data", "watch": "watch"}
NAME = {"stuck": "stuck on", "dropout": "goes silent", "chatter": "chattering", "rapid": "too-fast actuations",
        "volume": "too many in 5 min", "level": "count drops", "choppy": "erratic counts",
        "night_drop": "misses vehicles at night", "night_day": "busier at night", "corr": "doesn't follow traffic",
        "occspk": "erratic time ON", "prof": "time-of-day profile", "occ_hi": "ON longer than its kind",
        "no_yardstick": "no yardstick", "short_on": "too-short ONs"}
STUCK_BAD = 3600.0


# ------------------------------------------------------------------ small helpers
def fmt_t(t, day=False):
    t = pd.Timestamp(t)
    return t.strftime("%a %H:%M") if day else t.strftime("%H:%M")


def span_txt(a, b):
    a, b = pd.Timestamp(a), pd.Timestamp(b)
    e = "24:00" if b.hour == 0 and b.minute == 0 and b > a.normalize() else b.strftime("%H:%M")
    return f"{a:%a} {a:%H:%M}-{e}"


def mins(s):
    s = float(s)
    if s >= 3600:
        return f"{s / 3600:.1f} h"
    return f"{s / 60:.0f} min" if s >= 60 else f"{s:.0f} s"


def pct(x, nd=0):
    return f"{100 * float(x):.{nd}f} %"


def tname(fn, span):
    return f"{FN.get(fn, fn)} {'1-lane' if span == '1' else '2+-lane'}"


def hour_ranges(hs):
    hs = sorted(int(h) for h in hs)
    out, i = [], 0
    while i < len(hs):
        j = i
        while j + 1 < len(hs) and hs[j + 1] == hs[j] + 1:
            j += 1
        out.append(f"{hs[i]:02d}h" if i == j else f"{hs[i]:02d}-{hs[j]:02d}h")
        i = j + 1
    return ", ".join(out)


def lane_short(s):
    if not s:
        return ""
    k = s.split(",")
    return "L" + k[0] if len(k) == 1 else f"L{k[0]}-{k[-1]}"


def lane_long(s):
    if not s:
        return ""
    k = s.split(",")
    return f"lane {k[0]}" if len(k) == 1 else f"covers lanes {k[0]}-{k[-1]}"


def ticks(x):
    return np.rint(np.asarray(x, float) * 10.0)


def gaps_of(ev, t0, t1):
    """comms gaps (> 120 s without any event) inside [t0, t1): array of gap starts (s from t0) and ends."""
    e = ev[(ev.Timestamp >= t0) & (ev.Timestamp < t1)]
    ut = np.unique((e.Timestamp - t0).dt.total_seconds().to_numpy())
    gi = np.diff(ut) > hc.GAP_S
    return ut[:-1][gi], ut[1:][gi]


def det_events(ev, d, t0, t1):
    x = ev[ev.EventId.isin([81, 82]) & (ev.Parameter == d) & (ev.Timestamp >= t0) & (ev.Timestamp < t1)]
    x = x.assign(o=(x.EventId != 82).astype(int)).sort_values(["Timestamp", "o"])
    return (x.Timestamp - t0).dt.total_seconds().to_numpy(), x.EventId.to_numpy().astype(int)


# ------------------------------------------------------------------ package re-run (exact statistics + references)
def package(ev, w, dev):
    s, h = O.WIN[w]
    t0 = pd.Timestamp(s)
    t1 = t0 + pd.Timedelta(hours=h)
    ph, fn, ln, pc = HH.inputs(w.split("_")[0], dev)
    x1 = hc.health(ev, t0, t1, None, ph or None, fn or None, ln or None, pc or None)
    bad = set(x1.loc[x1.status == "bad", "detector"].astype(int))
    ph2 = {k: (np.nan if k in bad else v) for k, v in ph.items()}
    x2 = hc.health(ev, t0, t1, None, ph2, fn or None, ln or None, pc or None) if bad and ph else x1
    return x1, x2, ph, ph2, bad


def refs_for(B, phd, d):
    n, cov = B["n_on"].astype(float), B["cov"]
    dets = B["dets"]
    live = n[:, cov].sum(1) > 0
    phase = {int(k): v for k, v in phd.items() if v is not None and np.isfinite(v)}
    tw = hc.twins(n, cov)
    ref, kind, _ = hc._refs(dets, live, phase, tw)
    i = int(np.where(dets == d)[0][0])
    twins = [int(dets[j]) for j in np.where(tw[i])[0]]
    return i, ref[i], kind[i], twins


# ------------------------------------------------------------------ evidence builders
class Row:
    pass


def ev_stuck(c):
    r = c.r
    st, du, mo = V3.ons_pkg(c.ev, c.d, c.z0, c.z1)
    en = st + pd.to_timedelta(du, unit="s")
    ins = np.asarray((st >= c.t0) & (st < c.t1))
    lim = float(r.stuck_lim8)
    badv = max(STUCK_BAD, 2 * lim)
    j = int(np.argmax(np.where(ins, du, -1)))
    longest = float(du[j])
    c.auto.append(("stuck: longest ON (s)", longest, float(r.stuck_x)))
    over = ins & (du >= lim)
    order = np.argsort(-du * over)[: int(over.sum())]
    tn = tname(r.fn, r.span)
    lab = [f"{du[k] / 60:.0f} min" for k in order]
    w = (f"longest ON {longest / 60:.0f} min ({span_txt(st[j], en[j])}) vs limit for {tn} {lim / 60:.0f} min "
         f"(bad at {badv / 60:.0f} min)")
    if over.sum() > 1:
        w += f"; {int(over.sum())} ONs over {lim / 60:.0f} min in the sample ({', '.join(lab[:4])}{', ...' if len(lab) > 4 else ''}) - the longest decides"
    if float(r.s8_stuck) == 0.35 and longest >= lim and longest < badv * 10:
        co, rec = float(r.ep_co) if pd.notna(r.ep_co) else np.nan, float(r.ep_rec) if pd.notna(r.ep_rec) else np.nan
        if np.isfinite(co) and co >= hc.CO_STUCK_N:
            w += f"; held at suspect: {int(co)} other detectors were held ON at the same time"
        elif np.isfinite(rec):
            w += f"; held at suspect: afterwards it counted {rec:.0%} of its earlier share (usable again)"
    hold = w.split("; held at suspect: ")[1] if "; held at suspect: " in w else ""
    # panel: one stem per continuous ON >= 1 min
    big = du >= 60
    pan = dict(kind="stems", ylabel="length of each ON (minutes)", st=st[big], du=du[big] / 60, ins=ins[big],
               lim=lim / 60, badv=badv / 60, tn=tn,
               labels=[(st[k], du[k] / 60, f"{du[k] / 60:.0f} min\n{fmt_t(st[k])}-{fmt_t(en[k])}") for k in order[:6]],
               title=f"Stuck on: longest ON {longest / 60:.0f} min; limit {lim / 60:.0f} min, bad {badv / 60:.0f} min"
               + (f"; held at suspect: {hold}" if hold else ""))
    c.panels.append(pan)
    c.shade += [(st[k], en[k]) for k in order[:6]] if over.any() else [(st[j], en[j])]
    c.slab.append("the long ON(s)")
    c.why.append(w)
    c.claims += [("longest ON", f"{longest / 60:.0f} min"), ("limit", f"{lim / 60:.0f} min")]
    c.ep = (st[j], en[j])


def queue_context(c, a, b):
    """Q1 numbers from the note-104 bins (bs 15 min): healthy phase mates' mean % ON during the long ON vs their
    sample mean, and the 15-min correlation of own % ON with theirs outside it."""
    r = c.r
    Bc = c.bctx
    bs = int(Bc.bs.iloc[0])
    ia = int((a - c.t0).total_seconds() // bs)
    iz = int(np.ceil((b - c.t0).total_seconds() / bs))
    own = Bc[Bc.detector == c.d].set_index("b").sort_index()
    inep = own.index.to_series().between(ia, iz - 1).to_numpy()
    pp = Bc[(Bc.ph == r.phase) & (Bc.detector != c.d) & Bc.status.ne("bad")]
    keep = []
    for k, g in pp.groupby("detector"):
        ge = g[g.b.between(ia, iz - 1)]
        if len(ge) and (ge.occ >= .99).mean() >= .8:
            continue
        keep.append(int(k))
    pp = pp[pp.detector.isin(keep)]
    if not len(pp):
        return None
    pm = pp.groupby("b").occ.mean().reindex(own.index).to_numpy(float)
    during, usual = np.nanmean(pm[inep]), np.nanmean(pm)
    o = own.occ.to_numpy(float)
    m = ~inep & np.isfinite(o) & np.isfinite(pm)
    rr = float(np.corrcoef(o[m], pm[m])[0, 1]) if m.sum() >= 4 else np.nan
    tt = c.t0 + pd.to_timedelta(own.index.to_numpy() * bs + bs / 2, unit="s")
    c.auto.append(("Q1: mates during / usual", during / max(usual, 1e-9), float(r.ep_phx_h)))
    c.auto.append(("Q1: r own vs mates", rr, float(r.c_occ_hpeer)))
    return dict(mates=keep, during=during, usual=usual, x=during / max(usual, 1e-9), r=rr, tt=tt, pm=100 * pm)


def ev_queue_q1(c, verdict):
    r = c.r
    if not (pd.notna(r.ep_t0) and pd.notna(r.ep_t1)):
        if verdict != "ok":
            c.why.append("queue test not run: it needs one ON of 15 min or more")
        return
    a, b = pd.Timestamp(r.ep_t0), pd.Timestamp(r.ep_t1)
    q = queue_context(c, a, b)
    if q is None:
        c.why.append("queue test: no healthy phase mate")
        return
    ok_x, ok_r = q["x"] >= 1.5, np.isfinite(q["r"]) and q["r"] >= 0.70
    ms = "+".join(map(str, q["mates"]))
    txt = (f"healthy phase mate(s) det {ms}: {100 * q['during']:.0f} % ON during it vs {100 * q['usual']:.0f} % "
           f"over the sample ({q['x']:.1f}x; a queue needs 1.5x); outside it det {c.d}'s % ON vs theirs r = "
           f"{q['r']:.2f} (needs 0.70)")
    if verdict == "ok":
        w = "Q1 queue: " + txt + " -> ok"
    else:
        ed = float(r.ep_dur) if pd.notna(r.ep_dur) else np.nan
        fail = ([] if ok_x else ["mates not busier"]) + ([] if ok_r else ["does not track them"]) +             (["longer than 60 min, the most a queue can clear"] if ed >= 3600 else [])
        if not fail and pd.notna(r.light_full) and r.light_full > 0:
            fail.append(f"held ON in {int(r.light_full)} light-traffic 15-min periods elsewhere")
        w = "not a queue (" + ", ".join(fail) + "): " + txt
    c.why.append(w)
    c.occ_extra.append(dict(lab=f"healthy phase mate(s) det {ms}: average % ON", tt=q["tt"], y=q["pm"], col=INK,
                            ls="--", lw=2.0))
    c.occ_hl.append((100 * q["usual"], f"mates' average over the sample: {100 * q['usual']:.0f} %", INK))
    c.occ_title.append(f"mates {100 * q['during']:.0f} % ON during the long ON vs {100 * q['usual']:.0f} % usual "
                       f"({q['x']:.1f}x, needs 1.5x); r = {q['r']:.2f} (needs 0.70)")
    c.claims += [("mates ratio", f"{q['x']:.1f}x"), ("r", f"{q['r']:.2f}")]


def ev_dropout(c, xr, refi, refk, pass_lab=""):
    r = c.r
    B = c.B5
    i = c.i5
    b0, b1, lam = int(xr.drop_b0), int(xr.drop_b1), float(xr.drop_lam)
    a = c.t0 + pd.Timedelta(seconds=300 * b0)
    b = c.t0 + pd.Timedelta(seconds=300 * b1)
    n = B["n_on"].astype(float)
    cov = B["cov"]
    S = n[refi].sum(0) if refk == "phase" and len(refi) >= 2 and xr.drop_ref == "phase" else n.sum(0) - n[i]
    who = ("det " + "+".join(str(int(B["dets"][k])) for k in refi)) if xr.drop_ref == "phase" else \
        "the rest of the signal"
    x = n[i]
    out = np.ones(len(x), bool)
    out[b0:b1] = False
    sh = x[out & cov].sum() / max(S[out & cov].sum(), 1)
    lam2 = sh * S[b0:b1][cov[b0:b1]].sum()
    c.auto.append(("silent: expected in run", lam2, lam))
    own = x[b0:b1][cov[b0:b1]].sum()
    # 15-min expected inside the run, for the counts panel
    e5 = np.where(np.arange(len(x)) >= b0, 1, 0) * np.where(np.arange(len(x)) < b1, 1, 0) * sh * S
    tt5 = c.t0 + pd.to_timedelta(np.arange(len(x)) * 300 + 150, unit="s")
    k = len(x) // 3 * 3
    e15 = e5[:k].reshape(-1, 3).sum(1)
    tt15 = c.t0 + pd.to_timedelta(np.arange(k // 3) * 900 + 450, unit="s")
    e15 = np.where(e15 > 0, e15, np.nan)
    c.cnt_extra.append(dict(lab=f"expected det {c.d} while silent = its share of {who} outside that stretch "
                                f"({100 * sh:.1f} %) x their counts", tt=tt15, y=e15, col=THIS, ls="--", lw=2.2))
    dur = (b - a).total_seconds()
    # zoomed evidence panel: 5-min counts around the silent stretch (+-2 h), expected per 5 min inside it
    za, zb = max(c.z0, a - 2 * H1), min(c.z1, b + 2 * H1)
    ixz, ttz, nz, _ = RV.bins(c.ev, za, zb, 300)
    ttz = pd.to_datetime(ttz)
    ez = pd.Series(np.where(e5 > 0, e5, np.nan), index=tt5).reindex(ttz).to_numpy()
    lab_e = f"{pass_lab}expected det {c.d} per 5 min while silent ({100 * sh:.1f} % of {who})"
    old = [q for q in c.panels if q.get("tag") == "drop"]
    if old:
        old[0]["extra"].append((lab_e, ttz, ez, "#7a7a75", 2.0, ":"))
        old[0]["title"] += f"; {pass_lab}~{lam:.0f} expected"
    else:
        c.panels.append(dict(kind="lines", tag="drop", z=(za, zb), shade_only=[(a, b)], ylabel="actuations per 5 min",
                             series=[(f"det {c.d}: actual", ttz, nz[ixz[c.d]], THIS, 2.6, "-")],
                             extra=[(lab_e, ttz, ez, THIS, 2.0, "--")],
                             title=(f"Goes silent (zoom, 5-min counts): {own:.0f} actuations {span_txt(a, b)}; "
                                    f"{pass_lab}~{lam:.0f} expected (suspect from 30, bad from 100)")))
    w = (f"{pass_lab}silent {span_txt(a, b)} ({mins(dur)}): {own:.0f} actuations where ~{lam:.0f} expected from "
         f"{who} (suspect from 30, bad from 100)")
    c.why.append(w)
    c.cnt_title.append(f"{pass_lab.replace(', judged on the rest of the signal', '')}silent {span_txt(a, b)}: "
                       f"{own:.0f} counted, ~{lam:.0f} expected")
    c.shade.append((a, b))
    c.slab.append("silent stretch")
    c.claims += [("silent span", span_txt(a, b)), ("expected", f"~{lam:.0f}")]
    return who


def ev_level(c, xr):
    """count drops: package det_stats(B) with ref None = the rest of the signal; 15-min bins."""
    B = c.B5
    i = c.i5
    n15 = hc._agg(B["n_on"].astype(float))
    cov15 = hc._agg(B["cov"][None].astype(float))[0] == hc.AGG
    x = n15[i]
    S = n15.sum(0) - x
    xs, Ss = x[cov15], S[cov15]
    idx = np.where(cov15)[0]
    cut = int(xr.level_b) // hc.AGG          # the package's change point counts covered 15-min bins only
    x1, s1, x2, s2 = xs[:cut].sum(), Ss[:cut].sum(), xs[cut:].sum(), Ss[cut:].sum()
    sh1, sh2 = (x1 + .5) / max(s1, 1), (x2 + .5) / max(s2, 1)
    ratio = sh2 / sh1
    c.auto.append(("level: after/before", ratio, float(xr.level_ratio)))
    cp = c.t0 + pd.Timedelta(minutes=15 * int(idx[cut]))
    tt = c.t0 + pd.to_timedelta(np.arange(len(x)) * 900 + 450, unit="s")
    share = np.where(cov15 & (S > 0), 100 * x / np.maximum(S, 1), np.nan)
    top_ = 3.0 * 100 * max(sh1, sh2)
    off = np.nanmax(share) > top_
    pan = dict(kind="lines", ylabel=f"det {c.d}'s actuations as % of\nall other detectors' on the signal",
               ylim=(0, top_) if off else None,
               series=[(f"det {c.d}: its 15-min count as % of the rest of the signal's"
                        + (f" (cut at {top_:.0f} %, peaks up to {np.nanmax(share):.0f} %)" if off else ""), tt, share,
                        THIS, 2.4, "-")],
               segs=[(c.t0, cp, 100 * sh1, f"before {fmt_t(cp)}: {100 * sh1:.1f} %"),
                     (cp, c.t1, 100 * sh2, f"after: {100 * sh2:.1f} %")],
               title=(f"Count drops: share after {fmt_t(cp)} = {100 * sh2:.1f} % vs {100 * sh1:.1f} % before = "
                      f"{ratio:.0%} of before (suspect below 15 %, bad below 5 %)"))
    c.panels.append(pan)
    w = (f"its count as % of all other detectors on the signal fell from {100 * sh1:.1f} % before {fmt_t(cp, True)} to "
         f"{100 * sh2:.1f} % after = {ratio:.0%} of before (suspect below 15 %)")
    m1, m2 = Ss[:cut].mean() if cut else np.nan, Ss[cut:].mean() if cut < len(Ss) else np.nan
    if np.isfinite(m1) and np.isfinite(m2) and m1 < 0.2 * m2 and sh2 > 0.02:
        w += ("; the 'before' is the quiet night, when its phase mates counted almost nothing while it kept counting "
              "- a busy-at-night pattern rather than a drop")
    c.why.append(w)
    c.shade.append((cp, c.t1))
    c.slab.append("after the drop")
    c.claims += [("before", f"{100 * sh1:.1f} %"), ("after", f"{100 * sh2:.1f} %"), ("ratio", f"{ratio:.0%}")]
    return sh1, sh2, cp


def ev_choppy(c, xr, lim, limname):
    """erratic counts: package rel_disp on 15-min bins, reference = phase mates (>= 2, no twins) else the signal."""
    B = c.B5
    i, refi, kind, twins = c.i5, c.refi, c.refk, c.twins
    a = hc._agg(B["n_on"].astype(float), 3)
    ok = hc._agg(B["cov"][None].astype(float), 3)[0] == 3
    x = np.where(ok, a[i], 0.0)
    rr = np.where(ok, a[refi].sum(0), 0.0)
    h = 8
    X = hc._movsum(x, h) - x
    Rr = hc._movsum(rr, h) - rr
    s = X / np.maximum(Rr, 1e-9)
    e, v = s * rr, s * rr * (1 + s)
    m = ok & (Rr > 0) & (X + x > 0) & ((e + x) > 0)
    D = ((x - e) ** 2)[m].sum() / max(v[m].sum(), 1e-9)
    c.auto.append(("erratic chop15", D, float(xr.chop15)))
    tt = c.t0 + pd.to_timedelta(np.arange(len(x)) * 900 + 450, unit="s")
    dets = B["dets"]
    if kind == "phase":
        who = "det " + "+".join(str(int(dets[k])) for k in refi)
        why_ref = ""
    else:
        who = "all other detectors on the signal"
        mates = [k for k in c.dets if k != c.d]
        tw_ = [t for t in twins if t in mates]
        why_ref = ((f" (det {' and '.join(map(str, tw_))} count the same as det {c.d} - same zone on another input - "
                    f"so they are no yardstick, and the check needs 2 other phase mates)") if tw_ else
                   " (fewer than 2 other phase mates to compare with)") if mates else ""
    con = np.where(m, (x - e) ** 2, 0.0)
    top = np.argsort(con)[::-1][:3]
    sd = np.sqrt(np.maximum(v, 0))
    pan = dict(kind="lines", ylabel="actuations per 15 min",
               series=[(f"det {c.d}: actual", tt, np.where(ok, x, np.nan), THIS, 2.6, "-")],
               extra=[(f"expected det {c.d} = its share of {who} over the 2 h around each 15 min", tt,
                       np.where(m, e, np.nan), THIS, 1.8, "--")],
               band=(tt, np.where(m, np.maximum(e - 2 * sd, 0), np.nan), np.where(m, e + 2 * sd, np.nan),
                     "chance range (expected +- 2 x the usual random variation)"),
               marks=[(tt[k], x[k], f"{x[k] - e[k]:+.0f}") for k in top if m[k]],
               mark_lab="the 3 periods that add most to the score (actual - expected)",
               title=(f"Erratic counts = {D:.1f} (how far actual strays from expected, in units of normal chance "
                      f"variation; ~1 = normal); {limname} {lim:.1f}"))
    c.panels.append(pan)
    c.claims += [("erratic", f"{D:.1f}"), ("limit", f"{lim:.1f}")]
    c.window_only = True
    return D, who, why_ref, [(tt[k] - M15 / 2, tt[k] + M15 / 2) for k in top if m[k]]


def ev_volume(c):
    r = c.r
    ix, tt, n, _ = RV.bins(c.ev, c.z0, c.z1, 300)
    tts = pd.to_datetime(tt)
    ins = np.asarray((tts >= c.t0) & (tts < c.t1))
    y = n[ix[c.d]]
    j = int(np.nanargmax(np.where(ins, y, -1)))
    mx = float(y[j])
    c.auto.append(("volume max5", mx, float(r.max5)))
    lim = float(r.lim_max5)
    a = tts[j] - pd.Timedelta(seconds=150)
    tn = tname(r.fn, r.span)
    c.use5 = True
    c.cnt_hl += [(lim, f"limit for {tn}: {lim:.0f} per 5 min", MARK), (150, "old limit (all detectors): 150", MUTED)]
    c.cnt_marks.append((tts[j], mx, f"{mx:.0f} at {fmt_t(a)}-{fmt_t(a + pd.Timedelta(minutes=5))}"))
    c.cnt_title.append(f"Too many in 5 min: {mx:.0f} in {fmt_t(a)}-{fmt_t(a + pd.Timedelta(minutes=5))}; limit for "
                       f"{tn} {lim:.0f} (bad from {5 / 3 * lim:.0f})")
    c.why.append(f"{mx:.0f} actuations in one 5 min ({span_txt(a, a + pd.Timedelta(minutes=5))}) vs limit for {tn} "
                 f"{lim:.0f} (old limit 150)")
    c.shade.append((a, a + pd.Timedelta(minutes=5)))
    c.slab.append("the busiest 5 min")
    c.claims += [("max 5 min", f"{mx:.0f}"), ("limit", f"{lim:.0f}")]


def on_flags(ev, d, t0, t1, g0, g1):
    """per ON event of d in [t0, t1): time (s), chatter flag (< 3 ticks after the previous OFF), ON->ON interval
    (s, NaN over a comms gap or for the first), burst flag."""
    t, e = det_events(ev, d, t0, t1)
    on = e == 82
    prev_e = np.r_[0, e[:-1]]
    prev_t = np.r_[np.nan, t[:-1]]
    gap = np.where(on & (prev_e == 81), t - prev_t, np.nan)[on]
    ton = t[on]

    def spans(a, b):
        if not len(g0):
            return np.zeros(len(a), bool)
        k = np.searchsorted(g0, a)
        kk = np.minimum(k, len(g0) - 1)
        return (k < len(g0)) & (g0[kk] < b)
    gap[spans(ton - np.nan_to_num(gap, nan=0.0), ton)] = np.nan
    chat = ticks(np.nan_to_num(gap, nan=99)) < hc.CHAT_TICKS
    ioi = np.r_[np.nan, np.diff(ton)]
    if len(ton) > 1:
        bad = np.r_[False, spans(ton[:-1], ton[1:])]
        ioi[bad] = np.nan
    jt = ticks(np.nan_to_num(ioi, nan=99.0))
    fast = jt < 10
    # bursts: >= 4 successive intervals < 1 s; the ONs in such runs
    burst = np.zeros(len(ton), bool)
    k = 1
    while k < len(ton):
        if fast[k] and np.isfinite(ioi[k]):
            j = k
            while j + 1 < len(ton) and fast[j + 1] and np.isfinite(ioi[j + 1]):
                j += 1
            if j - k + 1 >= 4:
                burst[k - 1:j + 1] = True
            k = j + 1
        else:
            k += 1
    return ton, chat, ioi, jt, burst


def hourly(c, d, vals, ton):
    """per clock hour of the chart window: share (%) of ONs with flag; NaN with < 10 ONs."""
    hrs = pd.date_range(c.z0.floor("h"), c.z1, freq="h")
    hh = c.z0 + pd.to_timedelta(ton, unit="s")
    k = np.searchsorted(hrs.to_numpy(), hh.to_numpy(), side="right") - 1
    num = np.bincount(k, weights=vals.astype(float), minlength=len(hrs))[:len(hrs)]
    den = np.bincount(k, minlength=len(hrs))[:len(hrs)]
    return hrs, np.where(den >= 10, 100 * num / np.maximum(den, 1), np.nan)


def ev_fast(c, what):
    """chatter / rapid: whole-sample share (package definition, ticks) + per-hour bars over the chart window."""
    r = c.r
    g0, g1 = gaps_of(c.ev, c.t0, c.t1)
    ton, chat, ioi, jt, burst = on_flags(c.ev, c.d, c.t0, c.t1, g0, g1)
    tn = tname(r.fn, r.span)
    zg0, zg1 = gaps_of(c.ev, c.z0, c.z1)
    zton, zchat, zioi, zjt, zburst = on_flags(c.ev, c.d, c.z0, c.z1, zg0, zg1)
    if what == "chatter":
        v = float(chat.sum() / max(len(ton), 1))
        c.auto.append(("chatter share", v, float(r.chat_frac)))
        lim = float(r.lim_chat_frac)
        hrs, y = hourly(c, c.d, zchat, zton)
        mates = []
        for k in c.dets:
            if k == c.d:
                continue
            a = on_flags(c.ev, k, c.z0, c.z1, zg0, zg1)
            if len(a[0]) >= 50:
                mates.append((k, hourly(c, k, a[1], a[0])[1]))
        pan = dict(kind="bars", ylabel="% of its ONs that start < 0.3 s\nafter the previous OFF", hrs=hrs, y=y,
                   mates=mates, hl=[(100 * lim, f"limit for {tn}: {100 * lim:.1f} %", MARK)],
                   title=(f"Chattering: {100 * v:.1f} % of its {len(ton):,} ONs in the sample re-trigger < 0.3 s after "
                          f"the OFF; limit for {tn} {100 * lim:.1f} % (bad from {200 * lim:.0f} %)"),
                   barlab=f"det {c.d}, per hour")
        c.panels.append(pan)
        c.why.append(f"{100 * v:.1f} % of its ONs re-trigger < 0.3 s after the previous OFF vs limit for {tn} "
                     f"{100 * lim:.1f} %")
        c.claims += [("chatter", f"{100 * v:.1f} %"), ("limit", f"{100 * lim:.1f} %")]
        return
    f = np.isfinite(ioi)
    vals = {"ioi_lt05": float((jt[f] < 5).mean()), "ioi_lt1": float((jt[f] < 10).mean()),
            "burst_frac": float(burst.sum() / max(len(ton), 1))}
    lims = {k: float(r[f"lim_{k}"]) for k in vals}
    for k in vals:
        c.auto.append((f"rapid {k}", vals[k], float(r[k])))
    rat = {k: vals[k] / lims[k] for k in vals}
    drv = max(rat, key=rat.get)
    show = [k for k in ("ioi_lt1", "ioi_lt05", "burst_frac") if k == drv or rat[k] >= 1]
    desc = {"ioi_lt1": "start < 1 s after the previous ON", "ioi_lt05": "start < 0.5 s after the previous ON",
            "burst_frac": "are in bursts (5+ ONs each < 1 s apart)"}
    flags = {"ioi_lt1": np.r_[False, zjt[1:] < 10] & np.isfinite(zioi), "ioi_lt05": (zjt < 5) & np.isfinite(zioi),
             "burst_frac": zburst}
    cols = {"ioi_lt1": THIS, "ioi_lt05": "#2a9d8f", "burst_frac": "#a03ca0"}
    lines, hl, parts = [], [], []
    for k in show:
        hrs, y = hourly(c, c.d, flags[k], zton)
        lines.append((f"det {c.d}: % of ONs that {desc[k]}, per hour", hrs, y, cols[k]))
        hl.append((100 * lims[k], f"limit ({desc[k].split(' (')[0]}): {100 * lims[k]:.1f} %", cols[k]))
        parts.append(f"{100 * vals[k]:.1f} % of ONs {desc[k]} (limit {100 * lims[k]:.1f} %)")
    pan = dict(kind="hlines", ylabel="% of its ONs", lines=lines, hl=hl,
               title=f"Too-fast actuations, whole sample (limits for {tn}):\n" + "; ".join(parts))
    c.panels.append(pan)
    c.why.append(f"too fast (limits for {tn}): " + "; ".join(parts))
    c.claims += [(k, f"{100 * vals[k]:.1f} %") for k in show]


def ev_night(c, xr):
    B = c.B5
    i = c.i5
    n = B["n_on"].astype(float)
    cov = B["cov"]
    hr = B["hour"]
    nt = cov & ((hr >= 21) | (hr < 5))
    dy = cov & (hr >= 7) & (hr < 19)
    ref = str(xr.night_ref)
    dets = list(map(int, B["dets"]))
    if ref.startswith("d") and "which it tracks" in ref:
        p = int(ref.split(",")[0][1:])
        S = n[dets.index(p)]
        who = f"det {p}, which it tracks"
    elif ref.startswith("the other detectors"):
        S = n[c.refi].sum(0)
        who = "the other detectors on its phase"
    else:
        S = n.sum(0) - n[i]
        who = "the rest of the signal"
    sh = n[i, dy].sum() / max(S[dy].sum(), 1)
    exp_n = sh * S[nt].sum()
    xn = n[i, nt].sum()
    c.auto.append(("night expected", exp_n, float(xr.night_exp)))
    c.auto.append(("night counted", xn, float(xr.night_n)))
    k = len(S) // 3 * 3
    e15 = (sh * S)[:k].reshape(-1, 3).sum(1)
    tt15 = c.t0 + pd.to_timedelta(np.arange(k // 3) * 900 + 450, unit="s")
    c.cnt_extra.append(dict(lab=f"expected det {c.d} = {who} x det {c.d}'s daytime (07-19) share ({100 * sh:.0f} %)",
                            tt=tt15, y=e15, col=THIS, ls="--", lw=2.2))
    ratio = (xn + .5) / (exp_n + .5)
    c.cnt_title.append(f"Night 21:00-05:00: {xn:.0f} counted vs ~{exp_n:.0f} expected = {ratio:.2f} "
                       f"(suspect below 0.19, bad below 0.063)")
    day = c.t0.normalize()
    c.shade += [(day, day + 5 * H1), (day + 21 * H1, day + 24 * H1)]
    c.slab.append("night 21:00-05:00")
    c.why.append(f"night 21:00-05:00: {xn:.0f} actuations where ~{exp_n:.0f} expected from {who} at its daytime share "
                 f"({ratio:.2f}; suspect below 0.19)")
    c.claims += [("night counted", f"{xn:.0f}"), ("expected", f"~{exp_n:.0f}")]


def ev_n3(c):
    r = c.r
    S = pd.read_parquet(DCW / "health104" / "n3_spike_bins.parquet", filters=[("DeviceId", "==", c.dev),
                                                                            ("window", "==", c.w)])
    S = S[S.detector == c.d]
    bs = 900
    dbar = float(r.dbar)
    exc = float(S.exc.sum())
    c.auto.append(("N3 minutes", exc, float(r.n3_exc)))
    lim = float(r.lim_n3_exc)
    tn = tname(r.fn, r.span)
    ix, tt, n, occ = c.b15
    e = np.minimum(n[ix[c.d]] * dbar / bs * 100, 100)
    tms = c.t0 + pd.to_timedelta(S.b.to_numpy() * bs + bs / 2, unit="s")
    topb = set(S.sort_values("exc", ascending=False).head(5).b)
    marks = [(c.t0 + pd.Timedelta(seconds=int(b) * bs + bs / 2), 100 * o, f"+{x:.0f} min" if b in topb else "")
             for b, o, x in zip(S.b, S.occ, S.exc)]
    pan = dict(kind="lines", ylabel="% of each 15 min ON",
               series=[(f"det {c.d}: actual % ON", tt, occ[ix[c.d]], THIS, 2.6, "-")],
               extra=[(f"expected from its count x its usual ON length ({dbar:.1f} s)", tt, e, THIS, 1.8, "--")],
               marks=marks, mark_lab="15 min ON far longer than its count explains (extra minutes on the 5 largest)", ylim=(0, 102),
               title=(f"Erratic time ON: {exc:.0f} min ON beyond what its count explains in {len(S)} periods (marked) "
                      f"while its phase was not busy; limit for {tn} {lim:.0f} min"))
    c.panels.append(pan)
    c.shade += [(t - M15 / 2, t + M15 / 2) for t in tms]
    c.slab.append("ON far longer than its count explains")
    c.why.append(f"erratic time ON: {exc:.0f} min ON beyond its count x usual ON ({dbar:.1f} s) in {len(S)} 15-min "
                 f"periods, limit for {tn} {lim:.0f} min")
    c.claims += [("N3 minutes", f"{exc:.0f} min"), ("limit", f"{lim:.0f} min")]


def ev_d1(c):
    r = c.r
    st, du, mo = V3.ons_pkg(c.ev, c.d, c.t0, c.t1)
    en = st + pd.to_timedelta(du, unit="s")
    ok = np.asarray(en <= c.t1)
    rep = (mo > 0) & ok
    tot = float(du[ok].sum())
    rt = float(du[rep].sum())
    c.auto.append(("D1 rep time (s)", rt, float(r.rep_time_s)))
    c.auto.append(("D1 total ON (s)", tot, float(r.occ) * float(r.hours) * 3600))
    # 15-min minutes ON split
    ix, tt, n, occ = c.b15
    tts = pd.to_datetime(tt)
    edges = (tts - M15 / 2)
    cl = np.zeros(len(tts))
    rp = np.zeros(len(tts))
    st2, du2, mo2 = V3.ons_pkg(c.ev, c.d, c.z0, c.z1)
    for s, d_, m_ in zip(st2, du2, mo2):
        a, b = s, s + pd.Timedelta(seconds=float(d_))
        k0 = max(int((a - edges[0]) / M15), 0)
        k1 = min(int((b - edges[0]) / M15), len(tts) - 1)
        for k in range(k0, k1 + 1):
            ov = (min(b, edges[k] + M15) - max(a, edges[k])).total_seconds()
            if ov > 0:
                (rp if m_ > 0 else cl)[k] += ov / 60
    sh = rt / max(tot, 1e-9)
    pan = dict(kind="stack", ylabel="minutes ON per 15 min", tt=tts, a=cl, b=rp,
               alab=f"det {c.d}: ONs with a normal OFF", blab=f"det {c.d}: ONs logged again with no OFF between",
               title=(f"ON logged again without an OFF: {rt / 60:.0f} of its {tot / 60:.0f} min ON in the sample = "
                      f"{sh:.0%} (data note from 70 % and 10 min)"))
    c.panels.append(pan)
    c.why.append(f"data note: {rt / 60:.0f} of its {tot / 60:.0f} min ON ({sh:.0%}) are ONs logged again with no OFF "
                 f"between (extension-like; note from 70 %) - no status")
    c.claims += [("rep minutes", f"{rt / 60:.0f} of {tot / 60:.0f} min"), ("share", f"{sh:.0%}")]


def occ_hi_table():
    f = WK / "occhi_lim.parquet"
    if f.exists():
        return pd.read_parquet(f)
    Bn = pd.read_parquet(DCW / "health104" / "bins_ctx.parquet", columns=["DeviceId", "window", "detector", "fn",
                                                                         "b", "occ", "ref", "cong", "bs", "ph",
                                                                         "status"])
    X = pd.read_parquet(H8 / "resolved108_q998p995.parquet", columns=["DeviceId", "window", "detector", "cmode"])
    Bn = Bn.merge(X, on=["DeviceId", "window", "detector"])
    Bn = Bn[Bn.cmode.ne("unknown")]
    Bn["rmax"] = Bn.groupby(["DeviceId", "window", "detector"]).ref.transform("max")
    Bn["tl"] = np.clip(np.floor(5 * Bn.ref / Bn.rmax.where(Bn.rmax > 0)), 0, 4)
    lim = Bn[Bn.status.eq("ok")].groupby(["fn", "cmode", "bs", "tl"]).occ.quantile(.995).rename("lim").reset_index()
    lim.to_parquet(f)
    return lim


def ev_n1(c):
    r = c.r
    L = occ_hi_table()
    Bn = c.bctx.merge(pd.DataFrame({"detector": [c.d], "cmode": [r.cmode]}), on="detector")
    Bn["rmax"] = Bn.ref.max()
    Bn["tl"] = np.clip(np.floor(5 * Bn.ref / Bn.rmax.where(Bn.rmax > 0)), 0, 4)
    Bn = Bn.merge(L, on=["fn", "cmode", "bs", "tl"], how="left").sort_values("b")
    allb = c.bctx
    sel = allb[(allb.ph == r.phase)].groupby("b").occ.agg(["sum", "count"]).rename(columns={"sum": "ps", "count": "pc"})
    Bn = Bn.merge(sel, left_on="b", right_index=True, how="left")
    Bn["pocc"] = ((Bn.ps - Bn.occ.fillna(0)) / (Bn.pc - 1).where(Bn.pc > 1))
    Bn["hi"] = (Bn.occ > Bn.lim) & (Bn.cong < 1.5) & ~(Bn.pocc >= 0.40)
    nh = int(Bn.hi.sum())
    bs = int(Bn.bs.iloc[0])
    c.auto.append(("N1 fires (>= 30 min over)", float(nh * bs >= 1800), float(bool(r.occ_hi8))))
    tt = c.t0 + pd.to_timedelta(Bn.b.to_numpy() * bs + bs / 2, unit="s")
    mode = f" ({r.cmode})" if r.fn == "Count" else ""
    hb = Bn[Bn.hi]
    marks = [(c.t0 + pd.Timedelta(seconds=int(b) * bs + bs / 2), 100 * o, "") for b, o in zip(hb.b, hb.occ)]
    pan = dict(kind="lines", ylabel="% of each 15 min ON",
               series=[(f"det {c.d}: actual % ON", tt, 100 * Bn.occ.to_numpy(), THIS, 2.6, "-")],
               steps=[(f"limit: what 1 in 200 healthy {FN.get(r.fn, r.fn)}{mode} zones exceed at that traffic level",
                       tt, 100 * Bn.lim.to_numpy(), MARK)],
               marks=marks, mark_lab="over the limit while its phase mates were < 40 % ON",
               ylim=(0, max(5.0, 1.25 * np.nanmax(np.r_[100 * Bn.occ.to_numpy(), 100 * Bn.lim.to_numpy()]))),
               title=(f"ON longer than its kind: over the limit in {nh} 15-min periods = {nh * bs / 60:.0f} min "
                      f"(watch from 30 min)"))
    c.panels.append(pan)
    c.shade += [(t - M15 / 2, t + M15 / 2) for t, _, _ in marks]
    c.slab.append("over the limit for its kind")
    c.why.append(f"ON longer than other {FN.get(r.fn, r.fn)}{mode} zones at that traffic in {nh} 15-min periods "
                 f"({nh * bs / 60:.0f} min; watch from 30 min) -> watch")
    c.claims += [("periods over", f"{nh}"), ("minutes", f"{nh * bs / 60:.0f} min")]


def like_context(c):
    """Q2 / like-with-like: own 15-min % ON vs the mean of healthy same-class phase mates (else all healthy mates)."""
    r = c.r
    Bc = c.bctx
    Bc = Bc[Bc.ph == r.phase]
    hl = Bc.status.ne("bad") & Bc.occ.notna()
    same = Bc[hl & (Bc.fn == r.fn) & (Bc.detector != c.d)]
    alls = Bc[hl & (Bc.detector != c.d)]
    own = Bc[Bc.detector == c.d].set_index("b").occ.sort_index()
    use, lab = (same, f"healthy {FN.get(r.fn, r.fn)} mates") if len(same) else (alls, "healthy phase mates")
    pm = use.groupby("b").occ.mean().reindex(own.index)
    m = own.notna() & pm.notna()
    rr = float(np.corrcoef(own[m], pm[m])[0, 1]) if m.sum() >= 6 else np.nan
    c.auto.append(("like r", rr, float(r.c_occ_like)))
    bs = int(Bc.bs.iloc[0])
    tt = c.t0 + pd.to_timedelta(own.index.to_numpy() * bs + bs / 2, unit="s")
    ms = sorted(use.detector.unique().astype(int))
    c.occ_extra.append(dict(lab=f"{lab} det {'+'.join(map(str, ms))}: average % ON", tt=tt, y=100 * pm.to_numpy(),
                            col=INK, ls="--", lw=2.0))
    c.occ_title.append(f"det {c.d}'s % ON vs the average of {lab} det {'+'.join(map(str, ms))}: r = {rr:.2f}")
    return rr, lab, ms


def ev_queue_pattern(c):
    """Q2: busier half of the sample (phase traffic >= its median): slope of log count and log % ON vs log traffic."""
    Bn = pd.read_parquet(DCW / "health96" / "bins.parquet", filters=[("DeviceId", "==", c.dev), ("window", "==", c.w)])
    b = Bn[(Bn.detector == c.d) & Bn.n.notna() & Bn.ref.notna()].copy()
    hi = b.ref >= b.ref.median()
    out = {}
    for nm, y in (("cnt", np.log(b.n + 1.0)), ("occ", np.log(b.occ + 0.005))):
        xx = np.log(b.ref + 1.0)
        out[nm] = float(np.polyfit(xx[hi], y[hi], 1)[0])
    c.auto.append(("Q2 busier-half slope counts", out["cnt"], float(c.r.elhi_cnt)))
    c.auto.append(("Q2 busier-half slope time ON", out["occ"], float(c.r.elhi_occ)))
    I = c.I
    refd = [k for k in c.dets if k != c.d and I.fn.get(k) in ("Advance", "Count")]
    pan = dict(kind="scatter2", x=b.ref.to_numpy(), hi=hi.to_numpy(), y1=b.n.to_numpy(), y2=100 * b.occ.to_numpy(),
               s1=out["cnt"], s2=out["occ"], xlab=("phase traffic per 15 min (det " + "+".join(map(str, refd)) + ")") if refd else
               "traffic per 15 min (no Advance / Count on its phase: those of the whole signal)",
               title=(f"Queue pattern (busier half, filled dots): counts rise with traffic at slope {out['cnt']:.2f}, "
                      f"time ON at slope {out['occ']:.2f} - time ON grows faster = queue"
                      if out["occ"] > max(out["cnt"], 0) else
                      f"Busier half: counts slope {out['cnt']:.2f}, time ON slope {out['occ']:.2f} - no queue pattern"))
    c.panels.append(pan)
    c.claims += [("slope counts", f"{out['cnt']:.2f}"), ("slope time ON", f"{out['occ']:.2f}")]
    return out


PROF = None


def ev_prof(c, fired):
    global PROF
    if PROF is None:
        PROF = pd.read_parquet(H8 / "prof.parquet")
    r = c.r
    p = PROF[(PROF.DeviceId == c.dev) & (PROF.window == c.w) & (PROF.detector == c.d)].iloc[0]
    keys = p.ref_src.split("+")
    g = PROF[PROF.hl_prof]
    for k in keys:
        g = g[g[k] == p[k]]
    out = {}
    for nm, pre in (("cnt", "s"), ("occ", "o")):
        cols = [f"{pre}{h}" for h in range(24)]
        own = p[cols].to_numpy(float)
        med = g[cols].median().to_numpy()
        lo, hi = g[cols].quantile(.025).to_numpy(), g[cols].quantile(.975).to_numpy()
        dist = 0.5 * np.abs(own - med).sum()
        out[nm] = dict(own=own, med=med, lo=lo, hi=hi, d=dist, up=[h for h in range(24) if own[h] > hi[h]],
                       dn=[h for h in range(24) if own[h] < lo[h]])
    c.auto.append(("profile d_cnt", out["cnt"]["d"], float(r.d_cnt)))
    c.auto.append(("profile d_occ", out["occ"]["d"], float(r.d_occ)))
    lim = {"cnt": float(r.d_cnt_lim), "occ": float(r.d_occ_lim)}
    rat = {k: out[k]["d"] / lim[k] for k in out}
    drv = max(rat, key=rat.get)
    show = [k for k in ("cnt", "occ") if k == drv or rat[k] >= 1]
    vb = {"low": "low volume (< 20 per h)", "medium": "medium volume (20-100 per h)", "high": "high volume (> 100 per h)"}
    grp = (tname(r.fn, r.span) if "type" in keys else FN.get(r.fn, r.fn)) + ", " + vb[p.band]
    day = pd.Timestamp(c.t0).strftime("%a %d")
    parts = []
    for k in show:
        o = out[k]
        what = "actuations" if k == "cnt" else "time ON"
        hrs = []
        if o["up"]:
            hrs.append(f"above normal at {hour_ranges(o['up'])}")
        if o["dn"]:
            hrs.append(f"below at {hour_ranges(o['dn'])}")
        parts.append(f"{100 * o['d']:.0f} % of its day's {what} are in other hours than the normal day (limit "
                     f"{100 * lim[k]:.0f} %)" + ("; " + ", ".join(hrs) if hrs else "; no single hour outside the "
                                                                                   "normal range"))
        pan = dict(kind="hour", own=100 * o["own"], med=100 * o["med"], lo=100 * o["lo"], hi=100 * o["hi"],
                   up=o["up"], dn=o["dn"], what=what, grp=grp, n=len(g),
                   ylabel=f"% of the day's {what}\nin each hour",
                   title=(f"Time-of-day profile ({what}, {day}): {100 * o['d']:.0f} % of the day sits in other hours "
                          f"than the normal day of a {grp}; limit {100 * lim[k]:.0f} %"))
        c.panels.append(pan)
        c.claims.append((f"profile {what}", f"{100 * o['d']:.0f} %"))
    w = "profile: " + "; ".join(parts)
    if not fired:
        w += " -> not flagged"
    c.why.append(w)
    return out


# ------------------------------------------------------------------ one row
def build(i, rr, R, keep):
    c = Row()
    c.r = r = R[(R.DeviceId == rr.dev) & (R.window == rr.window) & (R.detector == rr.detector)].iloc[0]
    c.dev, c.w, c.d, c.kind, c.sig = rr.dev, rr.window, int(rr.detector), rr.kind, rr.signal
    c.ev = RV.events(c.dev)
    c.t0, c.t1, c.h = V3.wwin(r)
    cov0, cov1 = c.ev.Timestamp.min().floor("15min"), c.ev.Timestamp.max().ceil("15min")
    c.z0, c.z1 = max(cov0, c.t0 - 3 * H1), min(cov1, c.t1 + 3 * H1)
    x = R[(R.DeviceId == c.dev) & (R.window == c.w)].set_index("detector")
    c.I = I = pd.DataFrame({"phase": x.phase, "fn": x.fn, "status": x.st8, "status1": x.status})
    c.ln = V.lanes_of(r, c.ev, I, c.t0, c.t1)
    x1, x2, ph, ph2, bad1 = package(c.ev, c.w, c.dev)
    # the pass whose statistics the resolver stored (pass 2 unless the detector itself was bad in pass 1)
    use2 = c.d not in bad1 and x2 is not x1
    xp = (x2 if use2 else x1).set_index("detector").loc[c.d]
    c.xp, c.x1 = xp, x1.set_index("detector").loc[c.d]
    c.B5 = hc.events_to_bins(c.ev, c.t0, c.t1)
    c.i5, c.refi, c.refk, c.twins = refs_for(c.B5, ph2 if use2 else ph, c.d)
    c.b15 = RV.bins(c.ev, c.z0, c.z1, 900)
    ix = c.b15[0]
    c.dets = V.phase_dets(I, c.d, ix)
    c.bctx = pd.read_parquet(DCW / "health104" / "bins_ctx.parquet", filters=[("DeviceId", "==", c.dev),
                                                                            ("window", "==", c.w)])
    c.panels, c.why, c.claims, c.auto, c.shade, c.slab = [], [], [], [], [], []
    c.cnt_extra, c.cnt_hl, c.cnt_marks, c.cnt_title = [], [], [], []
    c.occ_extra, c.occ_hl, c.occ_title = [], [], []
    c.use5, c.window_only, c.ep, c.reply = False, False, None, ""
    left = [k for k in str(r.left8).split(",") if k and k != "nan"]
    watch = [k for k in str(r.watch8).split(",") if k and k != "nan"]
    cleared = [k for k in str(r.cleared8).split(",") if k and k != "nan"]
    rules = [k for k in str(r.rules8).split(",") if k and k != "nan"]
    s8 = {k: float(r[f"s8_{k}"]) if pd.notna(r.get(f"s8_{k}")) else 0.0 for k in
          ("stuck", "dropout", "chatter", "rapid", "volume", "level", "choppy", "night_drop", "occspk", "prof")}
    tn = tname(r.fn, r.span)
    k = c.kind
    # ---------- per row: evidence for every finding that decides the new status (strongest first)
    order = sorted(left, key=lambda q: -s8.get(q, 0))
    if k == "Y":
        x1r = c.x1
        i_, ref1, k1, _ = refs_for(c.B5, ph, c.d)
        who = ev_dropout(c, x1r, ref1, k1, pass_lab="first pass: ")
        mates = [int(c.B5["dets"][q]) for q in ref1]
        c.cnt_extra[-1]["lab"] = "first pass: " + c.cnt_extra[-1]["lab"]
        n_ = len(c.why)
        ev_dropout(c, c.xp, c.refi, "signal", pass_lab="second pass, judged on the rest of the signal: ")
        c.cnt_extra[-1].update(ls=":", col="#7a7a75", lab="second pass: " + c.cnt_extra[-1]["lab"])
        w2 = c.why.pop()
        c.why.append(f"det {'+'.join(map(str, mates))} are all flagged bad themselves (dotted), so without them: "
                     + w2.split(": ", 1)[1].split(" (suspect")[0] + " - below 30, and no healthy detector is left "
                     "on its phase to confirm or clear it -> watch ('no yardstick')")
    if k in ("Q2u", "Q2"):
        if k == "Q2u":
            lim = float(r.lim_chop15)
            D, who, why_ref, top = ev_choppy(c, c.xp, lim, f"limit for {tn}")
            c.why.append(f"erratic counts {D:.1f} vs limit for {tn} {lim:.1f} (expected = its share of {who}) -> "
                         f"not flagged")
            c.shade += top
            c.slab.append("the 3 periods that add most")
            rr_, lab, ms = like_context(c)
            c.why.append(f"its % ON follows the average of {lab} det {'+'.join(map(str, ms))}: r = {rr_:.2f}")
        else:
            sh1, sh2, cp = ev_level(c, c.xp)
            q = ev_queue_pattern(c)
            rr_, lab, ms = like_context(c)
            c.why[-1:] = [f"count drop ({c.why[-1]}) cleared by Q2: in its busier half its time ON rises with traffic "
                          f"(slope {q['occ']:.2f}) faster than its counts (slope {q['cnt']:.2f}) = queue pattern, and "
                          f"its % ON tracks {lab} det {'+'.join(map(str, ms))} (r = {rr_:.2f}, needs 0.80) -> ok"]
    for f in order:
        if f == "stuck":
            ev_stuck(c)
            ev_queue_q1(c, "fail")
        elif f == "dropout":
            ev_dropout(c, c.xp, c.refi, c.refk)
        elif f == "level":
            ev_level(c, c.xp)
        elif f == "choppy":
            lim = float(r.lim_chop15)
            D, who, why_ref, top = ev_choppy(c, c.xp, lim, f"limit for {tn}")
            c.why.append(f"erratic counts {D:.1f} vs limit for {tn} {lim:.1f}; expected = its share of {who}"
                         + why_ref)
            c.shade += top
            c.slab.append("the 3 periods that add most")
        elif f == "volume":
            ev_volume(c)
        elif f in ("chatter", "rapid"):
            ev_fast(c, f)
        elif f == "night_drop":
            ev_night(c, c.xp)
        elif f == "occspk":
            ev_n3(c)
        elif f == "prof":
            ev_prof(c, True)
    if k == "Q1":
        ev_stuck(c)
        c.why = [c.why[0].split("; held at suspect")[0]]
        c.panels[-1]["title"] = c.panels[-1]["title"].split("; held at suspect")[0]
        ev_queue_q1(c, "ok")
    if "rapid" in watch and "R7" in rules:
        ev_fast(c, "rapid")
        c.why[-1] += " - just over and its only finding -> watch"
    if "occ_hi" in watch:
        ev_n1(c)
    if k in ("old_only",):
        ev_prof(c, False)
        c.why[-1] += f" (the old 'doesn't follow traffic' / 'busier at night' check had fired)"
    if "D1" in str(r.dq8):
        ev_d1(c)
    if not c.why:
        c.why.append("nothing over a limit")
    # ---------- chart
    who_ = f"{c.sig} det {c.d} (P{int(r.phase)} {FN.get(r.fn, r.fn)}" + (f" {lane_short(c.ln.get(c.d, ''))}"
                                                                           if c.ln.get(c.d) else "") + ")"
    new = ST.get(r.st8, r.st8)
    old = ST.get(r.status, r.status)
    title = f"{who_}: {old} -> {new}"
    png = CH / f"{i:02d}_{c.kind}_{c.sig}_d{c.d}.png"
    draw(c, png, title)
    dpng = DC / png.name
    day_chart(c, dpng, f"{who_}: all saved data (Sat 26 16:15 - Mon 28 24:00), 15-min counts and % ON")
    smp = (f"24 h {c.t0:%a %d}" if c.h >= 24 else f"3 h {c.t0:%a %d} {c.t0:%H:%M}-{c.t1:%H:%M}")
    det_txt = f"det {c.d} · P{int(r.phase)} {FN.get(r.fn, r.fn)}" + (f" · {lane_long(c.ln.get(c.d, ''))}"
                                                                     if c.ln.get(c.d) else "")
    return dict(n=i, signal=c.sig, detector=c.d, det=det_txt, check=f"{KIND[c.kind]} · {smp}",
                oldnew=f"{old_txt(r)} -> {new_txt(r)}", why="; ".join(c.why), png=png, dpng=dpng, dev=c.dev,
                window=c.w, kind=c.kind, claims=c.claims, auto=c.auto, title=title)


def old_txt(r):
    s = {k: float(r[f"s_{k}"]) for k in ("dropout", "stuck", "chatter", "rapid", "volume", "level", "choppy",
                                         "night_drop", "night_day", "corr") if pd.notna(r.get(f"s_{k}"))
         and float(r[f"s_{k}"]) >= .35}
    ks = sorted(s, key=lambda q: -s[q])
    return ST.get(r.status, r.status) + (f" ({', '.join(NAME[q] for q in ks)})" if ks and r.status in
                                         ("suspect", "bad") else "")


def new_txt(r):
    left = [w for w in str(r.left8).split(",") if w and w != "nan"]
    wat = [w for w in str(r.watch8).split(",") if w and w != "nan"]
    st = ST.get(r.st8, r.st8)
    if left and r.st8 in ("suspect", "bad"):
        st += " (" + ", ".join(NAME.get(k, k) for k in left) + ")"
    elif wat and r.st8 == "watch":
        st += " (" + ", ".join(NAME.get(k, k) for k in wat) + ")"
    if "D1" in str(r.dq8):
        st += " + data note"
    return st


# ------------------------------------------------------------------ drawing
def dlabel(c, k):
    I = c.I
    p = I.phase.get(k, np.nan)
    s = f"det {k} " + (f"P{int(p)} " if np.isfinite(p) else "") + FN.get(I.fn.get(k), "?")
    if c.ln.get(k):
        s += " " + lane_short(c.ln[k])
    st = I.status.get(k)
    if k == c.d:
        return s + "  <- this row"
    return s + (f"  [{st}]" if st in ("suspect", "bad", "watch") else "")


def series(c, tt, Y, ix):
    out = []
    j = 0
    for k in c.dets:
        if k not in ix:
            continue
        if k == c.d:
            out.append((dlabel(c, k), tt, Y[ix[k]], THIS, 3.0, "-"))
        else:
            st = c.I.status.get(k)
            out.append((dlabel(c, k), tt, Y[ix[k]], PAL[j % len(PAL)], 1.3, ":" if st == "bad" else
                        ("--" if st == "suspect" else "-")))
            j += 1
    return out


def time_axis(ax, z0, z1, show=True):
    ax.set_xlim(z0, z1)
    span = z1 - z0
    step = 1 if span <= pd.Timedelta(hours=10) else 3
    tk = pd.date_range(z0.ceil(f"{step}h"), z1, freq=f"{step}h")
    ax.set_xticks(tk)
    ax.set_xticklabels([(f"{t:%a}\n{t:%H:%M}" if (t.hour == 0 or j == 0) else f"{t:%H:%M}") for j, t in enumerate(tk)]
                       if show else [])
    ax.grid(axis="x", color="#eeeeee", lw=0.8)


def style(ax):
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    ax.grid(axis="y", color="#e6e6e6", lw=0.8)
    ax.tick_params(labelsize=10)


def shade(ax, c, z0, z1, only=None):
    for a, b in (c.shade if only is None else only):
        a, b = max(pd.Timestamp(a), z0), min(pd.Timestamp(b), z1)
        if b > a:
            ax.axvspan(a, b, color=SHADE, lw=0, zorder=0)
    for t in (c.t0, c.t1):
        if z0 < t < z1:
            ax.axvline(t, color=MUTED, ls=":", lw=1.4, zorder=1)


def leg(ax, hs, ls, title=None):
    ax.legend(hs, ls, frameon=False, loc="upper left", bbox_to_anchor=(1.005, 1.0), fontsize=9.3, title=title,
              title_fontsize=9.5, alignment="left")


def ptitle(ax, t):
    ax.set_title(t, loc="left", fontsize=11, color=INK, fontweight="bold", wrap=True)


def draw_lines_panel(ax, c, p):
    hs, ls_ = [], []
    if "band" in p:
        tt, lo, hi, lab = p["band"]
        ax.fill_between(pd.to_datetime(tt), lo, hi, color="#d9d9d9", lw=0, zorder=1)
        hs.append(Patch(color="#d9d9d9"))
        ls_.append(lab)
    for lab, tt, y, col, lw, ls in p.get("series", []):
        l_, = ax.plot(pd.to_datetime(tt), y, color=col, lw=lw, ls=ls, zorder=4 if lw > 2 else 3)
        hs.append(l_)
        ls_.append(lab)
    for lab, tt, y, col, lw, ls in p.get("extra", []):
        l_, = ax.plot(pd.to_datetime(tt), y, color=col, lw=lw, ls=ls, zorder=4)
        hs.append(l_)
        ls_.append(lab)
    for lab, tt, y, col in p.get("steps", []):
        l_, = ax.step(pd.to_datetime(tt), y, where="mid", color=col, lw=1.6, ls="--", zorder=3)
        hs.append(l_)
        ls_.append(lab)
    for a, b, y, lab in p.get("segs", []):
        ax.plot([a, b], [y, y], color=MARK, ls="--", lw=2, zorder=5)
        ax.text(a + (b - a) * 0.02, y, lab, color=MARK, fontsize=10.5, fontweight="bold", va="bottom", zorder=6)
    if p.get("segs"):
        hs.append(Line2D([], [], color=MARK, ls="--", lw=2))
        ls_.append("average share before / after the change point")
    if p.get("marks"):
        for t, y, lab in p["marks"]:
            ax.scatter([t], [y], s=110, facecolors="none", edgecolors=MARK, linewidths=2, zorder=6)
            if lab:
                ax.annotate(lab, (t, y), xytext=(4, 8), textcoords="offset points", color=MARK, fontsize=10,
                            fontweight="bold", zorder=7)
        hs.append(Line2D([], [], marker="o", ls="", mfc="none", mec=MARK, mew=2, ms=9))
        ls_.append(p.get("mark_lab", ""))
    if p.get("ylim"):
        ax.set_ylim(*p["ylim"])
    else:
        ax.set_ylim(bottom=0)
    leg(ax, hs, ls_)


def draw_det_panel(ax, c, kind, first):
    ix, tt, n, occ = (RV.bins(c.ev, c.z0, c.z1, 300) if (kind == "cnt" and c.use5) else c.b15)
    Y = n if kind == "cnt" else occ
    ser = series(c, tt, Y, ix)
    hs, ls_ = [], []
    for lab, t_, y, col, lw, ls in ser:
        l_, = ax.plot(pd.to_datetime(t_), y, color=col, lw=lw, ls=ls, zorder=4 if lw > 2 else 3)
        if first:
            hs.append(l_)
            ls_.append(lab)
    if not first:
        hs.append(Line2D([], [], color=MUTED, lw=1.3))
        ls_.append("same detectors and colors as the actuations panel")
    ex = c.cnt_extra if kind == "cnt" else c.occ_extra
    for e in ex:
        l_, = ax.plot(pd.to_datetime(e["tt"]), e["y"], color=e["col"], ls=e["ls"], lw=e["lw"], zorder=5)
        hs.append(l_)
        ls_.append(e["lab"])
    hl = c.cnt_hl if kind == "cnt" else c.occ_hl
    for q, (y, lab, col) in enumerate(hl):
        ax.axhline(y, color=col, ls="--", lw=1.6, zorder=2)
        ax.text(c.z0 + (c.z1 - c.z0) * (0.005 if q % 2 == 0 else 0.995), y, f" {lab} ", color=col, fontsize=9.5,
                va="bottom", ha="left" if q % 2 == 0 else "right", zorder=6)
        hs.append(Line2D([], [], color=col, ls="--", lw=1.6))
        ls_.append(lab)
    if kind == "cnt":
        for t, y, lab in c.cnt_marks:
            ax.scatter([t], [y], s=120, facecolors="none", edgecolors=MARK, linewidths=2.2, zorder=7)
            ax.annotate(lab, (t, y), xytext=(6, 6), textcoords="offset points", color=MARK, fontsize=10.5,
                        fontweight="bold", zorder=8)
    if first and c.shade:
        hs.append(Patch(color=SHADE))
        ls_.append(" / ".join(dict.fromkeys(c.slab)) or "flagged period")
    if first and (c.z0 < c.t0 or c.t1 < c.z1):
        hs.append(Line2D([], [], color=MUTED, ls=":", lw=1.4))
        ls_.append(f"sample checked: {span_txt(c.t0, c.t1)}")
    bw = "5 min" if (kind == "cnt" and c.use5) else "15 min"
    ax.set_ylabel(f"actuations per {bw}" if kind == "cnt" else "% of each 15 min ON", fontsize=10.5)
    if kind == "occ":
        ax.set_ylim(0, 102)
    else:
        ax.set_ylim(bottom=0)
    t = c.cnt_title if kind == "cnt" else c.occ_title
    ptitle(ax, ("; ".join(t)) if t else ("Actuations of every detector on the phase" if kind == "cnt" else
                                         "% of time ON of every detector on the phase"))
    leg(ax, hs, ls_)


def draw(c, path, title):
    pans = c.panels
    kinds = [p["kind"] for p in pans]
    hr = [1.0 if k not in ("hour", "scatter2") else 0.95 for k in kinds] + [1.15, 0.95]
    fh = 3.1 * sum(hr) + 0.9
    fig = plt.figure(figsize=(16, fh))
    gs = fig.add_gridspec(len(hr), 1, height_ratios=hr, hspace=0.62, top=1 - 0.75 / fh, bottom=0.35 / fh)
    axes_t = []
    for j, p in enumerate(pans):
        if p["kind"] == "scatter2":
            sg = gs[j].subgridspec(1, 2, wspace=0.25)
            for q, (yy, s, ylab) in enumerate(((p["y1"], p["s1"], "actuations per 15 min"),
                                              (p["y2"], p["s2"], "% of each 15 min ON"))):
                ax = fig.add_subplot(sg[q])
                x, hi = p["x"], p["hi"]
                ax.scatter(x[~hi], np.maximum(yy[~hi], 0.3), s=14, color="#bbbbbb", label="quieter half")
                ax.scatter(x[hi], np.maximum(yy[hi], 0.3), s=18, color=THIS, label="busier half")
                xh = np.log(x[hi] + 1.0)
                yl = np.log(yy[hi] + 1.0) if q == 0 else np.log(yy[hi] / 100 + 0.005)
                k_, b_ = np.polyfit(xh, yl, 1)
                xs = np.linspace(xh.min(), xh.max(), 20)
                ys = np.exp(k_ * xs + b_) - 1.0 if q == 0 else 100 * (np.exp(k_ * xs + b_) - 0.005)
                ax.plot(np.exp(xs) - 1, ys, color=MARK, lw=2, label=f"fit, busier half: slope {s:.2f}")
                ax.set_xscale("log")
                ax.set_yscale("log")
                ax.set_xlabel(p["xlab"], fontsize=10)
                ax.set_ylabel(ylab, fontsize=10.5)
                style(ax)
                ax.legend(frameon=False, fontsize=9.5, loc="upper left")
                if q == 0:
                    ptitle(ax, p["title"])
            continue
        ax = fig.add_subplot(gs[j])
        if p["kind"] == "hour":
            h = np.arange(24)
            ax.fill_between(h, p["lo"], p["hi"], color="#d9d9d9", lw=0, step=None)
            l1, = ax.plot(h, p["med"], color=MUTED, ls="--", lw=1.6)
            l2, = ax.plot(h, p["own"], color=THIS, lw=3, marker="o", ms=4)
            hs = [Patch(color="#d9d9d9"), l1, l2]
            ls_ = [f"normal range: 95 % of {p['n']} healthy {p['grp']} detectors, same day",
                   "median of those healthy detectors", f"det {c.d}"]
            for hh in p["up"] + p["dn"]:
                ax.scatter([hh], [p["own"][hh]], s=140, facecolors="none", edgecolors=MARK, linewidths=2.2, zorder=6)
                ax.axvspan(hh - .5, hh + .5, color=SHADE, lw=0, zorder=0)
            if p["up"] or p["dn"]:
                hs.append(Line2D([], [], marker="o", ls="", mfc="none", mec=MARK, mew=2, ms=9))
                ls_.append("hour outside the normal range")
            ax.set_xlim(-0.5, 23.5)
            ax.set_xticks(range(24))
            ax.set_xticklabels([f"{x:02d}" for x in range(24)])
            ax.set_xlabel("hour of the day (each point = that clock hour, e.g. 07 = 07:00-08:00)", fontsize=10)
            ax.set_ylim(bottom=0)
            leg(ax, hs, ls_)
        elif p["kind"] == "stems":
            shade(ax, c, c.z0, c.z1)
            st, du, ins = p["st"], p["du"], p["ins"]
            for s_, d_, i_ in zip(st, du, ins):
                ax.plot([s_, s_], [0, d_], color=THIS if i_ else "#9aa9c9", lw=2.2 if i_ else 1.4)
            ax.scatter(st[ins], du[ins], s=22, color=THIS, zorder=4)
            ax.scatter(st[~ins], du[~ins], s=16, color="#9aa9c9", zorder=4)
            ax.axhline(p["lim"], color=MARK, ls="--", lw=1.6)
            ax.axhline(p["badv"], color="#7b1fa2", ls="--", lw=1.6)
            ax.text(c.z0 + (c.z1 - c.z0) * 0.005, p["lim"], f" limit for {p['tn']}: {p['lim']:.0f} min", color=MARK,
                    fontsize=9.5, va="bottom")
            ax.text(c.z0 + (c.z1 - c.z0) * 0.005, p["badv"], f" bad: {p['badv']:.0f} min", color="#7b1fa2",
                    fontsize=9.5, va="bottom")
            for q, (s_, d_, lab) in enumerate(p["labels"][:4]):
                ax.annotate(lab, (s_, d_), xytext=(5, 2 + 22 * (q % 2)), textcoords="offset points", color=MARK,
                            fontsize=9.5, fontweight="bold")
            top = max(np.max(du) if len(du) else 1, p["badv"]) * 1.15
            ax.set_ylim(0, top)
            hs = [Line2D([], [], color=THIS, lw=2.2), Line2D([], [], color="#9aa9c9", lw=1.4),
                  Line2D([], [], color=MARK, ls="--"), Line2D([], [], color="#7b1fa2", ls="--")]
            ls_ = [f"det {c.d}: one stem per continuous ON of 1 min or more (in the sample)",
                   "same, outside the sample (context only)", "limit (suspect)", "bad"]
            leg(ax, hs, ls_)
            time_axis(ax, c.z0, c.z1)
        elif p["kind"] in ("bars", "hlines"):
            shade(ax, c, c.z0, c.z1)
            hs, ls_ = [], []
            if p["kind"] == "bars":
                hrs, y = p["hrs"], p["y"]
                ax.bar(hrs + pd.Timedelta(minutes=30), y, width=pd.Timedelta(minutes=50), color=THIS, alpha=.85,
                       zorder=3)
                hs.append(Patch(color=THIS))
                ls_.append(p["barlab"])
                for j2, (k, ym) in enumerate(p["mates"]):
                    l_, = ax.plot(hrs + pd.Timedelta(minutes=30), ym, color=PAL[j2 % len(PAL)], lw=1.2, zorder=4)
                    hs.append(l_)
                    ls_.append(f"{dlabel(c, k)}, per hour")
            else:
                for lab, hrs, y, col in p["lines"]:
                    l_, = ax.plot(hrs + pd.Timedelta(minutes=30), y, color=col, lw=2.4, marker="o", ms=3.5, zorder=4)
                    hs.append(l_)
                    ls_.append(lab)
            for q, (y, lab, col) in enumerate(p["hl"]):
                ax.axhline(y, color=col, ls="--", lw=1.6, zorder=2)
                ax.text(c.z0 + (c.z1 - c.z0) * (0.005 if q % 2 == 0 else 0.995), y, f" {lab} ", color=col,
                        fontsize=9.5, va="bottom", ha="left" if q % 2 == 0 else "right")
                hs.append(Line2D([], [], color=col, ls="--", lw=1.6))
                ls_.append(lab)
            ax.set_ylim(bottom=0)
            leg(ax, hs, ls_)
            time_axis(ax, c.z0, c.z1)
        elif p["kind"] == "stack":
            shade(ax, c, c.z0, c.z1)
            ax.bar(p["tt"], p["a"], width=pd.Timedelta(minutes=13), color=THIS, zorder=3)
            ax.bar(p["tt"], p["b"], bottom=p["a"], width=pd.Timedelta(minutes=13), color=MARK, zorder=3)
            ax.set_ylim(0, 15.5)
            leg(ax, [Patch(color=THIS), Patch(color=MARK)], [p["alab"], p["blab"]])
            time_axis(ax, c.z0, c.z1)
        else:
            za, zb = p.get("z", (c.z0, c.z1))
            shade(ax, c, za, zb, p.get("shade_only"))
            draw_lines_panel(ax, c, p)
            time_axis(ax, za, zb)
        ax.set_ylabel(p.get("ylabel", ""), fontsize=10.5)
        style(ax)
        ptitle(ax, p["title"])
        axes_t.append(ax)
    for j, kind in enumerate(("cnt", "occ")):
        ax = fig.add_subplot(gs[len(pans) + j])
        shade(ax, c, c.z0, c.z1)
        draw_det_panel(ax, c, kind, first=(j == 0))
        time_axis(ax, c.z0, c.z1)
        style(ax)
    fig.suptitle(title, x=0.01, y=1 - 0.12 / fh, ha="left", va="top", fontsize=15, fontweight="bold", color=INK)
    fig.savefig(path, dpi=88, bbox_inches="tight")
    plt.close(fig)


def day_chart(c, path, title):
    ev = c.ev
    z0, z1 = ev.Timestamp.min().floor("15min"), ev.Timestamp.max().ceil("15min")
    ix, tt, n, occ = RV.bins(ev, z0, z1, 900)
    fig, axes = plt.subplots(2, 1, figsize=(16, 8.4), sharex=True)
    for j, (ax, Y, yl) in enumerate(zip(axes, (n, occ), ("actuations per 15 min", "% of each 15 min ON"))):
        shade(ax, c, z0, z1)
        hs, ls_ = [], []
        for lab, t_, y, col, lw, ls in series(c, tt, Y, ix):
            l_, = ax.plot(pd.to_datetime(t_), y, color=col, lw=lw if lw > 2 else 1.1, ls=ls)
            hs.append(l_)
            ls_.append(lab)
        if j == 0:
            if c.shade:
                hs.append(Patch(color=SHADE))
                ls_.append(" / ".join(dict.fromkeys(c.slab)))
            hs.append(Line2D([], [], color=MUTED, ls=":", lw=1.4))
            ls_.append(f"sample checked: {span_txt(c.t0, c.t1)}")
            leg(ax, hs, ls_)
        ax.set_ylabel(yl, fontsize=10.5)
        time_axis(ax, z0, z1)
        style(ax)
        if j == 1:
            ax.set_ylim(0, 102)
    fig.suptitle(title, x=0.01, ha="left", fontsize=14, fontweight="bold", color=INK)
    fig.tight_layout()
    fig.savefig(path, dpi=85, bbox_inches="tight")
    plt.close(fig)


# ------------------------------------------------------------------ sheet
HEADER = [
    "Health spot-check v108 (rebuilt). Is the NEW result right? Answer Y / N / ? in the yellow column. Each row's "
    "'Why' line uses only numbers you can read on its chart.",
    "Data: hi-res log Sat 26 Sep 16:15 - Mon 28 Sep 24:00, plus the classifier's phase / function / lanes. Old = the "
    "current package. New = proposed rules whose limits come from what HEALTHY detectors of the same type do. Type = "
    "classifier function x lanes covered (1 lane, or 2+ lanes). Status: ok / watch (a note only) / suspect / bad.",
    "Limit rule: for each type, the value only 1 in 500 healthy detectors of that type go past (1 in 200 for the "
    "time-of-day profile). Past the limit = suspect; far past it (stated per check), or two different kinds of "
    "problem together = bad. A single finding only just past its limit = watch.",
    "Charts: top panel(s) = the evidence for each finding, with the limit as a dashed line and the flagged value "
    "labelled. Then every detector on the same phase as its own line, e.g. 'det 12 P4 Advance L1' (L1 = lane 1, the "
    "busiest lane of the phase; L1-2 = one detector covering lanes 1 and 2); this row's detector = thick dark blue; "
    "a mate flagged bad = dotted, suspect = dashed, with its status in [ ]. Yellow = the flagged period; dotted "
    "vertical lines = start / end of the sample checked; 3 h of context either side where data exist. Real units, "
    "nothing scaled.",
    "THE CHECKS (each judged on the sample in the Check column):",
    "Stuck on = one continuous ON (ON to the next OFF) longer than the limit for its type (5-15 min); bad at 60 min. "
    "Q1: cleared as a queue only if its HEALTHY phase mates were ON at least 1.5x their usual during it AND its % ON "
    "tracks theirs (r >= 0.70) outside it.",
    "Goes silent = a stretch with no actuations where its usual share of the other detectors' counts says it should "
    "have had some: suspect from 30 expected, bad from 100. Y (no yardstick): if every detector it was compared with "
    "is itself bad, it cannot be judged -> watch.",
    "Count drops = its share of the signal's actuations falls at some point and stays low: after / before below 15 % "
    "= suspect, below 5 % = bad. Misses vehicles at night = 21:00-05:00 count / count expected from the detector it "
    "tracks (at its daytime share) below 0.19 = suspect, below 0.063 = bad.",
    "Erratic counts = how far its 15-min counts stray from 'expected' (its own share of its phase mates', or the "
    "signal's, counts over the 2 h around), in units of normal random variation (about 1 for a healthy detector); "
    "limit per type, bad at 2x. Q2: cleared if it shows the queue pattern (in the busier half of the day its time ON "
    "grows faster with traffic than its counts) AND its % ON tracks its healthy same-function mates (r >= 0.80).",
    "Too many in 5 min = most actuations in one 5 min vs the limit for its type (old rule: 150 for all). Too-fast "
    "actuations = share of ONs that start < 1 s (or < 0.5 s) after the previous ON, or come in bursts of 5+ ONs each "
    "< 1 s apart; limit per type and lane count (2+ lanes allow side-by-side cars). Chattering = share of ONs that "
    "start < 0.3 s after the previous OFF; limit per type, bad at 2x.",
    "Erratic time ON = minutes of ON time beyond what its count x its usual ON length explains, in 15-min periods "
    "when its phase was not busy; limit per type. ON longer than its kind (watch only) = 30 min or more of 15-min "
    "periods ON longer than 1 in 200 healthy zones of its function reach at that traffic level, phase mates < 40 % ON.",
    "Time-of-day profile (24-h samples) = the share of its day's actuations (and time ON) in each clock hour vs "
    "healthy detectors of the same type and volume (low < 20 / medium 20-100 / high > 100 per hour), same day. Score "
    "= the % of its day that sits in other hours than the normal day; limit per type and volume. It replaces "
    "'doesn't follow traffic' and 'busier at night'.",
    "Data note (D1) = most of its ON time is ONs logged again with no OFF between: reported, no status change.",
]


def old_entries(path):
    wb = openpyxl.load_workbook(path)
    ws = wb["Cases"]
    out = {}
    h0 = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1) if row[0] == "#")
    hdr = [str(x) for x in next(ws.iter_rows(min_row=h0, max_row=h0, values_only=True))]
    ci = {h: j for j, h in enumerate(hdr)}
    for row in ws.iter_rows(min_row=h0 + 1, values_only=True):
        if row[0] is None:
            continue
        det = int(str(row[ci["Det"]]).split("·")[0].replace("det", "").strip())
        out[(str(row[ci["Signal"]]), det)] = dict(answer=row[ci["Answer (Y / N / ?)"]], comment=row[ci["Comment"]],
                                                  earlier=row[ci["Your earlier comment"]],
                                                  reply=row[ci["Our reply"]] if "Our reply" in ci else None)
    return out


def write_xl(rows, keep, replies):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cases"
    for j, t in enumerate(HEADER, start=1):
        c = ws.cell(row=j, column=1, value=t)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.font = Font(bold=(j in (1, 5)), size=12 if j == 1 else 10)
        ws.merge_cells(start_row=j, start_column=1, end_row=j, end_column=12)
        ws.row_dimensions[j].height = 16 * max(1, -(-len(t) // 230)) + 3
    h0 = len(HEADER) + 2
    head = ["#", "Signal", "Det", "Check", "Old -> New", "Why", "Chart", "Day chart", "Your earlier comment",
            "Answer (Y / N / ?)", "Comment", "Our reply"]
    for j, h in enumerate(head, start=1):
        c = ws.cell(row=h0, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2A78D6")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for k, r in enumerate(rows, start=h0 + 1):
        e = keep.get((r["signal"], r["detector"]), {})
        vals = [r["n"], r["signal"], r["det"], r["check"], r["oldnew"], r["why"], "chart", "day chart",
                e.get("earlier"), e.get("answer"), e.get("comment"), replies.get(r["n"], "")]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=k, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=k, column=7).hyperlink = f"{CH.name}/{r['png'].name}"
        ws.cell(row=k, column=8).hyperlink = f"{DC.name}/{r['dpng'].name}"
        for j in (7, 8):
            ws.cell(row=k, column=j).font = Font(color="0563C1", underline="single")
        ws.cell(row=k, column=9).font = Font(italic=True, color="555555")
        ws.cell(row=k, column=10).fill = PatternFill("solid", fgColor="FFF8DC")
        ws.cell(row=k, column=12).font = Font(color="1F4E9C")
    for col, w in zip("ABCDEFGHIJKL", (4, 9, 24, 20, 28, 70, 7, 9, 36, 11, 36, 44)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(row=h0 + 1, column=1)
    wb.save(XL)
    return h0


def main():
    src = Path(sys.argv[1])
    keep = old_entries(src)
    WK.mkdir(parents=True, exist_ok=True)
    HH.init()
    R = pd.read_parquet(H8 / "resolved108_q998p995.parquet")
    rows = pd.read_csv(H8 / "review_rows108.csv")
    only = [int(a) for a in sys.argv[2].split(",")] if len(sys.argv) > 2 else None
    for p_ in (CH, DC):
        p_.mkdir(parents=True, exist_ok=True)
        if only is None:
            for f in p_.glob("*.png"):
                f.unlink()
    out, auto, claims = [], [], []
    for _, rr in rows.iterrows():
        if only is not None and int(rr.n) not in only:
            continue
        o = build(int(rr.n), rr, R, keep)
        print(o["n"], o["signal"], o["det"], "|", o["oldnew"], "|", o["why"], flush=True)
        for nm, a, b in o["auto"]:
            ok = (np.isfinite(a) and np.isfinite(b) and abs(a - b) <= max(0.02 * abs(b), 0.011)) or \
                 (not np.isfinite(a) and not np.isfinite(b))
            auto.append(dict(row=o["n"], signal=o["signal"], det=o["detector"], stat=nm, chart=a, resolver=b,
                             match=bool(ok)))
            if not ok:
                print(f"   AUTO MISMATCH {nm}: chart {a} vs resolver {b}")
        for cl, val in o["claims"]:
            claims.append(dict(row=o["n"], signal=o["signal"], det=o["detector"], claim=cl, value=val))
        out.append(o)
    if only is not None:
        return
    pd.DataFrame(auto).to_csv(WK / "auto_check.csv", index=False)
    pd.DataFrame(claims).to_csv(WK / "claims.csv", index=False)
    pd.DataFrame([{k: v for k, v in o.items() if k not in ("claims", "auto")} for o in out]).to_csv(
        WK / "rows_v3c.csv", index=False)
    replies = json.loads((WK / "replies.json").read_text(encoding="utf-8")) if (WK / "replies.json").exists() else {}
    replies = {int(k): v for k, v in replies.items()}
    need = {k for k, v in keep.items() if any(v.get(q) not in (None, "") for q in ("answer", "comment", "earlier"))}
    have = {(o["signal"], o["detector"]) for o in out}
    print("user entries:", len(need), "carried:", len(need & have), "lost:", sorted(need - have))
    write_xl(out, keep, replies)
    # cell-by-cell check of the user's entries
    chk = old_entries(XL)
    bad = [(k, q) for k in need for q in ("answer", "comment", "earlier") if keep[k].get(q) != chk.get(k, {}).get(q)]
    print("entries verified cell by cell:", "ALL MATCH" if not bad else f"MISMATCH {bad}")


if __name__ == "__main__":
    main()
