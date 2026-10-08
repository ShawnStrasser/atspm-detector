"""Note 47: follow-up spot-check of health v5 (review/spotcheck_health.xlsx), <= 6 rows, 66-h Sept-2026 window,
training signals only (never locked_v2): the user's three v4 comments (10090 d22 no-OFF heuristic, 10018 d2 short
ONs at night, 07035 d19 suspect vs bad) and cases of the v5 changes.  Charts reuse the note-43 layout
(h3_spotcheck.chart, continuous-ON plotting).  The v4 sheet's answers move to "earlier answers" (dated).

    python h5_spotcheck.py
"""
from __future__ import annotations

import sys
import urllib.parse
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
import h3_spotcheck as S  # noqa: E402
import h4_spotcheck as S4  # noqa: E402
import health_core as hc  # noqa: E402

ROWS = [  # signal, detector, chart kind, what was found (plain words for the user)
    ("10018", 2, "silent",
     "YOUR ANSWER (clearly bad). Now BAD. New check: at night (21:00-05:00) it counted 14 actuations where about 185 "
     "were expected from d4, which it tracks by day (8 %; healthy detectors stay above 19 %), and the few it sees at "
     "night are short. On your minimum-ON idea: it is built (a non-pulse detector whose ONs of 0.2 s or less are too "
     "many, overall or in its worst 3-h block; a 15-ft car over a 6-ft loop in 0.2 s = 72 mph), but it does not "
     "separate d2: d4 also gets 0.3-s ONs at night (free-flow ~48 mph), what differs is that d2 misses about 9 of 10 "
     "night vehicles. Over the whole population the short-ON check found no problem the other checks missed and raised "
     "false alarms (+0.5 pt), so it is reported as a note, not scored."),
    ("10090", 22, "stuck",
     "YOUR COMMENT (false heuristic). Fixed everywhere: an ON followed by another ON with no OFF is now treated as "
     "normal (extension); stuck-on is judged only from how long the detector stayed ON continuously, from an ON to the "
     "next OFF. d22 stays BAD for that reason alone: ON without a break 08:32-10:08 on Saturday and Sunday (96 min "
     "each, ~314 actuations expected meanwhile) and 36 min on Monday, and afterwards only 35 % of its usual share vs d23. "
     "All 568 earlier calls that quoted a missing OFF were re-checked the same way (none quotes it now)."),
    ("07035", 19, "rapid",
     "YOUR QUESTION (why suspect, not bad?). It has one finding only: 11 % of its ONs start within 1 s of the previous "
     "one, 1.3x the level that only 1 in 500 healthy detectors like it reach; on its own a finding must be about 1.6x "
     "that level to make a detector bad. Nothing else is wrong: counts follow d16, no stuck-on, no silence, no night "
     "drop; 14 % short ONs is within the healthy range. Several findings do add up (each at its limit takes the score "
     "down a step; about four such, or two well past the limit, make it bad). A stricter rule (any two different kinds "
     "of finding = bad) was tested: no extra known problem caught, slightly more healthy detectors called bad, so not "
     "adopted. Stays SUSPECT."),
    ("10041", 41, "silent",
     "NEW (short-ON check, a question for you). Predicted Count. 63 % of its ONs are 0.1-0.2 s, the rest 0.3-1 s and "
     "longer, and between 21:00 and 06:00 every ON is 0.2 s or less. Other detectors on this signal look the same. The data "
     "cannot tell a detector set to pulse that sometimes reports longer ONs from a normal one that stops seeing whole "
     "vehicles at night. Now OK with a note. Is this detector (and its neighbours d35-d48) set to pulse?"),
    ("11042", 22, "stuck",
     "YOUR EARLIER REMARK (should be flagged). Now SUSPECT: with continuous ON measured from an ON to the next OFF it "
     "stayed ON 48 min (Sat 06:26-07:14) and 15-40 min five more times; in the longest, 10 other detectors were held ON at "
     "the same time - a shared event, so the periods are listed and the detector is kept (was ok)."),
    ("04017", 9, "stuck",
     "NEW (what the fix costs). A pulse-mode Count detector ON for 15 min (Sun 05:16-05:31) together with 3 others, "
     "at an hour when about 2 actuations were expected. v4 ignored it because another ON came before the OFF; v5 counts "
     "it as one continuous ON, so it is now SUSPECT with the period listed. Ignoring long ONs in quiet hours instead "
     "would drop 10-35 % of the real problems caught, so every continuous ON of 15 min or more now counts."),
]


def main():
    wb0 = openpyxl.load_workbook(S.XL)
    if wb0.worksheets[0].title == "health follow-up v5":
        raise SystemExit("already rebuilt for v5 - restore the earlier workbook first")
    v3 = pd.read_parquet(S.H.REPO / "research" / "labels" / "function_labels_v3.parquet", columns=["DeviceId", "DeviceName"])
    nm = dict(zip(v3.DeviceName, v3.DeviceId.str.lower()))
    out = []
    for sig, d, kind, what in ROWS:
        dev = nm[sig]
        ev, h, B, ph, fn, ln = S4.load(dev)
        x = h[h.detector == d].iloc[0].copy()
        x["signal"] = sig
        png = S.CH / f"v5_{sig}_d{d}.png"
        S.chart(x, ev, h, B, ph, fn, kind, png)
        pdfs = sorted(S.PDF.glob(f"{sig}_*.pdf"))
        p = ph.get(d)
        sp = ln.get(d)
        out.append(dict(signal=sig, detector=d, phase=int(p) if p == p and p is not None else None,
                        function=fn.get(d, "-"), lanes=("-" if sp != sp or sp is None else
                                                        f"spans {int(sp)}" if sp >= 2 else "1 lane"),
                        status=x.status, what=what, png=png.name, pdf=pdfs[0] if pdfs else None))
        print(sig, d, x.status, "|", x.reason[:220])
    # earlier answers: the v4 sheet (answered 2026-09-29) first, then the existing "earlier answers" sheet
    ws0, w1 = wb0.worksheets[0], wb0.worksheets[1]
    prev = []
    hdr = [c.value for c in ws0[1]]
    for r in range(2, ws0.max_row + 1):
        row = {h_: ws0.cell(r, j + 1).value for j, h_ in enumerate(hdr)}
        if row.get("Signal") is None:
            continue
        links = {h_: ws0.cell(r, j + 1).hyperlink.target for j, h_ in enumerate(hdr) if ws0.cell(r, j + 1).hyperlink}
        what = next((v for k, v in row.items() if str(k).startswith("What was found")), None)
        prev.append(([ "2026-09-29", row["Signal"], row["Detector"], row["Phase"], row["Health status"], what,
                       "open chart", "open print", row["Your answer"], row["Comment"]],
                     {7: links.get("Chart link"), 8: links.get("Print link")}))
    for r in range(2, w1.max_row + 1):
        vals = [w1.cell(r, j).value for j in range(1, 11)]
        if vals[1] is None:
            continue
        prev.append((vals, {7: w1.cell(r, 7).hyperlink.target if w1.cell(r, 7).hyperlink else None,
                            8: w1.cell(r, 8).hyperlink.target if w1.cell(r, 8).hyperlink else None}))
    wb = Workbook()
    ws = wb.active
    ws.title = "health follow-up v5"
    ws.append(["Signal", "Detector", "Phase", "Predicted function", "Lanes", "Health status",
               "What was found (log Fri 18 Sep 16:15 - Mon 21 Sep 10:25, 2026)", "Chart link", "Print link",
               "Your answer", "Comment"])
    link = Font(color="0563C1", underline="single")
    for n, r in enumerate(out, start=2):
        ws.append([r["signal"], r["detector"], r["phase"], r["function"], r["lanes"], r["status"], r["what"],
                   "open chart", "open print" if r["pdf"] is not None else "no print on file", "", ""])
        ws.cell(n, 8).hyperlink = f"spotcheck_health_charts/{r['png']}"
        ws.cell(n, 8).font = link
        if r["pdf"] is not None:
            ws.cell(n, 9).hyperlink = "file:///" + urllib.parse.quote(str(r["pdf"]), safe=":\\/()_-.,'")
            ws.cell(n, 9).font = link
    w2 = wb.create_sheet("earlier answers")
    w2.append(["Answered", "Signal", "Detector", "Phase", "Health status then", "What was found then", "Chart link",
               "Print link", "Your answer", "Comment"])
    for n, (vals, lk) in enumerate(prev, start=2):
        w2.append(vals)
        for c, tgt in lk.items():
            if tgt:
                w2.cell(n, c).hyperlink = tgt
                w2.cell(n, c).font = link
    for sh, widths in ((ws, (9, 9, 7, 12, 10, 12, 95, 11, 11, 14, 40)), (w2, (11, 9, 9, 7, 12, 80, 11, 11, 14, 60))):
        for j, w in enumerate(widths):
            sh.column_dimensions[chr(65 + j)].width = w
        for c in sh[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
        for row in sh.iter_rows(min_row=1):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"
    wb.save(S.XL)
    print("saved", S.XL, len(out), "rows;", len(prev), "earlier answers kept")


if __name__ == "__main__":
    _ = hc
    main()
