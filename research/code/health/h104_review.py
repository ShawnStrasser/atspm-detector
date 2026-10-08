"""Note 104: review sheet review/health_review_v3.xlsx (+ _v3_charts/, _v3_day_charts/).

Rows = (a) the health_review_v2 rows the user could not judge (plus v2 row 9, where he doubted the scaling, and v2
row 3, which he rejected and which the rehoned R1 now keeps flagged), with charts redrawn for the specific
confusion, and (b) 3-4 examples for each new rule idea he gave (h104_resolve: R1', R1b, N1', N3).
The user's v2 answer + comment are shown in "Your earlier comment".  review/health_review_v2.xlsx is only READ.

Chart rules (answers to the v2 comments): NO line is scaled - every number is a real count or a real % of time ON;
the orange line is ONE named detector (the best-matching healthy detector on the same phase) unless the legend says
otherwise; the cases that are about ON lengths get a third panel with the length of every ON; 'erratic' rows mark the
15-min periods the check objects to with red circles; the R9 row zooms on the silent stretch the old check saw.
Saved w40 events only (Sat 26 16:15 - Mon 28 24:00); hi-res log + classifier outputs; seed 104.

    python h104_review.py
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

os.environ.setdefault("OMP_NUM_THREADS", "4")
sys.path.insert(0, str(Path(__file__).resolve().parent))
import h96_review as RV  # noqa: E402
import h104_resolve as M  # noqa: E402

import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
import pandas as pd  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

C3, O, hc = RV.C3, RV.O, RV.hc
THIS, PART, INK, MUTED, MARK = C3.THIS, C3.PART, C3.INK, C3.MUTED, "#c62828"
FN, CHK = RV.FN, dict(RV.CHK, occspk="erratic time ON", occ_hi="ON longer than its kind",
                      night_rel="relatively busy at night")
REPO = M.REPO
XL = REPO / "review" / "health_review_v3.xlsx"
XL2 = REPO / "review" / "health_review_v2.xlsx"
CH = REPO / "review" / "health_review_v3_charts"
DC = REPO / "review" / "health_review_v3_day_charts"
OUT = M.OUT
ST = {"ok": "ok", "suspect": "suspect", "bad": "bad", "not_enough_data": "too little data", "watch": "watch"}
KEEP = [3, 4, 5, 6, 7, 9, 12, 15, 16, 18, 19]          # v2 rows shown again
H1 = pd.Timedelta(hours=1)


# ------------------------------------------------------------------ data
def load():
    R = pd.read_parquet(OUT / "resolved104.parquet")
    R96 = pd.read_parquet(M.IN96 / "resolved.parquet", columns=["DeviceId", "window", "detector", "new_status",
                                                                "rules", "watch"])
    return R.merge(R96, on=["DeviceId", "window", "detector"], suffixes=("", "_96"))


def v2_answers():
    wb = openpyxl.load_workbook(XL2, read_only=True)
    ws = wb["Cases"]
    a = {r[0]: (r[10], r[11]) for r in ws.iter_rows(min_row=4, values_only=True) if r[0] is not None}
    v = pd.read_csv(M.IN96 / "review_rows.csv")
    v["ans"] = v.n.map(lambda k: a[k][0])
    v["com"] = v.n.map(lambda k: a[k][1])
    return v


def has(s, k):
    return s.fillna("").str.split(",").apply(lambda l: k in l)


def select(R, v2):
    rng = np.random.default_rng(104)
    R = R.assign(rnd=rng.random(len(R)))
    seen = set(zip(pd.read_csv(M.DCW / "health82" / "rows.csv").dev.str.lower(),
                   pd.read_csv(M.DCW / "health82" / "rows.csv").det)) | set(zip(v2.dev, v2.detector))
    used = set(v2.dev)
    base = R[R.wg.isin(["h3", "h24"]) & ~pd.Series(list(zip(R.DeviceId, R.detector))).isin(seen).to_numpy()]
    rows = []
    for k in KEEP:
        x = v2[v2.n == k].iloc[0]
        r = R[(R.DeviceId == x.dev) & (R.window == x.window) & (R.detector == x.detector)].iloc[0].copy()
        r["kind"], r["v2n"], r["v2rule"] = ("I1" if k == 3 else "KEEP"), k, x.rule
        r["earlier"] = f"v2 row {k}: {str(x.ans).strip()}" + (f" - {str(x.com).strip()}" if pd.notna(x.com) else "")
        rows.append(r)

    def take(c, n, tag, key=None):
        got = 0
        for _, r in c.sort_values("rnd").iterrows():
            if r.DeviceId in used:
                continue
            if key is not None and any(q.get("sub") == key(r) for q in rows if q.kind == tag):
                continue
            r = r.copy()
            r["kind"], r["v2n"], r["v2rule"], r["earlier"] = tag, None, None, ""
            r["sub"] = key(r) if key else ""
            rows.append(r)
            used.add(r.DeviceId)
            got += 1
            if got == n:
                break

    # I1: rehoned R1 keeps a stuck flag that the note-96 R1 cleared - one per failing condition
    c = base[has(base.rules_96, "R1") & ~has(base.rules, "R1") & ~has(base.rules, "R1b")]

    def why1(r):
        if r.light_full > 0:
            return "light"
        if r.n_hpeer > 0 and not (r.ep_phx_h >= 1.5):
            return "notbusy"
        return "corr"
    c = c[c.left.eq("stuck")]                     # the stuck finding is the only one: the rule decides the status
    take(c, 2, "I1", why1)
    # I2: pulse vs normal count zones: 2 normal-mode zones no longer flagged + 1 pulse zone newly flagged
    c = base[has(base.rules_96, "N1") & ~has(base.rules, "N1") & base["mode"].eq("normal") & base.fn.eq("Count")]
    take(c, 2, "I2")
    c = base[~has(base.rules_96, "N1") & has(base.rules, "N1") & base["mode"].eq("pulse") & base.fn.eq("Count")]
    take(c, 1, "I2")
    # I3: N3 erratic time ON: 3 clear (old ok -> suspect), different types, + 1 borderline
    c = base[has(base.rules, "N3") & base.status.eq("ok") & base.new_status.eq("suspect")]
    clear = c[c.n3_exc >= 2 * c.n3_lim]
    take(clear, 3, "I3", lambda r: r.fn if r.fn in ("Advance", "Count") else "other")
    bd = c[c.n3_exc < 1.15 * c.n3_lim]
    take(bd, 1, "I3b")
    # I5: N5 (busier at night relative to its day, count-type zones)
    c = base[has(base.rules, "N5") & base.status.eq("ok")]
    take(c, 2, "I5")
    # I4: R1b (long ON, could be congestion -> watch)
    c = base[has(base.rules, "R1b")]
    take(c, 3, "I4")
    return rows


# ------------------------------------------------------------------ helpers
def wwin(r):
    s0, h = O.WIN[r.window]
    t0 = pd.Timestamp(s0)
    return t0, t0 + pd.Timedelta(hours=h), h


def hm(t):
    return pd.Timestamp(t).strftime("%a %H:%M")


def ons(ev, d, z0, z1):
    """start time and length (s) of every ON of channel d (82 -> next 81)."""
    x = ev[ev.EventId.isin([81, 82]) & (ev.Parameter == d) & (ev.Timestamp >= z0) & (ev.Timestamp < z1)]
    x = x.sort_values(["Timestamp", "EventId"], ascending=[True, False])
    t, e = x.Timestamp.to_numpy(), x.EventId.to_numpy()
    k = np.where((e[:-1] == 82) & (e[1:] == 81))[0]
    return pd.to_datetime(t[k]), (t[k + 1] - t[k]) / np.timedelta64(1, "s")


def ons_pkg(ev, d, z0, z1):
    """the package's continuous ONs (an ON that starts it -> the next OFF; ON events in between = logged again
    without an OFF): start, length (s), number of ON events without their OFF inside."""
    x = ev[ev.EventId.isin([81, 82]) & (ev.Parameter == d) & (ev.Timestamp >= z0 - 2 * H1) & (ev.Timestamp < z1)]
    x = x.sort_values(["Timestamp", "EventId"], ascending=[True, False])
    t, e = x.Timestamp.to_numpy(), x.EventId.to_numpy()
    offs = np.where(e == 81)[0]
    st = np.where((e == 82) & np.r_[True, e[:-1] != 82])[0]
    k = np.searchsorted(offs, st)
    ok = k < len(offs)
    st, k = st[ok], k[ok]
    end = offs[k]
    du = (t[end] - t[st]) / np.timedelta64(1, "s")
    m = t[st] >= np.datetime64(z0)
    return pd.to_datetime(t[st][m]), du[m], (end - st - 1)[m]


def info(r):
    x = pd.read_parquet(OUT / "resolved104.parquet", columns=["DeviceId", "window", "detector", "phase", "fn",
                                                              "status"],
                        filters=[("DeviceId", "==", r.DeviceId), ("window", "==", r.window)])
    return x.set_index("detector")


def partner(I, d, n_w, ix, exclude=()):
    """best-matching healthy detector on the same phase (15-min counts in the sample), no twins."""
    p = I.phase.get(d, np.nan)
    best, bc = None, -2
    for k in I.index:
        if k == d or k in exclude or k not in ix or I.status.get(k) == "bad":
            continue
        if not (np.isfinite(p) and I.phase.get(k) == p):
            continue
        a, b = n_w[ix[d]], n_w[ix[k]]
        m = np.isfinite(a) & np.isfinite(b)
        if m.sum() < 4 or np.std(a[m]) == 0 or np.std(b[m]) == 0:
            continue
        if np.abs(a[m] - b[m]).sum() < 0.1 * (a[m] + b[m]).sum():
            continue                                                  # twin: same zone on two inputs
        c = np.corrcoef(a[m], b[m])[0, 1]
        if c > bc:
            best, bc = k, c
    return best


def lab(I, k):
    return f"det {k} ({FN.get(I.fn.get(k), '?')}, same phase)"


def choppy_bins(ev, r, I):
    """the package's chop15 per 15-min bin: expected = local share (+-2 h without the bin) x reference."""
    t0, t1, _ = wwin(r)
    ew = ev[(ev.Timestamp >= t0) & (ev.Timestamp < t1)]
    B = hc.events_to_bins(ew, t0, t1, bin_s=300)
    n = B["n_on"].astype(float)
    cov = B["cov"]
    dets = B["dets"]
    live = n[:, cov].sum(1) > 0
    ph = {int(k): v for k, v in I.phase.items()}
    tw = hc.twins(n, cov)
    ref, kind, _ = hc._refs(dets, live, ph, tw)
    a = hc._agg(n, 3)
    ok = hc._agg(cov[None].astype(float), 3)[0] == 3
    i = int(np.where(dets == r.detector)[0][0])
    x = np.where(ok, a[i], 0.0)
    rr = np.where(ok, a[ref[i]].sum(0), 0.0)
    h = 8
    X = hc._movsum(x, h) - x
    Rr = hc._movsum(rr, h) - rr
    s = X / np.maximum(Rr, 1e-9)
    e, v = s * rr, s * rr * (1 + s)
    m = ok & (Rr > 0) & (X + x > 0) & ((e + x) > 0)
    con = np.where(m, (x - e) ** 2, 0.0)
    D = con[m].sum() / max(v[m].sum(), 1e-9)
    tt = t0 + pd.to_timedelta(np.arange(len(x)) * 900 + 450, unit="s")
    refd = [int(dets[j]) for j in ref[i]]
    return dict(tt=tt, x=x, e=e, con=con, D=D, ref=refd, kind=kind[i])


# ------------------------------------------------------------------ charts
def panel_counts(ax, z0, z1, ser, shade, sample, marks=None):
    C3.line_chart(ax, z0, z1, ser, "actuations per 15 minutes", shade=shade, sample=sample)
    if marks is not None and len(marks[0]):
        ax.scatter(marks[0], marks[1], s=170, facecolors="none", edgecolors=MARK, linewidths=2.2, zorder=5)
        lg = ax.get_legend()
        h_, l_ = list(lg.legend_handles), [t.get_text() for t in lg.get_texts()]
        h_.append(plt.Line2D([], [], marker="o", ls="", mfc="none", mec=MARK, mew=2.2, ms=11))
        l_.append(marks[2])
        ax.legend(h_, l_, frameon=False, loc="lower left", bbox_to_anchor=(0, 1.0), ncol=3, fontsize=10.5)


def fig_n(npan):
    hr = [1, 1, 0.8][:npan]
    fig, axes = plt.subplots(npan, 1, figsize=(13, 3.9 * npan + 0.6), sharex=True,
                             gridspec_kw=dict(height_ratios=hr))
    return fig, np.atleast_1d(axes)


def finish(fig, axes, title, sub, path):
    for ax in axes:
        lg = ax.get_legend()
        if lg is not None:
            h_, l_ = list(lg.legend_handles), [t.get_text() for t in lg.get_texts()]
            ax.legend(h_, l_, ncol=2, fontsize=10.5, frameon=False, loc="lower left", bbox_to_anchor=(0, 1.0))
    for ax in axes[:-1]:
        ax.set_xlabel("")
    fig.suptitle(title, x=0.01, ha="left", fontsize=16, fontweight="bold", color=INK)
    fig.text(0.01, 1 - 0.55 / fig.get_figheight(), sub, ha="left", va="top", fontsize=12, color=MUTED)
    fig.tight_layout(rect=(0, 0, 1, 1 - (0.95 + 0.25 * sub.count("\n")) / fig.get_figheight()))
    fig.savefig(path, dpi=100)
    plt.close(fig)


def flagged(r, t0, t1, ev):
    k = r.kind
    if k in ("I1", "I4") or (k == "KEEP" and r.v2rule in ("R1", "V1")):
        return pd.Timestamp(r.ep_t0), pd.Timestamp(r.ep_t1), "the long ON"
    if k == "KEEP" and r.v2rule == "R3":
        return t0 + pd.Timedelta(minutes=5 * int(float(r.level_b))), t1, "after the drop"
    if k == "KEEP" and r.v2rule == "R4":
        ix, tt, n, _ = RV.bins(ev, t0, t1, 300)
        j = int(np.nanargmax(n[ix[int(r.detector)]]))
        a = t0 + pd.Timedelta(minutes=5 * j)
        return a, a + pd.Timedelta(minutes=5), "the busiest 5 minutes"
    return t0, t1, "the sample the check looked at"


def spike_marks(r, tt, occ_d):
    S = pd.read_parquet(OUT / "n3_spike_bins.parquet", filters=[("DeviceId", "==", r.DeviceId),
                                                                ("window", "==", r.window)])
    S = S[S.detector == r.detector]
    t0, _, _ = wwin(r)
    bs = 300 if r.window.startswith("m30") else 900
    times = t0 + pd.to_timedelta(S.b * bs + bs / 2, unit="s")
    tts = pd.to_datetime(tt)
    y = [occ_d[np.argmin(np.abs((tts - t).total_seconds()))] for t in times]
    return list(times), y, S


def chart(r, sig, png, dpng):
    d = int(r.detector)
    ev = RV.events(r.DeviceId)
    t0, t1, h = wwin(r)
    cov0, cov1 = ev.Timestamp.min().floor("15min"), ev.Timestamp.max().ceil("15min")
    I = info(r)
    p = I.phase.get(d, np.nan)
    model = f"P{int(p)} {FN.get(I.fn.get(d), '?')}" if np.isfinite(p) else f"no phase {FN.get(I.fn.get(d), '?')}"
    z0, z1 = (cov0, cov1) if h >= 24 else (max(cov0, t0 - 3 * H1), min(cov1, t1 + 3 * H1))
    ixw, _, nw, _ = RV.bins(ev, t0, t1, 900)
    ix, tt, n, occ = RV.bins(ev, z0, z1, 900)
    bad = set(I.index[I.status.eq("bad")]) - {d}
    pt = partner(I, d, nw, ixw)
    a, b, flab = flagged(r, t0, t1, ev)
    shade = [(a, b, flab)]
    extra = {}
    k, rule = r.kind, r.v2rule
    sub = f"Signal {sig}   |   det {d}: {model}"
    on_panel = k in ("I2", "I3", "I3b") or (k == "KEEP" and rule == "N1")
    if k == "KEEP" and rule == "R9":
        return chart_r9(r, sig, png, dpng, ev, I, model)
    fig, axes = fig_n(3 if on_panel else 2)
    ser = [(f"det {d}", tt, n[ix[d]], THIS, 2.6)]
    if pt is not None and pt in ix:
        ser.append((lab(I, pt), tt, n[ix[pt]], PART, 2.0))
    marks = None
    sc = pd.to_numeric(r.s_choppy, errors="coerce")
    if k == "KEEP" and rule in ("R2", "R7") and np.isfinite(sc) and sc >= .35:
        cb = choppy_bins(ev, r, I)
        top = np.argsort(cb["con"])[::-1][:3]
        marks = (list(cb["tt"][top]), list(cb["x"][top]), "the 3 periods the check objects to most")
        extra["cb"], extra["top"] = cb, top
        ser = [ser[0], (f"what the check expected det {d} to count (in the sample)", cb["tt"],
                        np.where(cb["e"] > 0, cb["e"], np.nan), PART, 2.0)]
    panel_counts(axes[0], z0, z1, ser, shade, (t0, t1), marks)
    # % ON panel
    if k in ("I1", "I4") or (k == "KEEP" and rule in ("R1", "V1")):
        peers = [q for q in I.index if q != d and q not in bad and np.isfinite(p) and I.phase.get(q) == p and q in ix]
        held = []
        for q in peers:
            ep = (pd.to_datetime(tt) >= a) & (pd.to_datetime(tt) <= b)
            if ep.any() and (occ[ix[q]][ep] >= 99).mean() >= .8:
                held.append(q)
        peers = [q for q in peers if q not in held]
        ser2 = [(f"det {d}", tt, occ[ix[d]], THIS, 2.6)]
        if peers:
            ser2.append((f"healthy detectors on its phase, average (det {', '.join(map(str, peers))})", tt,
                         np.nanmean(occ[[ix[q] for q in peers]], 0), PART, 2.0))
        extra["peers"], extra["held"] = peers, held
    else:
        ser2 = [(f"det {d}", tt, occ[ix[d]], THIS, 2.6)]
        if pt is not None and pt in ix:
            ser2.append((lab(I, pt), tt, occ[ix[pt]], PART, 2.0))
    C3.line_chart(axes[1], z0, z1, ser2, "% of each 15 min ON", shade=shade, sample=(t0, t1), ylim=(0, 102))
    if k in ("I3", "I3b") or (k == "KEEP" and rule == "N1" and r.n3_n >= 2 and r.fn in M.COUNT_T):
        tm, ym, S = spike_marks(r, tt, occ[ix[d]])
        if len(tm):
            axes[1].scatter(tm, ym, s=170, facecolors="none", edgecolors=MARK, linewidths=2.2, zorder=5)
            lg = axes[1].get_legend()
            h_, l_ = list(lg.legend_handles), [t.get_text() for t in lg.get_texts()]
            h_.append(plt.Line2D([], [], marker="o", ls="", mfc="none", mec=MARK, mew=2.2, ms=11))
            l_.append("ON much longer than its count explains")
            axes[1].legend(h_, l_, frameon=False, loc="lower left", bbox_to_anchor=(0, 1.0), ncol=2, fontsize=10.5)
        extra["S"] = S
        Om = pd.read_parquet(OUT / "n3_missing_off.parquet")
        Om = Om[(Om.DeviceId == r.DeviceId) & (Om.window == r.window) & (Om.detector == d)]
        if len(Om):
            extra["split"] = (Om.t_missoff.sum(), Om.t_longpair.sum())
    if on_panel:
        st, du, mo = ons_pkg(ev, d, z0, z1)
        ax = axes[2]
        ax.axvspan(max(a, z0), min(b, z1), color=C3.SHADE, lw=0, zorder=0)
        q = mo == 0
        ax.scatter(st[q], np.maximum(du[q], 0.05), s=6, color=THIS, alpha=.5, zorder=3,
                   label=f"det {d}: one dot per ON (ON to OFF)")
        if (~q).any():
            ax.scatter(st[~q], np.maximum(du[~q], 0.05), s=22, marker="x", color=MARK, zorder=4,
                       label="ON logged again before any OFF (missing OFF): counted as ON until the next OFF")
        ax.set_yscale("log")
        ax.set_ylim(0.05, max(10, np.nanmax(du) * 1.5) if len(du) else 10)
        ax.set_yticks([0.1, 0.2, 1, 10, 100, 1000])
        ax.set_yticklabels(["0.1", "0.2", "1", "10", "100", "1000"])
        ax.set_ylabel("length of each ON (s)")
        C3.time_axis(ax, z0, z1)
        C3.style(ax)
        ax.legend(frameon=False, loc="lower left", bbox_to_anchor=(0, 1.0), fontsize=10.5, ncol=2)
        extra["dur"] = (st, du, mo)
    title = title_of(r, d, extra, I)
    finish(fig, axes, title, sub, png)
    day_chart(r, sig, d, model, ev, I, pt, a, b, flab, dpng, t0, t1)
    return title, model, pt, extra


def chart_r9(r, sig, png, dpng, ev, I, model):
    """v2 row 15: zoom on the silent stretch the OLD check saw, with the old yardstick (its phase mates, two of which
    are broken) and a healthy detector."""
    d = int(r.detector)
    t0, t1, _ = wwin(r)
    p1 = pd.read_parquet(M.IN96 / "health.parquet", filters=[("DeviceId", "==", r.DeviceId),
                                                             ("window", "==", r.window)]).set_index("detector")
    q = p1.loc[d]
    a = t0 + pd.Timedelta(seconds=300 * int(q.drop_b0))
    b = t0 + pd.Timedelta(seconds=300 * int(q.drop_b1))
    mates = [k for k in p1.index if k != d and p1.pred_phase.get(k) == p1.pred_phase.get(d)]
    brok = [k for k in mates if p1.status.get(k) == "bad"]
    z0, z1 = a - 2 * H1, b + 2 * H1
    ix, tt, n, occ = RV.bins(ev, z0, z1, 300)
    fig, axes = fig_n(2)
    ser = [(f"det {d}", tt, n[ix[d]], THIS, 2.6),
           (f"its phase mates combined (det {', '.join(map(str, mates))}) - the old check's yardstick", tt,
            np.nansum(n[[ix[k] for k in mates if k in ix]], 0), PART, 2.0)]
    C3.line_chart(axes[0], z0, z1, ser, "actuations per 5 minutes", shade=[(a, b, "where the old check said 'silent'")])
    ser2 = [(f"det {k}", tt, n[ix[k]], c, 2.0) for k, c in zip(mates, ("#7a7a75", "#a03ca0", "#2a9d8f", "#c49a00"))
            if k in ix]
    C3.line_chart(axes[1], z0, z1, ser2, "actuations per 5 minutes", shade=[(a, b, "where the old check said 'silent'")])
    lam = float(q.drop_lam)
    title = f"Det {d} quiet {hm(a)}-{b:%H:%M}, zoomed in (5-min counts)"
    sub = (f"Signal {sig}   |   det {d}: {model}   |   bottom: each phase mate on its own; det "
           f"{', '.join(map(str, brok))} are themselves flagged bad")
    finish(fig, axes, title, sub, png)
    I2 = info(r)
    ixw, _, nw, _ = RV.bins(ev, t0, t1, 900)
    pt = partner(I2, d, nw, ixw, exclude=brok)
    day_chart(r, sig, d, model, ev, I2, pt, a, b, "where the old check said 'silent'", dpng, t0, t1)
    got = {k: float(np.nansum(n[ix[k]][(pd.to_datetime(tt) >= a) & (pd.to_datetime(tt) < b)])) for k in mates if k in ix}
    own = float(np.nansum(n[ix[d]][(pd.to_datetime(tt) >= a - H1) & (pd.to_datetime(tt) < a)]))
    return title, model, pt, dict(r9=(a, b, lam, mates, brok, got, own))


def day_chart(r, sig, d, model, ev, I, pt, a, b, flab, path, t0, t1):
    day = pd.Timestamp(a).normalize()
    y0, y1 = day, day + pd.Timedelta(days=1)
    ix2, tt2, n2, occ2 = RV.bins(ev, y0, y1, 900)
    cov1 = ev.Timestamp.max().ceil("15min")
    pl = lab(I, pt) if pt is not None and pt in ix2 else None
    sub2 = f"Signal {sig}   |   det {d}: {model}   |   real counts, nothing scaled"
    if cov1 < y1 - pd.Timedelta(minutes=15):
        sub2 += f"\nThe saved data end at {cov1:%H:%M}; nothing after that is shown."
    RV.two_panel(path, f"Det {d} - actuations and time ON per 15 min, {day:%a %d %b}", sub2, y0, y1, 900, ix2, tt2,
                 n2, occ2, d, n2[ix2[pt]] if pl else None, pl, occ2[ix2[pt]] if pl else None, pl,
                 [(max(a, y0), min(b, y1), flab)], (t0, t1), "actuations per 15 minutes")


# ------------------------------------------------------------------ text
def old_txt(r):
    if r.kind == "KEEP" and r.v2rule == "R9":
        return "suspect (goes silent)"
    s = {k: float(r[f"s_{k}"]) for k in RV.CHK if pd.notna(r.get(f"s_{k}")) and float(r[f"s_{k}"]) >= .35}
    mc = max(s, key=s.get) if s else None
    return ST.get(r.status, r.status) + (f" ({CHK[mc]})" if mc and r.status in ("suspect", "bad") else "")


def new_txt(r):
    st = ST.get(r.new_status, r.new_status)
    why = []
    left = [w for w in str(r.left).split(",") if w and w != "nan"]
    wat = [w for w in str(r.watch).split(",") if w and w != "nan"]
    if r.new_status in ("suspect", "bad") and left:
        why = [CHK.get(w, w) for w in left]
    elif r.new_status == "watch" and wat:
        why = [CHK.get(w, w) for w in wat]
    return st + (f" ({', '.join(why)})" if why else "")


def pct(x):
    return f"{100 * x:.0f} %"


def what_of(r):
    k, rule = r.kind, r.v2rule
    if k == "I1":
        return "Stuck on during a queue? (rule now stricter)"
    if k == "I2":
        return "Count zone set to pulse or normal (now judged separately)"
    if k in ("I3", "I3b"):
        return "NEW check: erratic time ON (counts fine, ON time spiky)"
    if k == "I4":
        return "Long ON in heavy traffic: possible congestion"
    if k == "I5":
        return "NEW note: busy at night compared with its day"
    return {"R2": "Counts called erratic, time ON follows traffic", "R3": "Count drops, but its time ON kept up",
            "R4": "Too many in 5 min, but it spans 2 lanes", "R7": "Just over a limit, nothing else wrong",
            "R9": "Flagged only because its neighbours are broken", "N1": "Count zone set to pulse or normal",
            "V1": "Long ON in heavy traffic: possible congestion"}[rule]


def title_of(r, d, x, I):
    k, rule = r.kind, r.v2rule
    if k in ("I1", "I4") or rule in ("R1", "V1"):
        return f"Det {d} held ON {float(r.ep_dur) / 60:.0f} min ({hm(r.ep_t0)}-{pd.Timestamp(r.ep_t1):%H:%M})"
    if "cb" in x:
        return f"Det {d}: counts vs its neighbour, with the periods the check objects to"
    if rule == "R3":
        return f"Det {d}: counts before and after {hm(pd.Timestamp(O.WIN[r.window][0]) + pd.Timedelta(minutes=5 * int(float(r.level_b))))}"
    if rule == "R4":
        return f"Det {d} had {float(r.max5):.0f} actuations in one 5 min"
    if k == "I2" or rule == "N1":
        st_, du, mo_ = x["dur"]
        t0_, t1_, _ = wwin(r)
        du = du[(st_ >= t0_) & (st_ < t1_)]
        return f"Det {d}: median ON {np.median(du):.1f} s - set to {'pulse' if np.median(du) <= .25 else 'normal'}"
    if k in ("I3", "I3b"):
        return f"Det {d}: counts look normal, time ON jumps"
    if k == "I5":
        return f"Det {d}: busy at night compared with its neighbour"
    return f"Det {d}"


def saw(r, x, I, pt):
    k, rule = r.kind, r.v2rule
    d = int(r.detector)
    ps = f"det {pt}" if pt is not None else "its neighbour"
    if k == "KEEP" and rule == "R9":
        a, b, lam, mates, brok, got, own = x["r9"]
        return (f"The old check saw det {d} silent {hm(a)}-{b:%H:%M} ({(b - a).total_seconds() / 60:.0f} min at night) "
                f"and expected about {lam:.0f} actuations, because its phase mates counted "
                f"{', '.join(f'det {k}: {got.get(k, 0):.0f}' for k in mates)} in that time - but det "
                f"{', '.join(map(str, brok))} are broken (flagged bad). In the hour before it counted {own:.0f}. "
                f"The chart now zooms on that night stretch (5-min counts). New logic: judged against healthy "
                f"detectors only, a 20-min lull at night is normal.")
    if "cb" in x:
        cb, top = x["cb"], x["top"]
        j = top[0]
        ex = "; ".join(f"{pd.Timestamp(cb['tt'][q]) - pd.Timedelta(minutes=7.5):%a %H:%M} counted {cb['x'][q]:.0f} where "
                       f"about {cb['e'][q]:.0f} were expected" for q in top)
        ee = cb["e"][cb["e"] > 0]
        nmed = float(np.median(ee)) if len(ee) else np.nan
        base = (f"How the 'erratic' number works: each 15 min is compared with what its usual share of the phase "
                f"traffic predicts (orange line: its share over the surrounding 2 h x the count of the other "
                f"detectors on its phase in that 15 min). At about {nmed:.0f} per 15 min chance alone moves a count "
                f"by about +-{np.sqrt(nmed):.0f}; here the typical miss is {np.sqrt(cb['D']):.1f}x that size. The "
                f"check squares this: {cb['D']:.1f} against a limit of {float(r.chop_lim):.1f} (the 'Nx' quoted last "
                f"time). The biggest misses (red circles): {ex}. ")
        if rule == "R2":
            return base + (f"So the line looks close by eye, yet it is statistically 'too jumpy'. It is a long zone "
                           f"and its time ON follows traffic ({r.c_occ_ref:.2f}), so the new logic calls it ok.")
        return base + "Only this one check fired, just over its limit, so the new logic gives a watch note only."
    if k == "KEEP" and rule == "R2":                  # corr row (v2 row 6)
        t0, t1, _ = wwin(r)
        ev = RV.events(r.DeviceId)
        ix, tt, n, _ = RV.bins(ev, t0, t1, 900)
        hr = pd.to_datetime(tt).hour
        nt = hr < 5
        own = np.nansum(n[ix[d]][nt]) / max(np.nansum(n[ix[d]]), 1)
        oth = np.nansum(n[ix[pt]][nt]) / max(np.nansum(n[ix[pt]]), 1) if pt is not None else np.nan
        return (f"Nothing is scaled now. 00:00-05:00 it made {pct(own)} of its day's actuations, {ps} {pct(oth)} - "
                f"so yes, it is relatively busier at night than its neighbour, which is why its counts correlate "
                f"poorly with the signal ({float(r['corr']):.2f}, normally about {float(r.corr_exp):.2f}). Its time ON "
                f"follows traffic ({r.c_occ_ref:.2f}). New logic: ok (long zone whose time ON follows traffic).")
    if k == "KEEP" and rule == "R3":
        t0, t1, _ = wwin(r)
        cp = t0 + pd.Timedelta(minutes=5 * int(float(r.level_b)))
        ev = RV.events(r.DeviceId)
        ix, tt, n, occ = RV.bins(ev, t0, t1, 900)
        tt = pd.to_datetime(tt)
        bf, af = tt < cp, tt >= cp
        hb, ha = bf.sum() / 4, af.sum() / 4
        o = lambda q, m: np.nansum(n[ix[q]][m])  # noqa: E731
        s = (f"Before {hm(cp)}: det {d} {o(d, bf) / hb:.0f} per hour, {ps} {o(pt, bf) / hb:.0f} per hour. After: det {d} "
             f"{o(d, af) / ha:.0f} per hour, {ps} {o(pt, af) / ha:.0f} per hour. ")
        return s + (f"So the 'drop' is really det {d} being very busy before {cp:%H:%M} (at night, when {ps} was "
                    f"quiet): relative to the phase traffic it counted {float(r.level_ratio):.0%} as much after as "
                    f"before (the check's limit is 15 %). Its time ON did keep up ({np.nanmean(occ[ix[d]][af]):.0f} % ON after vs "
                    f"{np.nanmean(occ[ix[d]][bf]):.0f} % before), so the new logic calls it ok.")
    if k == "KEEP" and rule == "R4":
        t0, t1, _ = wwin(r)
        ev = RV.events(r.DeviceId)
        z0 = t0 - 6 * H1
        ix, tt, n, _ = RV.bins(ev, z0, t0, 900)
        nt = np.nansum(n[ix[d]])
        no = np.nansum(n[ix[pt]]) if pt is not None and pt in ix else np.nan
        ix2, _, n2, _ = RV.bins(ev, t0, t1, 900)
        return (f"Nothing is scaled now. In the 6 h before the sample (00:00-06:00) det {d} counted {nt:.0f}, {ps} "
                f"{no:.0f}; during the 3-h sample det {d} {np.nansum(n2[ix2[d]]):.0f}, {ps} "
                f"{np.nansum(n2[ix2[pt]]):.0f}. Busiest 5 min: {float(r.max5):.0f} (limit 150; the classifier says it "
                f"spans {int(r.lanes)} lanes, healthy 2-lane detectors reach about 140). New logic: ok.")
    if k == "I2" or (k == "KEEP" and rule == "N1"):
        st, du, mo = x["dur"]
        t0, t1, _ = wwin(r)
        m = (st >= t0) & (st < t1)
        dd, mm = du[m], mo[m]
        md = np.median(dd)
        if md <= .25:
            lng = dd > 1
            nmo = int((lng & (mm > 0)).sum())
            return (f"Set to PULSE: {np.mean(dd <= .25):.0%} of its ONs in the sample last 0.2 s or less (median "
                    f"{md:.1f} s), as a pulse count zone should. But {lng.sum()} ONs lasted over 1 s (longest "
                    f"{dd.max():.0f} s" + (f"; {nmo} of them are ON events logged without an OFF, red x" if nmo else "")
                    + f") - a pulse zone should not hold ON. New logic compares pulse zones only with pulse zones: "
                    f"{new_txt(r)}.")
        return (f"Set to NORMAL (median ON {md:.1f} s, so ON for the whole vehicle). The old note compared it with "
                f"pulse-mode count zones too. New logic compares normal-mode count zones only with normal-mode ones: "
                f"{new_txt(r)}.")
    if k in ("I3", "I3b"):
        S = x.get("S")
        nb = int(r.n3_n)
        q = S.sort_values("exc", ascending=False).iloc[0] if S is not None and len(S) else None
        t0, _, _ = wwin(r)
        eg = ""
        if q is not None:
            tq = t0 + pd.Timedelta(seconds=int(q.b) * 900)
            eg = (f" E.g. {hm(tq)}: {q.n:.0f} actuations, ON {100 * q.occ:.0f} % of the 15 min where about "
                  f"{100 * q.n * float(r.dbar) / 900:.0f} % was expected.")
        mo_t, lp_t = x.get("split", (np.nan, np.nan))
        split = ""
        if np.isfinite(mo_t):
            split = (f" Where the long ON time comes from: ONs that really lasted long (ON to OFF) {lp_t / 60:.0f} min; "
                     f"ON events logged again without an OFF in between (missing OFF, red x) {mo_t / 60:.0f} min.")
        return (f"Its counts follow {ps}. But in {nb} 15-min periods it was ON at least 3x longer than its own count "
                f"explains (its usual ON {float(r.dbar):.1f} s), while the healthy detectors on its phase were not "
                f"busier than usual.{eg} Extra ON time not explained: {float(r.n3_exc):.0f} min (limit "
                f"{float(r.n3_lim):.0f} min = the most extreme 0.5 % of healthy count-type detectors).{split} "
                f"New: {new_txt(r)}.")
    if k == "I5":
        t0, t1, _ = wwin(r)
        ev = RV.events(r.DeviceId)
        ix, tt, n, _ = RV.bins(ev, t0, t1, 900)
        hr = pd.to_datetime(tt).hour
        nt, dy = hr < 5, (hr >= 7) & (hr < 19)
        f = lambda q: (np.nansum(n[ix[q]][nt]) / 5, np.nansum(n[ix[q]][dy]) / 12)  # noqa: E731
        a1, b1 = f(d)
        s = (f"00:00-05:00 it counted {a1:.0f} per hour, 07:00-19:00 {b1:.0f} per hour (night = {a1 / max(b1, 1e-9):.0%} "
             f"of day). ")
        if pt is not None and pt in ix:
            a2, b2 = f(pt)
            s += f"{ps}: {a2:.0f} vs {b2:.0f} per hour ({a2 / max(b2, 1e-9):.0%}). "
        return s + (f"For the signal as a whole night is about {float(r.sig_night_day):.0%} of day. The current check "
                    f"only looks when the night count is higher than the day count; the new note also looks at "
                    f"'much busier at night than the others' (3x). New: {new_txt(r)}.")
    if k in ("I1", "I4") or rule in ("R1", "V1"):
        peers = x.get("peers", [])
        held = x.get("held", [])
        s = (f"Held ON {float(r.ep_dur) / 60:.0f} min ({hm(r.ep_t0)}-{pd.Timestamp(r.ep_t1):%H:%M}); traffic on its phase "
             f"meanwhile {r.ep_refx:.1f}x its usual. ")
        if held:
            s += (f"Det {', '.join(map(str, held))} {'was' if len(held) == 1 else 'were'} held ON too and "
                  f"{'is' if len(held) == 1 else 'are'} not used as a yardstick. ")
        if k == "I4" or rule == "V1":
            return s + (f"In the hour before it was already {pct(r.pre_occ)} ON, its time ON follows the phase "
                        f"traffic ({r.c_occ_ref:.2f}) and it was never held ON in light traffic. New: watch only "
                        f"(possible congestion; was {old_txt(r)}).")
        if r.n_hpeer > 0:
            s += (f"Healthy detectors on its phase: {int(r.n_hpeer)}; they were {r.ep_phx_h:.1f}x their usual time ON "
                  f"(a queue needs 1.5x); its time ON matches theirs {r.c_occ_hpeer:.2f} (needs 0.7). ")
        else:
            s += "No healthy detector on its phase to compare with. "
        if r.light_full > 0:
            s += (f"It was also >= 90 % ON in {int(r.light_full)} other 15-min period(s) while traffic was light. ")
        cleared = has(pd.Series([r.rules]), "R1").iloc[0]
        return s + (f"New: {new_txt(r)}" + (" (queue)." if cleared else " - the old proposal called it ok (queue)."))
    return ""


# ------------------------------------------------------------------ main
def main():
    R = load()
    v2 = v2_answers()
    rows = select(R, v2)
    nm = RV.names()
    for p_ in (CH, DC):
        p_.mkdir(parents=True, exist_ok=True)
    out = []
    order = {"KEEP": 1, "I1": 2, "I4": 3, "I2": 4, "I3": 5, "I3b": 6, "I5": 7}
    rows = sorted(rows, key=lambda r: (order[r.kind], r.v2n if r.v2n else 99))
    for i, r in enumerate(rows, start=1):
        d = int(r.detector)
        sig = nm.get(r.DeviceId, r.DeviceId[:8])
        tag = r.kind if r.kind != "KEEP" else f"v2r{int(r.v2n)}"
        png = CH / f"{i:02d}_{tag}_{sig}_d{d}.png"
        dpng = DC / png.name
        title, model, pt, x = chart(r, sig, png, dpng)
        t0, t1, h = wwin(r)
        I = info(r)
        row = dict(n=i, what=what_of(r), signal=sig, det=f"det {d}: {model}",
                   sample=f"{'24 h' if h >= 24 else '3 h'}, {t0:%a %d %b %H:%M}", old=old_txt(r), new=new_txt(r),
                   saw=saw(r, x, I, pt), png=png, dpng=dpng, earlier=r.earlier, kind=r.kind, v2n=r.v2n,
                   dev=r.DeviceId, window=r.window, detector=d, title=title, partner=pt)
        out.append(row)
        print(i, tag, sig, d, r.window, "|", row["old"], "->", row["new"], "|", row["saw"][:160])
    pd.DataFrame(out).to_csv(OUT / "review_rows_v3.csv", index=False)
    write_xl(out)


def write_xl(rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cases"
    ws["A1"] = "Is this a real problem? Y / N / ?"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = ("Data: Sat 26 - Mon 28 Sep 2026. Rows 1-10: cases you could not judge last time, redrawn. Rows 11 on: "
                "examples of the rules you suggested. Nothing in the charts is scaled any more: every line is a real "
                "count (top) or a real % of time ON (middle). Blue = this detector; orange = ONE named detector on "
                "the same phase (or, where the legend says so, the average of the healthy ones). Some charts have a "
                "third panel with the length of every ON. Red circles = the periods the check objects to. Flagged "
                "period shaded. Day chart = the whole day. 'Old' = current package; 'New' = proposed logic.")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:M2")
    ws.row_dimensions[2].height = 76
    head = ["#", "What it is about", "Signal", "Det", "Sample", "Old result", "New result", "What the data shows",
            "Chart", "Day chart", "Your earlier comment", "Answer (Y / N / ?)", "Comment"]
    for j, h in enumerate(head, start=1):
        c = ws.cell(row=3, column=j, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2A78D6")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for k, r in enumerate(rows, start=4):
        vals = [r["n"], r["what"], r["signal"], r["det"], r["sample"], r["old"], r["new"], r["saw"], "chart",
                "day chart", r["earlier"] or None, None, None]
        for j, v in enumerate(vals, start=1):
            c = ws.cell(row=k, column=j, value=v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(row=k, column=9).hyperlink = f"{CH.name}/{r['png'].name}"
        ws.cell(row=k, column=10).hyperlink = f"{DC.name}/{r['dpng'].name}"
        for j in (9, 10):
            ws.cell(row=k, column=j).font = Font(color="0563C1", underline="single")
        ws.cell(row=k, column=11).font = Font(italic=True, color="555555")
        ws.cell(row=k, column=12).fill = PatternFill("solid", fgColor="FFF8DC")
    for col, w in zip("ABCDEFGHIJKLM", (4, 24, 8, 20, 16, 20, 22, 70, 8, 10, 40, 12, 36)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    wb.save(XL)
    print("saved", XL, len(rows), "rows")


if __name__ == "__main__":
    main()
