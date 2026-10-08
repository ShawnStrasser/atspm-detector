"""Note 38: 15-row spot-check of the health scorer for the user (review/spotcheck_health.xlsx).

Rebuilt 2026-09-28 for the scorer WITHOUT fault events (hb_nofault_eval.py -> nofault_eval.parquet).
One row per detector, training signals only (never locked_v2), whole Sept-2026 window
(Fri 18 16:15 - Mon 21 Sep 10:25).  Mix: bad (one per rule), suspect, healthy controls.
Each row gets a chart (review/spotcheck_health_charts/<signal>_d<det>.png): 15-min counts of the
detector (bold) vs the other detectors on its phase (thin) and the signal's total (scaled,
light), nights shaded; plus a second panel that shows the rule's evidence.
Phase and siblings come from the label table (evaluation only - the scorer never sees them).
"""
from __future__ import annotations

import sys
import urllib.parse
import warnings
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.dates as mdates  # noqa: E402
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
import pyarrow.dataset as ds  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font, PatternFill  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
import hb_data as H  # noqa: E402

warnings.filterwarnings("ignore")
OUT = H.REPO / "review" / "spotcheck_health.xlsx"
CH = H.REPO / "review" / "spotcheck_health_charts"
PDF = H.DCW / "cabinet" / "pdf"
EV = H.DCW / "official" / "stg" / "cache" / "events"
W0, WH = pd.Timestamp("2026-09-18 16:15"), 66.2
PREV = {"2C009", "2B358", "11026", "08154", "08056", "11001", "10022", "2B504", "04050", "03034", "2B036",
        "10015", "01063", "2B346", "03097"}             # signals of the first spot-check (already seen)
SIBC = ["#eb6834", "#1baf7a", "#eda100", "#e87ba4", "#4a3aa7", "#008300"]
DETC, TOTC, INK, MUTED = "#2a78d6", "#d9d8d3", "#0b0b0b", "#52514e"
RULE = {"s_dead": "dead", "s_dropout": "sudden silence", "s_stuck": "stuck on", "s_rapid": "rapid re-triggering",
        "s_chatter": "chatter", "s_corr": "does not follow traffic", "s_level": "drop in level",
        "s_night_day": "busier at night than by day", "s_erratic": "erratic counts", "s_volume": "volume"}


def fmt_t(t):
    return pd.Timestamp(t).strftime("%a %d %b %H:%M")


# ------------------------------------------------------------------------------------ data
def names_phases():
    v3 = pd.read_parquet(H.REPO / "research" / "labels" / "function_labels_v3.parquet",
                         columns=["DeviceId", "DeviceName", "detector", "phase_target"])
    v3["DeviceId"] = v3.DeviceId.str.lower()
    v3["detector"] = pd.to_numeric(v3.detector, errors="coerce")
    v3 = v3.dropna(subset=["detector"])
    v3["phase"] = pd.to_numeric(v3.phase_target.where(v3.phase_target.str.startswith("P")).str[1:], errors="coerce")
    return v3


def load_bins(dev):
    B = H.load_B("stg", dev)
    a, b = H.win_bins(B, str(W0), WH)
    return H.slice_B(B, a, b)


def events(dev, dets):
    t = ds.dataset(EV / f"DeviceId={dev}").to_table(
        filter=ds.field("EventId").isin([81, 82]) & ds.field("Parameter").isin([int(d) for d in dets]))
    e = t.to_pandas().drop_duplicates()
    e = e[(e.Timestamp >= W0) & (e.Timestamp < W0 + pd.Timedelta(hours=WH))]
    return e.sort_values(["Parameter", "Timestamp", "EventId"], ascending=[True, True, False])


def on_table(e, d):
    x = e[e.Parameter == d]
    on, off = x[x.EventId == 82].Timestamp.to_numpy(), x[x.EventId == 81].Timestamp.to_numpy()
    k = np.searchsorted(off, on, "left")
    dur = np.where(k < len(off), (off[np.minimum(k, len(off) - 1)] - on) / np.timedelta64(1, "s"), np.nan)
    return pd.DataFrame({"t": on, "dur": dur})


def same_ons(e, d1, d2, tol=0.5):
    """share of d1's ONs with a d2 ON within tol seconds (same zone on two inputs)."""
    x, y = on_table(e, d1).t.to_numpy(), np.sort(on_table(e, d2).t.to_numpy())
    if not len(x) or not len(y):
        return 0.0
    k = np.clip(np.searchsorted(y, x), 1, len(y) - 1)
    dt = np.minimum(np.abs(x - y[k - 1]), np.abs(y[k] - x)) / np.timedelta64(1, "s")
    return float((dt <= tol).mean())


def quiet_sibs(B, sib, b0, b1):
    """siblings whose rate inside [b0, b1) fell below a quarter of their rate in the 2 h before."""
    dets, out = list(B["dets"]), []
    for d in sib:
        x = B["n_on"][dets.index(d)]
        before = x[max(b0 - 24, 0):b0].mean() if b0 > 0 else x[b1:b1 + 24].mean()
        if before >= 1 and x[b0:b1].mean() < 0.25 * before:
            out.append(d)
    return out


# ------------------------------------------------------------------------------- sentences
def sentence(r, sib, B, P):
    """One plain sentence with numbers.  sib: {det: n_on} of the other detectors on the phase."""
    top = r.top
    ph = f"phase {int(r.phase)}"
    sibtxt = ", ".join(f"d{k} {v:,.0f}" for k, v in list(sib.items())[:3])
    if r.status == "ok":
        return (f"{r.n_on:,.0f} actuations in 66 h; rises and falls with the rest of the signal (15-min correlation "
                f"{r['corr']:.2f}) and with the other {ph} detectors ({sibtxt}); nothing unusual found.")
    if top == "s_dead":
        return (f"No actuations at all in 66 h, while the other {ph} detectors counted {sibtxt} "
                f"(about {r.dead_lam:.0f} would be expected even at a quarter of a typical detector's share).")
    if top == "s_dropout":
        i = list(B["dets"]).index(r.detector)
        b0, b1 = int(r.drop_b0), int(r.drop_b1)
        t0, t1 = B["start"] + pd.Timedelta(minutes=5 * b0), B["start"] + pd.Timedelta(minutes=5 * b1)
        inside = {k: B["n_on"][list(B["dets"]).index(k), b0:b1].sum() for k in list(sib)[:3]}
        end = "and stayed silent to the end of the log" if r.drop_to_end else f"until {fmt_t(t1)}"
        co = (f"; {int(r.co_silent)} other detector(s) on the signal went silent at the same time"
              if r.co_silent > 0 else "")
        q = quiet_sibs(B, sib, b0, b1)
        if q:
            co += ("; note that " + ", ".join(f"d{d}" for d in q) + " on the same phase also nearly stopped then "
                   "(a lull in that movement rather than a detector fault?)")
        return (f"Counted normally ({r.n_on:,.0f} actuations in total) but went silent at {fmt_t(t0)} {end} "
                f"({(b1 - b0) / 12:.1f} h), while in that time the other {ph} detectors counted ("
                + ", ".join(f"d{k} {v:,.0f}" for k, v in inside.items())
                + f"); about {r.drop_lam:.0f} actuations were expected from its usual share{co}.")
    if top == "s_stuck":
        sm = P["sib_max_on"]
        smt = f"{sm / 60:.1f} min" if sm >= 120 else (f"{sm:.0f} s" if sm >= 10 else f"{sm:.1f} s")
        return (f"One call was held ON for {r.dur_max / 60:.0f} min without a break (from {fmt_t(P['stuck_t'])}); "
                f"its usual ON lasts {P['med_on']:.1f} s; the longest ON among the other {ph} detectors was "
                f"{smt}.")
    if top in ("s_rapid", "s_chatter"):
        dup = (f" d{P['dup']} on the same phase carries almost the same actuations (same zone on two inputs?), "
               f"so it shows the same pattern; d{P['sib']} has {P['sib_lt1']:.0%}." if P.get("dup") else
               f" On the same phase d{P['sib']} has {P['sib_lt1']:.0%}.")
        return (f"{r.ioi_lt1:.0%} of its {r.n_on:,.0f} actuations start within 1 s of the previous one and "
                f"{r.burst_frac:.0%} come in bursts of 5 or more ({int(r.n_burst)} bursts); healthy detectors that "
                f"behave like this one (median ON {r.med_dur:.1f} s) stay under {P['lim1']:.0%} and "
                f"{P['limb']:.0%}.{dup}")
    if top == "s_corr":
        return (f"Its 15-min counts do not rise and fall with the traffic: correlation with the signal's total "
                f"{r['corr']:.2f} where about {r.corr_exp:.2f} is expected (d{P['sib']} on the same phase: "
                f"{P['sib_corr']:.2f}); {r.n_on:,.0f} actuations in 66 h.")
    if top == "s_level":
        tb = B["start"] + pd.Timedelta(minutes=5 * int(r.level_b))
        return (f"From {fmt_t(tb)} its share of the signal's traffic fell to {r.level_ratio:.0%} of what it was "
                f"before ({P['rate_before']:.0f} per hour before, {P['rate_after']:.0f} after); over the same split "
                + ", ".join(f"d{k}'s share went to {v:.0%}" for k, v in P["sib_ratio"].items())
                + " of its earlier share.")
    if top == "s_night_day":
        return (f"Counts more at night (00-05 h) than by day (07-19 h): {r.night_day:.1f}x its day rate, "
                f"while the rest of the signal runs at {r.sig_night_day:.2f}x - traffic cannot do that.")
    if top == "s_erratic":
        return (f"Its 15-min counts jump around {r.disp:.0f}x more than the signal's traffic would explain "
                f"(the other detectors are much steadier); {r.n_on:,.0f} actuations in 66 h.")
    return r.reason[:1].upper() + r.reason[1:] + "."


# ---------------------------------------------------------------------------------- charts
def shade_nights(ax, t0, t1):
    d = pd.Timestamp(t0).normalize() - pd.Timedelta(days=1)
    while d < t1:
        a, b = d + pd.Timedelta(hours=21), d + pd.Timedelta(hours=30)
        ax.axvspan(max(a, t0), min(b, t1), color="#eef0f5", zorder=0, lw=0) if b > t0 and a < t1 else None
        d += pd.Timedelta(days=1)


def style(ax):
    for s in ("top", "right"):
        ax.spines[s].set_visible(False)
    for s in ("left", "bottom"):
        ax.spines[s].set_color("#bbbbbb")
    ax.tick_params(colors=MUTED, labelsize=8)
    ax.grid(axis="y", color="#ececec", lw=0.6)


def chart(r, B, sib, e, P, path):
    dets = list(B["dets"])
    i = dets.index(r.detector)
    n15 = H.hc._agg(B["n_on"].astype(float))
    cov15 = H.hc._agg(B["cov"][None].astype(float))[0] > 0
    tt = B["start"] + pd.to_timedelta(np.arange(n15.shape[1]) * 15, unit="min")
    mask = lambda y: np.where(cov15, y, np.nan)  # noqa: E731
    second = r.status != "ok"
    fig, axs = plt.subplots(2 if second else 1, 1, figsize=(12, 7.2 if second else 4.2),
                            gridspec_kw={"height_ratios": [1.25, 1]} if second else None)
    ax = axs[0] if second else axs
    tot = n15.sum(0)
    ref = max(np.nanmax([n15[i].max()] + [n15[dets.index(k)].max() for k in sib]), 1)
    k = ref / max(tot.max(), 1)
    shade_nights(ax, tt[0], tt[-1])
    ax.fill_between(tt, 0, mask(tot * k), color=TOTC, lw=0, step="mid", zorder=1,
                    label=f"all {len(dets)} detectors on the signal (scaled x{k:.2f})")
    for j, (d, v) in enumerate(list(sib.items())[:6]):
        y = mask(n15[dets.index(d)])
        ax.plot(tt, y, color=SIBC[j], lw=1.0, zorder=2, label=f"d{d} (same phase)")
        jj = np.nanargmax(np.where(np.isfinite(y), y, -1))
        ax.annotate(f"d{d}", (tt[jj], y[jj]), xytext=(2, 2), textcoords="offset points", fontsize=7, color=MUTED)
    ax.plot(tt, mask(n15[i]), color=DETC, lw=2.4, zorder=3, label=f"d{r.detector} (this detector)")
    ax.set_ylabel("actuations per 15 min", fontsize=9, color=MUTED)
    ax.xaxis.set_major_locator(mdates.HourLocator(byhour=[0, 6, 12, 18]))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%a %H:%M"))
    ax.set_xlim(tt[0], tt[-1])
    ax.legend(fontsize=7.5, loc="upper left", frameon=False, ncol=3)
    ax.text(1.0, 1.01, "shaded = night 21:00-06:00, local time", transform=ax.transAxes, ha="right",
            fontsize=7.5, color=MUTED)
    style(ax)
    fig.suptitle(f"{r.signal}   detector {r.detector}   phase {int(r.phase)}   -   {r.status.upper()}: "
                 f"{P['short']}", fontsize=11.5, color=INK, x=0.01, ha="left", fontweight="bold")
    if second:
        a2 = axs[1]
        top = r.top
        if top == "s_dropout":
            b0, b1 = int(r.drop_b0), int(r.drop_b1)
            lo, hi = max(b0 - 36, 0), min(b1 + 36, B["n_on"].shape[1])
            t5 = B["start"] + pd.to_timedelta(np.arange(lo, hi) * 5, unit="min")
            a2.axvspan(t5[b0 - lo], t5[min(b1 - lo, len(t5) - 1)], color="#fde2d6", lw=0, zorder=0)
            for j, d in enumerate(list(sib)[:6]):
                a2.plot(t5, B["n_on"][dets.index(d), lo:hi], color=SIBC[j], lw=1.0, label=f"d{d}")
            a2.plot(t5, B["n_on"][i, lo:hi], color=DETC, lw=2.4, label=f"d{r.detector}")
            a2.set_title("zoom on the silence (shaded), 5-min counts", fontsize=9, loc="left", color=MUTED)
            a2.set_ylabel("actuations per 5 min", fontsize=9, color=MUTED)
            a2.xaxis.set_major_formatter(mdates.DateFormatter("%a %H:%M"))
        elif top == "s_dead":
            Bd = H.load_B("dec", r.DeviceId)
            ks = [r.detector] + list(sib)[:6]
            sept = [B["n_on"][dets.index(d)].sum() / (B["cov"].sum() / 12) for d in ks]
            dec = [(Bd["n_on"][list(Bd["dets"]).index(d)].sum() / (Bd["cov"].sum() / 12)
                    if Bd is not None and d in list(Bd["dets"]) else 0) for d in ks]
            x = np.arange(len(ks))
            if sum(dec) > 0:
                a2.bar(x - 0.2, dec, 0.38, color="#9bbbe6", label="Dec 2024 (older log)")
                a2.bar(x + 0.2, sept, 0.38, color=DETC, label="Sept 2026 (this window)")
            else:
                a2.bar(x, sept, 0.5, color=DETC, label="Sept 2026 (this window; no Dec 2024 log for this signal)")
            a2.set_xticks(x, [f"d{d}" + (" (this)" if d == r.detector else "") for d in ks])
            a2.set_ylabel("actuations per hour", fontsize=9, color=MUTED)
            a2.set_title("this detector and the others on its phase: actuations per hour", fontsize=9, loc="left",
                         color=MUTED)
        elif top == "s_stuck":
            for j, d in enumerate([P["sib"]]):
                o = on_table(e, d)
                a2.scatter(o.t, o.dur.clip(lower=0.1), s=4, color=SIBC[0], alpha=.5, label=f"d{d} (same phase)")
            o = on_table(e, r.detector)
            a2.scatter(o.t, o.dur.clip(lower=0.1), s=6, color=DETC, label=f"d{r.detector}")
            a2.set_yscale("log")
            a2.axhline(900, color="#e34948", lw=0.8, ls="--")
            a2.text(o.t.iloc[0], 1000, "15 min", fontsize=7, color="#e34948")
            a2.set_ylabel("length of each ON (s, log)", fontsize=9, color=MUTED)
            a2.set_xlim(tt[0], tt[-1])
            a2.set_title("how long each actuation stayed ON", fontsize=9, loc="left", color=MUTED)
            a2.xaxis.set_major_formatter(mdates.DateFormatter("%a %H:%M"))
        elif top in ("s_rapid", "s_chatter"):
            bins = np.logspace(-1, np.log10(600), 40)
            for d, c, lw in ((P["sib"], SIBC[0], 1.2), (r.detector, DETC, 2.4)):
                o = on_table(e, d)
                ioi = np.diff(o.t.to_numpy()) / np.timedelta64(1, "s")
                h, _ = np.histogram(ioi, bins)
                a2.step(bins[:-1], h / max(len(ioi), 1), where="post", color=c, lw=lw,
                        label=f"d{d}" + (" (this)" if d == r.detector else " (same phase)"))
            a2.axvspan(0.1, 1.0, color="#fde2d6", lw=0, zorder=0)
            a2.text(0.12, a2.get_ylim()[1] * 0.9, "< 1 s: too fast for\nseparate vehicles", fontsize=7.5,
                    color="#b04020")
            a2.set_xscale("log")
            a2.set_xlabel("time from one actuation to the next (s, log)", fontsize=9, color=MUTED)
            a2.set_ylabel("share of intervals", fontsize=9, color=MUTED)
            a2.set_title("gaps between successive ONs", fontsize=9, loc="left", color=MUTED)
        elif top in ("s_corr", "s_erratic"):
            for d, c, sz in ((P["sib"], SIBC[0], 8), (r.detector, DETC, 12)):
                y = n15[dets.index(d)][cov15]
                a2.scatter((tot - n15[i])[cov15], y, s=sz, color=c, alpha=.7,
                           label=f"d{d}" + (" (this)" if d == r.detector else " (same phase)"))
            a2.set_xlabel("all other detectors on the signal, per 15 min", fontsize=9, color=MUTED)
            a2.set_ylabel("this detector, per 15 min", fontsize=9, color=MUTED)
            a2.set_title("a healthy detector's counts climb with the signal's traffic (points along a line)",
                         fontsize=9, loc="left", color=MUTED)
        elif top == "s_level":
            sh = pd.Series(n15[i] / np.maximum(tot - n15[i], 1)).where(cov15).rolling(8, min_periods=4).mean()
            a2.plot(tt, 100 * sh, color=DETC, lw=2.4, label=f"d{r.detector}")
            for j, d in enumerate(list(sib)[:3]):
                s2 = pd.Series(n15[dets.index(d)] / np.maximum(tot - n15[dets.index(d)], 1)).where(cov15)
                a2.plot(tt, 100 * s2.rolling(8, min_periods=4).mean(), color=SIBC[j], lw=1, label=f"d{d}")
            a2.axvline(B["start"] + pd.Timedelta(minutes=5 * int(r.level_b)), color="#e34948", ls="--", lw=.8)
            a2.set_ylabel("% of the signal's other traffic (2-h mean)", fontsize=9, color=MUTED)
            a2.set_xlim(tt[0], tt[-1])
            a2.set_title("share of the signal's traffic over time", fontsize=9, loc="left", color=MUTED)
            a2.xaxis.set_major_formatter(mdates.DateFormatter("%a %H:%M"))
        elif top == "s_night_day":
            hr = tt.hour.to_numpy()
            for d, c, lw in [(r.detector, DETC, 2.4)] + [(d, SIBC[j], 1) for j, d in enumerate(list(sib)[:3])]:
                y = pd.Series(n15[dets.index(d)] * 4).where(cov15).groupby(hr).mean()
                a2.plot(y.index, y / max(y.max(), 1), color=c, lw=lw, label=f"d{d}")
            y = pd.Series(tot).where(cov15).groupby(hr).mean()
            a2.fill_between(y.index, 0, y / y.max(), color=TOTC, lw=0, label="whole signal")
            a2.set_xticks(range(0, 24, 3))
            a2.set_xlabel("hour of day", fontsize=9, color=MUTED)
            a2.set_ylabel("average per hour (each / its own peak)", fontsize=9, color=MUTED)
            a2.set_title("daily profile", fontsize=9, loc="left", color=MUTED)
        a2.legend(fontsize=7.5, frameon=False)
        style(a2)
    fig.tight_layout(rect=(0, 0, 1, 0.95))
    fig.savefig(path, dpi=110)
    plt.close(fig)


# -------------------------------------------------------------------------------- selection
def main():
    s = pd.read_parquet(H.HB / "nofault_eval.parquet")
    f = s[s.window.eq("full")].copy()
    locked = set(pd.read_csv(H.DCW / "official" / "locked_v2.csv").DeviceId.str.lower())
    folds = set(pd.read_csv(H.DCW / "folds_v4.csv").DeviceId.str.lower())
    f = f[f.DeviceId.isin(folds) & ~f.DeviceId.isin(locked)]
    v3 = names_phases()
    f = f.merge(v3[["DeviceId", "detector", "DeviceName", "phase"]], on=["DeviceId", "detector"], how="inner")
    f = f[f.phase.notna() & ~f.DeviceName.isin(PREV)].rename(columns={"DeviceName": "signal"})
    sc = [c for c in f if c.startswith("s_")]
    f["top"] = f[sc].fillna(0).idxmax(axis=1)
    # siblings on the same labelled phase that the log shows as live
    live = f[f.n_on > 0].groupby(["DeviceId", "phase"]).detector.apply(list).to_dict()
    f["n_live_sib"] = [len([d for d in live.get((a, p), []) if d != k]) for a, p, k in zip(f.DeviceId, f.phase, f.detector)]
    has_pdf = {q.name.split("_")[0] for q in PDF.glob("*.pdf")}
    f = f[(f.n_live_sib >= 1) & f.signal.isin(has_pdf)]          # the user needs the print to check
    rng = np.random.default_rng(3828)
    used, rows = set(), []

    def take(mask, k, kind):
        c = f[mask & ~f.signal.isin(used)]
        c = c.sample(frac=1, random_state=int(rng.integers(1e9))).drop_duplicates("signal")
        for _, r in c.head(k).iterrows():
            used.add(r.signal)
            rows.append((kind, r))

    bad, sus, ok = f.status.eq("bad"), f.status.eq("suspect"), f.status.eq("ok")
    take(bad & f.top.eq("s_dropout") & f.co_silent.eq(0) & (f.n_on > 300), 2, "bad")
    take(bad & f.top.eq("s_dead") & f.listed.astype(bool) & (f.dead_lam > 200), 1, "bad")
    take(bad & f.top.eq("s_stuck"), 1, "bad")
    take(bad & f.top.eq("s_rapid"), 1, "bad")
    take(bad & f.top.eq("s_corr"), 1, "bad")
    take(bad & f.top.eq("s_level"), 1, "bad")
    take(bad & f.top.eq("s_night_day"), 1, "bad")
    take(sus & f.top.eq("s_dropout") & (f.n_on > 300), 1, "suspect")
    take(sus & f.top.eq("s_rapid"), 1, "suspect")
    take(sus & f.top.eq("s_stuck"), 1, "suspect")
    take(sus & f.top.eq("s_erratic"), 1, "suspect")
    take(ok & (f.n_on > 1000) & (f["corr"] > .8), 3, "ok")
    CH.mkdir(parents=True, exist_ok=True)
    out = []
    for kind, r in rows:
        B = load_bins(r.DeviceId)
        dets = list(B["dets"])
        sibs = [d for d in v3[(v3.DeviceId == r.DeviceId) & (v3.phase == r.phase)].detector.astype(int)
                if d != r.detector and d in dets]
        sib = dict(sorted(((d, float(B["n_on"][dets.index(d)].sum())) for d in sibs), key=lambda x: -x[1]))
        sib = {k: v for k, v in sib.items() if v > 0} or sib
        e = events(r.DeviceId, [r.detector] + list(sib))
        P = {"sib": next(iter(sib))}
        P["short"] = {"ok": "nothing unusual"}.get(r.status, RULE.get(r.top, r.top))
        if r.status == "ok":
            r = r.copy()
            r["top"] = "none"
        if r.top == "s_stuck":
            o = on_table(e, r.detector)
            P["stuck_t"] = o.t[o.dur.idxmax()]
            P["med_on"] = o.dur.median()
            P["sib_max_on"] = max(on_table(e, d).dur.max() for d in sib)
            P["short"] += f" ({r.dur_max / 60:.0f} min single ON)"
        if r.top in ("s_rapid", "s_chatter"):
            dup = [d for d in sib if same_ons(e, r.detector, d) > 0.8]
            rest = [d for d in sib if d not in dup]
            if dup and rest:
                P["dup"], P["sib"] = dup[0], rest[0]
            g = H.hc.behaviour_group([r.pulse_frac], [r.med_dur])[0]
            lim = H.hc.RAPID_LIM[g]
            P["lim1"], P["limb"] = lim[1], lim[2]
            o = on_table(e, P["sib"])
            ioi = np.diff(o.t.to_numpy()) / np.timedelta64(1, "s")
            P["sib_lt1"] = float((ioi < 1).mean()) if len(ioi) else np.nan
            P["short"] += f" ({r.ioi_lt1:.0%} of ONs < 1 s apart)"
        if r.top in ("s_corr", "s_erratic"):
            n15 = H.hc._agg(B["n_on"].astype(float))
            c15 = H.hc._agg(B["cov"][None].astype(float))[0] == 3
            j = dets.index(P["sib"])
            P["sib_corr"] = float(np.corrcoef(n15[j][c15], (n15.sum(0) - n15[j])[c15])[0, 1])
        if r.top == "s_level":
            i = dets.index(r.detector)
            lb = int(r.level_b)
            P["rate_before"] = B["n_on"][i, :lb].sum() / max(B["cov"][:lb].sum() / 12, 1e-9)
            P["rate_after"] = B["n_on"][i, lb:].sum() / max(B["cov"][lb:].sum() / 12, 1e-9)
            P["short"] += f" (to {r.level_ratio:.0%} of its earlier share)"
            tot = B["n_on"].sum(0)
            P["sib_ratio"] = {}
            for d in list(sib)[:3]:
                x = B["n_on"][dets.index(d)]
                sh = lambda a, b: (x[a:b].sum() + .5) / max((tot - x)[a:b].sum(), 1)  # noqa: E731
                P["sib_ratio"][d] = sh(lb, None) / sh(0, lb)
        if r.top == "s_dropout":
            P["short"] += f" ({(r.drop_b1 - r.drop_b0) / 12:.1f} h)"
        what = sentence(r, sib, B, P)
        png = CH / f"{r.signal}_d{int(r.detector)}.png"
        chart(r, B, sib, e, P, png)
        pdfs = sorted(PDF.glob(f"{r.signal}_*.pdf"))
        out.append(dict(signal=r.signal, detector=int(r.detector), phase=int(r.phase), status=r.status, what=what,
                        png=png.name, pdf=pdfs[0] if pdfs else None, top=r.top, DeviceId=r.DeviceId))
        print(r.signal, r.detector, r.status, r.top, "|", what)
    df = pd.DataFrame(out).sample(frac=1, random_state=5).reset_index(drop=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "health spot-check"
    ws.append(["Signal", "Detector", "Phase", "Health status",
               "What was found (log Fri 18 Sep 16:15 - Mon 21 Sep 10:25, 2026)", "Chart", "Cabinet print",
               "Your answer (yes / no / ?)", "Comment"])
    link = Font(color="0563C1", underline="single")
    for n, r in enumerate(df.itertuples(index=False), start=2):
        ws.append([r.signal, r.detector, r.phase, r.status, r.what, "open chart",
                   "open print" if r.pdf is not None else "no print on file", "", ""])
        ws.cell(n, 6).hyperlink = f"spotcheck_health_charts/{r.png}"
        ws.cell(n, 6).font = link
        if r.pdf is not None:
            ws.cell(n, 7).hyperlink = "file:///" + urllib.parse.quote(str(r.pdf), safe=":\\/()_-.,'")
            ws.cell(n, 7).font = link
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDEBF7")
    for col, w in zip("ABCDEFGHI", (9, 9, 7, 13, 90, 12, 14, 14, 40)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=1):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    wb.save(OUT)
    df.drop(columns=["pdf"]).to_csv(H.HB / "spotcheck_rows.csv", index=False)


if __name__ == "__main__":
    main()
