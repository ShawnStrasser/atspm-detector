"""Note 46: follow-up spot-check of health v4 (review/spotcheck_health.xlsx), <= 8 rows, 66-h Sept-2026 window,
training signals only (never locked_v2).  Five rows answer the user's 2026-09-29 questions; three are new cases
of the v4 changes (recovery judged against the tracking partner, lane-aware rapid limits).  Charts reuse the
note-43 layout (h3_spotcheck.chart) with the v4 scorer and its out-of-fold inputs.  All earlier answers go to
one sheet "earlier answers" (with the date of the answer).

    python h4_spotcheck.py
"""
from __future__ import annotations

import sys
import urllib.parse
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
import h3_spotcheck as S  # noqa: E402
import h4_final as F  # noqa: E402
import hb_build as HB  # noqa: E402
import health_core as hc  # noqa: E402
import pyarrow.dataset as ds  # noqa: E402

ROWS = [  # signal, detector, chart kind, what was found (plain words for the user)
    ("2B146", 16, "choppy",
     "YOUR QUESTION (recurring spikes: traffic or fault?). The spikes come at the same times every day - about 06:30-07:30 "
     "and 17:30-18:30 - on Friday, Saturday, Sunday AND Monday, so the weekend has them too (unlikely for school traffic). "
     "The other phase-6 detectors do not show them. Inside the spikes the actuations look like normal vehicles: ON time "
     "0.2 s as usual, gaps median 4.6 s, only 1.4 % under 1 s (outside the spikes 1.0 %). No other check is triggered "
     "(no stuck, no silence, no rapid re-triggering). The function model, re-run on every 2-h window, calls it Count in "
     "all 33 windows. Now SUSPECT (was bad): a daily pattern with vehicle-like actuations is either traffic only its lane "
     "sees or a daily sensor effect such as low sun on a camera; with under 4 days of data the scorer will not clear it."),
    ("03033", 8, "stuck",
     "YOUR QUESTION (Sunday is not Saturday). Agreed: the recovery is now judged against d10, the detector it tracks "
     "1:1 (same lane), in the SAME hours instead of against the whole phase before vs after. After the 2.2-h stuck ON "
     "it counted 96 % of its usual share relative to d10. Now SUSPECT with only Sat 18:17-20:27 listed as bad (was bad)."),
    ("10045", 4, "stuck",
     "YOUR QUESTION (group failure). Every evening (Fri 19:20, Sat 18:00, Sun 18:15) d4 and d5 are held ON and d7 dips together while "
     "d6 on the same phase keeps counting normally (127 % of its usual share). A real queue would affect d6 too, so the "
     "reason now says: sensors failing together (e.g. one radar/camera unit or card), not a real queue; the periods are "
     "listed. It stays SUSPECT (drop the periods, keep the detector), as you suggested."),
    ("2B044", 14, "rapid",
     "YOUR QUESTION (fast actuations inside the choppy periods?). No: in its worst hour only 6-7 % of actuations "
     "come within 1 s of the previous one, well under the limit for detectors like it (16 %); over the whole sample "
     "2.4 %. Its counts are also not choppy against its phase partners (0.8, limit 5). It stays SUSPECT only for the "
     "15-min hold on Saturday 09:08, and afterwards it tracks d7 at its usual share."),
    ("07035", 19, "rapid",
     "YOUR QUESTION (does it span lanes?). The lane output says NO: d19 covers one lane (lane 2 of phase 6); the "
     "detector spanning lanes 1-2 on this phase is d18 (predicted Mid). So the single-lane limit applies. 11 % of its "
     "actuations come within 1 s of the previous one, and it is worst in the evening (24-30 % between 20:00 and "
     "midnight). Stays SUSPECT."),
    ("10018", 2, "silent",
     "NEW (recovery judged against the tracking partner). Silent Sat 22:45 - Sun 01:00 (2.2 h, about 200 actuations "
     "expected). Relative to d4, which it tracks, it came back at 52 % of its earlier share - just inside the 'clean "
     "recovery' band (50-200 %). Now SUSPECT with that period listed (was bad). Is 52 % really back to normal?"),
    ("10090", 22, "stuck",
     "NEW. Held ON with no OFF logged at about 08:30 on BOTH Saturday and Sunday (1.6 h each); after recovering it "
     "counted only 35 % of its usual share relative to d23, which it tracks. Now BAD (was suspect): a daily stuck "
     "loop that does not come back to its old level."),
    ("2B067", 42, "rapid",
     "NEW (lane-aware limits). Predicted Count spanning 2 lanes. Two side-by-side vehicles give two quick ONs, so it "
     "is now judged against healthy detectors that span 2 lanes: its quick re-actuations are within that range. Now OK "
     "(was suspect for rapid re-actuation)."),
]


def load(dev):
    ev = ds.dataset(S.EV / f"DeviceId={dev}").to_table(filter=ds.field("EventId").isin(list(hc.ALLOWED))).to_pandas()
    F.init(["final"])
    ph, fn, ln, pc = F.inputs("stg", dev)
    exp = HB.expected_channels().get(dev)
    h = hc.health(ev, S.T0, S.T1, exp, ph, fn, ln, pc)
    B = hc.events_to_bins(ev, S.T0, S.T1, exp)
    lab = {int(d): max(v, key=v.get) for d, v in fn.items()}
    return ev, h, B, ph, lab, ln


def main():
    if openpyxl.load_workbook(S.XL).worksheets[0].title == "health follow-up v4":
        raise SystemExit("already rebuilt for v4 - restore the earlier workbook first (it would count v4 rows as answers)")
    v3 = pd.read_parquet(S.H.REPO / "research" / "labels" / "function_labels_v3.parquet", columns=["DeviceId", "DeviceName"])
    nm = dict(zip(v3.DeviceName, v3.DeviceId.str.lower()))
    out = []
    for sig, d, kind, what in ROWS:
        dev = nm[sig]
        ev, h, B, ph, fn, ln = load(dev)
        x = h[h.detector == d].iloc[0].copy()
        x["signal"] = sig
        png = S.CH / f"v4_{sig}_d{d}.png"
        S.chart(x, ev, h, B, ph, fn, kind, png)
        pdfs = sorted(S.PDF.glob(f"{sig}_*.pdf"))
        p = ph.get(d)
        sp = ln.get(d)
        out.append(dict(signal=sig, detector=d, phase=int(p) if p == p and p is not None else None,
                        function=fn.get(d, "-"), lanes=("-" if sp != sp or sp is None else
                                                        f"spans {int(sp)}" if sp >= 2 else "1 lane"),
                        status=x.status, what=what, png=png.name, pdf=pdfs[0] if pdfs else None))
        print(sig, d, x.status, "|", x.reason[:160])
    old = openpyxl.load_workbook(S.XL)
    prev = []
    for ws0, date in ((old.worksheets[0], "2026-09-29"), (old.worksheets[1], "2026-09-28")):
        hdr = [c.value for c in ws0[1]]
        for r in range(2, ws0.max_row + 1):
            row = {h_: ws0.cell(r, j + 1).value for j, h_ in enumerate(hdr)}
            if row.get("Signal") is None:
                continue
            links = {h_: ws0.cell(r, j + 1).hyperlink.target for j, h_ in enumerate(hdr)
                     if ws0.cell(r, j + 1).hyperlink is not None}
            prev.append((date, row, links))
    wb = Workbook()
    ws = wb.active
    ws.title = "health follow-up v4"
    head = ["Signal", "Detector", "Phase", "Predicted function", "Lanes", "Health status",
            "What was found (log Fri 18 Sep 16:15 - Mon 21 Sep 10:25, 2026)", "Chart link", "Print link",
            "Your answer", "Comment"]
    ws.append(head)
    link = Font(color="0563C1", underline="single")
    for n, r in enumerate(out, start=2):
        ws.append([r["signal"], r["detector"], r["phase"], r["function"], r["lanes"], r["status"], r["what"],
                   "open chart", "open print" if r["pdf"] is not None else "no print on file", "", ""])
        ws.cell(n, 8).hyperlink = f"spotcheck_health_charts/{r['png']}"
        ws.cell(n, 8).font = link
        if r["pdf"] is not None:
            ws.cell(n, 9).hyperlink = "file:///" + urllib.parse.quote(str(r["pdf"]), safe=":\\/()_-.,'")
            ws.cell(n, 9).font = link
    w2 = wb.create_sheet("earlier answers")
    w2.append(["Answered", "Signal", "Detector", "Phase", "Health status then", "What was found then", "Chart link",
               "Print link", "Your answer", "Comment"])
    for n, (date, row, links) in enumerate(prev, start=2):
        g = lambda *ks: next((row[k] for k in ks if k in row and row[k] is not None), None)  # noqa: E731
        w2.append([date, g("Signal"), g("Detector"), g("Phase (predicted)", "Phase"), g("Health status"),
                   next((v for k, v in row.items() if str(k).startswith("What was found")), None),
                   "open chart", "open print", g("Your answer (yes / no / ?)"), g("Comment")])
        for c, key in ((7, "Chart"), (8, "Cabinet print")):
            if key in links:
                w2.cell(n, c).hyperlink = links[key]
                w2.cell(n, c).font = link
    for sh, widths in ((ws, (9, 9, 7, 12, 10, 12, 95, 11, 11, 14, 40)), (w2, (11, 9, 9, 7, 12, 80, 11, 11, 14, 60))):
        for j, w in enumerate(widths):
            sh.column_dimensions[chr(65 + j)].width = w
        for c in sh[1]:
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDEBF7")
        for row in sh.iter_rows(min_row=1):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        sh.freeze_panes = "A2"
    wb.save(S.XL)
    print("saved", S.XL, len(out), "rows;", len(prev), "earlier answers kept")


if __name__ == "__main__":
    main()
